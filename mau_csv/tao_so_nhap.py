"""Sinh các tệp Sổ nhập CLB (.xlsx) — mẫu trống và hai sổ ví dụ.

Chạy lại mỗi khi sửa `so_nhap.py` hoặc các bộ CSV ví dụ, để sổ và CSV không
lệch nhau:

    ./.venv/bin/python mau_csv/tao_so_nhap.py

Tạo ra:
    mau_csv/SO_NHAP_CLB.xlsx                         mẫu trống, phát cho trường
    mau_csv/vi_du_day_du/SO_NHAP_CLB_vi_du.xlsx      3 buổi · 9 CLB · 24 HS
    mau_csv/vi_du_ca_tuan/SO_NHAP_CLB_vi_du.xlsx     6 buổi · 18 CLB · 60 HS
    du_lieu_test/vi_du_huong_dan/SO_NHAP_VIDU.xlsx   bộ 10 em của HUONG_DAN_SU_DUNG.md
    du_lieu_test/bo_sach/SO_NHAP_SACH.xlsx           140 em, 0 cảnh báo
    du_lieu_test/SO_NHAP_TEST.xlsx                   120 em, cạnh tranh cao
    du_lieu_test/SO_NHAP_CO_LOI_CO_Y.xlsx            cố ý sai 5 chỗ: thử bảng lỗi

Các sổ ví dụ sinh THẲNG từ bộ ba tệp cũ cùng thư mục, nên nạp sổ hay nạp ba
tệp đều ra cùng một CSDL — `tests/test_so_nhap.py` canh điều đó.
"""

import csv
import os
import sys

import openpyxl

THU_MUC = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(THU_MUC))

import so_nhap  # noqa: E402

TEN_MAU = "SO_NHAP_CLB.xlsx"
TEN_VI_DU = "SO_NHAP_CLB_vi_du.xlsx"
BO_VI_DU = ("vi_du_day_du", "vi_du_ca_tuan")


GOC = os.path.dirname(THU_MUC)
DLT = os.path.join(GOC, "du_lieu_test")

# (thư mục, tiền tố bộ ba tệp cũ, đuôi, tên sổ)
BO_DU_LIEU_TEST = [
    (os.path.join(DLT, "vi_du_huong_dan"), "VIDU", ".csv", "SO_NHAP_VIDU.xlsx"),
    (os.path.join(DLT, "bo_sach"), "SACH", ".csv", "SO_NHAP_SACH.xlsx"),
    (DLT, "TEST", ".xlsx", "SO_NHAP_TEST.xlsx"),
]
TEN_BA_TEP = ("01_danh_sach_CLB", "02_chon_CLB_muon_thi", "03_xep_hang_nguyen_vong")


def _doc(duong_dan):
    """Một tệp cũ (.csv, hoặc .xlsx đọc trang đầu) -> [dict]."""
    if duong_dan.endswith(".csv"):
        with open(duong_dan, encoding="utf-8-sig", newline="") as f:
            return list(csv.DictReader(f))
    wb = openpyxl.load_workbook(duong_dan, read_only=True, data_only=True)
    dong = [[so_nhap.o_thanh_chu(v) for v in r]
            for r in wb.worksheets[0].iter_rows(values_only=True)]
    wb.close()
    dong = [r for r in dong if any(r)]
    return [dict(zip(dong[0], r)) for r in dong[1:]]


def _tu_ba_tep(thu_muc, ten, duoi=".csv", tien_to=""):
    clubs, students, ngoai_nv = so_nhap.tu_csv(*(
        _doc(os.path.join(thu_muc, tien_to + t + duoi)) for t in ten))
    if ngoai_nv:
        raise SystemExit(f"{thu_muc}: thi CLB không có trong nguyện vọng: {ngoai_nv}")
    return clubs, students


def tao_vi_du(ten_bo: str) -> str:
    goc = os.path.join(THU_MUC, ten_bo)
    clubs, students = _tu_ba_tep(goc, TEN_BA_TEP)
    ra = os.path.join(goc, TEN_VI_DU)
    so_nhap.ghi_so_nhap(ra, clubs, students)
    return ra


def tao_mau_trong() -> str:
    ra = os.path.join(THU_MUC, TEN_MAU)
    so_nhap.ghi_so_nhap(ra)
    return ra


def tao_du_lieu_test() -> list:
    ra = []
    for thu_muc, tien_to, duoi, ten in BO_DU_LIEU_TEST:
        clubs, students = _tu_ba_tep(thu_muc, TEN_BA_TEP, duoi, tien_to + "_")
        p = os.path.join(thu_muc, ten)
        so_nhap.ghi_so_nhap(p, clubs, students)
        ra.append(p)
    return ra


# Sổ cố ý sai: từ bộ 10 em, mỗi lỗi một chỗ, để thử bảng lỗi của màn hình
# nạp. (sheet, ô, giá trị mới, lỗi phần mềm phải báo)
LOI_CO_Y = [
    (so_nhap.SHEET_HS, "D3", "CLB Bóng chuyền", "so_nhap_clb_khong_co"),
    (so_nhap.SHEET_HS, "E4", "tám", "so_nhap_diem_sai"),
    (so_nhap.SHEET_HS, "A6", "HS04", "so_nhap_hs_trung"),
    (so_nhap.SHEET_HS, "F8", "=D8", "so_nhap_nv_trung"),
    (so_nhap.SHEET_CLB, "C3", 0, "so_nhap_chi_tieu_sai"),
]


def tao_so_co_loi() -> str:
    clubs, students = _tu_ba_tep(os.path.join(DLT, "vi_du_huong_dan"), TEN_BA_TEP,
                                 ".csv", "VIDU_")
    p = os.path.join(DLT, "SO_NHAP_CO_LOI_CO_Y.xlsx")
    so_nhap.ghi_so_nhap(p, clubs, students)
    wb = openpyxl.load_workbook(p)
    for sheet, o, gia_tri, _ in LOI_CO_Y:
        # "=D8": chép đúng tên ở NV1 sang NV2 — chọn một CLB hai lần.
        if isinstance(gia_tri, str) and gia_tri.startswith("="):
            gia_tri = wb[sheet][gia_tri[1:]].value
        wb[sheet][o] = gia_tri
    wb.save(p)
    return p


if __name__ == "__main__":
    for p in [tao_mau_trong()] + [tao_vi_du(bo) for bo in BO_VI_DU] \
            + tao_du_lieu_test() + [tao_so_co_loi()]:
        print("da tao", os.path.relpath(p, GOC))
