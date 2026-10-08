"""Sổ nhập CLB — MỘT tệp Excel thay cho ba tệp nạp cũ.

Ban giám khảo nhận xét bước nhập quá rối: ba tệp riêng, tên cột tiếng Anh,
gõ tay mã CLB (gõ sai một ký tự là mất cả học sinh), và mẫu .xlsx chỉ có bản
một buổi. Sổ nhập gom tất cả về một tệp, hai sheet, tiêu đề tiếng Việt:

  `1. CLB`        Tên CLB · Buổi · Chỉ tiêu · Suất ưu tiên · Nhóm ưu tiên · Mã CLB
  `2. Học sinh`   Mã HS · Họ tên · Nhóm ưu tiên · NV1 · Điểm 1 · NV2 · Điểm 2 · …

Ô NV là danh sách thả xuống tên CLB. Ô Điểm đứng ngay cạnh NV cùng số:
có điểm (hoặc chữ `thi`) nghĩa là em đã dự thi CLB đó; trống là không thi.
Vì thế tệp "chọn CLB muốn thi" cũ biến mất — nó chỉ còn là các ô điểm.

Một danh sách NV chung cho mọi buổi: buổi của từng nguyện vọng suy ra từ cột
Buổi của CLB, và thuật toán luôn lọc nguyện vọng theo buổi trước khi dùng
(xem `import_preferences_csv`). Nhờ vậy mẫu KHÔNG phụ thuộc trường có mấy buổi.

Tệp này là NGUỒN DUY NHẤT mô tả sổ nhập. Bốn nơi dùng chung nó: phần mềm
(api_nhap.py — đọc sổ), mẫu (mau_csv/tao_so_nhap.py), xuất dữ liệu đầu vào
(api_xuat.py) và hai skill sinh/xử lí dữ liệu. Sửa định dạng thì chỉ sửa ở đây.

Không phụ thuộc gì ngoài openpyxl và i18n_errors — để skill đóng gói được.
"""

from __future__ import annotations

import csv
import io
import math
import re
import unicodedata

from i18n_errors import err

# ---------------------------------------------------------------------------
# Tên sheet, tiêu đề cột
# ---------------------------------------------------------------------------
SHEET_CLB = "1. CLB"
SHEET_HS = "2. Học sinh"
SHEET_HD = "Hướng dẫn"

# Số cột NV mặc định của mẫu trống. Trần nguyện vọng của phần mềm là 10 MỖI
# BUỔI; 15 đủ cho trường một-hai buổi mà dòng vẫn vừa màn hình. Mẫu sinh từ
# dữ liệu có sẵn tự nới theo số buổi và số NV dài nhất.
SO_NV_MAC_DINH = 15
# Trần MỖI BUỔI của phần mềm (rbda_priority_pipeline.TRAN_NGUYEN_VONG_MOI_BUOI,
# TRAN_CLB_THI_MOI_BUOI). Chép lại ở đây vì tệp này phải chạy được một mình
# trong bản đóng gói của skill; tests/test_so_nhap.py canh hai bên bằng nhau.
TRAN_NV_MOI_BUOI = 10
TRAN_THI_MOI_BUOI = 5
SO_DONG_CLB_TRONG = 60
SO_DONG_HS_TRONG = 1500
TU_DANH_DAU_THI = "thi"

COT_CLB = ["Tên CLB", "Buổi", "Chỉ tiêu", "Suất ưu tiên", "Nhóm ưu tiên", "Mã CLB"]
COT_HS_DAU = ["Mã HS", "Họ tên", "Nhóm ưu tiên"]

# Nhãn buổi trong danh sách thả xuống -> mã buổi lưu trong CSDL.
NHAN_BUOI = [
    ("Thứ 2", "thu_2"), ("Thứ 3", "thu_3"), ("Thứ 4", "thu_4"),
    ("Thứ 5", "thu_5"), ("Thứ 6", "thu_6"), ("Thứ 7", "thu_7"),
    ("Chủ nhật", "chu_nhat"),
]


def bo_dau(chuoi) -> str:
    """'Mã HS *' -> 'ma hs'. Bỏ dấu, hạ chữ thường, gọn khoảng trắng.

    CHỈ dùng để so TIÊU ĐỀ CỘT (người dùng gõ lại tiêu đề không dấu, viết
    hoa cả cột…). Không dùng làm khoá tên CLB: nó bỏ cả phần trong ngoặc và
    mọi ký hiệu, nên "Bóng đá (nam)" với "Bóng đá (nữ)", hay "C++" với "C#",
    thành một — xem `khoa_ten`.
    """
    s = unicodedata.normalize("NFD", str(chuoi or ""))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = s.replace("đ", "d").replace("Đ", "D").lower()
    s = re.sub(r"\(.*?\)", " ", s)          # "(bắt buộc)" và chú thích kiểu đó
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return " ".join(s.split())


def khong_dau(chuoi) -> str:
    """'Bóng đá (Nữ)' -> 'bong da (nu)'. Chỉ bỏ dấu, hoa/thường, khoảng trắng
    thừa; GIỮ ngoặc và ký hiệu — khác `bo_dau`, vốn chỉ dành cho tiêu đề cột."""
    s = unicodedata.normalize("NFD", str(chuoi or ""))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = s.replace("đ", "d").replace("Đ", "D").casefold()
    return " ".join(s.split())


def khoa_long(chuoi) -> str:
    """Khoá so "gõ lỏng" cho ô NV: bỏ dấu, hoa/thường, và coi `_ - .` như
    khoảng trắng ("clb tin hoc", "CLB-Tin-Hoc" khớp mã `clb_tin_hoc`). GIỮ
    ngoặc và ký hiệu khác: "(nữ)" khác "(nam)", "C#" khác "C++"."""
    return " ".join(re.sub(r"[_.\-]+", " ", khong_dau(chuoi)).split())


def khoa_ten(ten) -> str:
    """Khoá so TÊN CLB: chỉ bỏ khác biệt hoa/thường và khoảng trắng thừa.

    Giữ dấu, ngoặc và ký hiệu: hai tên chỉ khác nhau ở đó là hai CLB khác
    nhau thật ("Bóng đá (nam)" / "Bóng đá (nữ)", "C++" / "C#", "Toán" / "Toan").
    """
    s = unicodedata.normalize("NFC", str(ten or "")).casefold()
    return " ".join(s.split())


def tao_ma_clb(ten: str) -> str:
    """'CLB Bóng đá (nữ)' -> 'clb_bong_da_nu'. Tất định: cùng tên, cùng mã.

    Giữ phần trong ngoặc; '+' -> 'plus', '#' -> 'sharp' để "C++" và "C#"
    ra hai mã khác nhau. Không tự thêm _2 khi trùng: số thêm vào phụ thuộc
    thứ tự dòng, sắp lại dòng là hai CLB đổi mã cho nhau. Trùng thì người
    dùng phải điền Mã CLB (lỗi `so_nhap_ma_tu_sinh_trung`).
    """
    s = unicodedata.normalize("NFD", str(ten or ""))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = s.replace("đ", "d").replace("Đ", "D").lower()
    s = s.replace("+", " plus ").replace("#", " sharp ")
    s = "_".join(re.sub(r"[^a-z0-9]+", " ", s).split()) or "clb"
    return s if s.startswith("clb") else "clb_" + s


_THU = {"2": 2, "hai": 2, "3": 3, "ba": 3, "4": 4, "tu": 4, "5": 5, "nam": 5,
        "6": 6, "sau": 6, "7": 7, "bay": 7}


def ma_buoi(nhan) -> str:
    """'Thứ 2' / 'thu 2' / 'T2' / 'Chủ nhật' -> 'thu_2' / 'chu_nhat'.

    Trống -> '' (trường một buổi). Chỉ đổi khi CẢ nhãn là một thứ trong
    tuần: 'thu_3_tiet_9' là một buổi riêng của trường, gộp nó vào 'thu_3'
    là biến hai buổi khác giờ thành trùng giờ. Nhãn khác giữ nguyên, chuẩn
    hoá đúng như `chuan_hoa_buoi` của phần mềm.
    """
    tho = " ".join(str(nhan or "").split()).strip()
    if not tho:
        return ""
    # Dấu câu là dấu tách ("Thứ-2", "T.2", "Thứ 2." đều là thứ Hai), nhưng
    # CHỮ trong ngoặc được giữ: bo_dau bỏ cả phần trong ngoặc nên "Thứ 3
    # (sáng)" và "Thứ 3 (chiều)" từng thành một buổi — hai giờ khác nhau bị
    # coi là trùng giờ.
    c = re.sub(r"[^a-z0-9]+", "_", khong_dau(tho)).strip("_")
    if c in ("cn", "chu_nhat"):
        return "chu_nhat"
    # "thu" + số hoặc chữ (thu_2, thu_hai); "t" CHỈ đi với số (t2, t_2):
    # "tba" (to be announced) hay "thai" không phải thứ Ba, thứ Hai.
    m = re.fullmatch(r"thu_?([a-z0-9]+)|t_?([0-9])", c)
    con_lai = (m.group(1) or m.group(2)) if m else None
    if con_lai in _THU:
        return f"thu_{_THU[con_lai]}"
    return tho.lower().replace(" ", "_")


def nhan_buoi(ma) -> str:
    """'thu_2' -> 'Thứ 2'. Mã lạ giữ nguyên; trống giữ trống."""
    for nhan, m in NHAN_BUOI:
        if m == ma:
            return nhan
    return ma or ""


def doc_diem(chuoi):
    """Một ô điểm -> float, hoặc None nếu không phải số.

    Nhận dấu phẩy thập phân (Excel tiếng Việt lưu 8,5). Chỉ đổi khi có ĐÚNG
    một dấu phẩy và không có dấu chấm — "1,234.5" là cách viết nghìn.
    """
    chuoi = str(chuoi if chuoi is not None else "").strip()
    if not chuoi:
        return None
    if chuoi.count(",") == 1 and "." not in chuoi:
        chuoi = chuoi.replace(",", ".")
    try:
        so = float(chuoi)
    except (TypeError, ValueError):
        return None
    return so if so == so and so not in (float("inf"), float("-inf")) else None


def o_thanh_chu(gia_tri) -> str:
    """Ô Excel -> chuỗi. 20.0 -> '20' (int('20.0') ném lỗi)."""
    if gia_tri is None:
        return ""
    if isinstance(gia_tri, bool):
        return "1" if gia_tri else ""
    if isinstance(gia_tri, float) and gia_tri.is_integer():
        return str(int(gia_tri))
    return str(gia_tri).strip()


# ---------------------------------------------------------------------------
# Nhận ra cột
# ---------------------------------------------------------------------------
_TEN_COT_CLB = {
    "name": {"ten clb", "ten cau lac bo", "ten", "clb", "cau lac bo"},
    "buoi": {"buoi", "buoi sinh hoat", "ngay", "thu"},
    "capacity": {"chi tieu", "so cho", "so luong", "chi tieu tong"},
    "reserve_capacity": {"suat uu tien", "chi tieu uu tien", "cho uu tien",
                         "suat du tru"},
    "reserve_group": {"nhom uu tien", "nhom du tru", "doi tuong uu tien"},
    "club_id": {"ma clb", "ma", "ma cau lac bo"},
}
_TEN_COT_HS = {
    "student_id": {"ma hs", "ma hoc sinh", "ma", "ma so", "ma so hoc sinh"},
    "name": {"ho ten", "ho va ten", "ten", "ten hoc sinh"},
    "reserve_group": {"nhom uu tien", "nhom du tru", "doi tuong uu tien"},
}
_MAU_NV = re.compile(r"^(?:nv|nguyen vong)\s*(\d+)$")
_MAU_DIEM = re.compile(r"^diem(?: nv| thi)?\s*(\d+)$")


def _tim_sheet(ten_cac_sheet, ten_chuan):
    """Tìm sheet theo tên đã bỏ dấu: '1. CLB', '1 clb', 'CLB' đều khớp."""
    muc = bo_dau(ten_chuan)
    muc_ngan = re.sub(r"^\d+ ", "", muc)
    for ten in ten_cac_sheet:
        b = bo_dau(ten)
        if b == muc or re.sub(r"^\d+ ", "", b) == muc_ngan:
            return ten
    return None


def la_so_nhap(ten_cac_sheet) -> bool:
    """Tệp có đủ hai sheet của sổ nhập không."""
    return bool(_tim_sheet(ten_cac_sheet, SHEET_CLB)
                and _tim_sheet(ten_cac_sheet, SHEET_HS))


def _khop_cot(tieu_de, bang_ten):
    """{khoá: chỉ số cột} theo bảng tên. Cột đầu khớp được giữ."""
    vi_tri = {}
    for i, t in enumerate(tieu_de):
        b = bo_dau(t)
        for khoa, cac_ten in bang_ten.items():
            if b in cac_ten and khoa not in vi_tri:
                vi_tri[khoa] = i
                break
    return vi_tri


def _cac_dong(ws):
    """(số dòng Excel, [ô đã thành chuỗi]) — bỏ dòng trống hoàn toàn."""
    for so_dong, hang in enumerate(ws.iter_rows(values_only=True), start=1):
        o = [o_thanh_chu(v) for v in hang]
        if any(o):
            yield so_dong, o


# ---------------------------------------------------------------------------
# ĐỌC sổ nhập
# ---------------------------------------------------------------------------
def doc_so_nhap(wb, clb_da_co=None, ma_da_co=None, clb_cu=None) -> dict:
    """Đọc sổ nhập (openpyxl Workbook) thành dữ liệu thuần, KHÔNG ghi gì.

    Trả về {"clubs": [...], "students": [...], "loi": [...]}.
    `loi` khác rỗng nghĩa là sổ có lỗi phải sửa — phần mềm KHÔNG nạp gì cả.
    Nạp một nửa sổ còn tệ hơn không nạp: người dùng tưởng đã xong.

    clubs:    {club_id, name, capacity, reserve_capacity, reserve_group, buoi, dong}
    students: {student_id, name, reserve_group, dong,
               nv: [{club_id, thi: bool, diem: str}]}

    clb_da_co: {khoa_ten(tên): club_id} của các CLB đã có trong phần mềm.
    Dòng CLB bỏ trống Mã CLB mà tên trùng một CLB đã có thì dùng mã của CLB
    đó — nạp lại sổ không đẻ ra CLB thứ hai.
    ma_da_co:  {club_id: tên} của mọi CLB đã có. Mã TỰ TẠO trùng (bỏ hoa/thường)
    mã một CLB đã có (mà tên khác) là lỗi: nạp vào sẽ đè lên CLB đó.
    clb_cu:    {club_id: {buoi, reserve_capacity, reserve_group}} của CLB đã có.
    Sheet CLB THIẾU HẲN cột Buổi / Suất ưu tiên / Nhóm ưu tiên thì CLB đã có
    giữ giá trị cũ (cột vắng mặt = "không biết"; cột có mà ô trống = "không có").
    """
    loi: list = []
    ten_cac_sheet = list(wb.sheetnames)
    ten_clb = _tim_sheet(ten_cac_sheet, SHEET_CLB)
    ten_hs = _tim_sheet(ten_cac_sheet, SHEET_HS)
    for ten, chuan in ((ten_clb, SHEET_CLB), (ten_hs, SHEET_HS)):
        if not ten:
            loi.append(err("so_nhap_thieu_sheet", sheet=chuan))
    if loi:
        return {"clubs": [], "students": [], "loi": loi}

    clubs = _doc_sheet_clb(wb[ten_clb], loi, clb_da_co or {}, ma_da_co or {}, clb_cu or {})
    students = _doc_sheet_hs(wb[ten_hs], clubs, loi)
    return {"clubs": [c for c in clubs if not c.get("hong")],
            "students": students, "loi": loi}


def _doc_sheet_clb(ws, loi, clb_da_co, ma_da_co, clb_cu) -> list:
    dong_iter = _cac_dong(ws)
    dau = next(dong_iter, None)
    if dau is None:
        loi.append(err("so_nhap_khong_co_clb", sheet=SHEET_CLB))
        return []
    _, tieu_de = dau
    cot = _khop_cot(tieu_de, _TEN_COT_CLB)
    for khoa, nhan in (("name", "Tên CLB"), ("capacity", "Chỉ tiêu")):
        if khoa not in cot:
            loi.append(err("so_nhap_thieu_cot", sheet=SHEET_CLB, cot=nhan))
    if loi:
        return []

    def lay(o, khoa):
        i = cot.get(khoa)
        return o[i].strip() if i is not None and i < len(o) else ""

    tho = []
    for so_dong, o in dong_iter:
        ten = lay(o, "name")
        if not ten and not lay(o, "club_id"):
            continue        # dòng chỉ có buổi/chỉ tiêu lạc: coi như trống
        tho.append((so_dong, o, ten))

    clubs, ten_da_gap, ma_da_gap = [], {}, {}
    for so_dong, o, ten in tho:
        if not ten:
            loi.append(err("so_nhap_clb_thieu_ten", sheet=SHEET_CLB, dong=so_dong))
            continue
        k_ten = khoa_ten(ten)
        if k_ten in ten_da_gap:
            loi.append(err("so_nhap_clb_trung_ten", sheet=SHEET_CLB, dong=so_dong,
                           dong_truoc=ten_da_gap[k_ten], ten=ten))
            continue
        ma = lay(o, "club_id")
        tu_sinh = False
        if not ma:
            # Mã lấy lại của CLB cùng tên đã có là mã THẬT trong phần mềm,
            # so như mã người dùng điền (phân biệt hoa/thường). Chỉ mã vừa
            # tạo từ tên mới là "tự sinh".
            ma = clb_da_co.get(k_ten)
            if ma is None:
                tu_sinh = True
                ma = tao_ma_clb(ten)
                # Mã tự tạo trùng mã một CLB KHÁC đã có trong phần mềm: nạp vào
                # là đổi tên, chỉ tiêu, buổi của CLB đó và chuyển cả nguyện vọng
                # của học sinh sang — không một cảnh báo nào. Bắt điền Mã CLB.
                # Gợi ý MÃ THẬT của CLB đó (CLB_TIN, không phải clb_tin vừa
                # tạo): mã phân biệt hoa/thường, điền clb_tin là tạo CLB mới.
                ma_cu = ma if ma in ma_da_co else next(
                    (k for k in ma_da_co if k.lower() == ma.lower()), None)
                if ma_cu is not None:
                    loi.append(err("so_nhap_ma_tu_sinh_trung_clb_da_co", sheet=SHEET_CLB,
                                   dong=so_dong, ten=ten, ma=ma, ma_cu=ma_cu,
                                   ten_cu=ma_da_co[ma_cu]))
                    continue
        # Phần mềm phân biệt hoa/thường ở mã (clb_a và CLB_A là hai CLB), nên
        # hai mã người dùng tự điền chỉ khác hoa/thường KHÔNG trùng — sổ xuất
        # từ một CSDL như thế phải nạp lại được. Mã TỰ TẠO thì so không phân
        # biệt hoa/thường: hàm nạp khớp mã lệch hoa/thường về mã đã có.
        trung = ma_da_gap.get(ma) or next(
            (v for k, v in ma_da_gap.items()
             if k.lower() == ma.lower() and (tu_sinh or v[1])), None)
        if trung:
            dong_truoc, truoc_tu_sinh = trung
            # Mã tự sinh đụng nhau: báo để người dùng điền Mã CLB, đừng tự
            # đánh số — số phụ thuộc thứ tự dòng.
            ma_loi = ("so_nhap_ma_tu_sinh_trung" if (tu_sinh or truoc_tu_sinh)
                      else "so_nhap_clb_trung_ma")
            loi.append(err(ma_loi, sheet=SHEET_CLB, dong=so_dong,
                           dong_truoc=dong_truoc, ma=ma, ten=ten))
            continue
        ten_da_gap[k_ten] = so_dong
        ma_da_gap[ma] = (so_dong, tu_sinh)

        cu = clb_cu.get(ma, {})

        def lay_hoac_cu(khoa, mac_dinh, o=o, cu=cu):
            # Cột VẮNG MẶT: giữ giá trị của CLB đã có (mới thì mặc định).
            # Không thế thì một sheet chỉ có Tên CLB + Chỉ tiêu xoá buổi và
            # suất ưu tiên của mọi CLB — trường nhiều buổi thành một buổi.
            return lay(o, khoa) if khoa in cot else cu.get(khoa, mac_dinh)

        chi_tieu = lay(o, "capacity")
        uu_tien = str(lay_hoac_cu("reserve_capacity", 0) or "0")
        try:
            f_chi_tieu, f_uu_tien = float(chi_tieu), float(uu_tien)
            # "inf", "1e999", "nan" đọc được thành float nhưng không thành
            # chỉ tiêu: int(inf) ném OverflowError, nan so sánh luôn sai.
            hop_le = (math.isfinite(f_chi_tieu) and math.isfinite(f_uu_tien)
                      and f_chi_tieu.is_integer() and f_uu_tien.is_integer())
            n_chi_tieu = int(f_chi_tieu) if hop_le else 0
            n_uu_tien = int(f_uu_tien) if hop_le else 0
            hop_le = hop_le and n_chi_tieu > 0 and 0 <= n_uu_tien <= n_chi_tieu
        except (ValueError, OverflowError):
            hop_le = False
        if not hop_le:
            loi.append(err("so_nhap_chi_tieu_sai", sheet=SHEET_CLB, dong=so_dong,
                           ten=ten, chi_tieu=chi_tieu or "0", uu_tien=uu_tien))
            # Vẫn giữ tên để tra NV: một lỗi chỉ tiêu không được kéo theo
            # một loạt lỗi "không có CLB" ở sheet học sinh.
            clubs.append({"club_id": ma, "name": ten, "hong": True})
            continue
        clubs.append({
            "club_id": ma, "name": ten, "capacity": n_chi_tieu,
            "reserve_capacity": n_uu_tien,
            "reserve_group": lay_hoac_cu("reserve_group", "") or "",
            "buoi": ma_buoi(lay(o, "buoi")) if "buoi" in cot else (cu.get("buoi") or ""),
            "dong": so_dong,
        })

    if not any(not c.get("hong") for c in clubs) and not loi:
        loi.append(err("so_nhap_khong_co_clb", sheet=SHEET_CLB))
    return clubs


def _doc_sheet_hs(ws, clubs, loi) -> list:
    dong_iter = _cac_dong(ws)
    dau = next(dong_iter, None)
    if dau is None:
        return []           # chưa có học sinh: nạp riêng danh sách CLB
    _, tieu_de = dau
    cot = _khop_cot(tieu_de, _TEN_COT_HS)
    cot_nv, cot_diem = {}, {}
    for i, t in enumerate(tieu_de):
        b = bo_dau(t)
        m = _MAU_NV.match(b)
        if m:
            cot_nv.setdefault(int(m.group(1)), i)
            continue
        m = _MAU_DIEM.match(b)
        if m:
            cot_diem.setdefault(int(m.group(1)), i)
    if "student_id" not in cot:
        loi.append(err("so_nhap_thieu_cot", sheet=SHEET_HS, cot="Mã HS"))
    if not cot_nv:
        loi.append(err("so_nhap_thieu_cot", sheet=SHEET_HS, cot="NV1"))
    if "student_id" not in cot or not cot_nv:
        return []

    # Ô NV nhận TÊN CLB (danh sách thả xuống) hoặc MÃ CLB. Khớp đúng tên
    # (bỏ hoa/thường) trước; gõ thiếu dấu thì chỉ nhận khi ra ĐÚNG MỘT CLB —
    # không bao giờ đoán giữa "Bóng đá (nam)" và "Bóng đá (nữ)".
    # Thứ tự: tên đúng TỪNG KÝ TỰ (danh sách thả xuống và sổ xuất ra đều ghi
    # tên), rồi mã đúng từng ký tự (clb_a và CLB_A là hai CLB khác nhau), rồi
    # mới bỏ hoa/thường. Khớp mã bỏ hoa/thường chỉ khi ra đúng một CLB.
    dung_ten = {c["name"]: c["club_id"] for c in clubs}
    dung_ma = {c["club_id"]: c["club_id"] for c in clubs}
    ma_bo_hoa: dict = {}
    for c in clubs:
        ma_bo_hoa.setdefault(khoa_ten(c["club_id"]), set()).add(c["club_id"])
    tra = {k: next(iter(v)) for k, v in ma_bo_hoa.items() if len(v) == 1}
    for c in clubs:
        tra[khoa_ten(c["name"])] = c["club_id"]
    # Khoá "gõ thiếu dấu" chỉ bỏ dấu, GIỮ ngoặc và ký hiệu: "Bóng đá (nữ)"
    # không được rơi vào "Bóng đá (nam)", "C#" không được rơi vào "C++".
    tra_long: dict = {}
    for c in clubs:
        for k in {khoa_long(c["name"]), khoa_long(c["club_id"])}:
            tra_long.setdefault(k, set()).add(c["club_id"])
    buoi_cua = {c["club_id"]: c.get("buoi", "") for c in clubs if not c.get("hong")}

    def tim_clb(gia_tri):
        ma = (dung_ten.get(gia_tri) or dung_ma.get(gia_tri)
              or tra.get(khoa_ten(gia_tri)))
        if ma is None:
            ung_vien = tra_long.get(khoa_long(gia_tri), set())
            ma = next(iter(ung_vien)) if len(ung_vien) == 1 else None
        return ma

    def lay(o, i):
        return o[i].strip() if i is not None and i < len(o) else ""

    students, ma_da_gap = [], {}
    for so_dong, o in dong_iter:
        sid = lay(o, cot["student_id"])
        if not sid:
            loi.append(err("so_nhap_thieu_ma_hs", sheet=SHEET_HS, dong=so_dong))
            continue
        # Trùng ĐÚNG từng ký tự mới là trùng. "hs01" và "HS01" là hai học
        # sinh với phần mềm; hàm nạp vẫn cảnh báo cặp mã như thế.
        if sid in ma_da_gap:
            loi.append(err("so_nhap_hs_trung", sheet=SHEET_HS, dong=so_dong,
                           dong_truoc=ma_da_gap[sid], student_id=sid))
            continue
        ma_da_gap[sid] = so_dong

        nv, da_chon = [], {}
        for k in sorted(set(cot_nv) | set(cot_diem)):
            o_nv = lay(o, cot_nv.get(k))
            o_diem = lay(o, cot_diem.get(k))
            if not o_nv:
                if o_diem:
                    loi.append(err("so_nhap_diem_khong_nv", sheet=SHEET_HS,
                                   dong=so_dong, student_id=sid, k=k))
                continue
            ma = tim_clb(o_nv)
            if ma is None:
                loi.append(err("so_nhap_clb_khong_co", sheet=SHEET_HS, dong=so_dong,
                               student_id=sid, k=k, gia_tri=o_nv))
                continue
            if ma in da_chon:
                loi.append(err("so_nhap_nv_trung", sheet=SHEET_HS, dong=so_dong,
                               student_id=sid, k=k, k_truoc=da_chon[ma], gia_tri=o_nv))
                continue
            da_chon[ma] = k
            thi = bool(o_diem)
            if thi and o_diem.lower() != TU_DANH_DAU_THI:
                so = doc_diem(o_diem)
                if so is None or so < 0:
                    loi.append(err("so_nhap_diem_sai", sheet=SHEET_HS, dong=so_dong,
                                   student_id=sid, k=k, gia_tri=o_diem))
                    continue
            diem = "" if (not thi or o_diem.lower() == TU_DANH_DAU_THI) else o_diem
            nv.append({"club_id": ma, "thi": thi, "diem": diem})
        # Trần MỖI BUỔI: soát ngay ở bước đọc thử. Để tới lúc nạp thì hàm nạp
        # chỉ bỏ cả buổi của em kèm cảnh báo, còn nguyện vọng cũ của em vẫn
        # nằm trong CSDL — sổ "sẵn sàng" mà kết quả lại khác.
        for khoa_dem, tran, ma_loi in (
                ("nv", TRAN_NV_MOI_BUOI, "so_nhap_qua_tran_nv"),
                ("thi", TRAN_THI_MOI_BUOI, "so_nhap_qua_tran_thi")):
            dem: dict = {}
            for x in nv:
                if khoa_dem == "thi" and not x["thi"]:
                    continue
                if x["club_id"] in buoi_cua:
                    b = buoi_cua[x["club_id"]]
                    dem[b] = dem.get(b, 0) + 1
            for b, n in dem.items():
                if n > tran:
                    loi.append(err(ma_loi, sheet=SHEET_HS, dong=so_dong, student_id=sid,
                                   buoi=nhan_buoi(b) or "-", n=n, tran=tran))
        students.append({
            "student_id": sid,
            "name": lay(o, cot.get("name")),
            # Cột Nhóm ưu tiên VẮNG MẶT -> None = "không biết", giữ nhóm cũ.
            # Cột có mà ô trống -> "" = em không thuộc nhóm nào.
            "reserve_group": (lay(o, cot["reserve_group"])
                              if "reserve_group" in cot else None),
            "nv": nv, "dong": so_dong,
        })
    return students


# ---------------------------------------------------------------------------
# Sổ -> ba văn bản CSV mà ba hàm nạp cũ đã hiểu (và đã được test kỹ)
# ---------------------------------------------------------------------------
def _csv(dong) -> str:
    buf = io.StringIO()
    csv.writer(buf, lineterminator="\n").writerows(dong)
    return buf.getvalue()


def thanh_csv(du_lieu: dict) -> dict:
    """{"clubs": csv, "test_selection": csv | "", "preferences": csv | ""}.

    MỌI học sinh có trong sổ đều có mặt ở cả hai tệp học sinh, kể cả em
    không thi CLB nào: sổ là bản đầy đủ, nạp lại sổ phải xoá được lựa chọn
    thi cũ của em đã bỏ thi.
    """
    clubs = [["club_id", "name", "capacity", "reserve_capacity", "reserve_group", "buoi"]]
    for c in du_lieu["clubs"]:
        clubs.append([c["club_id"], c["name"], c["capacity"], c["reserve_capacity"],
                      c["reserve_group"], c["buoi"]])

    hs = du_lieu["students"]
    if not hs:
        return {"clubs": _csv(clubs), "test_selection": "", "preferences": ""}

    n_thi = max([sum(1 for x in s["nv"] if x["thi"]) for s in hs] + [1])
    n_nv = max([len(s["nv"]) for s in hs] + [1])
    thi = [["student_id", "name", "reserve_group"]
           + [c for k in range(1, n_thi + 1) for c in (f"test_club_{k}", f"score_{k}")]]
    nv = [["student_id", "name", "reserve_group"]
          + [f"pref_{k}" for k in range(1, n_nv + 1)]]
    for s in hs:
        dau = [s["student_id"], s["name"], s["reserve_group"] or ""]
        o_thi = [v for x in s["nv"] if x["thi"] for v in (x["club_id"], x["diem"])]
        thi.append(dau + o_thi + [""] * (2 * n_thi - len(o_thi)))
        o_nv = [x["club_id"] for x in s["nv"]]
        nv.append(dau + o_nv + [""] * (n_nv - len(o_nv)))
    return {"clubs": _csv(clubs), "test_selection": _csv(thi), "preferences": _csv(nv)}


def tom_tat(du_lieu: dict) -> dict:
    hs = du_lieu["students"]
    return {
        "n_clb": len(du_lieu["clubs"]),
        "n_hoc_sinh": len(hs),
        "n_buoi": len({c["buoi"] for c in du_lieu["clubs"]}),
        "n_nguyen_vong": sum(len(s["nv"]) for s in hs),
        "n_thi": sum(1 for s in hs for x in s["nv"] if x["thi"]),
        "n_diem": sum(1 for s in hs for x in s["nv"] if x["diem"]),
    }


# ---------------------------------------------------------------------------
# GHI sổ nhập (mẫu trống, mẫu có sẵn CLB, xuất dữ liệu đầu vào, skill)
# ---------------------------------------------------------------------------
HUONG_DAN = [
    "SỔ NHẬP CLB — một tệp duy nhất để nạp vào phần mềm",
    "",
    "Bước 1. Sheet \"1. CLB\": mỗi câu lạc bộ một dòng.",
    "        Tên CLB và Chỉ tiêu là bắt buộc. Buổi chọn trong danh sách thả xuống",
    "        (trường chỉ có một buổi thì để trống cả cột).",
    "        Suất ưu tiên / Nhóm ưu tiên: chỉ điền nếu CLB giữ chỗ cho một nhóm",
    "        học sinh (ví dụ nhóm chinh_sach). Mã CLB để trống, phần mềm tự tạo.",
    "",
    "Bước 2. Sheet \"2. Học sinh\": mỗi học sinh một dòng.",
    "        Mã HS là bắt buộc và không được trùng.",
    "        Nhóm ưu tiên: chỉ điền với học sinh thuộc nhóm đó.",
    "",
    "Bước 3. NV1, NV2, NV3…: chọn tên CLB trong danh sách thả xuống,",
    "        theo thứ tự em thích nhất trước. Không cần điền kín.",
    "        Trường nhiều buổi: cứ điền chung một danh sách — phần mềm tự chia",
    "        theo buổi của từng CLB.",
    "",
    "Bước 4. Điểm 1, Điểm 2…: điểm thi của CLB ngay bên trái.",
    "        Có điểm = em đã thi CLB đó. Để trống = không thi.",
    "        Thi rồi mà chưa có điểm: gõ chữ  thi  rồi chấm sau trong phần mềm.",
    "",
    "Bước 5. Lưu tệp (giữ định dạng .xlsx).",
    "",
    "Bước 6. Mở phần mềm → màn hình Vận hành → kéo tệp này vào ô nạp",
    "        → xem lại phần tóm tắt → bấm Nhập.",
    "",
    "Sổ có lỗi (CLB trùng tên, NV ghi tên CLB không có, điểm không phải số…)",
    "thì phần mềm KHÔNG nạp gì cả và chỉ rõ sheet, dòng cần sửa.",
    "Nạp lại sổ đã sửa là cập nhật — không tạo trùng.",
    "",
    "Ví dụ một dòng học sinh:",
    "  Mã HS   Họ tên          NV1          Điểm 1   NV2           Điểm 2",
    "  HS001   Nguyễn Văn An   CLB Tin học  9,5      CLB Bóng rổ",
    "  → An thích Tin học nhất và đã thi được 9,5; nguyện vọng 2 là Bóng rổ,",
    "    CLB không thi tuyển.",
]


def ten_khong_trung(clubs) -> int:
    """Đổi tên CLB (tại chỗ) để không hai CLB nào trùng `khoa_ten`. Trả số tên đã đổi.

    Sổ chọn CLB theo TÊN, nên tên phải riêng. Dữ liệu từ nơi khác thì không
    chắc: trường mở "CLB Cờ vua" cả Thứ 2 lẫn Thứ 5, hay CSDL có "Toán" và
    "toán". CLB gặp đầu tiên giữ tên; CLB sau thêm buổi vào tên ("CLB Cờ vua
    (Thứ 5)"), cùng buổi thì thêm mã CLB. Tên trống lấy mã CLB.
    """
    da_dung, n_doi = set(), 0
    for c in clubs:
        c["name"] = (c.get("name") or "").strip() or c.get("club_id") or ""
    for c in clubs:
        goc = c["name"]
        if khoa_ten(goc) not in da_dung:
            da_dung.add(khoa_ten(goc))
            continue
        ung_vien = [nhan_buoi(c.get("buoi") or ""), c.get("club_id") or ""]
        ten = next((f"{goc} ({x})" for x in ung_vien
                    if x and khoa_ten(f"{goc} ({x})") not in da_dung), None)
        k = 2
        while ten is None or khoa_ten(ten) in da_dung:
            ten, k = f"{goc} ({k})", k + 1
        c["name"] = ten
        da_dung.add(khoa_ten(ten))
        n_doi += 1
    return n_doi


def ghi_so_nhap(duong_dan_hoac_tep, clubs=(), students=(), so_nv=None, co_nhom=True):
    """Ghi sổ nhập ra tệp .xlsx (đường dẫn hoặc tệp nhị phân đang mở).

    clubs:    dict có name, capacity, reserve_capacity, reserve_group, buoi, club_id
    students: dict có student_id, name, reserve_group,
              nv: [{club_id, thi, diem}] (diem: số, chuỗi, hoặc '')
    so_nv:    số cặp NV/Điểm; mặc định nới theo dữ liệu, tối thiểu SO_NV_MAC_DINH.
    co_nhom:  False thì bỏ hẳn cột Nhóm ưu tiên ở sheet học sinh. Dùng cho sổ
              sinh từ biểu mẫu: biểu mẫu không hỏi nhóm, và cột VẮNG MẶT nghĩa
              là phần mềm giữ nhóm đã gán; cột có mà ô trống là BỎ nhóm.
    """
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    # Bản sao: đổi tên cho khỏi trùng mà không đụng dữ liệu của bên gọi.
    clubs, students = [dict(c) for c in clubs], list(students)
    ten_khong_trung(clubs)
    so_buoi = len({c.get("buoi") or "" for c in clubs}) or 1
    so_nv = so_nv or max([SO_NV_MAC_DINH, 2 * so_buoi]
                         + [len(s.get("nv") or []) for s in students])

    NEN_TIEU_DE = PatternFill("solid", fgColor="1F3A5F")
    NEN_NHAP = PatternFill("solid", fgColor="FFF8DC")
    NEN_DIEM = PatternFill("solid", fgColor="EAF4EA")
    CHU_TIEU_DE = Font(bold=True, color="FFFFFF", size=11)
    VIEN = Border(bottom=Side(style="thin", color="D0D0D0"),
                  right=Side(style="thin", color="D0D0D0"))
    GIUA = Alignment(horizontal="center", vertical="center", wrap_text=True)

    wb = Workbook()

    # ---- 1. CLB ---------------------------------------------------------
    ws = wb.active
    ws.title = SHEET_CLB
    ws.append(COT_CLB)
    for c in clubs:
        ws.append([c.get("name", ""), nhan_buoi(c.get("buoi") or ""),
                   c.get("capacity", ""), c.get("reserve_capacity") or None,
                   c.get("reserve_group") or None, c.get("club_id") or None])
    so_dong_clb = max(len(clubs) + 1 + 20, SO_DONG_CLB_TRONG)
    for i, rong in enumerate([30, 12, 10, 13, 16, 18], start=1):
        ws.column_dimensions[get_column_letter(i)].width = rong
    for o in ws[1]:
        o.fill, o.font, o.alignment = NEN_TIEU_DE, CHU_TIEU_DE, GIUA
    ws.row_dimensions[1].height = 30
    for r in range(2, so_dong_clb + 1):
        for col in range(1, 7):
            o = ws.cell(row=r, column=col)
            o.fill, o.border = NEN_NHAP, VIEN
    ws.freeze_panes = "A2"

    dv_buoi = DataValidation(
        type="list", formula1='"' + ",".join(n for n, _ in NHAN_BUOI) + '"',
        allow_blank=True, showErrorMessage=False)
    dv_buoi.add(f"B2:B{so_dong_clb}")
    dv_chi_tieu = DataValidation(
        type="whole", operator="greaterThan", formula1="0", allow_blank=True,
        showErrorMessage=True, errorTitle="Chỉ tiêu",
        error="Chỉ tiêu phải là số nguyên lớn hơn 0.")
    dv_chi_tieu.add(f"C2:C{so_dong_clb}")
    dv_uu_tien = DataValidation(
        type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
        showErrorMessage=True, errorTitle="Suất ưu tiên",
        error="Suất ưu tiên là số nguyên, không lớn hơn Chỉ tiêu.")
    dv_uu_tien.add(f"D2:D{so_dong_clb}")
    for dv in (dv_buoi, dv_chi_tieu, dv_uu_tien):
        ws.add_data_validation(dv)
    ws.conditional_formatting.add(
        f"A2:A{so_dong_clb}",
        FormulaRule(formula=[f'AND(A2<>"",COUNTIF($A$2:$A${so_dong_clb},A2)>1)'],
                    fill=PatternFill("solid", fgColor="F4B6B6")))

    # ---- 2. Học sinh ----------------------------------------------------
    wh = wb.create_sheet(SHEET_HS)
    cot_dau = COT_HS_DAU if co_nhom else COT_HS_DAU[:2]
    nv1 = len(cot_dau) + 1                    # số thứ tự cột NV1 (tính từ 1)
    tieu_de = cot_dau + [c for k in range(1, so_nv + 1) for c in (f"NV{k}", f"Điểm {k}")]
    wh.append(tieu_de)
    for s in students:
        hang = [s.get("student_id", ""), s.get("name", "")]
        if co_nhom:
            hang.append(s.get("reserve_group") or None)
        for x in (s.get("nv") or [])[:so_nv]:
            ten = next((c["name"] for c in clubs if c.get("club_id") == x["club_id"]),
                       x["club_id"])
            diem = x.get("diem")
            if diem in (None, ""):
                diem = TU_DANH_DAU_THI if x.get("thi") else None
            elif isinstance(diem, str):
                so = doc_diem(diem)
                diem = (int(so) if so is not None and so.is_integer() else so) \
                    if so is not None else diem
            elif isinstance(diem, float) and diem.is_integer():
                diem = int(diem)
            hang += [ten, diem]
        wh.append(hang)
    so_dong_hs = max(len(students) + 1 + 200, SO_DONG_HS_TRONG)
    cot_cuoi = len(tieu_de)
    wh.column_dimensions["A"].width = 11
    wh.column_dimensions["B"].width = 24
    if co_nhom:
        wh.column_dimensions["C"].width = 13
    for col in range(nv1, cot_cuoi + 1):
        wh.column_dimensions[get_column_letter(col)].width = 22 if (col - nv1) % 2 == 0 else 8
    for o in wh[1]:
        o.fill, o.font, o.alignment = NEN_TIEU_DE, CHU_TIEU_DE, GIUA
    wh.row_dimensions[1].height = 30
    for r in range(2, so_dong_hs + 1):
        for col in range(1, cot_cuoi + 1):
            o = wh.cell(row=r, column=col)
            o.fill = NEN_DIEM if (col >= nv1 and (col - nv1) % 2 == 1) else NEN_NHAP
            o.border = VIEN
        # Mã HS là CHỮ: để Excel tự nhận thì 0012345 thành 12345.
        wh.cell(row=r, column=1).number_format = "@"
    wh.freeze_panes = f"{get_column_letter(nv1)}2"

    # Danh sách thả xuống: tên CLB lấy THẲNG từ sheet 1, nên thêm CLB ở
    # sheet 1 là danh sách cập nhật ngay, không phải sinh lại mẫu.
    dv_nv = DataValidation(
        type="list", formula1=f"'{SHEET_CLB}'!$A$2:$A${so_dong_clb}",
        allow_blank=True, showErrorMessage=True, errorStyle="warning",
        errorTitle="Tên CLB", error="Hãy chọn tên CLB trong danh sách (sheet 1. CLB).")
    wh.add_data_validation(dv_nv)
    for k in range(so_nv):
        c_nv, c_d = get_column_letter(nv1 + 2 * k), get_column_letter(nv1 + 1 + 2 * k)
        dv_nv.add(f"{c_nv}2:{c_nv}{so_dong_hs}")
        # Công thức tham chiếu CHÍNH cột điểm đó, nên mỗi cột một quy tắc.
        dv = DataValidation(
            type="custom", allow_blank=True, showErrorMessage=True,
            errorTitle="Điểm",
            error="Điểm là một số (vd 8,5) hoặc chữ  thi  nếu chưa có điểm.",
            formula1=f'OR(ISNUMBER({c_d}2),LOWER({c_d}2)="thi")')
        dv.add(f"{c_d}2:{c_d}{so_dong_hs}")
        wh.add_data_validation(dv)
    # Cùng một CLB chọn hai lần trên một dòng -> tô đỏ.
    for k in range(so_nv):
        c = get_column_letter(nv1 + 2 * k)
        dem = "+".join(f'(${get_column_letter(nv1 + 2 * j)}2={c}2)' for j in range(so_nv))
        wh.conditional_formatting.add(
            f"{c}2:{c}{so_dong_hs}",
            FormulaRule(formula=[f'AND({c}2<>"",({dem})>1)'],
                        fill=PatternFill("solid", fgColor="F4B6B6")))
    wh.conditional_formatting.add(
        f"A2:A{so_dong_hs}",
        FormulaRule(formula=[f'AND(A2<>"",COUNTIF($A$2:$A${so_dong_hs},A2)>1)'],
                    fill=PatternFill("solid", fgColor="F4B6B6")))

    # ---- Hướng dẫn ------------------------------------------------------
    hd = wb.create_sheet(SHEET_HD)
    hd.column_dimensions["A"].width = 90
    for i, dong in enumerate(HUONG_DAN, start=1):
        o = hd.cell(row=i, column=1, value=dong)
        if i == 1:
            o.font = Font(bold=True, size=14, color="1F3A5F")
        elif dong.startswith("Bước"):
            o.font = Font(bold=True)

    # Tên bắt đầu bằng "=" (vd học sinh tên "=1+1") phải là CHỮ, không thành
    # công thức Excel — vừa sai dữ liệu, vừa là lối chèn công thức độc.
    for trang in (ws, wh):
        for hang in trang.iter_rows(min_row=2, max_row=trang.max_row):
            for o in hang:
                if isinstance(o.value, str) and o.value.startswith("="):
                    o.data_type = "s"

    wb.active = 0
    wb.save(duong_dan_hoac_tep)


def tu_csv(clubs_rows, thi_rows, nv_rows):
    """Dữ liệu ba tệp CSV cũ (dict từ csv.DictReader) -> (clubs, students)
    cho `ghi_so_nhap`. Dùng để sinh sổ ví dụ từ các bộ CSV mẫu đã có, và để
    chuyển bộ ba tệp cũ của một trường sang sổ nhập.

    Nguyện vọng theo buổi (`thu_2_pref_1`…) được nối theo thứ tự buổi —
    đúng cách phần mềm vẫn ghi chúng vào CSDL. Trả thêm danh sách
    (student_id, club_id) thi một CLB KHÔNG có trong nguyện vọng: sổ nhập
    không chứa được những ô đó.
    """
    from rbda_priority_pipeline import sap_buoi

    clubs = [{
        "club_id": r["club_id"], "name": r.get("name") or r["club_id"],
        "capacity": int(float(r.get("capacity") or 0)),
        "reserve_capacity": int(float(r.get("reserve_capacity") or 0)),
        "reserve_group": r.get("reserve_group") or "",
        "buoi": (r.get("buoi") or "").strip(),
    } for r in clubs_rows]

    hs: dict = {}

    def lay_hs(r):
        sid = (r.get("student_id") or "").strip()
        s = hs.setdefault(sid, {"student_id": sid, "name": "", "reserve_group": "",
                                "nv": [], "_thi": {}})
        s["name"] = s["name"] or (r.get("name") or "")
        s["reserve_group"] = s["reserve_group"] or (r.get("reserve_group") or "")
        return s

    for r in nv_rows:
        if not (r.get("student_id") or "").strip():
            continue
        s = lay_hs(r)
        cot = list(r.keys())
        theo_buoi: dict = {}
        for c in cot:
            m = re.match(r"^(?P<b>.+)_pref_(?P<k>\d+)$", c or "")
            if m:
                theo_buoi.setdefault(m.group("b"), []).append((int(m.group("k")), c))
        if theo_buoi:
            thu_tu = [c for b in sap_buoi(theo_buoi) for _, c in sorted(theo_buoi[b])]
        else:
            thu_tu = sorted([c for c in cot if (c or "").startswith("pref_")],
                            key=lambda c: int(c.split("_")[1]))
        if "club_id" in cot:                    # dạng dài
            s.setdefault("_dai", []).append(
                (int(r.get("rank") or 999), r["club_id"]))
            continue
        for c in thu_tu:
            v = (r.get(c) or "").strip()
            if v and v not in [x["club_id"] for x in s["nv"]]:
                s["nv"].append({"club_id": v, "thi": False, "diem": ""})
    for s in hs.values():
        for _, cid in sorted(s.pop("_dai", [])):
            if cid not in [x["club_id"] for x in s["nv"]]:
                s["nv"].append({"club_id": cid, "thi": False, "diem": ""})

    for r in thi_rows:
        if not (r.get("student_id") or "").strip():
            continue
        s = lay_hs(r)
        if "club_id" in r:
            if r.get("club_id"):
                s["_thi"][r["club_id"]] = r.get("score") or ""
            continue
        for c, v in r.items():
            m = re.match(r"^test_club_(\d+)$", c or "")
            if m and (v or "").strip():
                s["_thi"][v.strip()] = (r.get(f"score_{m.group(1)}") or "").strip()

    ngoai_nv = []
    for s in hs.values():
        thi = s.pop("_thi")
        for x in s["nv"]:
            if x["club_id"] in thi:
                x["thi"], x["diem"] = True, thi[x["club_id"]]
        trong_nv = {x["club_id"] for x in s["nv"]}
        ngoai_nv += [(s["student_id"], c) for c in thi if c not in trong_nv]
    return clubs, list(hs.values()), ngoai_nv
