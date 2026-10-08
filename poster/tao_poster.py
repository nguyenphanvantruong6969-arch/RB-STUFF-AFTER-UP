"""Sinh KHUNG poster ba lá kiểu QUẢNG CÁO cho gian trưng bày.

Kích thước thật khi mở phẳng: 200 cm x 150 cm
    lá trái 50 cm | lá giữa 100 cm | lá phải 50 cm
Gập hai lá bên vào 90 độ: rộng 100 cm, sâu 50 cm, cao 150 cm.

Khung này có sẵn HÌNH (ảnh giao diện thật, sơ đồ có nhãn "Sơ đồ do AI tạo ra",
biểu đồ có nhãn "Dữ liệu mô phỏng") và SỐ LIỆU đã kiểm chứng. Mọi ô CHỮ để trống
dạng [ ... ] cho học sinh tự viết — Phụ lục 1 cấm AI viết bản thảo poster
(BAN_GIAO.md mục 6). Điền gì vào ô nào: xem poster/DU_LIEU_POSTER.md.

Chạy (vẽ hình trước):
    python poster/chup_giao_dien.py   # ảnh giao diện thật
    python poster/tao_hinh.py         # sơ đồ + biểu đồ
    python poster/tao_poster.py       # poster

Sinh ra trong poster/:
    poster_khung.html          trang in, đúng tỉ lệ 1:1 (đơn vị cm)
    poster_khung_200x150cm.pdf bản in tỉ lệ 1:1
    poster_huong_dan.pdf       cùng khung, phủ nếp gấp, lề an toàn, kích thước, mã ô
    poster_khung_1-2.pptx      bản SỬA ĐƯỢC HẾT, tỉ lệ 1:2 (100 x 75 cm) — in phóng 200%.
                               Sơ đồ = hình khối + chữ, biểu đồ = biểu đồ gốc PowerPoint
                               (cùng số liệu với tao_hinh.py); chỉ ảnh chụp giao diện là ảnh.
    xem_truoc_phang.png        ảnh xem trước khi mở phẳng
    xem_truoc_gap.png          ảnh minh hoạ khi đã gập
"""

from pathlib import Path

from PIL import Image

HERE = Path(__file__).resolve().parent
HINH = HERE / "hinh"

# ---------------------------------------------------------------- kích thước
W, H = 200.0, 150.0          # cm, mở phẳng
FOLDS = (50.0, 150.0)        # vị trí hai nếp gấp
MARGIN = 3.0                 # lề ngoài
FOLD_GAP = 2.5               # khoảng trống mỗi bên nếp gấp
PAD = 1.8                    # đệm trong ô

# Cỡ chữ ở tỉ lệ 1:1 (pt). Bản pptx 1:2 lấy một nửa.
PT_TITLE = 140
PT_KICKER = 40
PT_SUB = 48
PT_HEAD = 60
PT_BODY = 32
PT_SMALL = 24
PT_STAT = 110

# ---------------------------------------------------------------- bảng màu
# Token của giao diện phần mềm (style.css): giấy lạnh + ink + vàng đồng
INK = "#1C2321"
SLATE = "#5B6570"
PAPER = "#F3F4F1"
RAISED = "#FFFFFF"
HAIRLINE = "#DBDCD5"
GOLD = "#C98A1F"
GOLD_SOFT = "#F1DFB8"
MOSS = "#3F6B52"
MOSS_SOFT = "#DCE9E0"
GREEN = "#27805A"            # màu số liệu, trùng màu biểu đồ trong tao_hinh.py
RUST = "#B84A3E"

LX, LW = MARGIN, FOLDS[0] - FOLD_GAP - MARGIN                    # 3 → 47,5
CX, CW = FOLDS[0] + FOLD_GAP, FOLDS[1] - FOLDS[0] - 2 * FOLD_GAP  # 52,5 → 147,5
RX, RW = FOLDS[1] + FOLD_GAP, W - MARGIN - FOLDS[1] - FOLD_GAP    # 152,5 → 197
HEAD_H = 2.9                 # chiều cao dòng tiêu đề ô (bằng vòng số)

BANNED = ["Logo nhà trường", "Tên trường (kể cả viết tắt)", "Mã dự án",
          "Tên giáo viên hướng dẫn", "Tên học sinh tham gia"]


def img_h(name, w):
    """Chiều cao (cm) của ảnh khi đặt rộng w cm, giữ tỉ lệ gốc."""
    iw, ih = Image.open(HINH / name).size
    return w * ih / iw


# ================================================================ bố cục
# Mọi phần tử là một dict; HTML và PPTX cùng đọc danh sách này nên luôn khớp nhau.
#   card   : nền ô            {x,y,w,h,style}
#   head   : dòng tiêu đề ô   {x,y,w,num,title,code}
#   hint   : ô chữ chờ viết   {x,y,w,h,text,size}
#   note   : dòng chú thích cố định (nguồn, nhãn AI) {x,y,w,text}
#   image  : ảnh              {x,y,w,h,file}
#   stat   : ô số nổi bật     {x,y,w,h,value,meta,hint}
#   qr     : ô vuông QR       {x,y,w}
#   header : đầu lá giữa      {x,y,w,h}
def layout():
    E = []

    def card(x, y, w, h, style="plain"):
        E.append(dict(t="card", x=x, y=y, w=w, h=h, style=style))

    def head(x, y, w, num, title, code):
        E.append(dict(t="head", x=x + PAD, y=y + 1.6, w=w - 2 * PAD, num=num, title=title, code=code))
        return y + 1.6 + HEAD_H + 1.0

    def hint(x, y, w, h, text, size=PT_BODY):
        E.append(dict(t="hint", x=x, y=y, w=w, h=h, text=text, size=size))

    def note(x, y, w, text):
        E.append(dict(t="note", x=x, y=y, w=w, text=text))

    def image(x, y, w, file):
        h = img_h(file, w)
        E.append(dict(t="image", x=x, y=y, w=w, h=h, file=file))
        return y + h

    iw_side = LW - 2 * PAD

    # ------------------------------------------------ lá trái
    card(LX, 3.0, LW, 44.0)
    y = head(LX, 3.0, LW, "01", "Vấn đề", "T1")
    hint(LX + PAD, y, iw_side, 47.0 - PAD - y,
         "[ T1 — em viết 2–3 câu. Hiện nay xếp học sinh vào CLB thế nào? Ai thiệt, thiệt ra sao? ]")

    card(LX, 49.5, LW, 54.0)
    y = head(LX, 49.5, LW, "02", "Dùng thế nào", "T2")
    y = image(LX + PAD, y, iw_side, "giao_dien_van_hanh.png")
    note(LX + PAD, y + 0.4, iw_side, "Ảnh chụp giao diện thật · dữ liệu mô phỏng")
    hint(LX + PAD, y + 2.4, iw_side, 103.5 - PAD - y - 2.4,
         "[ T2 — em viết: người dùng làm mấy bước? Mỗi bước một dòng. ]")

    card(LX, 106.0, LW, 40.0)
    y = head(LX, 106.0, LW, "03", "Xem thử", "T3")
    q = 20.0
    E.append(dict(t="qr", x=LX + PAD, y=y, w=q))
    hint(LX + PAD + q + 1.5, y, iw_side - q - 1.5, q,
         "[ T3 — em viết 1 câu mời xem demo / dùng thử tại gian ]")

    # ------------------------------------------------ lá giữa
    E.append(dict(t="header", x=CX, y=MARGIN, w=CW, h=27.0))

    shot_w = 61.0
    card(CX, 32.5, shot_w, 57.0)
    y = image(CX + PAD, 32.5 + PAD, shot_w - 2 * PAD, "giao_dien_ket_qua.png")
    note(CX + PAD, y + 0.4, shot_w - 2 * PAD, "Ảnh chụp giao diện thật · tab Kết quả · dữ liệu mô phỏng (120 học sinh, 10 CLB)")
    hint(CX + PAD, y + 2.4, shot_w - 2 * PAD, 89.5 - PAD - y - 2.4,
         "[ G1 — em viết 1–2 câu: ảnh này cho thấy phần mềm làm được gì? ]")

    sx = CX + shot_w + 2.5
    sw = CX + CW - sx
    stats = [
        ("S1", "0", "cặp phá vỡ · mọi lần bốc thăm, mọi bộ dữ liệu mô phỏng"),
        ("S2", "0,14 giây", "xếp 2 000 học sinh · 40 CLB · dữ liệu mô phỏng"),
        ("S3", "816", "bài kiểm thử tự động"),
    ]
    th = (57.0 - 2 * 2.0) / 3
    for i, (code, value, meta) in enumerate(stats):
        ty = 32.5 + i * (th + 2.0)
        E.append(dict(t="stat", x=sx, y=ty, w=sw, h=th, value=value, meta=meta,
                      hint=f"[ {code} — em viết 1 dòng: con số này nghĩa là gì? ]"))

    card(CX, 92.0, CW, 54.0)
    y = head(CX, 92.0, CW, "04", "Cách hoạt động", "G2")
    y = image(CX + PAD, y, CW - 2 * PAD, "so_do_5_lop.png")
    hint(CX + PAD, y + 0.8, CW - 2 * PAD, 146.0 - PAD - y - 0.8,
         "[ G2 — em viết 1–2 câu tóm tắt sơ đồ. Nhãn \"Sơ đồ do AI tạo ra\" đã có sẵn trong hình — giữ nguyên. ]")

    # ------------------------------------------------ lá phải
    iw_r = RW - 2 * PAD
    card(RX, 3.0, RW, 58.0)
    y = head(RX, 3.0, RW, "05", "Lợi ích", "P1–P4")
    row = (61.0 - PAD - y) / 4
    for i in range(4):
        ry = y + i * row
        E.append(dict(t="bullet", x=RX + PAD, y=ry, n=i + 1))
        hint(RX + PAD + 3.2, ry, iw_r - 3.2, row - 0.8,
             f"[ P{i + 1} — một lợi ích, tối đa 2 dòng ]")

    card(RX, 63.5, RW, 66.0)
    y = head(RX, 63.5, RW, "06", "Đã kiểm chứng", "P5")
    y = image(RX + PAD, y, iw_r, "bieu_do_khai_gian.png")
    y = image(RX + PAD, y + 0.8, iw_r, "bieu_do_nguyen_vong.png")
    hint(RX + PAD, y + 0.8, iw_r, 129.5 - PAD - y - 0.8,
         "[ P5 — em viết 1–2 câu: hai biểu đồ này nói gì với học sinh? ]", PT_BODY)

    card(RX, 132.0, RW, 14.0, style="muted")
    E.append(dict(t="head", x=RX + PAD, y=132.0 + 1.2, w=iw_r, num="", title="Tuyên bố sử dụng AI",
                  code="P6", small=True))
    hint(RX + PAD, 132.0 + 1.2 + 2.6, iw_r, 14.0 - 1.2 - 2.6 - 1.0,
         "[ P6 — theo Phụ lục 1: phần nào có AI hỗ trợ ]", PT_SMALL)
    return E


# ================================================================ HTML
def _pos(e, h=None):
    hh = e["h"] if h is None else h
    return f"left:{e['x']}cm;top:{e['y']}cm;width:{e['w']}cm;height:{hh}cm;"


def el_html(e):
    t = e["t"]
    if t == "card":
        return f'<div class="card {e["style"]}" style="{_pos(e)}"></div>'
    if t == "head":
        num = f'<span class="num">{e["num"]}</span>' if e["num"] else ""
        cls = "head small" if e.get("small") else "head"
        return (f'<h2 class="{cls}" style="left:{e["x"]}cm;top:{e["y"]}cm;width:{e["w"]}cm">'
                f'{num}<span>{e["title"]}</span></h2>')
    if t == "hint":
        return f'<p class="hint" style="{_pos(e)}font-size:{e["size"]}pt">{e["text"]}</p>'
    if t == "note":
        return f'<p class="note" style="left:{e["x"]}cm;top:{e["y"]}cm;width:{e["w"]}cm">{e["text"]}</p>'
    if t == "image":
        return f'<img class="img" src="hinh/{e["file"]}" style="{_pos(e)}">'
    if t == "qr":
        return (f'<div class="qr" style="{_pos(e, e["w"])}"><span>[ QR tới video demo<br>'
                f'KHÔNG trỏ tới GitHub ]</span></div>')
    if t == "bullet":
        return f'<span class="bullet" style="left:{e["x"]}cm;top:{e["y"] + 0.1}cm">{e["n"]}</span>'
    if t == "stat":
        return (f'<div class="stat" style="{_pos(e)}"><div class="v">{e["value"]}</div>'
                f'<div class="m">{e["meta"]}</div><div class="h">{e["hint"]}</div></div>')
    if t == "header":
        return f"""<header class="top" style="{_pos(e)}">
  <p class="kicker">[ H1 — LĨNH VỰC DỰ THI ]</p>
  <h1>[ H2 — THÔNG ĐIỆP CHÍNH,<br>MỘT CÂU, TỐI ĐA HAI DÒNG ]</h1>
  <p class="sub">[ H3 — tên đề tài đầy đủ ]</p>
</header>"""
    raise ValueError(t)


def guide_html():
    """Lớp phủ hướng dẫn — chỉ có trong poster_huong_dan.pdf."""
    f1, f2 = FOLDS
    out = ['<div class="guide">']
    bands = [(0, 0, W, MARGIN), (0, H - MARGIN, W, MARGIN), (0, 0, MARGIN, H), (W - MARGIN, 0, MARGIN, H)]
    bands += [(f - FOLD_GAP, 0, 2 * FOLD_GAP, H) for f in FOLDS]
    for x, y, w, h in bands:
        out.append(f'<div class="g-safe" style="left:{x}cm;top:{y}cm;width:{w}cm;height:{h}cm"></div>')
    for f in FOLDS:
        out.append(f'<div class="g-fold" style="left:{f}cm"></div>')
        out.append(f'<div class="g-foldlabel" style="left:{f}cm">NẾP GẤP</div>')
    for x0, x1, label in [(0, f1, "50 cm"), (f1, f2, "100 cm"), (f2, W, "50 cm")]:
        out.append(f'<div class="g-dim" style="left:{x0}cm;width:{x1 - x0}cm"><span>{label}</span></div>')
    out.append('<div class="g-vdim"><span>150 cm</span></div>')
    banned = "".join(f"<li>{b}</li>" for b in BANNED)
    out.append(f"""
<div class="g-note" style="left:{f1 + 6}cm;top:36cm;width:88cm">
  <h3>Hướng dẫn in — lớp này KHÔNG in</h3>
  <ul>
    <li>Mở phẳng 200 × 150 cm. Gập hai lá bên vào 90° → rộng 100, sâu 50, cao 150 cm.</li>
    <li>Vùng gạch chéo: lề ngoài {MARGIN:g} cm và {FOLD_GAP:g} cm mỗi bên nếp gấp — không đặt chữ.</li>
    <li>Cỡ chữ 1:1: thông điệp {PT_TITLE} pt · tên ô {PT_HEAD} pt · nội dung {PT_BODY} pt · chú thích {PT_SMALL} pt.</li>
    <li>Mã ô (H1, T1, G1, S1, P1 …) khớp với bảng tra trong poster/DU_LIEU_POSTER.md.</li>
  </ul>
  <h3 class="ban">Không được xuất hiện trên poster</h3>
  <ul class="ban">{banned}</ul>
</div>""")
    out.append("</div>")
    return "\n".join(out)


def build_html(guide=False):
    body = "\n".join(el_html(e) for e in layout())
    g = guide_html() if guide else ""
    return f"""<!doctype html>
<html lang="vi"><head><meta charset="utf-8">
<title>Khung poster ba lá</title>
<link rel="stylesheet" href="../assets/fonts/fonts.css">
<style>
@page {{ size: {W}cm {H}cm; margin: 0; }}
* {{ box-sizing: border-box; margin: 0; padding: 0; }}
html, body {{ width:{W}cm; height:{H}cm; }}
body {{ background:{PAPER}; color:{INK}; font-family:"IBM Plex Sans",sans-serif;
       -webkit-print-color-adjust:exact; print-color-adjust:exact; }}
.board {{ position:relative; width:{W}cm; height:{H}cm; overflow:hidden; background:{PAPER}; }}
.board > * {{ position:absolute; }}
.strip {{ top:0; height:1.2cm; }}
.card {{ background:{RAISED}; border:0.08cm solid {HAIRLINE}; border-radius:0.6cm; }}
.card.muted {{ background:{MOSS_SOFT}; border-color:{MOSS_SOFT}; }}
.head {{ font-family:"Be Vietnam Pro",sans-serif; font-weight:700; font-size:{PT_HEAD}pt; line-height:1;
        height:{HEAD_H}cm; display:flex; align-items:center; gap:0.8cm; }}
.head .num {{ flex:none; display:inline-grid; place-items:center; width:{HEAD_H}cm; height:{HEAD_H}cm;
             border-radius:50%; background:{GOLD}; color:{RAISED}; font-size:{PT_HEAD * 0.62:.0f}pt; }}
.head.small {{ font-size:{PT_HEAD * 0.6:.0f}pt; height:2cm; color:{MOSS}; }}
.hint {{ color:{SLATE}; font-style:italic; line-height:1.4; border:0.1cm dashed {HAIRLINE};
        border-radius:0.4cm; padding:0.8cm 1cm; overflow:hidden; }}
.note {{ font-size:{PT_SMALL}pt; color:{SLATE}; line-height:1.2; }}
.img {{ border:0.08cm solid {HAIRLINE}; border-radius:0.3cm; object-fit:contain; background:#fff; }}
.qr {{ border:0.14cm dashed {SLATE}; border-radius:0.4cm; background:{PAPER}; display:grid; place-items:center;
      text-align:center; font-size:{PT_SMALL}pt; color:{SLATE}; }}
.bullet {{ width:2.4cm; height:2.4cm; border-radius:50%; background:{GREEN}; color:#fff; display:grid;
          place-items:center; font:700 {PT_HEAD * 0.5:.0f}pt "Be Vietnam Pro"; }}
.stat {{ background:{RAISED}; border:0.08cm solid {HAIRLINE}; border-left:0.6cm solid {GREEN};
        border-radius:0.6cm; padding:1.2cm 1.6cm; display:flex; flex-direction:column; gap:0.35cm; }}
.stat .v {{ font:700 {PT_STAT}pt/1 "Be Vietnam Pro"; color:{GREEN}; letter-spacing:-0.02em; }}
.stat .m {{ font-size:{PT_SMALL}pt; color:{INK}; font-weight:600; }}
.stat .h {{ font-size:{PT_SMALL}pt; color:{SLATE}; font-style:italic; margin-top:auto;
           border-top:0.08cm dashed {HAIRLINE}; padding-top:0.4cm; }}
.top {{ background:{INK}; color:{PAPER}; border-radius:0.6cm; padding:2.2cm 3cm; display:flex;
       flex-direction:column; justify-content:center; gap:0.9cm; }}
.top .kicker {{ font-weight:600; font-size:{PT_KICKER}pt; letter-spacing:0.12em; color:{GOLD}; }}
.top h1 {{ font-family:"Be Vietnam Pro",sans-serif; font-weight:700; font-size:{PT_TITLE}pt; line-height:1.02; }}
.top .sub {{ font-size:{PT_SUB}pt; color:{HAIRLINE}; }}

.guide {{ inset:0; pointer-events:none; }}
.guide > * {{ position:absolute; }}
.g-safe {{ background:repeating-linear-gradient(45deg, rgba(184,74,62,.16) 0 0.5cm, transparent 0.5cm 1.2cm); }}
.g-fold {{ top:0; height:{H}cm; border-left:0.2cm dashed {RUST}; transform:translateX(-0.1cm); }}
.g-foldlabel {{ top:{H / 2}cm; transform:translate(-50%,-50%) rotate(-90deg); background:{RUST}; color:#fff;
               font:600 30pt "IBM Plex Sans"; padding:0.3cm 1.2cm; border-radius:0.3cm; }}
.g-dim {{ top:{H - 2.4}cm; height:1.8cm; border-left:0.12cm solid {RUST}; border-right:0.12cm solid {RUST};
         display:grid; place-items:center; }}
.g-dim::before {{ content:""; position:absolute; left:0; right:0; top:50%; border-top:0.12cm solid {RUST}; }}
.g-dim span, .g-vdim span {{ position:relative; background:{RUST}; color:#fff; font:600 34pt "IBM Plex Sans";
                            padding:0.1cm 1cm; border-radius:0.3cm; }}
.g-vdim {{ left:0.6cm; top:0; width:1.8cm; height:{H}cm; border-top:0.12cm solid {RUST};
          border-bottom:0.12cm solid {RUST}; display:grid; place-items:center; }}
.g-vdim::before {{ content:""; position:absolute; top:0; bottom:0; left:50%; border-left:0.12cm solid {RUST}; }}
.g-vdim span {{ transform:rotate(-90deg); white-space:nowrap; }}
.g-note {{ background:rgba(255,255,255,.96); border:0.2cm solid {RUST}; border-radius:0.6cm;
          padding:2cm 2.4cm; font-size:30pt; line-height:1.4; color:{INK}; }}
.g-note h3 {{ font:700 40pt "Be Vietnam Pro"; color:{RUST}; margin-bottom:0.8cm; }}
.g-note h3.ban {{ margin-top:1.4cm; }}
.g-note ul {{ padding-left:1.4cm; }}
.g-note ul.ban {{ display:flex; flex-wrap:wrap; gap:0.4cm 2.4cm; list-style:"✕  "; font-weight:600; }}
</style></head>
<body><div class="board">
<div class="strip" style="left:0;width:{FOLDS[0]}cm;background:{MOSS}"></div>
<div class="strip" style="left:{FOLDS[0]}cm;width:{FOLDS[1] - FOLDS[0]}cm;background:{GOLD}"></div>
<div class="strip" style="left:{FOLDS[1]}cm;width:{W - FOLDS[1]}cm;background:{MOSS}"></div>
{body}
{g}
</div></body></html>"""


# ================================================================ PPTX 1:2
def build_pptx(path, s=0.5):
    from pptx import Presentation
    from pptx.util import Cm, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.enum.dml import MSO_LINE_DASH_STYLE

    def rgb(h):
        return RGBColor.from_string(h.lstrip("#"))

    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(W * s), Cm(H * s)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    sh = slide.shapes

    def rect(x, y, w, h, fill, line=None, lw=0.08, rounded=True, dash=False, shape=None, into=None, radius=0.6):
        kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE)
        r = (into or sh).add_shape(kind, Cm(x * s), Cm(y * s), Cm(w * s), Cm(h * s))
        if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            r.adjustments[0] = min(radius / min(w, h), 0.5)
        r.fill.solid()
        r.fill.fore_color.rgb = rgb(fill)
        if line:
            r.line.color.rgb = rgb(line)
            r.line.width = Cm(lw * s)
            if dash:
                r.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        else:
            r.line.fill.background()
        r.shadow.inherit = False
        return r

    def text(x, y, w, h, runs, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, frame=None, into=None, wrap=True):
        tf = frame or (into or sh).add_textbox(Cm(x * s), Cm(y * s), Cm(w * s), Cm(h * s)).text_frame
        tf.word_wrap = wrap
        tf.vertical_anchor = anchor
        for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, m, 0)
        for i, (t, pt, color, font, bold, italic) in enumerate(runs):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run()
            r.text = t
            r.font.size = Pt(pt * s)
            r.font.color.rgb = rgb(color)
            r.font.name = font
            r.font.bold = bold
            r.font.italic = italic
        return tf

    # ---------- hình SỬA ĐƯỢC: sơ đồ = hình khối + chữ, biểu đồ = biểu đồ gốc PowerPoint
    import tao_hinh as th
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION

    PT_PER_CM = 72 / 2.54

    def ghi(g, x, y, k, e_units_x, e_units_y, chu, co, mau, dam, font):
        """Chữ đặt theo chân chữ (đơn vị khung hình) → hộp chữ không ngắt dòng."""
        top = y + (e_units_y - co * 0.9) * k
        text(x + e_units_x * k, top, 1, co * 1.3 * k,
             [(chu, co * k * PT_PER_CM, mau, font, dam, False)], into=g.shapes, wrap=False)

    def fam(f):
        return "Be Vietnam Pro" if "Vietnam" in f else "IBM Plex Sans"

    def so_do(e):
        k = e["w"] / th.SO_DO_W
        x0, y0 = e["x"], e["y"]
        g = sh.add_group_shape()
        g.name = "So do 5 lop"
        for p in th.so_do_spec():
            if p[0] == "rect":
                _, x, y, w, h, fill, stroke, sw = p
                rect(x0 + x * k, y0 + y * k, w * k, h * k, fill, stroke, sw * k, into=g.shapes, radius=18 * k)
            elif p[0] == "circle":
                _, cx, cy, r, fill, label = p
                c = rect(x0 + (cx - r) * k, y0 + (cy - r) * k, 2 * r * k, 2 * r * k, fill,
                         shape=MSO_SHAPE.OVAL, into=g.shapes)
                text(0, 0, 0, 0, [(label, 30 * k * PT_PER_CM, RAISED, "Be Vietnam Pro", True, False)],
                     MSO_ANCHOR.MIDDLE, PP_ALIGN.CENTER, frame=c.text_frame)
            elif p[0] == "text":
                _, x, y, chu, co, mau, dam, f = p
                ghi(g, x0, y0, k, x, y, chu, co, mau, dam, fam(f))
            elif p[0] == "arrow":
                _, x1, y, x2 = p
                rect(x0 + x1 * k, y0 + (y - 16) * k, (x2 - x1) * k, 32 * k, SLATE,
                     shape=MSO_SHAPE.RIGHT_ARROW, into=g.shapes)
            elif p[0] == "darrow":
                _, x, y1, y2 = p
                rect(x0 + (x - 16) * k, y0 + y1 * k, 32 * k, (y2 - y1) * k, SLATE,
                     shape=MSO_SHAPE.DOWN_ARROW, into=g.shapes)

    def chu_bieu_do(g, e, k, tieu_de, phu, nguon, y_nguon):
        ghi(g, e["x"], e["y"], k, 60, 80, tieu_de, 44, INK, True, "Be Vietnam Pro")
        ghi(g, e["x"], e["y"], k, 60, 128, phu, 28, SLATE, False, "IBM Plex Sans")
        ghi(g, e["x"], e["y"], k, 60, y_nguon, nguon, 24, "#8A929A", False, "IBM Plex Sans")

    def ve_bieu_do(g, e, k, kieu, du_lieu, y1, y2, mau, nhan, vmax, buoc):
        cd = CategoryChartData()
        cd.categories = [c for c, _ in du_lieu]
        cd.add_series("Số em", [v for _, v in du_lieu])
        gf = g.shapes.add_chart(kieu, Cm((e["x"] + 40 * k) * s), Cm((e["y"] + y1 * k) * s),
                                Cm(1520 * k * s), Cm((y2 - y1) * k * s), cd)
        ch = gf.chart
        ch.has_legend = False
        ch.has_title = False
        ch.font.size = Pt(26 * k * PT_PER_CM * s)
        ch.font.name = "IBM Plex Sans"
        ch.font.color.rgb = rgb(INK)
        plot = ch.plots[0]
        plot.gap_width = 70
        ser = plot.series[0]
        for i, pt in enumerate(ser.points):
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = rgb(mau[i])
            dl = pt.data_label
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
            dl.text_frame.text = nhan[i]
            r = dl.text_frame.paragraphs[0].runs[0]
            r.font.size = Pt(34 * k * PT_PER_CM * s)
            r.font.bold = True
            r.font.color.rgb = rgb(INK)
        va = ch.value_axis
        va.minimum_scale, va.maximum_scale, va.major_unit = 0, vmax, buoc
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = rgb(HAIRLINE)
        va.format.line.fill.background()
        va.tick_labels.font.color.rgb = rgb("#8A929A")
        ca = ch.category_axis
        ca.format.line.color.rgb = rgb(SLATE)
        ca.tick_labels.font.bold = True
        ca.tick_labels.font.size = Pt(32 * k * PT_PER_CM * s)
        return ch

    def bd_nguyen_vong(e):
        k = e["w"] / 1600
        g = sh.add_group_shape()
        g.name = "Bieu do nguyen vong"
        tong = sum(v for _, v in th.NV_DATA)
        chu_bieu_do(g, e, k, "Được xếp vào nguyện vọng thứ mấy",
                    f"{tong} em được xếp · bộ 120 học sinh, 10 CLB · hạt giống bốc thăm 42",
                    "Dữ liệu mô phỏng — nguồn: du_lieu_test/SO_LIEU_DA_KIEM_CHUNG.md, mục 2", 960)
        ve_bieu_do(g, e, k, XL_CHART_TYPE.COLUMN_CLUSTERED, th.NV_DATA, 160, 900,
                   [th.GREEN] + ["#7FAE95"] * 3,
                   [f"{v} em · {round(100 * v / tong)}%" for _, v in th.NV_DATA], 70, 20)

    def bd_khai_gian(e):
        k = e["w"] / 1600
        g = sh.add_group_shape()
        g.name = "Bieu do khai gian"
        chu_bieu_do(g, e, k, "Số em khai sai nguyện vọng mà được lợi",
                    "Vét cạn mọi cách khai của từng em · 200 thể hiện, 1 400 lượt học sinh",
                    "Dữ liệu mô phỏng — nguồn: NGHIEN_CUU_TOI_UU.md, TN3", 700)
        ch = ve_bieu_do(g, e, k, XL_CHART_TYPE.BAR_CLUSTERED, th.KHAI_GIAN_DATA, 170, 650,
                        [th.GREEN, th.GRAY_BAR],
                        ["0 / 1 400 em", "258 / 1 400 em (18,43%)"], 400, 100)
        ch.category_axis.reverse_order = True

    NATIVE = {"so_do_5_lop.png": so_do, "bieu_do_nguyen_vong.png": bd_nguyen_vong,
              "bieu_do_khai_gian.png": bd_khai_gian}

    rect(0, 0, W, H, PAPER, rounded=False)
    rect(0, 0, FOLDS[0], 1.2, MOSS, rounded=False)
    rect(FOLDS[0], 0, FOLDS[1] - FOLDS[0], 1.2, GOLD, rounded=False)
    rect(FOLDS[1], 0, W - FOLDS[1], 1.2, MOSS, rounded=False)

    for e in layout():
        t = e["t"]
        x, y, w = e["x"], e["y"], e.get("w", 0)
        if t == "card":
            muted = e["style"] == "muted"
            rect(x, y, w, e["h"], MOSS_SOFT if muted else RAISED, MOSS_SOFT if muted else HAIRLINE)
        elif t == "head":
            small = e.get("small")
            tx = x
            if e["num"]:
                c = rect(x, y, HEAD_H, HEAD_H, GOLD, shape=MSO_SHAPE.OVAL)
                text(0, 0, 0, 0, [(e["num"], PT_HEAD * 0.62, RAISED, "Be Vietnam Pro", True, False)],
                     MSO_ANCHOR.MIDDLE, PP_ALIGN.CENTER, frame=c.text_frame)
                tx = x + HEAD_H + 0.8
            text(tx, y, x + w - tx, 2.0 if small else HEAD_H,
                 [(e["title"], PT_HEAD * (0.6 if small else 1), MOSS if small else INK, "Be Vietnam Pro", True, False)],
                 MSO_ANCHOR.MIDDLE)
        elif t == "hint":
            b = rect(x, y, w, e["h"], RAISED, HAIRLINE, 0.1, dash=True)
            b.fill.background()
            text(x + 1, y + 0.8, w - 2, e["h"] - 1.6,
                 [(e["text"], e["size"], SLATE, "IBM Plex Sans", False, True)])
        elif t == "note":
            text(x, y, w, 1.6, [(e["text"], PT_SMALL, SLATE, "IBM Plex Sans", False, False)])
        elif t == "image" and e["file"] in NATIVE:
            NATIVE[e["file"]](e)
        elif t == "image":
            sh.add_picture(str(HINH / e["file"]), Cm(x * s), Cm(y * s), Cm(w * s), Cm(e["h"] * s))
        elif t == "qr":
            b = rect(x, y, w, w, PAPER, SLATE, 0.14, dash=True)
            text(0, 0, 0, 0, [("[ QR tới video demo — KHÔNG trỏ tới GitHub ]", PT_SMALL, SLATE, "IBM Plex Sans", False, False)],
                 MSO_ANCHOR.MIDDLE, PP_ALIGN.CENTER, frame=b.text_frame)
        elif t == "bullet":
            c = rect(x, y + 0.1, 2.4, 2.4, GREEN, shape=MSO_SHAPE.OVAL)
            text(0, 0, 0, 0, [(str(e["n"]), PT_HEAD * 0.5, RAISED, "Be Vietnam Pro", True, False)],
                 MSO_ANCHOR.MIDDLE, PP_ALIGN.CENTER, frame=c.text_frame)
        elif t == "stat":
            rect(x, y, w, e["h"], RAISED, HAIRLINE)
            rect(x, y, 0.6, e["h"], GREEN, rounded=False)
            text(x + 1.6, y + 1.2, w - 2.8, e["h"] - 2.4, [
                (e["value"], PT_STAT, GREEN, "Be Vietnam Pro", True, False),
                (e["meta"], PT_SMALL, INK, "IBM Plex Sans", True, False),
                (e["hint"], PT_SMALL, SLATE, "IBM Plex Sans", False, True),
            ])
        elif t == "header":
            rect(x, y, w, e["h"], INK)
            text(x + 3, y + 2.2, w - 6, e["h"] - 4.4, [
                ("[ H1 — LĨNH VỰC DỰ THI ]", PT_KICKER, GOLD, "IBM Plex Sans", True, False),
                ("[ H2 — THÔNG ĐIỆP CHÍNH, MỘT CÂU, TỐI ĐA HAI DÒNG ]", PT_TITLE, PAPER, "Be Vietnam Pro", True, False),
                ("[ H3 — tên đề tài đầy đủ ]", PT_SUB, HAIRLINE, "IBM Plex Sans", False, False),
            ], MSO_ANCHOR.MIDDLE)

    slide.notes_slide.notes_text_frame.text = (
        "Tỉ lệ 1:2 — slide 100 x 75 cm, in phóng 200% ra 200 x 150 cm.\n"
        "Nếp gấp ở 25 cm và 75 cm trên slide (50 cm và 150 cm khi in).\n"
        "Mã ô (H1, T1, G1, S1, P1…) khớp bảng tra poster/DU_LIEU_POSTER.md.\n"
        "Không đặt: " + ", ".join(BANNED) + ".")
    prs.save(path)


# ================================================================ ảnh gập
def folded_html(png_name):
    # Mỗi lá lấy đúng phần của ảnh phẳng; hai lá bên xoay quanh nếp gấp.
    return f"""<!doctype html><html><head><meta charset="utf-8"><style>
body {{ margin:0; width:1600px; height:1100px; background:#E6E7E2; font-family:sans-serif; overflow:hidden; }}
.scene {{ position:absolute; left:800px; top:140px; width:0; height:0; perspective:2800px;
         perspective-origin:0 250px; }}
.board {{ position:absolute; left:-500px; top:0; width:1000px; height:1500px; transform-style:preserve-3d;
         transform-origin:50% 0; transform:scale(.5); }}
.p {{ position:absolute; top:0; height:1500px; background:url({png_name}) no-repeat;
     background-size:2000px 1500px; box-shadow:0 0 0 1px rgba(0,0,0,.08); }}
.c {{ left:0; width:1000px; background-position:-500px 0; }}
.l {{ left:-500px; width:500px; background-position:0 0; transform-origin:right center;
     transform:rotateY(62deg); filter:brightness(.93); }}
.r {{ left:1000px; width:500px; background-position:-1500px 0; transform-origin:left center;
     transform:rotateY(-62deg); filter:brightness(.93); }}
.floor {{ position:absolute; left:-300px; top:1480px; width:1600px; height:60px;
         background:radial-gradient(ellipse at center, rgba(0,0,0,.28), transparent 70%); }}
.cap {{ position:fixed; bottom:24px; left:0; right:0; text-align:center; color:#5B6570; font-size:22px; }}
</style></head><body>
<div class="scene"><div class="board"><div class="floor"></div>
<div class="p l"></div><div class="p c"></div><div class="p r"></div></div></div>
<div class="cap">Minh hoạ khi gập — hai lá bên gập vào 90° khi trưng bày thật: rộng 100 cm, sâu 50 cm, cao 150 cm</div>
</body></html>"""


def _chromium():
    """Dùng Chromium có sẵn trên máy nếu bản Playwright chưa tải trình duyệt riêng."""
    import glob
    for pat in ("/opt/pw-browsers/chromium", "/opt/pw-browsers/chromium-*/chrome-linux/chrome"):
        for c in sorted(glob.glob(pat)):
            if Path(c).is_file():
                return c
    return None


def main():
    from playwright.sync_api import sync_playwright

    thieu = [f for f in ("giao_dien_van_hanh.png", "giao_dien_ket_qua.png", "so_do_5_lop.png",
                         "bieu_do_khai_gian.png", "bieu_do_nguyen_vong.png") if not (HINH / f).exists()]
    if thieu:
        raise SystemExit("Thiếu hình: %s — chạy chup_giao_dien.py và tao_hinh.py trước." % ", ".join(thieu))

    (HERE / "poster_khung.html").write_text(build_html(), encoding="utf-8")
    (HERE / "_huong_dan.html").write_text(build_html(guide=True), encoding="utf-8")

    with sync_playwright() as p:
        br = p.chromium.launch(executable_path=_chromium())
        pg = br.new_page()
        for src, pdf in [("poster_khung.html", "poster_khung_200x150cm.pdf"),
                         ("_huong_dan.html", "poster_huong_dan.pdf")]:
            pg.goto((HERE / src).as_uri())
            pg.evaluate("document.fonts.ready")
            pg.pdf(path=str(HERE / pdf), width=f"{W}cm", height=f"{H}cm",
                   print_background=True, prefer_css_page_size=True)
        pg2 = br.new_page(viewport={"width": 2000, "height": 1500})
        pg2.goto((HERE / "poster_khung.html").as_uri())
        pg2.evaluate("document.fonts.ready")
        z = 2000 / (W * 96 / 2.54)
        pg2.evaluate(f"document.body.style.zoom='{z}'")
        pg2.wait_for_timeout(300)
        pg2.screenshot(path=str(HERE / "xem_truoc_phang.png"), clip={"x": 0, "y": 0, "width": 2000, "height": 1500})
        (HERE / "_gap.html").write_text(folded_html("xem_truoc_phang.png"), encoding="utf-8")
        pg3 = br.new_page(viewport={"width": 1600, "height": 1100})
        pg3.goto((HERE / "_gap.html").as_uri())
        pg3.wait_for_timeout(300)
        pg3.screenshot(path=str(HERE / "xem_truoc_gap.png"))
        br.close()

    for tmp in ("_huong_dan.html", "_gap.html"):
        (HERE / tmp).unlink()
    build_pptx(HERE / "poster_khung_1-2.pptx")
    print("Đã sinh tệp trong", HERE)


if __name__ == "__main__":
    main()
