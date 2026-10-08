"""Sinh tệp Word giải thích biểu mẫu 6 phần: THIET_KE_6_PHAN.docx.

    python mau_forms_thi_clb/tao_word_6_phan.py

Danh sách CLB và đề thi đọc thẳng từ THIET_KE_6_PHAN.html (bản đã duyệt),
nên tệp Word và trang chạy thử không lệch nhau. Toàn bộ chữ trong tệp Word
viết bằng tiếng Việt đơn giản và không dùng gạch ngang dài.

Cần python-docx (requirements-dev.txt).
"""

import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

THU_MUC = os.path.dirname(os.path.abspath(__file__))
HTML = os.path.join(THU_MUC, "THIET_KE_6_PHAN.html")
RA = os.path.join(THU_MUC, "THIET_KE_6_PHAN.docx")

MAU = RGBColor(0x0D, 0x69, 0x64)
CO_BUOI, KHONG_BUOI = "Có", "Không, em bận buổi này"
CO, KHONG, DUNG = "Có, em thi CLB này", "Không", "Em không thi thêm CLB nào buổi này"
DIEM_CAU, TRAN_THI = 5, 5


def doc_du_lieu():
    """[(mã buổi, tên buổi, [(club_id, tên, [(câu, đáp án đúng, [lựa chọn])])])]."""
    sys.path.insert(0, THU_MUC)
    import tao_google_forms
    return tao_google_forms.doc_du_lieu()


def so_cau(du_lieu):
    """Đánh số câu liên tục như Forms: {mã câu: số}."""
    so, n = {"student_id": 1, "name": 2}, 2
    for ma, _ten, ds in du_lieu:
        n += 1; so["di-" + ma] = n
        n += 1; so[ma] = n
        for cid, _name, de in ds:
            n += 1; so["thi-" + cid] = n
            for j in range(len(de)):
                n += 1; so["%s-%d" % (cid, j + 1)] = n
    return so, n


# ------------------------------------------------------------- định dạng

def to_nen(o, mau_hex):
    tc = o._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), mau_hex)
    tc.append(shd)


def bang(doc, tieu_de, dong, dam_cot_dau=False):
    t = doc.add_table(rows=1, cols=len(tieu_de))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(tieu_de):
        o = t.rows[0].cells[i]
        o.text = ""
        r = o.paragraphs[0].add_run(h)
        r.bold, r.font.size = True, Pt(10)
        to_nen(o, "DBECEA")
    for d in dong:
        cells = t.add_row().cells
        for i, v in enumerate(d):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(10)
            r.bold = dam_cot_dau and i == 0
    doc.add_paragraph()
    return t


def doan(doc, chu, dam=None, nghieng=False):
    p = doc.add_paragraph()
    if dam:
        r = p.add_run(dam)
        r.bold = True
    r = p.add_run(chu)
    r.italic = nghieng
    return p


def gach(doc, chu, dam=None, kieu="List Bullet"):
    p = doc.add_paragraph(style=kieu)
    if dam:
        p.add_run(dam).bold = True
    p.add_run(chu)
    return p


# ------------------------------------------------------------- nội dung

def viet(du_lieu):
    so, tong = so_cau(du_lieu)
    n_buoi = len(du_lieu)
    n_clb = sum(len(ds) for _, _, ds in du_lieu)
    n_cau_buoi = 2 + sum(1 + len(de) for _, _, de in du_lieu[0][2])
    so_re = n_buoi * (1 + 2 * len(du_lieu[0][2]))
    so_re_gon = n_buoi * (1 + len(du_lieu[0][2]))
    so_dap_an = sum(len(de) for _, _, ds in du_lieu for _, _, de in ds)

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    for ten in ("Title", "Heading 1", "Heading 2"):
        doc.styles[ten].font.color.rgb = MAU
        doc.styles[ten].font.name = "Calibri"

    doc.add_heading("Biểu mẫu đăng ký CLB và thi tuyển gồm 6 phần", 0)
    doan(doc, "Bản thiết kế đã duyệt. Tài liệu này giải thích cách biểu mẫu hoạt động "
              "và minh hoạ bằng một học sinh điền thử.", nghieng=True)
    doan(doc, "Tệp này để đọc. Đừng nhập tệp này vào Forms. Tệp để nhập là "
              "NHAP_FORMS_6_PHAN.docx.", dam="Lưu ý: ")

    # 1
    doc.add_heading("1. Mục đích", 1)
    doan(doc, "Học sinh dùng một biểu mẫu Microsoft Forms duy nhất để làm ba việc: "
              "đăng ký CLB theo từng buổi trong tuần, chọn CLB muốn thi, và làm luôn đề thi "
              "của CLB đó. Bài thi được chấm tự động. Giáo viên tải kết quả ra Excel rồi "
              "nạp vào phần mềm xếp CLB.")

    # 2
    doc.add_heading("2. Cấu trúc 6 phần", 1)
    doan(doc, "Phần 1 là thông tin học sinh. Mỗi buổi học trong tuần là một phần riêng.")
    dong = [["Phần 1", "Thông tin học sinh", "Mã học sinh, họ và tên", "2"]]
    for k, (ma, ten, ds) in enumerate(du_lieu):
        dong.append(["Phần %d" % (k + 2), ten, "%d CLB, mỗi CLB 2 câu đề thi" % len(ds),
                     str(n_cau_buoi)])
    dong.append(["", "Tổng cộng", "%d CLB" % n_clb, str(tong)])
    bang(doc, ["Phần", "Tên", "Nội dung", "Số câu"], dong, dam_cot_dau=True)
    doan(doc, "Mỗi buổi có %d câu: 1 câu mở buổi, 1 câu xếp hạng, %d câu chọn CLB để thi "
              "và %d câu đề thi. Học sinh chỉ thấy đề của CLB mình chọn thi."
         % (n_cau_buoi, len(du_lieu[0][2]), n_cau_buoi - 2 - len(du_lieu[0][2])))

    # 3
    doc.add_heading("3. Một buổi đi thế nào", 1)
    doan(doc, "Lấy Thứ Hai làm ví dụ. Bốn buổi còn lại giống hệt, chỉ khác danh sách CLB.")
    gach(doc, " “Em có đăng ký sinh hoạt Thứ Hai không?” Chọn Có thì đi tiếp. "
              "Chọn “Không, em bận buổi này” thì sang ngay Phần 3 (Thứ Ba). "
              "Học sinh bận buổi này chỉ cần bấm một lần.",
         dam="Bước 1. Câu mở buổi.", kieu="List Paragraph")
    gach(doc, " Học sinh kéo thả 10 CLB, CLB thích nhất lên đầu. Bước này làm trước khi chọn "
              "CLB để thi.", dam="Bước 2. Xếp hạng 10 CLB.", kieu="List Paragraph")
    gach(doc, " Mỗi CLB một câu “Thi CLB ...?”, theo đúng thứ tự trong danh sách. Mỗi câu "
              "có ba lựa chọn:", dam="Bước 3. Chọn CLB để thi.", kieu="List Paragraph")
    gach(doc, " đề thi của CLB đó hiện ra ngay bên dưới.", dam="“%s”:" % CO)
    gach(doc, " sang câu của CLB kế tiếp.", dam="“%s”:" % KHONG)
    gach(doc, " sang thẳng buổi sau, không phải bấm Không cho các CLB còn lại.",
         dam="“%s”:" % DUNG)
    gach(doc, " Mỗi đề 2 câu, mỗi câu %d điểm. Làm xong 2 câu, biểu mẫu tự đi tới câu "
              "chọn của CLB kế tiếp." % DIEM_CAU,
         dam="Bước 4. Làm đề thi.", kieu="List Paragraph")
    doan(doc, "Mỗi buổi học sinh được thi tối đa %d CLB." % TRAN_THI, dam="Lưu ý: ")

    # 4
    ma2, ten2, ds2 = du_lieu[0]
    ma3, ten3, _ = du_lieu[1]
    c1, c2, c3 = ds2[0], ds2[1], ds2[2]
    doc.add_heading("4. Ví dụ: một học sinh điền thử", 1)
    doan(doc, "Bạn Nguyễn Minh An (mã HS024) muốn thi %s và %s vào %s, còn %s thì bận."
         % (c1[1], c3[1], ten2, ten3))
    dong = [
        [so["student_id"], "Mã học sinh", "Ghi HS024", "Câu kế tiếp"],
        [so["name"], "Họ và tên", "Ghi Nguyễn Minh An", "Phần 2 · %s" % ten2],
        [so["di-" + ma2], "Em có đăng ký sinh hoạt %s không?" % ten2, CO_BUOI,
         "Câu %d (xếp hạng)" % so[ma2]],
        [so[ma2], "Xếp hạng 10 CLB", "Kéo %s lên số 1, %s lên số 2" % (c1[1], c3[1]),
         "Câu %d" % so["thi-" + c1[0]]],
        [so["thi-" + c1[0]], "Thi %s?" % c1[1], CO, "Hiện 2 câu đề thi (câu %d, %d)"
         % (so[c1[0] + "-1"], so[c1[0] + "-2"])],
        [so[c1[0] + "-1"], c1[2][0][0], "Chọn “%s” (đúng)" % c1[2][0][1], "Câu kế tiếp"],
        [so[c1[0] + "-2"], c1[2][1][0], "Chọn “%s” (sai)" % c1[2][1][2][1],
         "Câu %d" % so["thi-" + c2[0]]],
        [so["thi-" + c2[0]], "Thi %s?" % c2[1], KHONG,
         "Bỏ qua đề, sang câu %d" % so["thi-" + c3[0]]],
        [so["thi-" + c3[0]], "Thi %s?" % c3[1], CO, "Hiện 2 câu đề thi"],
        [so[c3[0] + "-1"], c3[2][0][0], "Chọn “%s” (đúng)" % c3[2][0][1], "Câu kế tiếp"],
        [so[c3[0] + "-2"], c3[2][1][0], "Chọn “%s” (đúng)" % c3[2][1][1],
         "Câu %d" % so["thi-" + ds2[3][0]]],
        [so["thi-" + ds2[3][0]], "Thi %s?" % ds2[3][1], DUNG, "Phần 3 · %s" % ten3],
        [so["di-" + ma3], "Em có đăng ký sinh hoạt %s không?" % ten3, KHONG_BUOI,
         "Phần 4 · %s" % du_lieu[2][1]],
    ]
    bang(doc, ["Câu", "Câu hỏi", "An chọn", "Biểu mẫu đi tới"], dong)
    doan(doc, "Kết quả của An sau khi nộp:", dam="")
    bang(doc, ["Buổi", "Nguyện vọng 1", "CLB đã thi", "Điểm (thang 10)"], [
        [ten2, c1[1], c1[1], "5 (đúng 1/2 câu)"],
        ["", "", c3[1], "10 (đúng 2/2 câu)"],
        [ten3, "Bận", "Không thi", ""],
    ])
    n_tra_loi = sum(1 for d in dong if so["di-" + ma2] <= d[0] < so["di-" + ma3])
    doan(doc, "Ở %s, An chỉ trả lời %d câu thay vì phải lướt qua cả %d câu. Ở %s, An chỉ "
              "bấm một lần." % (ten2, n_tra_loi, n_cau_buoi, ten3))

    # 5
    doc.add_heading("5. Bảng rẽ nhánh của Phần 2 · %s" % ten2, 1)
    doan(doc, "Rẽ nhánh đặt trên từng câu, ngay trong cùng một phần. Lựa chọn Có luôn để "
              "mặc định, vì câu kế tiếp chính là câu đề thi. Bốn buổi còn lại đặt theo cùng "
              "quy tắc.")
    sau = "Phần 3 · %s" % ten3
    dong = [[so["di-" + ma2], "Em có đăng ký sinh hoạt %s không?" % ten2,
             "Câu %d" % so[ma2], sau, ""],
            [so[ma2], "Xếp hạng 10 CLB", "Không rẽ nhánh", "", ""]]
    for i, (cid, name, _de) in enumerate(ds2):
        ke = ds2[i + 1][0] if i + 1 < len(ds2) else None
        dong.append([so["thi-" + cid], "Thi %s?" % name,
                     "Câu %d (mặc định)" % (so["thi-" + cid] + 1),
                     ("Câu %d" % so["thi-" + ke]) if ke else sau, sau])
    bang(doc, ["Câu", "Nội dung", "Có", "Không", "Không thi thêm"], dong)

    # 6
    doc.add_heading("6. Chấm điểm", 1)
    gach(doc, "Mỗi đề 2 câu, mỗi câu %d điểm. Điểm mỗi CLB theo thang 10." % DIEM_CAU)
    gach(doc, "Chấm theo nội dung đáp án học sinh chọn, không theo vị trí. Vì vậy có thể "
              "bật xáo trộn lựa chọn cho từng câu đề thi.")
    gach(doc, "Không cho học sinh xem đáp án ngay sau khi nộp, để bạn làm trước không báo "
              "đáp án cho bạn làm sau.")
    gach(doc, "Bộ chuyển đổi đọc tệp Excel xuất từ Forms, chấm lại theo đáp án gốc, và báo "
              "nếu điểm Forms chấm khác đáp án gốc.")

    # 7
    doc.add_heading("7. So với bản cũ", 1)
    bang(doc, ["", "Bản cũ", "Bản 6 phần"], [
        ["Số phần", "26", "6"],
        ["CLB mỗi buổi", "3", "10"],
        ["CLB có thi", "10 (2 mỗi buổi)", "%d (tất cả)" % n_clb],
        ["Số câu mỗi đề", "5 câu, mỗi câu 2 điểm", "2 câu, mỗi câu %d điểm" % DIEM_CAU],
        ["Tổng số câu hỏi", "67", str(tong)],
        ["Học sinh bận một buổi", "Bấm Không ở 2 câu dự thi", "Bấm một lần"],
        ["Học sinh chỉ thi 1 CLB trong buổi", "Bấm Không cho CLB còn lại",
         "Bấm “Không thi thêm” là xong buổi"],
        ["Chỗ rẽ nhánh phải cài tay", "10", str(so_re)],
        ["Câu phải đánh dấu đáp án", "50", str(so_dap_an)],
    ], dam_cot_dau=True)

    # 8
    doc.add_heading("8. Những điều cần biết", 1)
    gach(doc, " Câu chọn nhiều ô không rẽ nhánh được. Vì vậy mỗi CLB là một câu chọn một: "
              "Có, Không, hoặc Không thi thêm.",
         dam="Microsoft Forms không cho tick nhiều ô rồi hiện đề.")
    gach(doc, " Forms không đổi thứ tự câu hỏi theo câu trả lời, nên câu chọn CLB đi theo "
              "danh sách, không theo thứ tự học sinh vừa xếp.",
         dam="Câu chọn CLB có thứ tự cố định.")
    gach(doc, " Phần mềm xếp CLB chỉ nhận %d. Forms không tự chặn được, nên câu hỏi ghi rõ "
              "“tối đa %d” và bộ chuyển đổi sẽ báo nếu học sinh chọn quá." % (TRAN_THI, TRAN_THI),
         dam="Tối đa %d CLB dự thi mỗi buổi." % TRAN_THI)
    gach(doc, " Nếu bỏ trống, Forms coi như đi tiếp và đề thi vẫn hiện ra.",
         dam="Câu chọn CLB và câu đề thi phải là câu bắt buộc.")
    gach(doc, " CLB học sinh không muốn thì để cuối.", dam="Câu xếp hạng buộc xếp đủ 10 CLB.")
    gach(doc, " Có %d chỗ rẽ nhánh và %d câu phải đánh dấu đáp án. Bỏ lựa chọn “Không thi "
              "thêm” thì còn %d chỗ rẽ nhánh, nhưng học sinh phải bấm nhiều hơn."
         % (so_re, so_dap_an, so_re_gon), dam="Phải cài tay khá nhiều.")
    gach(doc, " Bản này có %d câu, dưới mức khoảng 200 câu mà Forms cho phép. Thêm CLB hoặc "
              "thêm câu đề thi thì phải tính lại." % tong, dam="Giới hạn số câu.")

    # 9
    doc.add_heading("9. Các bước dựng trên Microsoft Forms", 1)
    for chu in [
        "Trên forms.office.com chọn Quick import, chọn tệp NHAP_FORMS_6_PHAN.docx, chọn "
        "dạng Quiz. Không nhập tệp tài liệu này.",
        "Kiểm tra lại: đủ 6 phần, 162 câu, câu nào có lựa chọn thì vẫn là câu chọn một.",
        "Đổi 5 câu xếp hạng (mã [thu_2] tới [thu_6]) sang dạng Ranking, giữ nguyên chữ và "
        "các lựa chọn.",
        "Đặt bắt buộc cho câu mở buổi, câu xếp hạng, câu chọn CLB và câu đề thi.",
        "Với mỗi câu đề thi: đánh dấu đáp án đúng, đặt %d điểm." % DIEM_CAU,
        "Đặt rẽ nhánh theo bảng ở mục 5. Các buổi khác đặt theo cùng quy tắc.",
        "Làm như vậy cho cả 5 buổi, Phần 2 tới Phần 6.",
        "Giữ nguyên phần [mã] ở đầu mỗi câu hỏi. Bộ chuyển đổi dựa vào mã này để đọc tệp Excel.",
        "Cài đặt: chỉ người trong trường được điền, mỗi người một lần, không xáo trộn câu "
        "hỏi, không tự hiện kết quả cho học sinh.",
        "Tự điền thử hai phiếu, tải Excel và chạy bộ chuyển đổi trước khi gửi cho học sinh.",
    ]:
        gach(doc, chu, kieu="List Number")

    # phụ lục
    doc.add_page_break()
    doc.add_heading("Phụ lục. Danh sách CLB và đề thi mẫu", 1)
    doan(doc, "Tên CLB và câu hỏi là nội dung mẫu. Trường thay bằng đề thật trước khi dùng.",
         nghieng=True)
    for ma, ten, ds in du_lieu:
        doc.add_heading(ten, 2)
        dong = []
        for cid, name, de in ds:
            for j, (q, dung, _ops) in enumerate(de):
                dong.append([name if j == 0 else "", "Câu %d" % so["%s-%d" % (cid, j + 1)],
                             q, dung])
        bang(doc, ["CLB", "Số câu", "Câu hỏi", "Đáp án đúng"], dong, dam_cot_dau=True)

    return doc


def kiem_chu(doc):
    """Không được có gạch ngang dài hay gạch ngang vừa ở bất cứ đâu."""
    chu = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            chu += [c.text for c in r.cells]
    sai = [x for x in chu if "—" in x or "–" in x]
    if sai:
        raise SystemExit("Còn gạch ngang dài: %r" % sai[:3])


def main():
    du_lieu = doc_du_lieu()
    doc = viet(du_lieu)
    kiem_chu(doc)
    # Mẫu mặc định của python-docx thiếu thuộc tính percent ở w:zoom (sai
    # chuẩn OOXML). Word bỏ qua được, nhưng thêm vào cho tệp hợp lệ.
    for z in doc.settings.element.findall(qn("w:zoom")):
        z.set(qn("w:percent"), "100")
    doc.save(RA)
    print("Đã ghi %s" % RA)


if __name__ == "__main__":
    sys.exit(main())
