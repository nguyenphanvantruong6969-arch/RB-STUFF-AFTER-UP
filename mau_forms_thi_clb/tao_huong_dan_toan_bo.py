"""Sinh tệp Word hướng dẫn toàn bộ quy trình: HUONG_DAN_TOAN_BO_QUY_TRINH.docx.

    python mau_forms_thi_clb/tao_huong_dan_toan_bo.py

Từ soạn đề thi, cài skill, sinh mã Google Apps Script, tạo biểu mẫu, chấm và
xuất dữ liệu, xử lí raw, tới nạp vào phần mềm xếp CLB; cộng phần sửa mã và bảo
trì. Tiếng Việt đơn giản, không gạch ngang dài (kiem_chu chặn khi sinh).

Cần python-docx (requirements-dev.txt).
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Pt, RGBColor  # noqa: E402

import tao_word_6_phan as tw  # noqa: E402
from tao_word_6_phan import bang, doan, gach  # noqa: E402

RA = os.path.join(tw.THU_MUC, "HUONG_DAN_TOAN_BO_QUY_TRINH.docx")
KHO = "nguyenphanvantruong6969-arch/rbda-kiosk"
NHANH = "claude/microsoft-form-project-collection-p71hzh"

_DEM = [0]


def muc(doc, ten, cap=1):
    _DEM[0] = 0
    return doc.add_heading(ten, cap)


def buoc(doc, chu, dam=None):
    """Bước đánh số, đếm lại từ 1 sau mỗi tiêu đề (List Number của Word đánh nối tiếp)."""
    _DEM[0] += 1
    p = doc.add_paragraph(style="List Paragraph")
    p.add_run("%d. " % _DEM[0]).bold = True
    if dam:
        p.add_run(dam).bold = True
    p.add_run(chu)
    return p


def lenh(doc, *dong):
    """Khối lệnh: chữ đơn cách, nền xám nhạt."""
    for d in dong:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before, pf.space_after = Pt(0), Pt(0)
        pf.left_indent = Pt(18)
        ppr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "F2F4F4")
        ppr.append(shd)
        r = p.add_run(d)
        r.font.name = "Consolas"
        r.element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        r.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def luu_y(doc, chu, dam="Lưu ý: "):
    p = doan(doc, chu, dam=dam)
    for r in p.runs:
        r.font.color.rgb = RGBColor(0x8A, 0x5A, 0x00)
    return p


def viet():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    for ten in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        doc.styles[ten].font.color.rgb = tw.MAU
        doc.styles[ten].font.name = "Calibri"

    doc.add_heading("Hướng dẫn toàn bộ quy trình: từ đề thi tới dữ liệu đầu vào phần mềm xếp CLB", 0)
    doan(doc, "Tài liệu này đi từ đầu tới cuối: cần chuẩn bị gì, cài skill thế nào, soạn đề, "
              "sinh mã, tạo biểu mẫu Google Forms, chấm và xuất dữ liệu, biến dữ liệu thành "
              "một Sổ nhập CLB nạp vào phần mềm, và cách sửa mã khi cần. Làm theo đúng thứ tự các phần.",
         nghieng=True)

    # ------------------------------------------------------------------ 0
    muc(doc, "0. Nhìn toàn cảnh")
    doan(doc, "Có năm chặng. Mỗi chặng ra một thứ để chặng sau dùng.")
    bang(doc, ["Chặng", "Làm gì", "Dùng gì", "Ra gì"], [
        ["1. Soạn đề", "Điền buổi, CLB, đề thi, đáp án", "Excel", "MAU_DE_THI.xlsx"],
        ["2. Sinh mã", "Biến tệp Excel thành mã tạo biểu mẫu", "Skill tao-google-form-clb",
         "TAO_GOOGLE_FORM.gs, HUONG_DAN_TUNG_BUOC.md"],
        ["3. Tạo biểu mẫu", "Chạy mã trên Google, cài đặt, gửi học sinh", "Google Apps Script, Google Forms",
         "Biểu mẫu và phiếu trả lời"],
        ["4. Chấm và xuất", "Chấm điểm từng CLB, xuất dữ liệu", "Hàm chamTheoCLB trong mã",
         "Thư mục Kết quả CLB (4 tệp Excel)"],
        ["5. Xử lí raw", "Soát, chấm lại, ra sổ nhập", "Skill xu-li-raw-forms",
         "SO_NHAP_CLB.xlsx và báo cáo bất thường"],
    ], dam_cot_dau=True)
    doan(doc, "Sau chặng 5 là kéo Sổ nhập CLB vào phần mềm xếp CLB và chạy sắp xếp (mục 10).")

    # ------------------------------------------------------------------ 1
    muc(doc, "1. Cần chuẩn bị gì")
    bang(doc, ["Thứ cần có", "Để làm gì", "Ghi chú"], [
        ["Tài khoản Google (nên dùng tài khoản của trường)", "Chạy Apps Script, chứa biểu mẫu và kết quả",
         "Học sinh cũng cần đăng nhập Google nếu bật giới hạn một lần trả lời"],
        ["Trình duyệt (Chrome, Edge)", "Mở script.google.com, Google Forms, Google Drive", ""],
        ["Claude có Skills: claude.ai hoặc Claude Code", "Chạy hai skill",
         "Trên claude.ai cần bật chạy mã (code execution) và tải skill lên"],
        ["Hai tệp skill: tao-google-form-clb.zip, xu-li-raw-forms.zip", "Sinh mã, xử lí raw",
         "Hoặc dùng thẳng trong kho mã rbda-kiosk (đã có sẵn trong .claude/skills/)"],
        ["Excel hoặc LibreOffice", "Điền MAU_DE_THI.xlsx, điền chỉ tiêu CLB", ""],
        ["Phần mềm xếp CLB (PhanBoCauLacBo.exe) trên Windows 10 hoặc 11", "Nạp dữ liệu, chạy sắp xếp",
         "Không cần cài đặt, giải nén là chạy"],
        ["Chỉ khi tự chạy lệnh: Python 3.10 trở lên và thư viện openpyxl", "Chạy script của skill trên máy",
         "Cài bằng: pip install openpyxl"],
        ["Chỉ khi sửa mã: kho mã rbda-kiosk, Git, pytest", "Sửa khuôn mẫu, chạy kiểm tra, đóng gói lại skill",
         "Kho: %s, nhánh %s" % (KHO, NHANH)],
    ], dam_cot_dau=True)

    muc(doc, "1.1. Các tệp và vai trò của chúng", 2)
    bang(doc, ["Tệp", "Vai trò"], [
        ["MAU_DE_THI.xlsx", "Đầu vào của mọi thứ: buổi, CLB, đề thi, đáp án, cài đặt"],
        ["TAO_GOOGLE_FORM.gs", "Mã Google Apps Script. Có hai hàm: taoBieuMau (tạo biểu mẫu) và chamTheoCLB (chấm, xuất Excel)"],
        ["HUONG_DAN_TUNG_BUOC.md", "Hướng dẫn riêng cho đúng bộ đề vừa sinh (số buổi, số CLB, ví dụ làm thử)"],
        ["Thư mục Kết quả CLB <ngày giờ>", "Trên Google Drive, do chamTheoCLB tạo, chứa 4 tệp Excel (mục 7)"],
        ["RAW_PHIEU.xlsx, DAP_AN.xlsx", "Dữ liệu raw đã chấm từng câu, và đáp án. Đầu vào của skill xu-li-raw-forms"],
        ["SO_NHAP_CLB.xlsx", "Sổ nhập CLB: tệp DUY NHẤT phần mềm xếp CLB nạp được (trang 1. CLB, trang 2. Học sinh)"],
        ["BAO_CAO_BAT_THUONG.md", "Những điều bất thường trong phiếu, kèm mã học sinh"],
        ["DINH_DANG_RAW.md", "Mô tả từng cột của dữ liệu raw. Skill xử lí raw đọc theo tài liệu này"],
    ], dam_cot_dau=True)

    # ------------------------------------------------------------------ 2
    muc(doc, "2. Cài skill để Claude dùng được")
    doan(doc, "Skill là một thư mục có tệp SKILL.md (hướng dẫn cho Claude) và các script. Claude "
              "đọc SKILL.md để biết khi nào dùng và chạy lệnh nào. Có ba cách cài, chọn một.")

    muc(doc, "2.1. Trên claude.ai", 2)
    buoc(doc, "Đăng nhập claude.ai. Vào Settings (Cài đặt).")
    buoc(doc, "Tìm mục bật chạy mã và tạo tệp (Code execution and file creation), bật lên. "
              "Skill cần chạy script Python nên phải bật mục này.")
    buoc(doc, "Vào mục Skills (thường nằm trong Capabilities), bấm Upload skill, chọn "
              "tao-google-form-clb.zip. Làm lại với xu-li-raw-forms.zip.")
    buoc(doc, "Kiểm tra hai skill đã hiện trong danh sách và đang bật.")
    luu_y(doc, "Tên các mục trong Settings có thể khác một chút tuỳ phiên bản và gói tài khoản. "
               "Nếu không thấy mục Skills, tài khoản có thể chưa có tính năng này.")

    muc(doc, "2.2. Trên Claude Code, cho mọi dự án", 2)
    buoc(doc, "Giải nén hai tệp .zip vào thư mục ~/.claude/skills/ trên máy "
              "(Windows: C:\\Users\\<tên>\\.claude\\skills\\).")
    buoc(doc, "Sau khi giải nén phải có: ~/.claude/skills/tao-google-form-clb/SKILL.md và "
              "~/.claude/skills/xu-li-raw-forms/SKILL.md.")
    buoc(doc, "Mở lại Claude Code. Skill tự được nhận.")

    muc(doc, "2.3. Trong kho mã rbda-kiosk", 2)
    doan(doc, "Hai skill đã nằm sẵn trong .claude/skills/ của kho. Mở Claude Code trong thư mục kho "
              "là dùng được, không cần cài. Đây là cách đầy đủ nhất: skill xử lí raw còn nạp thử "
              "kết quả vào phần mềm thật để kiểm tra.")

    muc(doc, "2.4. Skill cần gì để chạy được", 2)
    gach(doc, " Python 3.10 trở lên và thư viện openpyxl (đọc, ghi Excel). Trên claude.ai "
              "môi trường chạy mã thường có sẵn; trên máy mình thì cài bằng pip install openpyxl.",
         dam="Phần mềm:")
    gach(doc, " mỗi tệp .zip đã mang theo mọi tệp phụ thuộc (khuôn mẫu mã, bộ soát, tài liệu "
              "định dạng, tệp Excel mẫu). Không cần tải thêm gì.", dam="Tệp đi kèm:")
    gach(doc, " với tao-google-form-clb là tệp Excel đề thi; với xu-li-raw-forms là "
              "RAW_PHIEU.xlsx, DAP_AN.xlsx và SO_NHAP_CLB.xlsx đã điền chỉ tiêu.",
         dam="Đầu vào:")
    gach(doc, " nói đúng việc cần làm, ví dụ “tạo Google Form từ tệp này” hoặc “xử lí raw”. "
              "Claude nhận ra và dùng skill.", dam="Cách gọi:")

    doan(doc, "Bên trong mỗi tệp .zip:")
    lenh(doc,
         "tao-google-form-clb/",
         "  SKILL.md                       hướng dẫn cho Claude",
         "  scripts/tao_form.py            lệnh chính: mau, tao",
         "  scripts/tao_google_forms.py    khuôn mẫu sinh tệp .gs",
         "  scripts/sinh_du_lieu.py        bộ soát theo giới hạn của phần mềm",
         "  scripts/so_nhap.py             định nghĩa Sổ nhập CLB (đọc, ghi)",
         "  scripts/i18n_errors.py         câu báo lỗi của sổ nhập",
         "  scripts/THIET_KE_6_PHAN.html   bộ đề ví dụ",
         "  assets/MAU_DE_THI.xlsx         tệp Excel mẫu",
         "  references/DINH_DANG_RAW.md    định dạng dữ liệu raw",
         "",
         "xu-li-raw-forms/",
         "  SKILL.md",
         "  scripts/xu_li_raw.py           lệnh chính",
         "  scripts/sinh_du_lieu.py        bộ soát",
         "  scripts/so_nhap.py             ghi Sổ nhập CLB",
         "  scripts/i18n_errors.py",
         "  references/DINH_DANG_RAW.md")

    # ------------------------------------------------------------------ 3
    muc(doc, "3. Chặng 1: soạn đề trong MAU_DE_THI.xlsx")
    doan(doc, "Lấy tệp mẫu: có sẵn trong assets/ của skill, trong kho ở mau_forms_thi_clb/MAU_DE_THI.xlsx, "
              "hoặc nhờ Claude “tạo tệp Excel mẫu đề thi”. Tệp mẫu đã điền sẵn 50 CLB ví dụ để thấy "
              "cách điền. Phải thay bằng dữ liệu thật của trường.")

    muc(doc, "3.1. Trang Cai_dat", 2)
    bang(doc, ["khoa", "gia_tri", "Ý nghĩa"], [
        ["tieu_de", "Đăng ký CLB và thi tuyển", "Tên biểu mẫu học sinh thấy"],
        ["diem_moi_cau", "5", "Điểm mỗi câu đề thi (số nguyên lớn hơn 0)"],
        ["toi_da_clb_thi_moi_buoi", "5", "Mỗi buổi học sinh được tick tối đa bao nhiêu CLB thi (1 tới 5)"],
    ], dam_cot_dau=True)

    muc(doc, "3.2. Trang CLB", 2)
    doan(doc, "Mỗi CLB một dòng, theo đúng thứ tự muốn hiện trên biểu mẫu.")
    bang(doc, ["Cột", "Điền gì", "Ví dụ"], [
        ["buoi", "Mã buổi: thu_2, thu_3, thu_4, thu_5, thu_6, thu_7, chu_nhat", "thu_2"],
        ["ten_buoi", "Tên hiện cho học sinh. Bỏ trống thì tự điền theo mã buổi", "Thứ Hai"],
        ["club_id", "Mã CLB: chữ không dấu, số, gạch dưới. Không trùng", "clb_covua"],
        ["ten_clb", "Tên hiện cho học sinh", "CLB Cờ vua"],
    ], dam_cot_dau=True)
    gach(doc, "Tối đa 10 CLB mỗi buổi (giới hạn nguyện vọng của phần mềm).")
    gach(doc, "Mã CLB phải giữ nguyên qua mọi năm nếu muốn so sánh số liệu. Đổi mã là thành CLB khác.")

    muc(doc, "3.3. Trang De_thi", 2)
    doan(doc, "Mỗi câu hỏi một dòng.")
    bang(doc, ["Cột", "Điền gì"], [
        ["club_id", "Mã CLB của câu này, phải có ở trang CLB"],
        ["cau_hoi", "Nội dung câu hỏi"],
        ["lua_chon_a, b, c, d", "Các lựa chọn. Ít nhất 2, điền liền từ a. Không để trống ở giữa. Không có hai lựa chọn giống nhau"],
        ["dap_an", "Một chữ a, b, c hoặc d (viết hoa cũng được)"],
    ], dam_cot_dau=True)
    gach(doc, "CLB không có dòng nào ở trang này là CLB không tổ chức thi: vẫn có trong danh sách "
              "Top để học sinh chọn nguyện vọng, nhưng không có trong câu tick và không có đề.")
    gach(doc, "Nên để khoảng 2 tới 5 câu mỗi CLB. Biểu mẫu trên khoảng 400 câu có thể làm Apps Script "
              "chạy quá 6 phút (mục 11).")

    # ------------------------------------------------------------------ 4
    muc(doc, "4. Chặng 2: sinh mã TAO_GOOGLE_FORM.gs")
    muc(doc, "4.1. Cách dễ: nhờ Claude", 2)
    buoc(doc, "Mở cuộc trò chuyện với Claude (đã cài skill, mục 2).")
    buoc(doc, "Đính kèm MAU_DE_THI.xlsx đã điền, gõ: “tạo Google Form từ tệp này”.")
    buoc(doc, "Claude chạy skill và gửi lại hai tệp: TAO_GOOGLE_FORM.gs và HUONG_DAN_TUNG_BUOC.md, "
              "cùng tóm tắt số buổi, số CLB, số câu.")
    buoc(doc, "Nếu Claude báo “KHÔNG GHI TỆP NÀO”, tệp Excel có chỗ sai. Claude liệt kê từng lỗi kèm "
              "số dòng (ví dụ “De_thi dòng 12: dap_an ‘e’ phải là một chữ trong a, b, c, d”). Sửa "
              "trong Excel rồi gửi lại.")
    doan(doc, "Có đề ở dạng khác (văn bản, Word)? Gửi cho Claude và nói rõ muốn tạo biểu mẫu. Claude "
              "chép nguyên văn vào tệp mẫu và hỏi lại chỗ thiếu đáp án; không tự bịa câu hỏi.")

    muc(doc, "4.2. Cách tự chạy lệnh (trong kho mã hoặc thư mục skill đã giải nén)", 2)
    doan(doc, "Trong kho mã:")
    lenh(doc,
         "pip install openpyxl",
         "python .claude/skills/tao-google-form-clb/scripts/tao_form.py mau --ra MAU_DE_THI.xlsx",
         "python .claude/skills/tao-google-form-clb/scripts/tao_form.py tao --vao MAU_DE_THI.xlsx --ra form_moi")
    doan(doc, "Ở thư mục skill đã giải nén, thay đường dẫn đầu bằng tao-google-form-clb/scripts/tao_form.py. "
              "Lệnh mau tạo tệp Excel mẫu; lệnh tao đọc tệp đã điền và ghi hai tệp vào thư mục form_moi.")

    muc(doc, "4.3. Biểu mẫu sinh ra có gì", 2)
    doan(doc, "Phần đầu là mã học sinh và họ tên. Mỗi buổi sau đó:")
    buoc(doc, "“Em có đăng ký sinh hoạt <buổi> không?” Chọn Không thì sang buổi sau.")
    buoc(doc, "Top 1 tới Top N (N là số CLB của buổi): mỗi Top là một danh sách thả xuống, chọn một CLB "
              "hoặc “(Bỏ trống)”. Không bắt xếp đủ.")
    buoc(doc, "Tick các CLB muốn thi. Google chặn cứng, không cho tick quá số tối đa đã đặt.")
    buoc(doc, "“Em thi CLB nào trước?”, làm đề CLB đó, rồi “CLB nào tiếp theo?” (chỉ các CLB đứng sau "
              "trong danh sách). Học sinh chỉ thấy đề của CLB mình chọn.")
    doan(doc, "Mọi câu hỏi có mã trong ngoặc vuông ở đầu, ví dụ [thu_2-top-1], [clb_covua-1]. Mã này "
              "giúp hàm chấm và skill đọc đúng dữ liệu. Không xoá mã khi sửa chữ trên Google Forms.")

    # ------------------------------------------------------------------ 5
    muc(doc, "5. Chặng 3: tạo biểu mẫu trên Google")
    muc(doc, "5.1. Chạy mã (làm một lần, khoảng 5 phút)", 2)
    buoc(doc, "Mở script.google.com bằng tài khoản Google của trường, bấm “Dự án mới”. Cách khác: mở một "
              "biểu mẫu Google Forms trống, bấm dấu ba chấm góc trên bên phải, chọn “Apps Script”.")
    buoc(doc, "Xoá hết chữ có sẵn trong khung soạn thảo (dòng function myFunction...).")
    buoc(doc, "Mở TAO_GOOGLE_FORM.gs bằng Notepad, chọn tất cả (Ctrl + A), sao chép (Ctrl + C), dán vào "
              "khung soạn thảo (Ctrl + V).")
    buoc(doc, "Bấm biểu tượng đĩa mềm để lưu. Bấm vào tên dự án ở góc trên bên trái, đổi thành “Đăng ký CLB”.")
    buoc(doc, "Ở thanh trên, chọn hàm taoBieuMau trong ô chọn hàm, bấm “Chạy”.")
    buoc(doc, "Google hỏi quyền: bấm “Xem lại quyền”, chọn tài khoản.")
    buoc(doc, "Màn hình “Google chưa xác minh ứng dụng này”: bấm “Nâng cao” (hoặc “Hiển thị cài đặt nâng "
              "cao”), rồi “Đi tới Đăng ký CLB (không an toàn)”. Cảnh báo này hiện vì mã do chính mình "
              "tạo, chưa qua Google duyệt; người phát triển là email của chính mình.")
    buoc(doc, "Màn hình chọn quyền: tick “Chọn tất cả” (Google Drive, Google Trang tính, Biểu mẫu, Kết nối "
              "với dịch vụ bên ngoài), bấm “Tiếp tục”. Bỏ quyền nào thì hàm cần quyền đó sẽ báo lỗi.")
    buoc(doc, "Đợi 1 tới 3 phút. Bấm “Nhật ký thực thi” nếu khung chưa mở. Chép hai đường dẫn: đường dẫn "
              "SỬA biểu mẫu và đường dẫn gửi HỌC SINH. Khung nhật ký không bấm được đường dẫn: bôi đen, "
              "Ctrl + C, dán vào tab mới.")
    luu_y(doc, "mỗi lần chạy taoBieuMau là tạo thêm một biểu mẫu mới. Chỉ chạy một lần. Lỡ chạy hai lần thì "
               "xoá biểu mẫu thừa trong Google Drive.")

    doan(doc, "Bốn quyền dùng để làm gì:")
    bang(doc, ["Quyền", "Mã dùng để"], [
        ["Google Drive", "Tạo thư mục Kết quả CLB, lưu tệp Excel, dọn bảng tính tạm vào thùng rác"],
        ["Google Trang tính", "Tạo bảng điểm và các trang 01, 02, 03, raw, đáp án"],
        ["Biểu mẫu", "Tạo biểu mẫu, đọc phiếu học sinh nộp để chấm"],
        ["Kết nối với dịch vụ bên ngoài", "Chuyển bảng tính thành tệp .xlsx qua trang xuất tệp của chính Google"],
    ], dam_cot_dau=True)
    doan(doc, "Thu hồi quyền lúc nào cũng được: myaccount.google.com, mục Bảo mật, Ứng dụng và dịch vụ bên thứ ba.")

    muc(doc, "5.2. Cài đặt và làm thử (bắt buộc trước khi gửi học sinh)", 2)
    doan(doc, "Mở đường dẫn SỬA biểu mẫu.")
    buoc(doc, "Nếu nút “Xuất bản” góc trên còn màu tím, bấm để xuất bản. Chưa xuất bản thì học sinh không mở được.")
    buoc(doc, "Tab Cài đặt, mục Bài kiểm tra: “Công bố điểm” chọn “Sau khi xem xét thủ công”. Bỏ tick "
              "“Câu hỏi bị trả lời sai” và “Câu trả lời đúng”. Như vậy học sinh nộp xong không thấy đáp án "
              "để báo cho bạn làm sau.")
    buoc(doc, "Tab Cài đặt, mục Câu trả lời: bật “Giới hạn 1 câu trả lời”. Học sinh phải đăng nhập Google.")
    buoc(doc, "Bấm biểu tượng con mắt (Xem trước), làm thử hai phiếu: một phiếu bận mọi buổi; một phiếu ở một "
              "buổi chọn vài CLB ở Top 1, Top 2, tick 2 CLB, làm đề CLB thứ nhất rồi CLB thứ hai, rồi chọn "
              "“Em không thi thêm CLB nào buổi này”. Kiểm tra chỉ hiện đề của đúng hai CLB đó.")
    buoc(doc, "Chạy chamTheoCLB (mục 7) trên hai phiếu thử, rồi xử lí raw (mục 9). Đi trọn một vòng rồi "
              "mới gửi học sinh. Xong thì xoá hai phiếu thử: tab Câu trả lời, dấu ba chấm, “Xoá tất cả "
              "câu trả lời”.")

    # ------------------------------------------------------------------ 6
    muc(doc, "6. Gửi học sinh và đóng khảo sát")
    buoc(doc, "Gửi đường dẫn HỌC SINH qua nhóm lớp, email hoặc mã QR (Google Forms có nút Gửi, chọn biểu "
              "tượng đường dẫn để lấy đường dẫn ngắn).")
    buoc(doc, "Nhắc học sinh bốn điều: buổi bận thì chọn “Không, em bận buổi này”; bấm vào Top 1 chọn CLB "
              "thích nhất, rồi Top 2...; mỗi CLB chỉ chọn một lần, Top không dùng để “(Bỏ trống)”; tick "
              "trước các CLB muốn thi (tối đa 5), rồi chọn lần lượt để làm bài. Chỉ bài của CLB đã tick "
              "và đã chọn ở một Top mới được tính điểm.")
    buoc(doc, "Theo dõi số phiếu ở tab Câu trả lời.")
    buoc(doc, "Hết hạn: tab Câu trả lời, tắt “Chấp nhận câu trả lời”.")

    # ------------------------------------------------------------------ 7
    muc(doc, "7. Chặng 4: chấm và xuất dữ liệu")
    buoc(doc, "Mở lại dự án Apps Script ở mục 5 (script.google.com, mục Dự án của tôi, hoặc từ biểu mẫu: dấu "
              "ba chấm, Apps Script).")
    buoc(doc, "Chọn hàm chamTheoCLB, bấm “Chạy”. Lần đầu Google có thể hỏi lại quyền: làm như mục 5.1.")
    buoc(doc, "Nhật ký thực thi hiện: “Đã chấm N học sinh”, đường dẫn “Thư mục kết quả”, số cảnh báo nếu "
              "có, và lời nhắc điền chỉ tiêu.")
    buoc(doc, "Mở đường dẫn thư mục (chép từ nhật ký, dán vào tab mới). Hoặc vào drive.google.com, tìm "
              "thư mục “Kết quả CLB <ngày giờ>”.")
    buoc(doc, "Tải về: quay ra thư mục cha, bấm chuột phải vào thư mục “Kết quả CLB ...”, chọn “Tải "
              "xuống”. Google nén thành một tệp .zip. Giải nén.")
    doan(doc, "Bốn tệp trong thư mục:")
    bang(doc, ["Tệp", "Nội dung", "Dùng để"], [
        ["BANG_DIEM.xlsx", "Bảng điểm từng CLB (mỗi em một dòng) và trang Cảnh báo",
         "Giáo viên xem"],
        ["SO_NHAP_CLB.xlsx", "Sổ nhập CLB: trang 1. CLB (cột Chỉ tiêu để trống), trang 2. Học sinh "
         "(nguyện vọng, điểm cạnh CLB đã thi)", "Điền chỉ tiêu (mục 8), rồi nạp thẳng hoặc qua skill"],
        ["RAW_PHIEU.xlsx", "Mọi câu trả lời của mọi phiếu, đã chấm từng câu", "Đầu vào skill xu-li-raw-forms"],
        ["DAP_AN.xlsx", "Đáp án 100 câu (hoặc số câu của đề)", "Đầu vào skill xu-li-raw-forms"],
    ], dam_cot_dau=True)
    gach(doc, "Có phiếu nộp muộn thì chạy lại chamTheoCLB. Mỗi lần chạy tạo một thư mục mới, thư mục cũ giữ nguyên.")
    gach(doc, "Cách khác để lấy câu trả lời thô: trong Google Forms, tab Câu trả lời, dấu ba chấm, “Tải các "
              "câu trả lời xuống (.csv)”. Tệp này chưa tách điểm từng CLB và skill không đọc tệp này; "
              "chỉ dùng để lưu trữ.")
    gach(doc, "Biểu mẫu tạo bằng bản mã cũ vẫn chấm được bằng chamTheoCLB bản mới: dán mã mới vào dự án "
              "(không chạy taoBieuMau), điền đường dẫn SỬA biểu mẫu vào dòng const LINK_BIEU_MAU = '...' "
              "ở đầu mã nếu nhật ký báo “Chưa biết chấm biểu mẫu nào”.")

    # ------------------------------------------------------------------ 8
    muc(doc, "8. Điền chỉ tiêu CLB")
    buoc(doc, "Mở SO_NHAP_CLB.xlsx bằng Excel, trang “1. CLB”.")
    buoc(doc, "Điền cột Chỉ tiêu cho mọi CLB: số học sinh tối đa, số nguyên lớn hơn 0.")
    buoc(doc, "Có suất dành cho nhóm ưu tiên thì điền Suất ưu tiên (không lớn hơn Chỉ tiêu) và "
              "Nhóm ưu tiên (ví dụ chinh_sach). Không có thì để trống cả hai.")
    buoc(doc, "Không sửa Mã CLB và Buổi. Lưu lại, giữ dạng .xlsx.")

    # ------------------------------------------------------------------ 9
    muc(doc, "9. Chặng 5: xử lí raw thành Sổ nhập CLB")
    muc(doc, "9.1. Cách dễ: nhờ Claude", 2)
    buoc(doc, "Đính kèm ba tệp: RAW_PHIEU.xlsx, DAP_AN.xlsx, SO_NHAP_CLB.xlsx (đã điền chỉ tiêu).")
    buoc(doc, "Gõ: “xử lí raw”.")
    buoc(doc, "Claude chạy skill xu-li-raw-forms, gửi lại một SO_NHAP_CLB.xlsx mới (điểm đã chấm lại) và "
              "BAO_CAO_BAT_THUONG.md, cùng tóm tắt những điều cần xem.")
    buoc(doc, "Nếu chạy trong kho mã rbda-kiosk, Claude còn nạp thử sổ vào phần mềm thật và báo số "
              "cảnh báo.")

    muc(doc, "9.2. Cách tự chạy lệnh", 2)
    lenh(doc,
         "python .claude/skills/xu-li-raw-forms/scripts/xu_li_raw.py \\",
         "    --raw    \"Kết quả CLB .../RAW_PHIEU.xlsx\" \\",
         "    --dap-an \"Kết quả CLB .../DAP_AN.xlsx\" \\",
         "    --clb    \"Kết quả CLB .../SO_NHAP_CLB.xlsx\" \\",
         "    --ra     nap")
    doan(doc, "Ở thư mục skill đã giải nén, dùng xu-li-raw-forms/scripts/xu_li_raw.py. Trên Windows, viết "
              "cả lệnh trên một dòng và bỏ dấu \\ ở cuối mỗi dòng.")
    doan(doc, "Không có --clb thì skill vẫn chạy, nhưng cột Chỉ tiêu của sổ để trống và phải điền "
              "trước khi nạp. Mã thoát 1 và dòng “KHÔNG GHI TỆP NÀO” nghĩa là đầu vào hỏng (thiếu cột, "
              "tệp chỉ tiêu thiếu CLB, buổi không khớp): đọc lý do, sửa, chạy lại.")

    muc(doc, "9.3. Skill làm gì với dữ liệu", 2)
    gach(doc, "Chấm lại mọi câu theo DAP_AN.xlsx, không tin điểm có sẵn.")
    gach(doc, "Mỗi mã học sinh chỉ giữ phiếu nộp đầu tiên (đây là bài thi, nộp lại không thành làm lại).")
    gach(doc, "Chỉ tính bài của CLB đã tick và đã chọn ở một Top, tối đa 5 CLB mỗi buổi.")
    gach(doc, "Một CLB chọn ở nhiều Top: giữ Top cao nhất. Bỏ trống một Top ở giữa: dồn lên.")
    gach(doc, "Soát theo 9 điều bắt buộc của phần mềm. Có lỗi thì không ghi tệp nào.")

    muc(doc, "9.4. Đọc BAO_CAO_BAT_THUONG.md trước khi nạp", 2)
    bang(doc, ["Mục trong báo cáo", "Nghĩa là", "Nên làm"], [
        ["Đáp án trên Google Forms lệch đáp án gốc", "Nhiều bài lệch ở cùng một câu: đáp án trên Forms bị tick sai",
         "Sửa đáp án trên Forms. Điểm trong sổ đã theo đáp án gốc"],
        ["Làm bài mà không tick CLB đó", "Bài không được tính", "Liên hệ học sinh nếu cần; không tự cộng điểm"],
        ["Tick CLB mà không làm bài", "Em tick nhưng không làm đề", "Chỉ cần biết"],
        ["Phiếu nộp trùng, phiếu không có mã", "Đã giữ phiếu đầu, bỏ phiếu không mã", "Kiểm với danh sách lớp"],
        ["Mã học sinh khác dạng với số đông, một email nhiều mã", "Có thể gõ nhầm mã", "Sửa mã cho đúng rồi chạy lại"],
        ["Bỏ trống Top ở giữa, chọn một CLB ở nhiều Top", "Đã tự dồn lên, giữ Top cao nhất", "Chỉ cần biết"],
        ["Thiếu xếp hạng, câu hỏi lạ", "Biểu mẫu có thể bị sửa tay", "Kiểm lại biểu mẫu"],
    ], dam_cot_dau=True)

    # ------------------------------------------------------------------ 10
    muc(doc, "10. Nạp vào phần mềm xếp CLB")
    buoc(doc, "Mở PhanBoCauLacBo.exe. Lần đầu Windows có thể hiện “Windows protected your PC”: bấm More info, "
              "rồi Run anyway.")
    buoc(doc, "Thẻ “01 · Vận hành sắp xếp”: kéo SO_NHAP_CLB.xlsx thả vào ô kéo thả. Phần mềm đọc thử và "
              "hiện dòng tóm tắt (số CLB, buổi, học sinh); sổ có lỗi thì liệt kê từng lỗi kèm trang, dòng. "
              "Đối chiếu con số rồi bấm “Nhập sổ”.")
    buoc(doc, "Đọc mục “Cảnh báo dữ liệu” ngay dưới. Cảnh báo không chặn nhưng có thể âm thầm làm đổi kết "
              "quả. Hiểu hết từng cảnh báo rồi mới đi tiếp.")
    buoc(doc, "Bấm “Chạy sắp xếp”, xác nhận. Ghi lại số seed đang dùng: cùng dữ liệu và cùng seed thì luôn ra "
              "cùng kết quả, để kiểm chứng sau này.")
    buoc(doc, "Bấm “Xuất kết quả”. Tệp ra ở thư mục Tải xuống (Downloads): một tệp .xlsx mang đi họp, tệp "
              "tổng .csv, và thư mục mỗi CLB một tệp kèm danh sách em chưa được xếp.")
    doan(doc, "Chi tiết phần mềm: tệp HUONG_DAN_SU_DUNG.md trong kho mã.")

    # ------------------------------------------------------------------ 11
    muc(doc, "11. Giới hạn cần biết")
    bang(doc, ["Giới hạn", "Vì sao", "Đã xử lí thế nào"], [
        ["Câu tick không tự hiện đề", "Google Forms không rẽ nhánh theo câu chọn nhiều",
         "Tick trước để bị chặn ở mức tối đa, rồi chọn lần lượt CLB để làm bài"],
        ["Chọn làm bài theo thứ tự trong danh sách", "Google không nhớ em đã làm CLB nào",
         "Trang “tiếp theo” chỉ liệt kê CLB đứng sau"],
        ["Chọn trùng CLB ở hai Top", "Các danh sách thả xuống độc lập", "Khi chấm giữ Top cao hơn, có cảnh báo"],
        ["Chỉ có một tổng điểm trên Forms", "Chế độ Quiz cộng mọi câu", "chamTheoCLB tách điểm từng CLB"],
        ["Apps Script chạy tối đa 6 phút mỗi lần", "Giới hạn của Google",
         "Trên khoảng 400 câu thì chia thành hai biểu mẫu"],
        ["Tối đa 10 CLB mỗi buổi, 5 CLB thi mỗi buổi", "Giới hạn của phần mềm xếp CLB", "Skill chặn ngay khi sinh mã"],
    ], dam_cot_dau=True)

    # ------------------------------------------------------------------ 12
    muc(doc, "12. Lỗi thường gặp")
    bang(doc, ["Hiện tượng", "Nguyên nhân", "Cách xử lí"], [
        ["“Google chưa xác minh ứng dụng này”", "Mã tự tạo, chưa qua Google duyệt",
         "Nâng cao, rồi Đi tới ... (không an toàn), rồi chọn tất cả quyền"],
        ["Biểu mẫu hiện hết mọi câu dù chưa chọn", "Đang ở màn hình sửa", "Bấm Xem trước để thấy như học sinh"],
        ["“Chưa biết chấm biểu mẫu nào”", "Dự án không nhớ biểu mẫu (tạo từ nơi khác hoặc bản cũ)",
         "Dán đường dẫn SỬA biểu mẫu vào const LINK_BIEU_MAU ở đầu mã, lưu, chạy lại chamTheoCLB"],
        ["“Không xuất được ... (mã lỗi ...)”", "Mạng hoặc Google bận", "Chạy lại chamTheoCLB"],
        ["Chạy quá 6 phút, bị ngắt", "Biểu mẫu quá dài", "Xoá biểu mẫu dở, chia đề thành hai tệp Excel"],
        ["Có hai biểu mẫu giống nhau", "Chạy taoBieuMau hai lần", "Xoá một, chỉ gửi một đường dẫn"],
        ["Học sinh không mở được", "Chưa xuất bản, đã tắt nhận câu trả lời, hoặc chưa đăng nhập Google",
         "Kiểm tra ba điều đó"],
        ["Skill báo “KHÔNG GHI TỆP NÀO”", "Tệp Excel hoặc raw có chỗ sai", "Đọc danh sách lỗi có số dòng, sửa, chạy lại"],
        ["Phần mềm báo Chỉ tiêu phải lớn hơn 0", "Chưa điền cột Chỉ tiêu ở trang 1. CLB", "Điền rồi kéo sổ vào lại"],
        ["Mở CSV bằng Excel bị vỡ dấu", "Excel đoán sai bảng mã", "Nạp thẳng vào phần mềm (đọc đúng), hoặc mở bằng Data, From Text, chọn UTF-8"],
        ["Không thấy mục Skills trên claude.ai", "Gói tài khoản hoặc chưa bật chạy mã",
         "Bật Code execution; hoặc dùng Claude Code, hoặc tự chạy lệnh (mục 4.2, 9.2)"],
    ], dam_cot_dau=True)

    # ------------------------------------------------------------------ 13
    muc(doc, "13. Sửa mã và bảo trì (cho người phụ trách kỹ thuật)")
    muc(doc, "13.1. Lấy mã nguồn", 2)
    lenh(doc,
         "git clone https://github.com/%s.git" % KHO,
         "cd rbda-kiosk",
         "git checkout %s" % NHANH,
         "pip install -r requirements-dev.txt     (hoặc tối thiểu: pip install openpyxl python-docx pytest)")

    muc(doc, "13.2. Tệp nào làm gì", 2)
    bang(doc, ["Tệp trong kho", "Vai trò"], [
        ["mau_forms_thi_clb/tao_google_forms.py", "Khuôn mẫu mã Apps Script (biến MAU) và hàm noi_dung() đổ dữ liệu vào"],
        [".claude/skills/tao-google-form-clb/scripts/tao_form.py", "Đọc, soát tệp Excel, gọi noi_dung(), viết hướng dẫn"],
        [".claude/skills/xu-li-raw-forms/scripts/xu_li_raw.py", "Đọc raw, chấm lại, soát, ghi Sổ nhập CLB và báo cáo"],
        [".claude/skills/sinh-du-lieu-clb/scripts/sinh_du_lieu.py", "Bộ soát 9 điều và bộ ghi sổ nhập (dùng chung)"],
        ["so_nhap.py", "Định nghĩa Sổ nhập CLB: phần mềm, skill và bộ chuyển Forms cùng dùng"],
        ["mau_forms_thi_clb/DINH_DANG_RAW.md", "Định dạng raw: hợp đồng giữa mã Apps Script và skill xử lí raw"],
        ["mau_forms_thi_clb/THIET_KE_6_PHAN.html", "Bộ đề ví dụ và trang chạy thử thiết kế"],
        ["mau_forms_thi_clb/dong_goi_skill.py", "Đóng gói hai skill thành .zip tự chạy được"],
        ["tests/test_skill_*.py, tests/test_mau_forms_thi_clb.py, tests/test_dong_goi_skill.py", "Kiểm tra tự động"],
    ], dam_cot_dau=True)

    muc(doc, "13.3. Muốn đổi gì thì sửa ở đâu", 2)
    bang(doc, ["Muốn đổi", "Sửa ở đâu"], [
        ["Đề thi, CLB, buổi, điểm, số CLB thi tối đa", "Chỉ sửa MAU_DE_THI.xlsx, sinh lại mã. Không sửa mã"],
        ["Chữ trên biểu mẫu, cách rẽ nhánh, cách chấm", "Biến MAU trong tao_google_forms.py (mã JavaScript nằm trong chuỗi Python)"],
        ["Cột dữ liệu raw", "Sửa cùng lúc: COT_RAW trong MAU, DINH_DANG_RAW.md, COT_RAW trong xu_li_raw.py. Test báo nếu lệch"],
        ["Quy tắc soát, loại bất thường trong báo cáo", "xu_li_raw.py"],
        ["Giới hạn 10 CLB, 5 CLB thi", "Là giới hạn của phần mềm, nằm trong sinh_du_lieu.py. Không đổi nếu phần mềm không đổi"],
    ], dam_cot_dau=True)
    doan(doc, "Trong biến MAU, dấu % phải viết thành %% (Python dùng % để đổ dữ liệu vào), và dấu \\ trong "
              "biểu thức chính quy giữ nguyên vì MAU là chuỗi thô.")

    muc(doc, "13.4. Sau mỗi lần sửa", 2)
    buoc(doc, "Sinh lại tệp mẫu trong kho:")
    lenh(doc, "python mau_forms_thi_clb/tao_google_forms.py")
    buoc(doc, "Kiểm cú pháp JavaScript (cần Node.js): chép TAO_GOOGLE_FORM.gs thành một tệp .js rồi chạy node --check.")
    buoc(doc, "Chạy kiểm tra. Các test dùng bản giả lập Google (FormApp) bằng Node để kiểm rẽ nhánh và chấm; "
              "máy không có Node thì các test đó tự bỏ qua.")
    lenh(doc, "python -m pytest tests/test_mau_forms_thi_clb.py tests/test_skill_tao_google_form.py \\",
              "    tests/test_skill_xu_li_raw.py tests/test_dong_goi_skill.py -q",
              "python -m pytest tests -q        (toàn bộ, khoảng 3 phút)")
    buoc(doc, "Đóng gói lại skill nếu dùng ngoài kho, rồi tải lại lên claude.ai:")
    lenh(doc, "python mau_forms_thi_clb/dong_goi_skill.py --ra skill_zip")
    buoc(doc, "Sinh lại các tài liệu Word nếu nội dung đổi:")
    lenh(doc, "python mau_forms_thi_clb/tao_huong_dan_toan_bo.py",
              "python mau_forms_thi_clb/tao_huong_dan_google.py")
    buoc(doc, "Commit và đẩy lên nhánh. Google không tự cập nhật: phải dán mã mới vào dự án Apps Script. "
              "Đổi cách tạo biểu mẫu thì phải chạy lại taoBieuMau (ra biểu mẫu mới); chỉ đổi cách chấm thì "
              "chỉ cần dán mã và chạy lại chamTheoCLB.")

    muc(doc, "13.5. Không chạy được Apps Script ở máy tính", 2)
    doan(doc, "Mã Apps Script chỉ chạy thật trên Google. Các test trong kho chạy mã đó trên một bản giả lập "
              "FormApp viết bằng Node: đủ để bắt lỗi cú pháp, rẽ nhánh sai, chấm sai. Mỗi khi đổi mã, vẫn "
              "phải làm thử trên Google thật (mục 5.2) trước khi dùng với học sinh.")

    # ------------------------------------------------------------------ 14
    muc(doc, "14. Bảo vệ dữ liệu học sinh")
    gach(doc, "Phiếu trả lời là dữ liệu thật của học sinh, phần lớn chưa đủ 18 tuổi. Không đưa vào kho mã "
              "nguồn, không gửi qua kênh công khai.")
    gach(doc, "Khi nhờ Claude xử lí raw với dữ liệu thật, ưu tiên chạy trên máy trong kho mã (Claude Code) "
              "hoặc tự chạy lệnh (mục 9.2), thay vì tải tệp lên nơi khác.")
    gach(doc, "Xoá thư mục Kết quả CLB cũ trên Drive khi không còn cần.")
    gach(doc, "Các bộ dữ liệu ví dụ trong kho đều là dữ liệu mô phỏng, tên và điểm là bịa.")

    # ------------------------------------------------------------------ 15
    muc(doc, "15. Tóm tắt một trang")
    for dong in [
        "Điền MAU_DE_THI.xlsx.",
        "Claude: “tạo Google Form từ tệp này” → nhận TAO_GOOGLE_FORM.gs.",
        "script.google.com → dán mã → chạy taoBieuMau → cho phép quyền → lấy hai đường dẫn.",
        "Xuất bản, cài đặt công bố điểm thủ công và giới hạn 1 câu trả lời, làm thử trọn một vòng.",
        "Gửi đường dẫn học sinh. Hết hạn thì tắt nhận câu trả lời.",
        "Chạy chamTheoCLB → tải thư mục Kết quả CLB về, giải nén.",
        "Điền cột Chỉ tiêu ở trang 1. CLB của SO_NHAP_CLB.xlsx.",
        "Claude: “xử lí raw” kèm RAW_PHIEU, DAP_AN, SO_NHAP_CLB → nhận sổ nhập mới và báo cáo.",
        "Đọc báo cáo bất thường.",
        "Phần mềm: kéo SO_NHAP_CLB.xlsx vào → Nhập sổ → đọc cảnh báo → Chạy sắp xếp (ghi lại seed) → Xuất kết quả.",
    ]:
        buoc(doc, dong)

    for z in doc.settings.element.findall(qn("w:zoom")):
        z.set(qn("w:percent"), "100")
    return doc


def main():
    doc = viet()
    tw.kiem_chu(doc)
    doc.save(RA)
    print("Đã ghi %s" % RA)


if __name__ == "__main__":
    sys.exit(main())
