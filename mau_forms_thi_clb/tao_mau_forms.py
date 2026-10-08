"""Sinh toàn bộ mẫu Microsoft Forms từ de_thi.json.

    python mau_forms_thi_clb/tao_mau_forms.py

Ghi ra, cạnh tệp này:

    NHAP_FORMS_THI_CLB.docx   đem nhập bằng Quick Import của Forms
    CAI_DAT_FORMS.md          những gì Quick Import KHÔNG mang theo: đổi
                              câu Ranking, rẽ nhánh, đáp án + điểm, cài đặt
    MAU_XUAT_TU_FORMS.xlsx    tệp Forms xuất ra, dựng giả 60 học sinh — để
                              chạy thử bộ chuyển đổi trước khi có dữ liệu thật

Sửa đề, thêm CLB, đổi buổi: sửa de_thi.json rồi chạy lại. Không sửa tay ba
tệp trên — lần sinh sau sẽ ghi đè, và sửa tay là cách biểu mẫu lệch khỏi
đáp án dùng để chấm.

Cần python-docx (requirements-dev.txt) để ghi .docx; thiếu thì bỏ qua riêng
tệp đó, hai tệp kia vẫn sinh.
"""

import os
import random
import sys
from datetime import datetime, timedelta

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import bieu_mau as bm  # noqa: E402

THU_MUC = bm.THU_MUC
CHU_CAI = "abcdefghijklmnopqrstuvwxyz"


# ------------------------------------------------------------ Quick Import

def ghi_docx(de, duong):
    """Bố cục giống KHAO_SAT_NHAP_FORMS.docx — dạng Quick Import đọc tốt:
    tiêu đề phần in đậm, câu hỏi đánh số, lựa chọn 'a. …' thụt lề."""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)

    def dong(chu, co=11, dam=False, truoc=0, sau=4, thut=0):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before, pf.space_after = Pt(truoc), Pt(sau)
        if thut:
            pf.left_indent = Pt(thut)
        r = p.add_run(chu)
        r.bold, r.font.size = dam, Pt(co)

    dong(de["tieu_de"], co=16, dam=True, sau=10)
    so = 0
    for i, (ten, cac_cau) in enumerate(bm.cac_phan(de), 1):
        dong("PHẦN %d — %s" % (i, ten), co=14, dam=True, truoc=18, sau=8)
        for tieu_de, _loai, lua_chon, _bb in cac_cau:
            so += 1
            dong("%d. %s" % (so, tieu_de), co=12, dam=True, truoc=12, sau=4)
            for j, lc in enumerate(lua_chon):
                dong("%s. %s" % (CHU_CAI[j], lc), thut=18, sau=2)
    doc.save(duong)


# -------------------------------------------------------- bảng cài đặt tay

def ghi_cai_dat(de):
    phan = bm.cac_phan(de)
    so_cau = sum(len(c) for _, c in phan)
    n_thi = len(bm.clb_co_thi(de))
    L = []
    w = L.append
    w("# Cài đặt Microsoft Forms cho mẫu \"%s\"" % de["tieu_de"])
    w("")
    w("> **Tệp này do `tao_mau_forms.py` sinh từ `de_thi.json`.** Đừng sửa tay —")
    w("> sửa `de_thi.json` rồi chạy lại. Hướng dẫn đầy đủ: `README.md` cùng thư mục.")
    w("")
    w("Biểu mẫu có **%d phần, %d câu hỏi**: %d buổi, %d CLB, trong đó **%d CLB "
      "tổ chức thi** (mỗi đề %d câu × %s điểm = %s điểm)."
      % (len(phan), so_cau, len(de["buoi"]), len(de["cau_lac_bo"]), n_thi,
         len(bm.clb_co_thi(de)[0]["de_thi"]), bm.ghi_so(de["diem_moi_cau"]),
         bm.ghi_so(bm.diem_toi_da(de, bm.clb_co_thi(de)[0]))))
    w("")
    w("## 1. Bố cục — thứ tự phần CHÍNH LÀ luồng đi của học sinh")
    w("")
    w("| Phần | Tên phần | Câu hỏi | Loại | Bắt buộc |")
    w("|---|---|---|---|---|")
    so = 0
    for i, (ten, cac_cau) in enumerate(phan, 1):
        for j, (td, loai, _lc, bb) in enumerate(cac_cau):
            so += 1
            ma = bm.ma_cua_cau(td)
            w("| %s | %s | %d · `[%s]` | %s | %s |" % (
                i if j == 0 else "", ten if j == 0 else "", so, ma,
                {"text": "Text", "ranking": "**Ranking**", "choice": "Choice"}[loai],
                "✅" if bb else "—"))
    w("")
    w("## 2. Đổi %d câu nguyện vọng sang Ranking" % len(de["buoi"]))
    w("")
    w("Quick Import không biết loại **Ranking** — các câu `[thu_…]` về tới Forms")
    w("sẽ là câu Choice. Với từng câu: bấm **+ Add new → ⌄ → Ranking** ngay dưới")
    w("nó, **dán nguyên văn** tiêu đề và các lựa chọn (giữ cả phần `[mã]` và phần")
    w("`(club_id)`), để **không bắt buộc**, rồi xoá câu Choice cũ.")
    w("")
    w("## 3. Rẽ nhánh — chỉ %d chỗ" % n_thi)
    w("")
    w("Chọn câu cổng → **…** → **Add branching**. Đáp án **\"%s\"** để nguyên "
      "(đi tiếp sang phần đề thi ngay sau). Đáp án **\"%s\"** đặt **Go to**:" % (bm.CO, bm.KHONG))
    w("")
    w("| Câu cổng | \"%s\" → Go to |" % bm.KHONG)
    w("|---|---|")
    for clb, toi in bm.re_nhanh(de):
        w("| `[thi-%s]` %s | **%s** |" % (clb["club_id"], clb["name"], toi))
    w("")
    w("Làm xong đề thi CLB này thì Forms **tự** sang phần kế tiếp — đúng là câu")
    w("cổng của CLB sau, hoặc nguyện vọng của buổi sau. Không cần đặt thêm nhánh.")
    w("")
    w("## 4. Đáp án và điểm (chế độ Quiz)")
    w("")
    w("Chỉ cần nếu tạo dạng **Quiz**. Mỗi câu đề thi: bấm vào đáp án đúng →")
    w("dấu ✔ **Correct answer**, ô **Points** = **%s**. Câu cổng, câu nguyện vọng,"
      % bm.ghi_so(de["diem_moi_cau"]))
    w("mã và tên học sinh: **không** cho điểm.")
    w("")
    w("Bộ chuyển đổi **chấm lại** theo chính bảng này, và báo nếu điểm Forms")
    w("chấm lệch — tức là có câu tick nhầm đáp án trên Forms.")
    w("")
    for clb in bm.clb_co_thi(de):
        w("**%s** (`%s`, %s)" % (clb["name"], clb["club_id"], bm.ten_buoi(de, clb["buoi"])))
        w("")
        w("| Câu | Đáp án đúng |")
        w("|---|---|")
        for i, cau in enumerate(clb["de_thi"], 1):
            w("| `[%s-%d]` %s | %s. %s |" % (clb["club_id"], i, cau["cau"].replace("|", "\\|"),
                                          CHU_CAI[cau["dap_an"]], cau["lua_chon"][cau["dap_an"]]))
        w("")
    w("## 5. Settings")
    w("")
    w("| Cài đặt | Giá trị | Vì sao |")
    w("|---|---|---|")
    w("| Who can fill out this form | **Only people in my organization** | Học sinh đăng nhập tài khoản trường |")
    w("| **One response per person** | **BẬT** | Đây là bài thi: nộp lại là làm lại tới khi được điểm cao. Bộ chuyển đổi cũng chỉ giữ phiếu **đầu** của mỗi mã học sinh |")
    w("| Shuffle questions | **TẮT** | Xáo câu là hỏng rẽ nhánh |")
    w("| Shuffle options (từng câu đề thi) | Bật được | Chấm theo nội dung lựa chọn, không theo vị trí |")
    w("| Show results automatically (Quiz) | **TẮT** | Bật là em làm trước biết đáp án, chuyền cho em làm sau |")
    w("| Start date / End date | Đặt đủ cả hai | Hết giờ tự khoá, không phải canh |")
    w("")
    return "\n".join(L)


# ------------------------------------------------ tệp Forms xuất ra, dựng giả

HO = ["Nguyễn", "Trần", "Lê", "Phạm", "Hoàng", "Vũ", "Đặng", "Bùi", "Đỗ", "Ngô"]
DEM = ["Văn", "Thị", "Minh", "Thu", "Gia", "Quốc", "Ngọc", "Hải", "Bảo", "Khánh"]
TEN = ["An", "Bình", "Châu", "Dũng", "Hà", "Khoa", "Linh", "Nam", "Phúc",
       "Quân", "Trang", "Vy", "Yến", "Tùng", "Mai", "Long"]


def ghi_mau_xuat(de, duong, so_hoc_sinh=60, hat_giong=2026):
    """Giả đúng hình dạng Forms xuất ở chế độ Quiz: cột hệ thống trước, mỗi
    câu một cột; câu có điểm kèm cột 'Points - …' và 'Feedback - …'.
    Ranking và Choice nhiều lựa chọn là MỘT ô, nối bằng dấu chấm phẩy."""
    from openpyxl import Workbook

    rd = random.Random(hat_giong)
    phan = bm.cac_phan(de)
    co_diem = {bm.cau_de_thi(c, i) for c in bm.clb_co_thi(de)
               for i in range(1, len(c["de_thi"]) + 1)}

    tieu_de = ["ID", "Start time", "Completion time", "Email", "Name",
               "Total points", "Quiz feedback"]
    for _ten, cac_cau in phan:
        for td, *_ in cac_cau:
            tieu_de.append(td)
            if td in co_diem:
                tieu_de += ["Points - " + td, "Feedback - " + td]

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(tieu_de)
    gio = datetime(2026, 9, 7, 7, 30)
    for n in range(1, so_hoc_sinh + 1):
        o = {"ID": n, "Email": "anonymous", "Name": ""}
        o[bm.cau_ma_hoc_sinh()] = "HS%03d" % n
        o[bm.cau_ho_ten()] = "%s %s %s" % (rd.choice(HO), rd.choice(DEM), rd.choice(TEN))
        nang_luc = rd.uniform(0.7, 0.95)
        tong = 0
        # Mỗi em xếp 2–4 buổi; buổi còn lại bỏ trống = em bận.
        buoi_ranh = rd.sample(de["buoi"], rd.randint(2, 4))
        for b in de["buoi"]:
            o_buoi = bm.clb_cua_buoi(de, b["ma"])
            xep = []
            if b in buoi_ranh:
                # Ranking của Forms bắt xếp ĐỦ mọi lựa chọn.
                xep = rd.sample(o_buoi, len(o_buoi))
                o[bm.cau_xep_hang(de, b["ma"])] = "".join(bm.lua_chon_clb(c) + ";" for c in xep)
            for clb in o_buoi:
                if not clb["de_thi"]:
                    continue
                thi = clb in xep[:2] and rd.random() < 0.75
                o[bm.cau_cong_thi(de, clb)] = bm.CO if thi else bm.KHONG
                if not thi:
                    continue
                dung_ca = [rd.random() < nang_luc for _ in clb["de_thi"]]
                # Ít nhất 2 câu đúng: bài 0–2 điểm là THẬT với đề 5 câu, nhưng
                # phần mềm báo "điểm lạ" khi dưới 1/3 trung vị (README mục 5).
                # Tệp mẫu phải nạp sạch để người đọc không lẫn cảnh báo cố ý.
                for k in rd.sample(range(len(dung_ca)), len(dung_ca)):
                    if sum(dung_ca) >= 2:
                        break
                    dung_ca[k] = True
                for i, cau in enumerate(clb["de_thi"], 1):
                    dung = dung_ca[i - 1]
                    sai = [x for k, x in enumerate(cau["lua_chon"]) if k != cau["dap_an"]]
                    td = bm.cau_de_thi(clb, i)
                    o[td] = cau["lua_chon"][cau["dap_an"]] if dung else rd.choice(sai)
                    o["Points - " + td] = de["diem_moi_cau"] if dung else 0
                    tong += o["Points - " + td]
        o["Total points"] = tong
        bat_dau = gio + timedelta(minutes=3 * n + rd.randint(0, 2))
        o["Start time"] = bat_dau
        o["Completion time"] = bat_dau + timedelta(minutes=rd.randint(6, 25))
        ws.append([o.get(c) for c in tieu_de])
    wb.save(duong)


def main():
    de = bm.doc_de_thi()
    with open(os.path.join(THU_MUC, "CAI_DAT_FORMS.md"), "w", encoding="utf-8") as f:
        f.write(ghi_cai_dat(de))
    ghi_mau_xuat(de, os.path.join(THU_MUC, "MAU_XUAT_TU_FORMS.xlsx"))
    try:
        ghi_docx(de, os.path.join(THU_MUC, "NHAP_FORMS_THI_CLB.docx"))
    except ImportError:
        print("Thiếu python-docx — BỎ QUA NHAP_FORMS_THI_CLB.docx "
              "(pip install -r requirements-dev.txt)", file=sys.stderr)
    print("Đã sinh mẫu Forms: %d phần, %d câu hỏi."
          % (len(bm.cac_phan(de)), sum(len(c) for _, c in bm.cac_phan(de))))


if __name__ == "__main__":
    main()
