"""Sinh tệp Quick Import cho biểu mẫu 6 phần: NHAP_FORMS_6_PHAN.docx.

    python mau_forms_thi_clb/tao_nhap_forms_6_phan.py

Tệp này để NHẬP vào Microsoft Forms (Quick Import), khác với
THIET_KE_6_PHAN.docx là tài liệu để đọc. Trong tệp chỉ có tiêu đề phần, câu
hỏi đánh số và lựa chọn: đoạn văn nào khác Forms cũng sẽ biến thành câu hỏi.

Dữ liệu CLB và đề thi đọc từ THIET_KE_6_PHAN.html qua tao_word_6_phan.py.
Thứ tự lựa chọn của câu đề thi xáo đúng như trang HTML, nên biểu mẫu trên
Forms trông giống hệt bản chạy thử.

Cần python-docx (requirements-dev.txt).
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Pt  # noqa: E402

import tao_word_6_phan as tw  # noqa: E402
from tao_google_forms import xao  # noqa: E402,F401  (dùng chung, test gọi tn.xao)

RA = os.path.join(tw.THU_MUC, "NHAP_FORMS_6_PHAN.docx")
CHU_CAI = "abcdefghijklmnopqrstuvwxyz"


def cac_phan(du_lieu):
    """[(tiêu đề phần, [(mã, câu hỏi, [lựa chọn])])], 6 phần, 162 câu."""
    phan = [("PHẦN 1: Thông tin học sinh", [
        ("student_id", "Mã học sinh", []),
        ("name", "Họ và tên", []),
    ])]
    for k, (ma, ten, ds) in enumerate(du_lieu):
        cau = [
            ("di-" + ma, "Em có đăng ký sinh hoạt %s không?" % ten,
             [tw.CO_BUOI, tw.KHONG_BUOI]),
            (ma, "Xếp hạng các CLB %s. Kéo CLB em thích nhất lên đầu." % ten,
             ["%s (%s)" % (name, cid) for cid, name, _de in ds]),
        ]
        for cid, name, de in ds:
            cau.append(("thi-" + cid, "Thi %s? Tối đa %d CLB mỗi buổi." % (name, tw.TRAN_THI),
                        [tw.CO, tw.KHONG, tw.DUNG]))
            for j, (q, _dung, ops) in enumerate(de):
                cau.append(("%s-%d" % (cid, j + 1), q, xao(ops, cid + str(j))))
        phan.append(("PHẦN %d: %s" % (k + 2, ten), cau))
    return phan


def viet(du_lieu):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    def dong(chu, co=11, dam=False, truoc=0, sau=4, thut=0):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before, pf.space_after = Pt(truoc), Pt(sau)
        if thut:
            pf.left_indent = Pt(thut)
        r = p.add_run(chu)
        r.bold, r.font.size = dam, Pt(co)

    dong("Đăng ký CLB và thi tuyển", co=16, dam=True, sau=10)
    so = 0
    for ten_phan, cac_cau in cac_phan(du_lieu):
        dong(ten_phan, co=14, dam=True, truoc=18, sau=8)
        for ma, cau, lua_chon in cac_cau:
            so += 1
            dong("%d. [%s] %s" % (so, ma, cau), co=12, dam=True, truoc=12, sau=4)
            for j, lc in enumerate(lua_chon):
                dong("%s. %s" % (CHU_CAI[j], lc), thut=18, sau=2)
    for z in doc.settings.element.findall(qn("w:zoom")):
        z.set(qn("w:percent"), "100")
    return doc


def main():
    du_lieu = tw.doc_du_lieu()
    doc = viet(du_lieu)
    tw.kiem_chu(doc)
    doc.save(RA)
    print("Đã ghi %s" % RA)


if __name__ == "__main__":
    sys.exit(main())
