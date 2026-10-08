"""Nạp dữ liệu: đọc CSV/Excel, nhận diện loại tệp, ba hàm nhập
(CLB, chọn thi, nguyện vọng) và các bước soát dùng chung.

Một phần của PipelineAPI (xem api.py) — tách riêng cho dễ đọc; mọi hàm
public ở đây vẫn là hàm của PipelineAPI và giao diện gọi được như cũ.
"""

import csv
import io
import re
import so_nhap
from i18n_errors import err
from i18n_errors import phan_hoi_loi as _fail
from i18n_errors import phan_hoi_ok as _ok
from rbda_priority_pipeline import (
    khoa_sap_buoi,
    sap_buoi,
    BUOI_MAC_DINH,
    TRAN_CLB_THI_MOI_BUOI,
    TRAN_NGUYEN_VONG_MOI_BUOI,
    gom_theo_buoi,
    loi_suc_chua,
)


class _HuyNapSo(Exception):
    """Một bước nạp Sổ nhập CLB trả về lỗi: huỷ cả giao dịch, trả lỗi đó ra."""

    def __init__(self, ket_qua):
        super().__init__("nap so nhap bi huy")
        self.ket_qua = ket_qua


class NhapDuLieuMixin:
    # -----------------------------------------------------------------
    # NHẬP DỮ LIỆU TỪ MICROSOFT FORMS (CSV đã chuẩn hoá bởi
    # 06_ms_forms_transform.py) — trước đây KHÔNG có đường nào để CSV
    # này vào app.db, đây là mảnh còn thiếu duy nhất trong luồng dữ liệu.
    #
    # JS đọc file bằng FileReader.readAsText() ở phía trình duyệt rồi
    # gửi NGUYÊN VĂN nội dung CSV (chuỗi text) qua đây — không cần
    # main.py hỗ trợ thêm gì, không phụ thuộc hộp thoại chọn file gốc
    # của hệ điều hành (vốn không ổn định trong mọi bản pywebview).
    #
    # Hỗ trợ 2 định dạng cho MỖI loại CSV — tự nhận diện theo header,
    # không cần người dùng chọn định dạng:
    #   (a) "dài" (long) — đúng 1-1 với cấu trúc bảng DB, do
    #       06_ms_forms_transform.py xuất ra:
    #         nguyện vọng:      student_id,name,club_id,rank
    #         chọn club thi:    student_id,club_id
    #   (b) "rộng" (wide) — 1 dòng/học sinh, tiện đọc bằng mắt:
    #         nguyện vọng:      student_id,name,pref_1,pref_2,...,pref_10
    #         chọn club thi:    student_id,name,test_club_1,test_club_2,...
    # -----------------------------------------------------------------

    @staticmethod
    def _parse_csv_rows(csv_text: str):
        """Tự nhận diện dấu phân cách (, hoặc ;) và trả về (fieldnames, rows)."""
        # BOM (\ufeff) — Excel LUÔN thêm ký tự này vào đầu file khi lưu
        # "CSV UTF-8", và Microsoft Forms xuất ra cũng vậy. Nó vô hình
        # khi mở file bằng mắt, nhưng dính liền vào tên cột đầu tiên:
        # "student_id" đọc lên thành "\ufeffstudent_id". Không cắt bỏ thì
        # phần mềm báo THIẾU CỘT student_id trên một file hoàn toàn đúng —
        # lỗi cực khó hiểu với giáo viên. .strip() KHÔNG cắt được BOM vì
        # nó không phải khoảng trắng.
        csv_text = csv_text.lstrip("\ufeff")
        sample = csv_text[:4096]
        # CHI lay dau phan cach tu Sniffer, KHONG lay ca dialect.
        #
        # csv.Sniffer doan luon quy uoc trich dan, va no doan SAI: voi
        # 'HS1,"Tran ""Bo"" Van A, Jr.",clb_a' no tra ve doublequote=False,
        # tuc la bo quy uoc "" = mot dau nhay. Hau qua: o ten bi cat ngay
        # dau phay, phan duoi ('Jr."') roi sang cot ke ben va bi hieu la
        # ma club. Ca dong bi bo qua, chi kem mot canh bao "club khong ton
        # tai" khong lien quan gi toi nguyen nhan that.
        # Ten CLB tieng Viet rat hay co dau nhay: CLB "Vi Cong Dong".
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
        except csv.Error:
            delimiter = ","
        reader = csv.DictReader(io.StringIO(csv_text), delimiter=delimiter)
        fieldnames = [ (f or "").strip().lower() for f in (reader.fieldnames or []) ]
        def o_that(v):
            """Gỡ đúng lớp bảo vệ `_an_toan_cho_excel` đã thêm khi XUẤT.

            Tệp "Dữ liệu đầu vào hiện tại" phải nạp lại được y nguyên. Tên
            "=1+1" xuất ra thành "'=1+1"; không gỡ ở đây thì nạp lại là tên
            đổi thành "'=1+1". Chỉ gỡ dấu nháy đứng TRƯỚC = + - @ — đúng và
            chỉ đúng thứ phía xuất thêm vào.
            """
            if not isinstance(v, str):
                return v
            v = v.strip()
            if v[:1] == "'" and v[1:2] in ("=", "+", "-", "@", "\t", "\r"):
                return v[1:]
            return v

        rows = []
        for raw_row in reader:
            row = { (k or "").strip().lower(): o_that(v)
                    for k, v in raw_row.items() if k is not None }
            rows.append(row)
        return fieldnames, rows

    # -----------------------------------------------------------------
    # TỰ NHẬN DIỆN LOẠI FILE
    #
    # Giao diện cũ có HAI ô nạp file và người dùng phải tự chọn đúng ô.
    # Kéo nhầm ô KHÔNG hề báo lỗi: file nguyện vọng dạng dài
    # (student_id, name, club_id, rank) nạp vào ô "chọn club thi" vẫn
    # khớp đủ cột, nên nó ghi thẳng vào club_test_selection và báo
    # "thành công". Nguyện vọng thật mất sạch mà không một cảnh báo —
    # lỗi làm sai kết quả phân bổ của cả trường mà không ai biết.
    #
    # Nên bỏ hẳn việc bắt người dùng chọn: đọc dòng tiêu đề là biết file
    # gì. Nhưng CHỈ kết luận khi CHẮC CHẮN — bộ cột
    # (student_id, club_id) không kèm rank vừa có thể là chọn club thi
    # dạng dài, vừa có thể là nguyện vọng dạng dài thiếu cột rank. Gặp
    # trường hợp đó thì HỎI LẠI, tuyệt đối không đoán: đoán sai là đúng
    # lại cái bug vừa chữa.
    # -----------------------------------------------------------------

    # -----------------------------------------------------------------
    # ĐỌC THẲNG FILE EXCEL (.xlsx)
    #
    # Microsoft Forms xuất kết quả ra .xlsx. Trước đây người vận hành
    # phải mở Excel -> File -> Save As -> chọn đúng "CSV UTF-8" -> rồi
    # mới nạp được. Bước thừa đó lại là bước dễ sai nhất: chọn nhầm
    # "CSV (Comma delimited)" thì tên tiếng Việt hỏng hết dấu, mà chọn
    # đúng thì Excel chèn BOM (bản cũ báo thiếu cột student_id vì thế).
    # Đọc thẳng .xlsx là bỏ được cả lớp lỗi đó.
    #
    # openpyxl là Python thuần, không cần .NET hay thư viện hệ thống —
    # khác hẳn pythonnet, nên không kéo theo rủi ro đóng gói đã gặp.
    # -----------------------------------------------------------------

    @staticmethod
    def _o_excel_thanh_chu(gia_tri) -> str:
        """Đưa một ô Excel về đúng chuỗi mà phần CSV đang chờ.

        Quan trọng nhất là số: Excel lưu chỉ tiêu 20 dưới dạng số thực,
        đọc thô ra thành "20.0" và int("20.0") ném ValueError — cả dòng
        CLB sẽ bị bỏ qua dù file hoàn toàn đúng.
        """
        if gia_tri is None:
            return ""
        if isinstance(gia_tri, bool):
            return "1" if gia_tri else ""
        if isinstance(gia_tri, float) and gia_tri.is_integer():
            return str(int(gia_tri))
        return str(gia_tri).strip()

    # Tệp .xlsx là một tệp zip. Tệp nén vài trăm KB có thể bung ra hàng
    # GB (zip bomb) và openpyxl sẽ đọc hết vào bộ nhớ — máy kiosk treo
    # cứng giữa buổi. Một trường vài nghìn học sinh chưa tới 1 MB, nên
    # hai trần dưới đây rộng hàng chục lần mức cần mà vẫn chặn được.
    _TRAN_TEP_XLSX_MB = 20
    _TRAN_GIAI_NEN_XLSX_MB = 200

    @classmethod
    def _soat_co_xlsx(cls, du_lieu: bytes):
        """Trả về lỗi `err(...)` nếu tệp quá lớn, None nếu đọc được.

        Soát cỡ SAU GIẢI NÉN bằng mục lục của zip, trước khi openpyxl
        mở tệp — mục lục chỉ vài KB nên soát không tốn gì. Tệp không
        phải zip thì để openpyxl tự báo lỗi như cũ.
        """
        import zipfile

        mb = 1024 * 1024
        if len(du_lieu) > cls._TRAN_TEP_XLSX_MB * mb:
            return err("file_too_large",
                       size_mb=round(len(du_lieu) / mb, 1),
                       max_mb=cls._TRAN_TEP_XLSX_MB)
        try:
            with zipfile.ZipFile(io.BytesIO(du_lieu)) as z:
                tong = sum(i.file_size for i in z.infolist())
        except zipfile.BadZipFile:
            return None
        if tong > cls._TRAN_GIAI_NEN_XLSX_MB * mb:
            return err("file_too_large",
                       size_mb=round(tong / mb, 1),
                       max_mb=cls._TRAN_GIAI_NEN_XLSX_MB)
        return None

    def xlsx_to_csv_text(self, file_base64: str, sheet_name: str = ""):
        """Đọc file .xlsx (gửi lên dạng base64) và trả về text CSV.

        Phần còn lại của luồng nhập giữ nguyên: detect_csv_kind và
        import_csv_auto vẫn làm việc trên text, không cần biết dữ liệu
        đến từ .csv hay .xlsx.
        """
        try:
            import base64 as _b64

            try:
                import openpyxl
            except ImportError:
                return _fail(err("xlsx_support_missing"))

            try:
                du_lieu = _b64.b64decode(file_base64 or "", validate=False)
            except Exception as e:
                return _fail(err("xlsx_read_failed", detail=str(e)))
            loi_co = self._soat_co_xlsx(du_lieu)
            if loi_co:
                return _fail(loi_co)

            try:
                wb = openpyxl.load_workbook(
                    io.BytesIO(du_lieu), read_only=True, data_only=True
                )
            except Exception as e:
                return _fail(err("xlsx_read_failed", detail=str(e)))

            ten_cac_sheet = list(wb.sheetnames)
            if not ten_cac_sheet:
                return _fail(err("xlsx_empty"))
            if sheet_name and sheet_name in ten_cac_sheet:
                ws = wb[sheet_name]
            else:
                # Sheet DAU TIEN, khong phai sheet dang active: file Forms
                # xuat ra luon de du lieu o sheet dau, con active co the la
                # sheet nguoi dung xem cuoi cung truoc khi luu.
                ws = wb[ten_cac_sheet[0]]

            dong = []
            for hang in ws.iter_rows(values_only=True):
                o = [self._o_excel_thanh_chu(v) for v in hang]
                # Excel hay de lai vai hang rong o cuoi bang.
                if any(x for x in o):
                    dong.append(o)
            wb.close()

            if not dong:
                return _fail(err("xlsx_empty"))

            buf = io.StringIO()
            csv.writer(buf, lineterminator="\n").writerows(dong)
            return _ok({
                "csv_text": buf.getvalue(),
                "sheet_name": ws.title,
                "sheet_names": ten_cac_sheet,
                "n_rows": len(dong),
            })
        except Exception as e:
            return _fail(self._loi_co_ghi_vet(err("xlsx_read_failed", detail=str(e))))

    # -----------------------------------------------------------------
    # SỔ NHẬP CLB — MỘT tệp Excel thay cho ba tệp (xem so_nhap.py)
    #
    # Giao diện chỉ nhận sổ nhập. Sổ được dịch thành đúng ba văn bản CSV
    # mà ba hàm nạp bên dưới đã hiểu và đã được test kỹ, rồi nạp theo thứ
    # tự bắt buộc CLB -> chọn thi -> nguyện vọng. Mọi luật soát của ba hàm
    # đó (trần mỗi buổi, nhóm dự trữ lạ, mã lệch hoa/thường…) vẫn áp dụng.
    # -----------------------------------------------------------------
    def _doc_so_nhap(self, file_base64: str):
        """-> (du_lieu, None) hoặc (None, lỗi). Không ghi gì."""
        import base64 as _b64

        try:
            import openpyxl
        except ImportError:
            return None, err("xlsx_support_missing")
        try:
            du_lieu = _b64.b64decode(file_base64 or "", validate=False)
        except Exception as e:
            return None, err("xlsx_read_failed", detail=str(e))
        if not du_lieu:
            return None, err("xlsx_empty")
        loi_co = self._soat_co_xlsx(du_lieu)
        if loi_co:
            return None, loi_co
        try:
            wb = openpyxl.load_workbook(io.BytesIO(du_lieu), read_only=True,
                                        data_only=True)
        except Exception as e:
            return None, err("xlsx_read_failed", detail=str(e))
        try:
            if not so_nhap.la_so_nhap(wb.sheetnames):
                return None, err("so_nhap_khong_phai_so")
            return so_nhap.doc_so_nhap(wb, self._clb_theo_ten(), self._ma_clb_da_co(),
                                       self._clb_cu()), None
        finally:
            wb.close()

    def _clb_theo_ten(self) -> dict:
        """{khoa_ten(tên): club_id} của CLB đã có — để dòng sổ bỏ trống Mã CLB
        mà trùng tên một CLB đã có thì cập nhật CLB đó, không tạo CLB thứ hai.
        Tên mà CSDL có hai CLB trùng khoá thì bỏ: không đoán."""
        dem: dict = {}
        with self._ket_noi_doc() as cur:
            for cid, ten in cur.execute("SELECT club_id, name FROM clubs"):
                dem.setdefault(so_nhap.khoa_ten(ten or cid), []).append(cid)
        return {k: v[0] for k, v in dem.items() if len(v) == 1}

    def _clb_cu(self) -> dict:
        """{club_id: {buoi, reserve_capacity, reserve_group}} — giá trị giữ lại
        khi sheet CLB thiếu hẳn cột tương ứng (xem doc_so_nhap)."""
        with self._ket_noi_doc() as cur:
            return {r[0]: {"buoi": "" if (r[1] or BUOI_MAC_DINH) == BUOI_MAC_DINH else r[1],
                           "reserve_capacity": r[2] or 0, "reserve_group": r[3] or ""}
                    for r in cur.execute(
                        "SELECT club_id, buoi, reserve_capacity, reserve_group FROM clubs")}

    def _ma_clb_da_co(self) -> dict:
        """{club_id: tên} của mọi CLB đã có (xem doc_so_nhap)."""
        with self._ket_noi_doc() as cur:
            return {cid: (ten or cid)
                    for cid, ten in cur.execute("SELECT club_id, name FROM clubs")}

    def _canh_bao_so_nhap(self, du_lieu) -> tuple:
        """Những điều nạp sổ này SẼ làm với dữ liệu đang có — không chặn nạp,
        nhưng người vận hành phải thấy trước khi bấm Nhập.

        Trả về (cảnh báo, cặp (mã HS, CLB) ghi `thi`, mã HS bỏ trống nhóm).
        """
        cap_thi = [(s["student_id"], x["club_id"]) for s in du_lieu["students"]
                   for x in s["nv"] if x["thi"] and x["diem"] == ""]
        # None = sổ không có cột Nhóm ưu tiên (sổ sinh từ biểu mẫu): giữ nhóm.
        # So nhãn ĐÃ CHUẨN HOÁ: ô "-" chuẩn hoá ra rỗng, tức là không nhóm.
        bo_nhom = [s["student_id"] for s in du_lieu["students"]
                   if s["reserve_group"] is not None
                   and not self.chuan_hoa_nhom_du_tru(s["reserve_group"])]
        ma_trong_so = {c["club_id"] for c in du_lieu["clubs"]}
        canh_bao = []
        with self._ket_noi_doc() as cur:
            diem = {(r[0], r[1]) for r in cur.execute(
                "SELECT student_id, club_id FROM club_scores")}
            n_xoa = sum(1 for cap in cap_thi if cap in diem)
            co_nhom = {r[0] for r in cur.execute(
                "SELECT student_id FROM students "
                "WHERE reserve_group IS NOT NULL AND reserve_group <> ''")}
            n_bo = sum(1 for sid in bo_nhom if sid in co_nhom)
            ngoai = [r[1] or r[0] for r in cur.execute(
                "SELECT club_id, name FROM clubs ORDER BY name")
                if r[0] not in ma_trong_so]
        if n_xoa:
            canh_bao.append(err("so_nhap_canh_bao_xoa_diem", n=n_xoa))
        if n_bo:
            canh_bao.append(err("so_nhap_canh_bao_bo_nhom", n=n_bo))
        if ngoai and du_lieu["clubs"]:
            canh_bao.append(err("so_nhap_canh_bao_clb_ngoai_so", n=len(ngoai),
                                ten=", ".join(ngoai[:5])))
        return canh_bao, cap_thi, bo_nhom

    def xem_truoc_so_nhap(self, file_base64: str):
        """Đọc sổ và tóm tắt, KHÔNG ghi gì vào CSDL.

        Trả về {n_clb, n_hoc_sinh, n_buoi, n_nguyen_vong, n_thi, n_diem,
        loi: [...]}. `loi` khác rỗng thì giao diện không cho bấm Nhập.
        """
        try:
            du_lieu, loi = self._doc_so_nhap(file_base64)
            if loi:
                return _fail(loi)
            kq = so_nhap.tom_tat(du_lieu)
            kq["loi"] = du_lieu["loi"]
            kq["canh_bao"] = [] if du_lieu["loi"] else self._canh_bao_so_nhap(du_lieu)[0]
            return _ok(kq)
        except Exception as e:
            return _fail(self._loi_co_ghi_vet(err("xlsx_read_failed", detail=str(e))))

    def import_so_nhap(self, file_base64: str):
        """Nạp cả sổ nhập: CLB, rồi CLB dự thi kèm điểm, rồi nguyện vọng.

        Sổ có lỗi (xem so_nhap.doc_so_nhap) thì KHÔNG ghi gì và trả về danh
        sách lỗi kèm sheet, dòng. Sổ sạch thì ba bước nạp chạy nối nhau; cảnh
        báo của cả ba gom về một danh sách.
        """
        try:
            du_lieu, loi = self._doc_so_nhap(file_base64)
            if loi:
                return _fail(loi)
            if du_lieu["loi"]:
                return _fail(du_lieu["loi"])

            csv_text = so_nhap.thanh_csv(du_lieu)
            kq = so_nhap.tom_tat(du_lieu)
            canh_bao, cap_thi, bo_nhom = self._canh_bao_so_nhap(du_lieu)

            def buoc(res):
                # Một bước hỏng -> huỷ CẢ giao dịch, kể cả các bước đã xong.
                if not res["ok"]:
                    raise _HuyNapSo(res)
                canh_bao.extend(res["data"]["warnings"])
                return res["data"]

            # MỘT giao dịch cho cả sổ: CLB đã ghi mà bước học sinh hỏng (CSDL
            # bị khoá, lỗi bất ngờ) thì CLB cũng được huỷ. Màn hình nói "chưa
            # nạp gì cả" là đúng sự thật.
            with self._mot_giao_dich():
                d = buoc(self.import_clubs_csv(csv_text["clubs"]))
                kq.update(n_clubs_created=d["n_clubs_created"],
                          n_clubs_updated=d["n_clubs_updated"],
                          n_students_created=0, n_students_written=0,
                          n_students_skipped=0, n_scores_written=0)

                if csv_text["preferences"]:
                    # Chọn thi trước nguyện vọng: đúng thứ tự giao diện cũ vẫn
                    # nạp, nên kết quả trùng khít với bộ ba tệp CSV.
                    d = buoc(self.import_test_selection_csv(csv_text["test_selection"]))
                    kq["n_students_created"] += d["n_students_created"]
                    kq["n_scores_written"] = d["n_scores_written"]

                    d = buoc(self.import_preferences_csv(csv_text["preferences"]))
                    kq["n_students_created"] += d["n_students_created"]
                    kq["n_students_written"] = d["n_students_with_preferences_written"]
                    kq["n_students_skipped"] = d["n_students_skipped"]

                    # Sổ là bản ĐẦY ĐỦ cho mỗi em có trong sổ. Hàm nạp CSV giữ
                    # nguyên ô trống (đúng với CSV, nơi cột trống là "không
                    # biết"); với sổ, ô trống là "không có":
                    #   `thi`              -> chưa có điểm: xoá điểm cũ của CLB đó;
                    #   Nhóm ưu tiên trống -> em không còn thuộc nhóm nào.
                    with self._ket_noi_ghi() as cur:
                        cur.executemany(
                            "DELETE FROM club_scores WHERE student_id = ? AND club_id = ?",
                            cap_thi)
                        cur.executemany(
                            "UPDATE students SET reserve_group = NULL WHERE student_id = ?",
                            [(sid,) for sid in bo_nhom])
                        # Họ tên sửa trong sổ phải vào phần mềm. Hàm nạp CSV chỉ
                        # điền tên còn trống; với sổ, ô Họ tên có chữ là tên đúng.
                        cur.executemany(
                            "UPDATE students SET name = ? WHERE student_id = ?",
                            [(s["name"], s["student_id"]) for s in du_lieu["students"]
                             if (s["name"] or "").strip()])

            kq["warnings"] = canh_bao
            return _ok(kq)
        except _HuyNapSo as huy:
            return huy.ket_qua
        except Exception as e:
            return _fail(self._loi_co_ghi_vet(err("loi_nap_so_nhap", detail=str(e))))

    def detect_csv_kind(self, csv_text: str):
        """Đọc dòng tiêu đề để biết đây là file gì.

        Trả về {kind, format, confident, candidates, fieldnames}.
        confident=False nghĩa là giao diện PHẢI hỏi lại người dùng.
        """
        try:
            fieldnames, rows = self._parse_csv_rows(csv_text)
            if not rows:
                return _fail(err("csv_empty"))

            co = set(fieldnames)

            def ket_luan(kind, fmt):
                return _ok({"kind": kind, "format": fmt, "confident": True,
                            "candidates": [kind], "fieldnames": fieldnames})

            # capacity chi co nghia voi file danh sach CLB.
            if "capacity" in co and "club_id" in co:
                return ket_luan("clubs", "")
            if any(f.startswith("pref_") for f in fieldnames):
                return ket_luan("preferences", "wide")
            # Bo cot nguyen vong theo tung buoi: thu_3_pref_1, thu_5_pref_1...
            # Phai xet SAU `pref_` o tren de mot tep co ca hai kieu van vao
            # nhanh cu — nhung phai xet TRUOC nhanh `test_club_`, vi mot ten
            # buoi co the tinh co bat dau bang "test".
            if self._cot_nguyen_vong_theo_buoi(fieldnames):
                return ket_luan("preferences", "wide_theo_buoi")
            if any(f.startswith("test_club_") for f in fieldnames):
                return ket_luan("test_selection", "wide")
            # rank chi co nghia voi nguyen vong — chon club thi khong xep hang.
            if "rank" in co and "club_id" in co and "student_id" in co:
                return ket_luan("preferences", "long")

            if "student_id" in co and "club_id" in co:
                return _ok({
                    "kind": "", "format": "long", "confident": False,
                    "candidates": ["test_selection", "preferences"],
                    "fieldnames": fieldnames,
                })

            return _ok({"kind": "unknown", "format": "", "confident": False,
                        "candidates": [], "fieldnames": fieldnames})
        except Exception as e:
            return _fail(err("error_detecting_csv_kind", detail=str(e)))

    def import_csv_auto(self, csv_text: str, kind: str = "",
                        create_missing_students: bool = True):
        """Nạp một file CSV bất kỳ — tự nhận diện loại, không cần chọn ô.

        kind chỉ dùng khi tự nhận diện KHÔNG chắc: giao diện hỏi lại
        người dùng rồi truyền câu trả lời vào đây.
        """
        try:
            if not kind:
                nhan_dien = self.detect_csv_kind(csv_text)
                if not nhan_dien["ok"]:
                    return nhan_dien
                d = nhan_dien["data"]
                if not d["confident"]:
                    # KHONG doan. Tha khong nhap con hon nhap vao sai bang.
                    return _fail(err(
                        "csv_kind_ambiguous",
                        candidates=d["candidates"], fieldnames=d["fieldnames"],
                    ))
                kind = d["kind"]

            if kind == "clubs":
                res = self.import_clubs_csv(csv_text)
            elif kind == "preferences":
                res = self.import_preferences_csv(csv_text, create_missing_students)
            elif kind == "test_selection":
                res = self.import_test_selection_csv(csv_text, create_missing_students)
            else:
                return _fail(err("csv_kind_unknown"))

            if res["ok"]:
                res["data"]["kind"] = kind
            return res
        except Exception as e:
            return _fail(self._loi_co_ghi_vet(err("error_importing_csv_auto", detail=str(e))))

    @staticmethod
    def _khop_buoi_da_co(cur, buoi, nhan_goc=None):
        """Dùng lại nhãn buổi ĐÃ CÓ trong CSDL nếu nó cùng một buổi với `buoi`.

        CSDL tạo trước bản chuẩn hoá thứ có thể còn nhãn `thứ_3`. CLB mới khai
        "Thứ 3" thành `thu_3`; nếu không quy về nhãn cũ thì hai CLB cùng giờ
        nằm ở hai buổi, và một em trúng cả hai. Không đổi dữ liệu đã lưu, nên
        lịch sử chạy và số bốc thăm đã khoá vẫn khớp.

        `nhan_goc`: nhãn đúng như người dùng / sổ ghi. Trùng TỪNG KÝ TỰ một nhãn
        đã lưu thì giữ nguyên nhãn đó (vd sổ không có cột Buổi nên giữ buổi cũ
        `thứ_3`): đọc thử đã soát trần theo đúng nhãn ấy, nạp không được gộp
        khác đi.
        """
        if not buoi:
            return buoi
        cu = [r[0] for r in cur.execute(
            "SELECT buoi, COUNT(*) FROM clubs WHERE buoi IS NOT NULL AND buoi <> '' "
            "GROUP BY buoi ORDER BY COUNT(*) DESC, buoi COLLATE THU_TU_BUOI")]
        goc = (nhan_goc or "").strip()
        if goc and goc in cu:
            return goc
        if buoi in cu:
            return buoi
        # Mọi nhãn cũ khớp đều CÙNG một ngày (ma_buoi giữ chữ trong ngoặc, nên
        # "Thứ 3 (sáng)" không khớp "Thứ 3"). Dùng nhãn nhiều CLB nhất — `cu`
        # đã sắp theo số CLB — thay vì đẻ thêm nhãn thứ ba cho cùng ngày.
        return next((b for b in cu if so_nhap.ma_buoi(b) == buoi), buoi)

    @staticmethod
    def _ghi_clb(cur, club_id, name, capacity, reserve_capacity, reserve_group, buoi):
        """Thêm hoặc cập nhật MỘT CLB — câu SQL duy nhất cho cả form lẫn tệp CSV."""
        cur.execute(
            """
            INSERT INTO clubs (club_id, name, capacity, reserve_capacity, reserve_group, buoi)
            VALUES (?, ?, ?, ?, ?, ?)
            ON CONFLICT(club_id) DO UPDATE SET
                name=excluded.name,
                capacity=excluded.capacity,
                reserve_capacity=excluded.reserve_capacity,
                reserve_group=excluded.reserve_group,
                buoi=excluded.buoi
            """,
            (club_id, name, capacity, reserve_capacity, reserve_group, buoi),
        )

    def import_clubs_csv(self, csv_text: str):
        """Nhập danh sách CLB bằng CSV.

        Đây là nút thắt ĐẦU TIÊN người dùng gặp: mẫu CSV học sinh bắt
        buộc club phải tồn tại trước (không thì cả học sinh bị bỏ qua),
        mà club lại chỉ tạo được bằng cách gõ form từng cái.

        Cột: club_id, name, capacity — bắt buộc.
             reserve_capacity, reserve_group — tuỳ chọn (trường không
             dùng dự trữ thì bỏ trống cả hai).
             buoi — tuỳ chọn. Buổi sinh hoạt trong tuần, vd `thu_3`. Hai
             CLB cùng giá trị `buoi` là trùng giờ. Bỏ trống cả cột thì
             mọi CLB thuộc một buổi và phần mềm chạy y hệt bản cũ.
        """
        try:
            fieldnames, rows = self._parse_csv_rows(csv_text)
            if not rows:
                return _fail(err("csv_empty"))
            for bat_buoc in ("club_id", "name", "capacity"):
                if bat_buoc not in fieldnames:
                    return _fail(err("csv_missing_columns", fieldnames=fieldnames))

            with self._ket_noi_ghi() as cur:
                da_co = {r[0] for r in cur.execute("SELECT club_id FROM clubs")}

                n_tao, n_sua, n_bo = 0, 0, 0
                canh_bao = []
                # Dem CLB co khai buoi va khong khai, de canh bao bo du lieu
                # TRON hai kieu — xem cho raise o cuoi vong lap.
                co_buoi, khong_buoi = [], []
                for i, row in enumerate(rows, start=2):   # dong 1 la tieu de
                    club_id = (row.get("club_id") or "").strip()
                    if not club_id:
                        canh_bao.append(err("csv_club_row_invalid", line=i,
                                            reason="club_id"))
                        n_bo += 1
                        continue
                    try:
                        capacity = int(row.get("capacity") or 0)
                        reserve_capacity = int(row.get("reserve_capacity") or 0)
                    except ValueError:
                        canh_bao.append(err("csv_club_row_invalid", line=i,
                                            reason="capacity"))
                        n_bo += 1
                        continue
                    # Cung dung dieu kien nhu create_or_update_club — khong de
                    # duong CSV lot qua thu duong UI da chan.
                    if loi_suc_chua(capacity, reserve_capacity):
                        canh_bao.append(err("csv_club_row_invalid", line=i,
                                            reason="capacity"))
                        n_bo += 1
                        continue

                    buoi_value = self._khop_buoi_da_co(
                        cur, self.chuan_hoa_buoi(row.get("buoi")), row.get("buoi")) or None
                    (co_buoi if buoi_value else khong_buoi).append(club_id)

                    self._ghi_clb(cur, club_id, (row.get("name") or "").strip(),
                                  capacity, reserve_capacity,
                                  self.chuan_hoa_nhom_du_tru(row.get("reserve_group")) or None,
                                  buoi_value)
                    if club_id in da_co:
                        n_sua += 1
                    else:
                        n_tao += 1
                        da_co.add(club_id)

                # Tep khai buoi cho MOT SO CLB va bo trong so con lai la
                # tinh huong nguy hiem nhat cua ca buoc nhap nay: nhung CLB
                # bo trong roi hết vào cùng một buổi mặc định, tức là chúng
                # bị coi là TRÙNG GIỜ với nhau — người nhập không hề định
                # thế và không có gì trên màn hình nói ra. Phải kêu.
                if co_buoi and khong_buoi:
                    canh_bao.append(err(
                        "clb_thieu_buoi",
                        n=len(khong_buoi),
                        club_ids=sorted(khong_buoi)[:5],
                    ))

            return _ok({
                "n_clubs_created": n_tao,
                "n_clubs_updated": n_sua,
                "n_rows_skipped": n_bo,
                "n_clubs_co_buoi": len(co_buoi),
                "warnings": canh_bao,
            })
        except Exception as e:
            return _fail(self._loi_co_ghi_vet(err("error_importing_clubs_csv", detail=str(e))))

    def preview_import_csv(self, csv_text: str, kind: str):
        """
        Xem trước trước khi nhập thật (không ghi DB) — cho UI hiện
        'phát hiện định dạng X, Y dòng, Z học sinh sẽ được tạo mới'
        trước khi người dùng bấm xác nhận nhập.
        kind: "preferences" hoặc "test_selection"
        """
        try:
            fieldnames, rows = self._parse_csv_rows(csv_text)
            if not rows:
                return _fail(err("csv_empty"))

            is_wide = any(f.startswith("pref_") or f.startswith("test_club_") for f in fieldnames)
            fmt = "wide" if is_wide else "long"

            with self._ket_noi_doc() as cur:
                existing_ids = {r[0] for r in cur.execute("SELECT student_id FROM students")}

            row_student_ids = {r.get("student_id", "") for r in rows if r.get("student_id")}
            new_students = [sid for sid in row_student_ids if sid not in existing_ids]

            return _ok({
                "format": fmt,
                "kind": kind,
                "fieldnames": fieldnames,
                "n_rows": len(rows),
                "n_students_detected": len(row_student_ids),
                "n_new_students": len(new_students),
                "sample_row": rows[0] if rows else None,
            })
        except Exception as e:
            return _fail(err("error_reading_csv_preview", detail=str(e)))

    # -----------------------------------------------------------------
    # CHUAN HOA NHAN NHOM DU TRU
    #
    # reserve_group la chuoi TU DO, go o HAI noi (file danh sach CLB va
    # file hoc sinh) va phai khop nhau thi co che du tru moi chay. Truoc
    # day so khop bang chuoi chinh xac, nen CLB khai "chinh_sach" con
    # giao vien go "Chinh sach" la HAI NHOM KHAC NHAU: hoc sinh dien
    # chinh sach vao theo dien general, mat suat du tru. Pipeline chay
    # het va khong bao loi — chi co hai dong trong muc Canh bao du lieu,
    # rat de luot qua.
    #
    # Nay moi nhan deu di qua day TRUOC KHI ghi vao DB, nen hai ben go
    # kieu nao cung quy ve cung mot ma.
    #
    # Chuan hoa luc GHI chu khong phai luc SO SANH la co y:
    # rbda_priority_pipeline.py van so khop chuoi chinh xac nhu cu, khong
    # phai sua gi — cho nhay cam nhat cua du an khong bi dung toi.
    # -----------------------------------------------------------------

    @staticmethod
    def chuan_hoa_nhom_du_tru(raw) -> str:
        """"Chinh sach", "CHINH SACH", "Chinh-Sach" -> "chinh_sach"."""
        import unicodedata

        chu = (raw or "").strip()
        if not chu:
            return ""
        # NFD tach dau ra khoi nguyen am roi loc bo — nhung KHONG tach
        # duoc chu D gach ngang, phai thay tay truoc.
        chu = chu.replace("Đ", "D").replace("đ", "d")
        chu = "".join(
            c for c in unicodedata.normalize("NFD", chu)
            if not unicodedata.combining(c)
        )
        chu = re.sub(r"[^0-9a-zA-Z]+", "_", chu.lower())
        return chu.strip("_")

    # Cot reserve_group la TUY CHON trong ca hai file hoc sinh. Truoc day
    # hoc sinh tao bang CSV luon co nhom du tru rong, phai vao man hinh 04
    # gan tay tung em — ma nhom du tru CHINH LA co che uu tien cua RB-DA,
    # quen buoc do thi phan du tru vo hieu hoan toan va pipeline van chay
    # tron tru khong bao gi.
    #
    # Quy tac (ghi ca trong HUONG_DAN_CSV.md):
    #   o co gia tri -> GHI DE nhom hien co (nguoi nhap chu dong dua vao)
    #   o trong      -> GIU NGUYEN, khong xoa (file thieu cot khong duoc
    #                   lam mat du lieu da gan truoc do)
    @staticmethod
    def _doc_diem(chuoi: str):
        """Đọc một ô điểm. Trả về float, hoặc None nếu không phải số.

        Chấp nhận cả dấu phẩy thập phân: Excel bản tiếng Việt lưu 8,5
        chứ không phải 8.5. Chỉ đổi khi có ĐÚNG một dấu phẩy và không có
        dấu chấm nào — "1,234.5" là cách viết nghìn, không đụng vào.
        """
        chuoi = (chuoi or "").strip()
        if not chuoi:
            return None
        if chuoi.count(",") == 1 and "." not in chuoi:
            chuoi = chuoi.replace(",", ".")
        try:
            so = float(chuoi)
        except (TypeError, ValueError):
            return None
        return so if so == so and so not in (float("inf"), float("-inf")) else None

    @staticmethod
    def _soat_dong_trung(rows, is_wide: bool) -> list:
        """Cảnh báo khi một student_id xuất hiện nhiều lần trong file DẠNG RỘNG.

        Dạng rộng là mỗi học sinh một dòng, nên dòng thứ hai của cùng một
        mã sẽ âm thầm ghi đè dòng đầu — học sinh điền form hai lần, hoặc
        người nhập dán nhầm, là mất nguyện vọng của một dòng mà không ai
        biết. Dạng DÀI thì nhiều dòng mỗi học sinh là bình thường, không
        cảnh báo gì.
        """
        if not is_wide:
            return []
        dem: dict = {}
        for row in rows:
            sid = (row.get("student_id") or "").strip()
            if sid:
                dem[sid] = dem.get(sid, 0) + 1
        return [
            err("csv_duplicate_student_rows", student_id=sid, n=n)
            for sid, n in sorted(dem.items()) if n > 1
        ]

    @staticmethod
    def _soat_ma_trung_hoa_thuong(cur, ma_trong_file) -> list:
        """Cảnh báo khi mã học sinh mới chỉ khác hoa/thường với mã đã có.

        `hs001` và `HS001` tạo ra HAI học sinh riêng biệt: file tick chọn
        dùng kiểu này, file nguyện vọng dùng kiểu kia, thế là hai hồ sơ
        rời rạc mỗi cái thiếu một nửa, và pipeline xử lý như hai người.

        KHÔNG tự gộp — gộp nhầm hai em có thật là hỏng nặng hơn nhiều.
        Chỉ báo để người nhập tự quyết.
        """
        da_co = {
            r[0] for r in cur.execute("SELECT student_id FROM students").fetchall()
        }
        theo_thuong = {}
        for m in sorted(da_co):
            theo_thuong.setdefault(m.lower(), m)

        # Hai cach viet nam TRONG CUNG MOT FILE cung phai bat duoc: nguoi
        # nhap go tay 'hs201' o dong nay va 'HS201' o dong kia thi ca hai
        # deu la ma moi, khong ma nao co san trong CSDL de doi chieu.
        canh_bao = []
        for m in sorted(set(ma_trong_file)):
            if m in da_co:
                theo_thuong.setdefault(m.lower(), m)
                continue
            cu = theo_thuong.get(m.lower())
            if cu:
                canh_bao.append(
                    err("csv_student_id_case_conflict", student_id=m, da_co=cu)
                )
            else:
                theo_thuong[m.lower()] = m
        return canh_bao

    @staticmethod
    def _soat_ma_nghi_bi_cat(ma_trong_file) -> list:
        """Cảnh báo mã học sinh nghi bị Excel cắt mất số 0 đứng đầu.

        `0012345` để Excel tự nhận định dạng thì thành số `12345`. Phần
        mềm nhận đúng thứ Excel đã lưu nên KHÔNG cứu được — nhưng PHÁT
        HIỆN được: trong cùng một file, mã toàn chữ số mà ngắn hơn hẳn
        những mã còn lại gần như chắc chắn đã bị cắt.

        Chỉ cảnh báo, KHÔNG chặn nhập và KHÔNG tự thêm số 0 vào: phần
        mềm không có cách nào biết mã gốc dài bao nhiêu, đoán thêm là tự
        bịa dữ liệu.

        Ba điều kiện để tránh báo nhiễu:
          - chỉ xét mã TOÀN CHỮ SỐ (mã có chữ thì Excel không đụng tới)
          - cần ít nhất 3 mã như vậy mới đủ cơ sở nói cái nào bất thường
          - mã ngắn phải là THIỂU SỐ; phần lớn mã đều ngắn thì đó là quy
            ước của trường, không phải lỗi Excel
        """
        toan_so = [m for m in ma_trong_file if m.isdigit()]
        if len(toan_so) < 3:
            return []

        dem_do_dai: dict = {}
        for m in toan_so:
            dem_do_dai[len(m)] = dem_do_dai.get(len(m), 0) + 1
        # Do dai pho bien nhat; hoa thi lay do dai LON hon (an toan hon —
        # gia thiet ma day du moi la chuan).
        pho_bien = max(dem_do_dai.items(), key=lambda kv: (kv[1], kv[0]))[0]

        nghi = sorted({m for m in toan_so if len(m) < pho_bien})
        # Ma ngan chiem da so -> quy uoc cua truong, khong phai loi.
        if len(nghi) * 2 >= len(toan_so):
            return []

        return [
            err("csv_student_id_maybe_truncated",
                student_id=m, do_dai=len(m), do_dai_pho_bien=pho_bien)
            for m in nghi
        ]

    @staticmethod
    def _khop_club_id(cur):
        """Trả về hàm đưa club_id người dùng gõ về đúng mã gốc trong DB.

        Khớp CHÍNH XÁC trước; chỉ khi không thấy mới thử bỏ qua hoa/thường.
        Nếu trường thật sự tạo cả `clb_a` lẫn `CLB_A` thì mã khớp chính
        xác vẫn thắng — phần mềm không tự đoán hộ.

        Tha thứ hoa/thường KHÔNG phải tha thứ mọi thứ: mã sai hẳn vẫn bị
        bỏ qua kèm cảnh báo như cũ.
        """
        that = [r[0] for r in cur.execute("SELECT club_id FROM clubs").fetchall()]
        chinh_xac = set(that)
        theo_thuong = {}
        for m in that:
            theo_thuong.setdefault(m.lower(), m)

        def khop(cid: str):
            cid = (cid or "").strip()
            if cid in chinh_xac:
                return cid
            return theo_thuong.get(cid.lower())

        return khop

    @staticmethod
    def _soat_nhom_du_tru_la(cur, nhom_da_ghi: dict) -> list:
        """Cảnh báo NGAY LÚC NHẬP nếu nhãn nhóm không CLB nào nhận.

        Mục "Cảnh báo dữ liệu" cũng bắt được ca này, nhưng nằm ở panel
        khác và rất dễ lướt qua. Sai ở đây thì học sinh diện ưu tiên mất
        suất dự trữ mà pipeline vẫn chạy trơn, nên phải đập vào mắt ngay
        khi vừa nạp file.

        nhom_da_ghi: { nhãn đã chuẩn hoá: số học sinh mang nhãn đó }
        """
        import difflib

        if not nhom_da_ghi:
            return []
        nhom_clb = {
            r[0] for r in cur.execute(
                "SELECT DISTINCT reserve_group FROM clubs "
                "WHERE reserve_group IS NOT NULL AND TRIM(reserve_group) <> ''"
            ).fetchall()
        }
        canh_bao = []
        for nhan, n in sorted(nhom_da_ghi.items()):
            if nhan in nhom_clb:
                continue
            # Go nham mot hai ky tu la ca gap nhat — chi thang nhom dinh go.
            gan_giong = difflib.get_close_matches(nhan, sorted(nhom_clb), n=1, cutoff=0.6)
            canh_bao.append(err(
                "csv_reserve_group_unknown",
                reserve_group=nhan, n=n,
                goi_y=gan_giong[0] if gan_giong else "",
            ))
        return canh_bao

    @staticmethod
    def _ghi_nhom_du_tru(cur, student_id: str, reserve_group: str) -> str:
        """Trả về nhãn đã chuẩn hoá (chuỗi rỗng nếu không ghi gì) để bên
        gọi đếm được từng nhãn phục vụ _soat_nhom_du_tru_la."""
        nhom = NhapDuLieuMixin.chuan_hoa_nhom_du_tru(reserve_group)
        if not nhom:
            return ""
        cur.execute(
            "UPDATE students SET reserve_group = ? WHERE student_id = ?",
            (nhom, student_id),
        )
        return nhom

    @classmethod
    def _buoi_qua_tran(cls, cur, ds_clb, tran):
        """Buổi ĐẦU TIÊN vượt trần, dạng `(buoi, số lượng)` — None nếu không có.

        Dùng cho hai đường giao diện, nơi cả danh sách tới cùng một lúc và
        người dùng sửa được ngay, nên từ chối cả lời gọi là đúng. Hai đường
        NẠP TỆP thì khác: ở đó bỏ riêng buổi vi phạm, vì người nạp không ngồi
        sửa từng em được.
        """
        buoi_cua_clb = cls._buoi_cua_clb(cur)
        _giu, bo = cls._cat_buoi_qua_tran(ds_clb, buoi_cua_clb, tran)
        if not bo:
            return None
        b = sorted(bo)[0]
        return b, bo[b]

    @staticmethod
    def _buoi_cua_clb(cur) -> dict:
        """{club_id: buoi} — CLB chưa khai buổi thuộc buổi mặc định.

        Một chỗ duy nhất đọc quan hệ CLB→buổi, để hai bước nạp và bước soát
        lệch buổi không trôi ra hai luật khác nhau.
        """
        return {
            r[0]: (r[1] or BUOI_MAC_DINH)
            for r in cur.execute("SELECT club_id, buoi FROM clubs")
        }

    @staticmethod
    def _cat_buoi_qua_tran(ds_clb, buoi_cua_clb: dict, tran: int):
        """Bỏ những BUỔI vượt trần, giữ nguyên phần còn lại đúng thứ tự cũ.

        Trả `(giu, bo)` với `bo = {buoi: số lượng đã khai}` — rỗng nghĩa là
        không buổi nào vượt.

        Giữ nguyên thứ tự là bắt buộc chứ không phải tiện tay: danh sách
        nguyện vọng LÀ một thứ hạng, xáo nó là đổi nguyện vọng của học sinh.

        Bỏ theo BUỔI chứ không bỏ trọn học sinh: các buổi độc lập với nhau,
        nên một lỗi ở thứ Ba không có lý do gì làm em mất cả tuần.

        LƯU Ý cho người gọi: `ds_clb` phải là mã CLB ĐÃ CHUẨN HOÁ. Mã lệch
        hoa/thường tra không ra buổi, rơi vào buổi mặc định, và gom nhầm nhóm.
        """
        theo_buoi = gom_theo_buoi(
            ds_clb, {c: {"buoi": b} for c, b in buoi_cua_clb.items()})
        bo = {b: len(ds) for b, ds in theo_buoi.items() if len(ds) > tran}
        if not bo:
            return list(ds_clb), {}
        giu = [c for c in ds_clb
               if buoi_cua_clb.get(c, BUOI_MAC_DINH) not in bo]
        return giu, bo

    _MAU_COT_NGUYEN_VONG_THEO_BUOI = re.compile(r"^(?P<buoi>.+)_pref_(?P<k>\d+)$")

    @classmethod
    def _cot_nguyen_vong_theo_buoi(cls, fieldnames) -> dict:
        """Tách bộ cột `<buoi>_pref_<k>` thành {buoi: [cột theo thứ tự k]}.

        Rỗng nghĩa là tệp dùng bộ cột cũ (`pref_1`, `pref_2`…) hoặc dạng dài.
        Cột `pref_1` KHÔNG khớp mẫu này vì mẫu đòi có phần tên buổi đứng
        trước `_pref_`.
        """
        theo_buoi: dict = {}
        for f in fieldnames or []:
            m = cls._MAU_COT_NGUYEN_VONG_THEO_BUOI.match(f.strip())
            if m:
                theo_buoi.setdefault(m.group("buoi"), []).append((int(m.group("k")), f))
        return {
            b: [ten for _k, ten in sorted(cot)]
            for b, cot in theo_buoi.items()
        }

    # -----------------------------------------------------------------
    # BA BƯỚC DÙNG CHUNG CHO HAI TỆP HỌC SINH (nguyện vọng / chọn thi)
    #
    # Trước đây mỗi bước được viết hai lần, mỗi hàm nhập một bản. Sửa một
    # bản mà quên bản kia là hai tệp cùng loại dữ liệu cư xử khác nhau —
    # đúng loại lỗi khó thấy nhất.
    # -----------------------------------------------------------------

    @staticmethod
    def _bo_trung_giu_thu_tu(ds: list) -> tuple:
        """(danh sách đã bỏ trùng, giữ lần xuất hiện ĐẦU; có trùng không)."""
        gon = list(dict.fromkeys(ds))
        return gon, len(gon) != len(ds)

    @staticmethod
    def _khop_ds_clb(khop_club, ds: list) -> tuple:
        """Đổi mã người dùng gõ về mã GỐC trong CSDL (chấp nhận lệch hoa/thường).

        Trả về (ds mã gốc, ds mã không tìm thấy, [(mã gõ, mã gốc|None)]).
        """
        da_khop = [(c, khop_club(c)) for c in ds]
        khong_co = [c for c, that in da_khop if that is None]
        return [that for _, that in da_khop if that is not None], khong_co, da_khop

    @staticmethod
    def _tao_hoac_bo_sung_hoc_sinh(cur, sid, name, da_co: set, duoc_tao: bool) -> str:
        """Học sinh chưa có thì tạo (nếu được phép); đã có mà thiếu tên thì
        điền tên. Trả về "tao_moi", "da_co" hoặc "bo_qua"."""
        if sid not in da_co:
            if not duoc_tao:
                return "bo_qua"
            cur.execute(
                "INSERT INTO students (student_id, name, stb_number, reserve_group) "
                "VALUES (?, ?, NULL, NULL)",
                (sid, name or sid),
            )
            da_co.add(sid)
            return "tao_moi"
        if name:
            cur.execute(
                "UPDATE students SET name = ? WHERE student_id = ? AND (name IS NULL OR name = '')",
                (name, sid),
            )
        return "da_co"

    def import_preferences_csv(self, csv_text: str, create_missing_students: bool = True):
        """
        Nhập CSV nguyện vọng (Bước 2 — xếp hạng) từ Microsoft Forms.
        Ghi đè TOÀN BỘ nguyện vọng cũ của TỪNG học sinh xuất hiện trong
        file (giống hành vi submit_preferences ở kiosk) — không đụng
        tới học sinh không có trong file.
        """
        try:
            fieldnames, rows = self._parse_csv_rows(csv_text)
            if not rows:
                return _fail(err("csv_empty"))

            cot_theo_buoi = self._cot_nguyen_vong_theo_buoi(fieldnames)
            is_wide = bool(cot_theo_buoi) or any(
                f.startswith("pref_") for f in fieldnames)

            # Gom thanh { student_id: (name, [club_id_theo_thu_tu], nhom) }
            grouped: dict = {}
            canh_bao_buoi: list = []
            if cot_theo_buoi:
                # Bo cot theo buoi: thu_3_pref_1, thu_3_pref_2, thu_5_pref_1...
                # Moi buoi mot cau Ranking rieng tren Microsoft Forms.
                #
                # Cac buoi duoc NOI LIEN NHAU theo thu tu ten buoi, va rank
                # ghi xuong CSDL van la 1..n chay suot. Nghe nhu mat thong
                # tin buoi, nhung khong: thuat toan luon LOC danh sach nay
                # theo buoi truoc khi dung, ma loc mot day da sap thi phan
                # con lai van dung thu tu. Nho vay du lieu nhap bang bo cot
                # cu va bo cot moi song chung duoc trong cung mot CSDL.
                with self._ket_noi_doc() as cur:
                    buoi_that = self._buoi_cua_clb(cur)
                for i, row in enumerate(rows, start=2):
                    sid = row.get("student_id")
                    if not sid:
                        continue
                    thu_tu = []
                    for buoi_cot in sap_buoi(cot_theo_buoi):
                        for cot in cot_theo_buoi[buoi_cot]:
                            cid = (row.get(cot) or "").strip()
                            if not cid:
                                continue
                            that = buoi_that.get(cid)
                            # CLB khong ton tai -> de buoc kiem tra chung phia
                            # sau bao loi; o day chi soat chuyen LECH BUOI.
                            # So theo nhãn CHUẨN: CLB khai "t3" được lưu thành
                            # thu_3, cột vẫn tên t3_pref_1 — cùng một buổi.
                            if (that is not None and that != buoi_cot
                                    and so_nhap.ma_buoi(that) != so_nhap.ma_buoi(buoi_cot)):
                                canh_bao_buoi.append(err(
                                    "nguyen_vong_lech_buoi", line=i, student_id=sid,
                                    club_id=cid, buoi_cot=buoi_cot, buoi_that=that))
                                continue
                            thu_tu.append(cid)
                    grouped[sid] = (row.get("name", ""), thu_tu,
                                    row.get("reserve_group", ""))
            elif is_wide:
                pref_cols = sorted(
                    [f for f in fieldnames if f.startswith("pref_")],
                    key=lambda f: int(f.split("_")[1]) if f.split("_")[1].isdigit() else 999,
                )
                for row in rows:
                    sid = row.get("student_id")
                    if not sid:
                        continue
                    ordered = [row[c] for c in pref_cols if row.get(c)]
                    grouped[sid] = (row.get("name", ""), ordered,
                                    row.get("reserve_group", ""))
            else:
                if "student_id" not in fieldnames or "club_id" not in fieldnames:
                    return _fail(err("csv_missing_columns", fieldnames=fieldnames))
                has_rank = "rank" in fieldnames
                by_sid_rows: dict = {}
                for row in rows:
                    sid = row.get("student_id")
                    if not sid or not row.get("club_id"):
                        continue
                    by_sid_rows.setdefault(sid, []).append(row)
                for sid, sid_rows in by_sid_rows.items():
                    if has_rank:
                        # `or ""`: DictReader dien None vao o thieu cua dong ngan, va
                        # None.isdigit() lam hong CA lan nhap vi mot dong.
                        sid_rows.sort(key=lambda r: int(r["rank"]) if (r.get("rank") or "").isdigit() else 999)
                    name = next((r.get("name") for r in sid_rows if r.get("name")), "")
                    nhom = next((r.get("reserve_group") for r in sid_rows
                                 if r.get("reserve_group")), "")
                    grouped[sid] = (name, [r["club_id"] for r in sid_rows], nhom)

            with self._ket_noi_ghi() as cur:
                khop_club = self._khop_club_id(cur)
                existing_students = {r[0] for r in cur.execute("SELECT student_id FROM students")}

                n_created, n_updated, n_skipped = 0, 0, 0
                # Dem RIENG voi n_skipped. "Bo mot buoi cua mot em" va "bo
                # tron mot em" la hai chuyen khac han; gop chung la lap lai
                # dung kieu nham lan da de ra con so "45 em trang tay".
                n_buoi_bo_qua = 0
                buoi_cua_clb = self._buoi_cua_clb(cur)
                row_errors = self._soat_dong_trung(rows, is_wide)
                row_errors += self._soat_ma_trung_hoa_thuong(cur, grouped.keys())
                row_errors += self._soat_ma_nghi_bi_cat(grouped.keys())
                # Điểm chỉ thuộc file CHỌN CLB THI. Gặp cột điểm ở đây là
                # người nhập tưởng đã nạp điểm rồi — im lặng bỏ qua đúng là
                # loại lỗi im lặng dự án này đã phải sửa nhiều lần.
                if any(f == "score" or f.startswith("score_") for f in fieldnames):
                    row_errors.append(err("csv_scores_ignored_here"))
                nhom_da_ghi: dict = {}

                for sid, (name, ordered_clubs, nhom_du_tru) in grouped.items():
                    deduped, co_trung = self._bo_trung_giu_thu_tu(ordered_clubs)
                    if co_trung:
                        row_errors.append(err("csv_pref_duplicate_deduped", student_id=sid))

                    deduped, invalid_clubs, _ = self._khop_ds_clb(khop_club, deduped)
                    if invalid_clubs:
                        row_errors.append(err("csv_unknown_clubs_skipped", student_id=sid, club_ids=invalid_clubs))
                        n_skipped += 1
                        continue

                    # TRAN TINH THEO TUNG BUOI, va kiem O DAY chu khong som
                    # hon: tra buoi phai dung ma DA CHUAN HOA, ma lech
                    # hoa/thuong tra khong ra buoi thi roi vao buoi mac dinh
                    # va gom nham nhom.
                    #
                    # Tran cu dem ca tuan va BO TRON hoc sinh vuot tran. No
                    # viet cho truong MOT buoi; voi sau buoi thi em nao chon
                    # 2 CLB moi buoi da la 12 nguyen vong va roi ra ngoai. Do
                    # duoc: 31/180 em cua bo `bo_sau_buoi` va 4/160 em cua
                    # `bo_nhieu_buoi` bi gat lang le khoi dot xep.
                    giu, buoi_bo = self._cat_buoi_qua_tran(
                        deduped, buoi_cua_clb, TRAN_NGUYEN_VONG_MOI_BUOI)
                    for b, n in sorted(buoi_bo.items(), key=lambda bn: khoa_sap_buoi(bn[0])):
                        # Truong MOT buoi co nhan buoi la BUOI_MAC_DINH — mot
                        # chuoi noi bo, khong phai thu doc len duoc. Neu no ra
                        # la bat nguoi dung doc mot khai niem truong ho khong
                        # co. Hai ma loi vi the, va truong mot buoi thay dung
                        # cau nhu ban cu.
                        row_errors.append(err(
                            "csv_pref_too_many_skipped"
                            if b == BUOI_MAC_DINH else "csv_pref_bo_buoi_qua_tran",
                            student_id=sid, buoi=b, count=n,
                            tran=TRAN_NGUYEN_VONG_MOI_BUOI))
                    n_buoi_bo_qua += len(buoi_bo)
                    if buoi_bo and not giu:
                        # Moi buoi deu vuot tran -> em nay that su trang tay.
                        n_skipped += 1
                        continue
                    deduped = giu

                    ket = self._tao_hoac_bo_sung_hoc_sinh(
                        cur, sid, name, existing_students, create_missing_students)
                    if ket == "bo_qua":
                        row_errors.append(err("csv_student_missing_skipped", student_id=sid))
                        n_skipped += 1
                        continue
                    n_created += ket == "tao_moi"

                    nhan = self._ghi_nhom_du_tru(cur, sid, nhom_du_tru)
                    if nhan:
                        nhom_da_ghi[nhan] = nhom_da_ghi.get(nhan, 0) + 1

                    cur.execute("DELETE FROM preferences WHERE student_id = ?", (sid,))
                    cur.executemany(
                        "INSERT INTO preferences (student_id, club_id, rank) VALUES (?, ?, ?)",
                        [(sid, cid, i + 1) for i, cid in enumerate(deduped)],
                    )
                    n_updated += 1

                row_errors.extend(self._soat_nhom_du_tru_la(cur, nhom_da_ghi))

            # Nguyen vong lech buoi: bao TRUOC cac canh bao khac. Day la loi
            # nguoi nhap sua duoc ngay va sua xong thi ket qua doi that, nen
            # no khong duoc lan giua mot danh sach dai.
            row_errors = canh_bao_buoi + row_errors

            return _ok({
                "n_students_created": n_created,
                "n_students_with_preferences_written": n_updated,
                "n_students_skipped": n_skipped,
                "n_buoi_bo_qua": n_buoi_bo_qua,
                "n_nguyen_vong_lech_buoi": len(canh_bao_buoi),
                "warnings": row_errors,
            })
        except Exception as e:
            return _fail(self._loi_co_ghi_vet(err("error_importing_preferences_csv", detail=str(e))))

    def import_test_selection_csv(self, csv_text: str, create_missing_students: bool = True):
        """
        Nhập CSV chọn club muốn thi/xét (Bước 1 — tick-box) từ
        Microsoft Forms. Ghi đè toàn bộ lựa chọn thi cũ của từng học
        sinh xuất hiện trong file, giống hành vi submit_test_selection.
        """
        try:
            fieldnames, rows = self._parse_csv_rows(csv_text)
            if not rows:
                return _fail(err("csv_empty"))

            is_wide = any(f.startswith("test_club_") for f in fieldnames)

            grouped: dict = {}
            diem_tho: dict = {}     # sid -> {club_id thô: chuỗi điểm}
            loi_som: list = []      # lỗi phát hiện trước khi mở CSDL

            def _hau_to(ten: str, tien_to: str):
                duoi = ten[len(tien_to):]
                return duoi if duoi.isdigit() else None

            if is_wide:
                # Ghép test_club_N với score_N theo HẬU TỐ SỐ trong tên
                # cột, KHÔNG theo vị trí. Ô CLB bỏ trống bị lọc đi, nên
                # ghép theo vị trí sẽ lệch: em bỏ trống test_club_2 mà
                # điền test_club_3 thì điểm sẽ gán nhầm cho club khác.
                cot_clb = {}
                cot_diem = {}
                for f in fieldnames:
                    if f.startswith("test_club_"):
                        n = _hau_to(f, "test_club_")
                        if n:
                            cot_clb[n] = f
                    elif f.startswith("score_"):
                        n = _hau_to(f, "score_")
                        if n:
                            cot_diem[n] = f

                for row in rows:
                    sid = row.get("student_id")
                    if not sid:
                        continue
                    selected, diem_hs = [], {}
                    for n in sorted(cot_clb, key=int):
                        cid = (row.get(cot_clb[n]) or "").strip()
                        cot_d = cot_diem.get(n)
                        d = (row.get(cot_d) or "").strip() if cot_d else ""
                        if cid:
                            selected.append(cid)
                            if d:
                                diem_hs[cid] = d
                        elif d:
                            # Có điểm mà không có club -> gần như chắc là
                            # gõ lệch cột. Báo, đừng đoán club nào.
                            loi_som.append(err("csv_score_without_club",
                                               student_id=sid, cot=n))
                    grouped[sid] = (row.get("name", ""), selected,
                                    row.get("reserve_group", ""))
                    if diem_hs:
                        diem_tho[sid] = diem_hs
            else:
                if "student_id" not in fieldnames or "club_id" not in fieldnames:
                    return _fail(err("csv_missing_columns", fieldnames=fieldnames))
                for row in rows:
                    sid = row.get("student_id")
                    if not sid or not row.get("club_id"):
                        continue
                    name, clubs_list, nhom = grouped.get(sid, ("", [], ""))
                    grouped[sid] = (row.get("name") or name,
                                    clubs_list + [row["club_id"]],
                                    row.get("reserve_group") or nhom)
                    d = (row.get("score") or "").strip()
                    if d:
                        diem_tho.setdefault(sid, {})[row["club_id"]] = d

            with self._ket_noi_ghi() as cur:
                khop_club = self._khop_club_id(cur)
                existing_students = {r[0] for r in cur.execute("SELECT student_id FROM students")}

                n_created, n_updated, n_skipped, n_diem = 0, 0, 0, 0
                n_diem_xoa = 0
                n_buoi_bo_qua = 0
                buoi_cua_clb = self._buoi_cua_clb(cur)
                row_errors = list(loi_som)
                row_errors += self._soat_dong_trung(rows, is_wide)
                row_errors += self._soat_ma_trung_hoa_thuong(cur, grouped.keys())
                row_errors += self._soat_ma_nghi_bi_cat(grouped.keys())
                nhom_da_ghi: dict = {}

                for sid, (name, club_ids, nhom_du_tru) in grouped.items():
                    deduped, _ = self._bo_trung_giu_thu_tu(club_ids)
                    deduped, invalid_clubs, da_khop = self._khop_ds_clb(khop_club, deduped)
                    if invalid_clubs:
                        row_errors.append(err("csv_unknown_clubs_skipped", student_id=sid, club_ids=invalid_clubs))
                        n_skipped += 1
                        continue

                    # TRAN CLB THI TINH THEO TUNG BUOI. Thi qua nam CLB trong
                    # MOT buoi thi khong truong nao cham xue; con ca tuan sau
                    # buoi, nam CLB moi buoi la hoan toan binh thuong.
                    #
                    # Kiem O DAY, sau khi chuan hoa ma: ma lech hoa/thuong tra
                    # khong ra buoi thi roi vao buoi mac dinh va gom nham nhom.
                    giu, buoi_bo = self._cat_buoi_qua_tran(
                        deduped, buoi_cua_clb, TRAN_CLB_THI_MOI_BUOI)
                    for b, n in sorted(buoi_bo.items(), key=lambda bn: khoa_sap_buoi(bn[0])):
                        row_errors.append(err(
                            "csv_thi_too_many_skipped"
                            if b == BUOI_MAC_DINH else "csv_thi_bo_buoi_qua_tran",
                            student_id=sid, buoi=b, count=n,
                            tran=TRAN_CLB_THI_MOI_BUOI))
                    n_buoi_bo_qua += len(buoi_bo)
                    clb_bi_bo = set(deduped) - set(giu)
                    if buoi_bo and not giu:
                        n_skipped += 1
                        continue
                    deduped = giu

                    ket = self._tao_hoac_bo_sung_hoc_sinh(
                        cur, sid, name, existing_students, create_missing_students)
                    if ket == "bo_qua":
                        row_errors.append(err("csv_student_missing_skipped", student_id=sid))
                        n_skipped += 1
                        continue
                    n_created += ket == "tao_moi"

                    nhan = self._ghi_nhom_du_tru(cur, sid, nhom_du_tru)
                    if nhan:
                        nhom_da_ghi[nhan] = nhom_da_ghi.get(nhan, 0) + 1

                    cur.execute("DELETE FROM club_test_selection WHERE student_id = ?", (sid,))
                    cur.executemany(
                        "INSERT INTO club_test_selection (student_id, club_id) VALUES (?, ?)",
                        [(sid, cid) for cid in deduped],
                    )
                    n_updated += 1

                    # Điểm ghi SAU lựa chọn thi, và ghi bằng club_id ĐÃ KHỚP
                    # (khop_club chấp nhận lệch hoa/thường), không phải chuỗi
                    # thô trong file — ghi thô sẽ tạo một dòng điểm mồ côi mà
                    # không màn hình nào đọc tới.
                    diem_hs = diem_tho.get(sid)
                    if diem_hs:
                        anh_xa = {tho: that for tho, that in da_khop}
                        for tho, chuoi in diem_hs.items():
                            that = anh_xa.get(tho)
                            if that is None:
                                row_errors.append(err("csv_score_for_unselected_club",
                                                      student_id=sid, club_id=tho))
                                continue
                            # CLB bi bo vi buoi vuot tran: bo luon diem cua no.
                            # Khong bo thi day la mot dong diem mo coi cho mot
                            # CLB em khong con du thi. Va KHONG canh bao lan
                            # hai — canh bao ve buoi vuot tran o tren da noi
                            # dung nguyen nhan; them mot dong "cham diem CLB
                            # khong dang ky" chi lam nguoi doc tuong co hai loi.
                            if that in clb_bi_bo:
                                continue
                            so = self._doc_diem(chuoi)
                            if so is None:
                                # Chỉ bỏ RIÊNG ô điểm này. Bỏ cả học sinh vì
                                # một ô gõ sai là mất nhiều hơn được.
                                row_errors.append(err("csv_score_not_a_number",
                                                      student_id=sid, club_id=that, score=chuoi))
                                continue
                            if so < 0:
                                row_errors.append(err("csv_score_negative",
                                                      student_id=sid, club_id=that, score=chuoi))
                                continue
                            cur.execute(
                                "INSERT INTO club_scores (student_id, club_id, score) VALUES (?, ?, ?) "
                                "ON CONFLICT(student_id, club_id) DO UPDATE SET score = excluded.score",
                                (sid, that, so),
                            )
                            n_diem += 1

                    # Tệp mới không còn CLB nào đó cho em này -> điểm cũ của
                    # CLB ấy đi theo ô tick. Gọi SAU khi ghi điểm từ tệp: mọi
                    # điểm vừa ghi đều thuộc CLB đang tick nên không bị đụng.
                    n_xoa = self._xoa_diem_khong_con_tick(cur, sid)
                    if n_xoa:
                        n_diem_xoa += n_xoa
                        row_errors.append(err("csv_score_removed_unselected",
                                              student_id=sid, n=n_xoa))

                row_errors.extend(self._soat_nhom_du_tru_la(cur, nhom_da_ghi))


            return _ok({
                "n_students_created": n_created,
                "n_students_with_selection_written": n_updated,
                "n_scores_written": n_diem,
                "n_scores_removed": n_diem_xoa,
                "n_students_skipped": n_skipped,
                "n_buoi_bo_qua": n_buoi_bo_qua,
                "warnings": row_errors,
            })
        except Exception as e:
            return _fail(self._loi_co_ghi_vet(err("error_importing_test_selection_csv", detail=str(e))))
