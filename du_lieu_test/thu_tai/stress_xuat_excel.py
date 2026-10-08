"""Thử sức tệp Excel kết quả (`export_ket_qua`) — sáu nhóm phép thử.

    ./.venv/bin/python du_lieu_test/thu_tai/stress_xuat_excel.py [--nhanh]

⚠️  DỮ LIỆU MÔ PHỎNG do máy sinh, không phải khảo sát học sinh có thật.

VÌ SAO CẦN. Từ khi xuất kết quả gom về MỘT tệp Excel, hỏng sổ Excel là hỏng
cả lần xuất — trước đó sổ hỏng thì lặng lẽ mất, các tệp .csv vẫn ra. Nên
mọi đầu vào làm openpyxl ngã giờ chặn đứng nhà trường khỏi kết quả.

Sáu nhóm:
  1. Quy mô và tốc độ — tới 5.000 em, 100 CLB, 5 buổi
  2. Giới hạn cứng của Excel — ô 32.767 ký tự, 1.026 ngắt trang mỗi trang
  3. Dữ liệu lạ và độc — ký tự điều khiển, công thức, emoji, mã số 0 đầu...
  4. Trạng thái lạ — chưa chạy, kết quả cũ, chạy một phần buổi, xoá sau khi chạy
  5. Hệ thống tệp — thư mục không ghi được, xuất 50 lần, hai luồng cùng xuất
  6. Nhất quán ngẫu nhiên — 30 bộ dữ liệu × 3 hạt giống, đối chiếu từng trang

Mỗi kịch bản qua CÙNG một bộ kiểm `kiem_so` (nhóm 6), nên "xuất được" luôn
đi kèm "và nói đúng". Cuối cùng mở lại mọi tệp đáng ngờ bằng LibreOffice
(nếu máy có) — openpyxl đọc được chưa chắc là bảng tính hợp lệ.

Kết quả: `ket_qua_stress_xuat.csv` cạnh tệp này, và bảng tóm tắt in ra màn hình.
"""

import collections
import csv
import glob
import hashlib
import os
import random
import re
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import threading
import time
import tracemalloc
import zipfile

GOC = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, GOC)

from openpyxl import load_workbook  # noqa: E402

from api import PipelineAPI  # noqa: E402

THU_MUC = os.path.dirname(os.path.abspath(__file__))
SEED = 2026
NHANH = "--nhanh" in sys.argv

TRANG = ["Hướng dẫn", "Danh sách học sinh", "Theo CLB",
         "Chưa có chỗ", "Thống kê", "Dấu vết (kỹ thuật)"]
O_TOI_DA = 32767           # ky tu toi da mot o Excel
NGAT_TOI_DA = 1026         # ngat trang thu cong toi da mot trang tinh Excel
KY_TU_CAM = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

KQ: list[dict] = []
TEP_SOAT_LO: list[str] = []    # tep se mo lai bang LibreOffice


def ghi(nhom, kich_ban, loi, **so):
    """Ghi một dòng kết quả. `loi` rỗng = ĐẠT."""
    ok = not loi
    KQ.append({"nhom": nhom, "kich_ban": kich_ban, "ket_qua": "ĐẠT" if ok else "HỎNG",
               "chi_tiet": "; ".join(loi)[:400], **so})
    print("  [%s] %-58s %s" % ("ĐẠT " if ok else "HỎNG", kich_ban[:58],
                               "" if ok else "; ".join(loi)[:160]))


def lam_sach(v):
    """Giá trị một ô SAU KHI xuất: bỏ ký tự Excel cấm, cắt ở 32.767 ký tự.
    Bộ kiểm so dữ liệu CSDL đã qua hàm này với nội dung tệp."""
    if v is None:
        return None
    s = KY_TU_CAM.sub("", str(v))
    return s[:O_TOI_DA]


# --------------------------------------------------------------------- #
# SINH DỮ LIỆU
# --------------------------------------------------------------------- #

def sinh_csv(n_hs, n_clb, n_buoi, seed, ty_le=1.0, ten_hs=None, ten_clb=None,
             ma_clb=None, ten_buoi=None, ma_hs=None):
    """(csv CLB, csv thi, csv nguyện vọng). Nhu cầu lệch kiểu Zipf như
    chay_thu_tai.py; mỗi em 3 nguyện vọng mỗi buổi, thi 4 CLB đầu."""
    rng = random.Random(seed)
    buoi = ten_buoi or (["thu_%d" % (i + 2) for i in range(n_buoi)] if n_buoi > 1 else [None])
    ma = ma_clb or ["c%04d" % i for i in range(n_clb)]
    ten = ten_clb or ["CLB %d" % i for i in range(n_clb)]
    b_cua = {m: buoi[i % len(buoi)] for i, m in enumerate(ma)}
    theo_buoi = collections.defaultdict(list)
    for m in ma:
        theo_buoi[b_cua[m]].append(m)
    cho = max(1, int(round(n_hs * ty_le / max(1, len(ma) / len(buoi)))))

    def q(x):
        x = "" if x is None else str(x)
        return '"%s"' % x.replace('"', '""') if any(c in x for c in ',"\n\r') else x

    dong_clb = ["club_id,name,capacity,reserve_capacity,reserve_group"
                + (",buoi" if n_buoi > 1 else "")]
    for i, m in enumerate(ma):
        r = max(1, cho // 5) if i % 5 == 0 else 0
        dong_clb.append(",".join(q(x) for x in [m, ten[i], cho, r, "chinh_sach" if r else ""]
                                 + ([b_cua[m]] if n_buoi > 1 else [])))

    cot_nv = []
    for b in buoi:
        for k in range(1, 4):
            cot_nv.append(("%s_pref_%d" % (b, k)) if b else ("pref_%d" % k))
    dong_thi = ["student_id,name,reserve_group," + ",".join(
        "test_club_%d,score_%d" % (k, k) for k in range(1, 5))]
    dong_nv = ["student_id,name,reserve_group," + ",".join(cot_nv)]
    for i in range(n_hs):
        sid = (ma_hs[i] if ma_hs else "S%06d" % i)
        tn = (ten_hs[i] if ten_hs else "HS %d" % i)
        nhom = "chinh_sach" if rng.random() < 0.2 else ""
        nv = []
        for b in buoi:
            ds = theo_buoi[b]
            ts = [1.0 / (j + 1) for j in range(len(ds))]
            chon = []
            con = list(range(len(ds)))
            for _ in range(min(3, len(ds))):
                k = rng.choices(con, weights=[ts[j] for j in con])[0]
                con.remove(k)
                chon.append(ds[k])
            nv += chon + [""] * (3 - len(chon))
        thi = [x for x in nv if x][:4]
        o = [sid, tn, nhom]
        for k in range(4):
            o += [thi[k], "%.1f" % rng.uniform(4, 10)] if k < len(thi) else ["", ""]
        dong_thi.append(",".join(q(x) for x in o))
        dong_nv.append(",".join(q(x) for x in [sid, tn, nhom] + nv))
    return "\n".join(dong_clb), "\n".join(dong_thi), "\n".join(dong_nv)


def dung(n_hs, n_clb, n_buoi=1, seed=SEED, chay=True, **kw):
    d = tempfile.mkdtemp(prefix="stress_")
    api = PipelineAPI(os.path.join(d, "app.db"))
    api.thu_muc_xuat = os.path.join(d, "TaiXuong")
    os.makedirs(api.thu_muc_xuat)
    for text in sinh_csv(n_hs, n_clb, n_buoi, seed, **kw):
        r = api.import_csv_auto(text)
        if not r["ok"]:
            raise RuntimeError("nap that bai: %r" % r["errors"][:2])
    if chay:
        r = api.run_pipeline(seed=seed)
        if not r["ok"]:
            raise RuntimeError("chay that bai: %r" % r["errors"][:2])
    return api


def sua_csdl(api, sql, *tham):
    with sqlite3.connect(api.db_path) as c:
        c.execute(sql, tham)


def van_tay_csdl(api) -> str:
    with sqlite3.connect(api.db_path) as c:
        return hashlib.sha256("\n".join(c.iterdump()).encode()).hexdigest()


# --------------------------------------------------------------------- #
# BỘ KIỂM DÙNG CHUNG
# --------------------------------------------------------------------- #

def _bang_duoi_khoa(ws):
    if not ws.freeze_panes:
        return [], []
    h = ws[ws.freeze_panes].row
    tieu_de = [c.value for c in ws[h - 1]]
    dong = [list(r) for r in ws.iter_rows(min_row=h, values_only=True)
            if any(o not in (None, "") for o in r)]
    return tieu_de, dong


def doc_theo_clb(ws):
    ra, hien, cot = collections.Counter(), None, None
    clb = []
    for r in ws.iter_rows(values_only=True):
        a = r[0]
        if isinstance(a, str) and "Mã CLB: " in a:
            hien = next(p[len("Mã CLB: "):] for p in a.split(" · ") if p.startswith("Mã CLB: "))
            clb.append(hien)
            cot = None
        elif hien is not None and "Mã học sinh" in r:
            cot = list(r).index("Mã học sinh")
        elif hien is not None and cot is not None and r[cot] not in (None, ""):
            ra[(hien, str(r[cot]))] += 1
    return ra, clb


def kiem_gioi_han(wb, path):
    loi = []
    for ws in wb.worksheets:
        if len(ws.title) > 31:
            loi.append("ten trang %r > 31" % ws.title)
        if len(ws.row_breaks.brk) > NGAT_TOI_DA:
            loi.append("%s: %d ngat trang > %d" % (ws.title, len(ws.row_breaks.brk), NGAT_TOI_DA))
        for r in ws.iter_rows():
            for c in r:
                if isinstance(c.value, str):
                    if len(c.value) > O_TOI_DA:
                        loi.append("%s!%s dai %d" % (ws.title, c.coordinate, len(c.value)))
                    if KY_TU_CAM.search(c.value):
                        loi.append("%s!%s co ky tu cam" % (ws.title, c.coordinate))
                if c.hyperlink and c.hyperlink.location:
                    dich = c.hyperlink.location.split("!")[0].strip("'")
                    if dich not in wb.sheetnames:
                        loi.append("lien ket hong %r" % c.hyperlink.location)
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if n.startswith("xl/worksheets/") and re.search(rb"<f[ >]", z.read(n)):
                loi.append("co cong thuc trong %s" % n)
    return loi[:8]


def kiem_so(api, path, kiem_cap=True):
    """Mọi trang có nói ĐÚNG điều CSDL đang chứa không. Trả về danh sách lỗi."""
    loi = []
    if not path or not os.path.isfile(path):
        return ["khong co tep"]
    wb = load_workbook(path)
    if wb.sheetnames != TRANG:
        return ["sai trang: %r" % wb.sheetnames]
    loi += kiem_gioi_han(wb, path)

    with api._ket_noi_doc() as cur:
        kq = [(r[0], r[1], r[2]) for r in cur.execute(
            "SELECT student_id, buoi, club_id FROM match_results")]
        so_buoi = len(api._ds_buoi(cur))
    ds_hs = {lam_sach(s) for s, _, _ in kq}
    co_cho = {lam_sach(s) for s, _, c in kq if c}
    chua = ds_hs - co_cho
    cap = collections.Counter((lam_sach(c), lam_sach(s)) for s, _, c in kq if c)

    # Danh sach
    td, dong = _bang_duoi_khoa(wb["Danh sách học sinh"])
    if kq:
        mong = len(ds_hs) if so_buoi > 1 else len(kq)
        if len(dong) != mong:
            loi.append("Danh sach %d dong, mong %d" % (len(dong), mong))
        i = td.index("Mã học sinh")
        if {str(r[i]) for r in dong} != ds_hs:
            loi.append("Danh sach: tap ma hoc sinh lech")
    # Theo CLB
    doc, clb = doc_theo_clb(wb["Theo CLB"])
    if doc != cap:
        thieu = cap - doc
        thua = doc - cap
        loi.append("Theo CLB lech: thieu %d, thua %d (vd %s)" % (
            sum(thieu.values()), sum(thua.values()), list((thieu or thua).items())[:2]))
    if len(clb) != len(set(clb)):
        loi.append("Theo CLB: mot CLB co hai khoi")
    # Chua co cho
    _, dong = _bang_duoi_khoa(wb["Chưa có chỗ"])
    ma_chua = {str(r[0]) for r in dong if r[0] and not str(r[0]).startswith("(")}
    if kq and ma_chua != chua:
        loi.append("Chua co cho %d em, mong %d" % (len(ma_chua), len(chua)))
    # Huong dan
    hd = {r[0]: r[1] for r in wb["Hướng dẫn"].iter_rows(values_only=True) if r and r[0]}
    if kq:
        for nhan, mong in (("Số học sinh", len(ds_hs)), ("Số em chưa có chỗ", len(chua)),
                           ("Số em đã có câu lạc bộ", len(co_cho))):
            if hd.get(nhan) != mong:
                loi.append("Huong dan %s=%r, mong %r" % (nhan, hd.get(nhan), mong))
        if "Chỗ đã xếp" in hd and hd["Chỗ đã xếp"] != sum(cap.values()):
            loi.append("Huong dan Cho da xep=%r, mong %d" % (hd["Chỗ đã xếp"], sum(cap.values())))
        # Thong ke: tong cot Da xep
        dong = [list(r) for r in wb["Thống kê"].iter_rows(values_only=True)]
        try:
            h = next(n for n, r in enumerate(dong) if "Tỉ lệ chọi" in r and "Đã xếp" in r)
            c = dong[h].index("Đã xếp")
            tong = 0
            for r in dong[h + 1:]:
                if not any(r):
                    break
                tong += r[c] or 0
            if tong != sum(cap.values()):
                loi.append("Thong ke tong Da xep=%d, mong %d" % (tong, sum(cap.values())))
        except StopIteration:
            loi.append("Thong ke: khong thay bang tung CLB")
    return loi


# --------------------------------------------------------------------- #
# 1. QUY MÔ VÀ TỐC ĐỘ
# --------------------------------------------------------------------- #

def do(ham):
    """(kết quả, giây, MB đỉnh). Đo giờ và đo bộ nhớ ở HAI lần chạy riêng:
    tracemalloc làm Python chậm đi 4-5 lần, đo chung thì con số giờ sai —
    lần thử đầu báo 24 giây cho một lần xuất thật ra chỉ mất 5 giây."""
    t = time.perf_counter()
    r = ham()
    giay = time.perf_counter() - t
    tracemalloc.start()
    ham()
    _, dinh = tracemalloc.get_traced_memory()
    tracemalloc.stop()
    return r, round(giay, 3), round(dinh / 2**20, 1)


def nhom_1():
    print("\n1. QUY MÔ VÀ TỐC ĐỘ")
    luoi = [(200, 10), (1000, 10), (1000, 50), (2000, 50), (2000, 100),
            (5000, 50), (5000, 100)]
    if NHANH:
        luoi = [(200, 10), (1000, 50)]
    for n_hs, n_clb in luoi:
        for n_buoi in (1, 5):
            api = dung(n_hs, n_clb, n_buoi)
            r, giay, mb = do(lambda: api.export_ket_qua(""))
            loi = [] if r["ok"] else ["xuat hong: %r" % r["errors"][:1]]
            if r["ok"]:
                loi += kiem_so(api, r["data"]["path"])
                if giay >= 10:
                    loi.append("cham: %.1f giay" % giay)
                if mb >= 500:
                    loi.append("ton bo nho: %.0f MB" % mb)
            kb = round(os.path.getsize(r["data"]["path"]) / 1024) if r["ok"] else 0
            r2, giay_csv, mb_csv = do(lambda: api.export_ket_qua("", True))
            r3, giay_cu, _ = do(lambda: api.export_csv(""))
            if not (r2["ok"] and r3["ok"]):
                loi.append("xuat kem CSV / export_csv hong")
            ghi(1, "%d em, %d CLB, %d buổi" % (n_hs, n_clb, n_buoi), loi,
                giay_excel=giay, mb_dinh=mb, kb_tep=kb,
                giay_excel_kem_csv=giay_csv, giay_export_csv_cu=giay_cu)
            if r["ok"] and n_hs == max(x for x, _ in luoi) and n_buoi == 5:
                TEP_SOAT_LO.append(r["data"]["path"])


# --------------------------------------------------------------------- #
# 2. GIỚI HẠN CỨNG CỦA EXCEL
# --------------------------------------------------------------------- #

def nhom_2():
    print("\n2. GIỚI HẠN CỨNG CỦA EXCEL")
    api = dung(5, 2)
    sua_csdl(api, "UPDATE students SET name = ? WHERE student_id = 'S000000'", "A" * 40000)
    sua_csdl(api, "UPDATE clubs SET name = ? WHERE club_id = 'c0000'", "B" * 40000)
    r = api.export_ket_qua("")
    loi = [] if r["ok"] else ["xuat hong: %r" % r["errors"][:1]]
    if r["ok"]:
        loi += kiem_so(api, r["data"]["path"])
        TEP_SOAT_LO.append(r["data"]["path"])
    ghi(2, "tên học sinh và tên CLB dài 40.000 ký tự", loi)

    n = 300 if NHANH else 1100
    api = dung(60, n, chay=True)
    r = api.export_ket_qua("")
    loi = [] if r["ok"] else ["xuat hong: %r" % r["errors"][:1]]
    if r["ok"]:
        loi += kiem_so(api, r["data"]["path"])
        TEP_SOAT_LO.append(r["data"]["path"])
    ghi(2, "%d CLB (ngắt trang Excel tối đa %d)" % (n, NGAT_TOI_DA), loi)


# --------------------------------------------------------------------- #
# 3. DỮ LIỆU LẠ VÀ ĐỘC
# --------------------------------------------------------------------- #

def mot_ca(nhom, ten, chuan_bi, them=None, n_buoi=1, **kw):
    try:
        api = dung(kw.pop("n_hs", 12), kw.pop("n_clb", 3), n_buoi, chay=False, **kw)
        if chuan_bi:
            chuan_bi(api)
        api.run_pipeline(seed=SEED)
        r = api.export_ket_qua("", True)
        if not r["ok"]:
            ghi(nhom, ten, ["xuat hong: %s" % (r["errors"][0].get("params") or r["errors"][0])])
            return None
        loi = kiem_so(api, r["data"]["path"])
        if them:
            loi += them(api, r["data"]) or []
        TEP_SOAT_LO.append(r["data"]["path"])
        ghi(nhom, ten, loi)
        return api
    except Exception as e:      # mot ca hong khong duoc lam dung ca bo thu
        ghi(nhom, ten, ["NGOAI LE %s: %s" % (type(e).__name__, e)])
        return None


def nhom_3():
    print("\n3. DỮ LIỆU LẠ VÀ ĐỘC")
    for ch in ["\x00", "\x01", "\x0b", "\x0c", "\x1f"]:
        mot_ca(3, "ký tự điều khiển %r trong tên học sinh (qua nạp CSV)" % ch, None,
               ten_hs=["An%sBình %d" % (ch, i) for i in range(12)])
    mot_ca(3, "ký tự điều khiển trong tên CLB", lambda a: sua_csdl(
        a, "UPDATE clubs SET name = ? WHERE club_id = 'c0000'", "Bóng\x0bđá"))
    mot_ca(3, "ký tự điều khiển trong mã CLB", None,
           ma_clb=["clb\x0ba", "clb\x01b", "clb_c"])
    mot_ca(3, "ký tự điều khiển trong mã học sinh", None,
           ma_hs=["HS\x0b%02d" % i for i in range(12)])
    mot_ca(3, "ký tự điều khiển trong tên buổi", None, n_buoi=2,
           ten_buoi=["thu\x0b2", "thu_3"], n_clb=4)
    for dau in ["=", "+", "-", "@", "\t", "\r"]:
        mot_ca(3, "tên bắt đầu bằng %r (công thức)" % dau, lambda a, d=dau: (
            sua_csdl(a, "UPDATE students SET name = ? WHERE student_id = 'S000000'",
                     d + 'HYPERLINK("http://x","bam")'),
            sua_csdl(a, "UPDATE clubs SET name = ? WHERE club_id = 'c0000'", d + "1+1")))
    mot_ca(3, "emoji, NFD, chữ Ả Rập, nháy, ·, ] * trong mã và tên", None,
           ma_clb=["clb·a", "clb]b*", "clb'c\""],
           ten_clb=["⚽ Bóng đá 🏀", "Tiếng Việt NFD", "نادي القراءة"],
           ten_hs=["Nguyễn \"Văn\" O'Brien 😀 %d" % i for i in range(12)])

    def so_0(api, d):
        ws = load_workbook(d["path"])["Danh sách học sinh"]
        _, dong = _bang_duoi_khoa(ws)
        ma = {r[0] for r in dong}
        return [] if "0012345" in ma and all(isinstance(m, str) for m in ma) else [
            "mat so 0 dau hoac ma thanh so: %r" % sorted(ma)[:3]]
    mot_ca(3, "mã học sinh có số 0 đứng đầu", None, them=so_0,
           ma_hs=["0012345"] + ["%07d" % i for i in range(1, 12)])

    def trung_cot(api, d):
        ws = load_workbook(d["path"])["Danh sách học sinh"]
        td, _ = _bang_duoi_khoa(ws)
        return [] if len(td) == len(set(td)) else ["tieu de cot trung nhau: %r" % td]
    mot_ca(3, "buổi tên 'Họ tên' (trùng tên cột thời khoá biểu)", None, n_buoi=2,
           ten_buoi=["Họ tên", "thu_3"], n_clb=4, them=trung_cot)
    mot_ca(3, "tên học sinh rỗng / NULL", lambda a: (
        sua_csdl(a, "UPDATE students SET name = NULL WHERE student_id = 'S000000'"),
        sua_csdl(a, "UPDATE students SET name = '' WHERE student_id = 'S000001'")))
    mot_ca(3, "trường chỉ 1 học sinh", None, n_hs=1, n_clb=1)
    mot_ca(3, "mọi em đều có chỗ (dư chỗ)", None, ty_le=3.0)
    mot_ca(3, "không em nào có chỗ ở CLB đông (thiếu chỗ nặng)", None, n_hs=60, ty_le=0.1)


# --------------------------------------------------------------------- #
# 4. TRẠNG THÁI LẠ
# --------------------------------------------------------------------- #

def chu_huong_dan(path):
    ws = load_workbook(path)["Hướng dẫn"]
    return " ".join(str(o) for r in ws.iter_rows(values_only=True) for o in r if o).lower()


def nhom_4():
    print("\n4. TRẠNG THÁI LẠ")
    d = tempfile.mkdtemp()
    api = PipelineAPI(os.path.join(d, "app.db"))
    api.thu_muc_xuat = d
    r = api.export_ket_qua("")
    ghi(4, "CSDL trống hoàn toàn", [] if r["ok"] and load_workbook(r["data"]["path"]).sheetnames == TRANG
        else ["hong: %r" % r.get("errors")])

    api.create_or_update_club("clb_a", "CLB A", 5, 0, "")
    r = api.export_ket_qua("")
    ghi(4, "có CLB, chưa có học sinh", [] if r["ok"] else ["hong: %r" % r.get("errors")])

    api = dung(30, 4, chay=False)
    r = api.export_ket_qua("")
    loi = [] if r["ok"] else ["hong"]
    if r["ok"] and "chưa chạy" not in chu_huong_dan(r["data"]["path"]):
        loi.append("Huong dan khong noi chua chay")
    ghi(4, "có dữ liệu, chưa chạy lần nào", loi)

    api = dung(30, 4)
    api.create_or_update_club("c0001", "CLB 1 sửa", 99, 0, "")
    r = api.export_ket_qua("")
    loi = kiem_so(api, r["data"]["path"]) if r["ok"] else ["hong"]
    if r["ok"] and "chạy lại" not in chu_huong_dan(r["data"]["path"]):
        loi.append("khong canh bao ket qua cu")
    if r["ok"] and not r["data"].get("ket_qua_cu"):
        loi.append("phan hoi khong bao ket_qua_cu")
    ghi(4, "chạy rồi sửa dữ liệu (kết quả cũ)", loi)

    api = dung(80, 10, n_buoi=5)
    api.run_pipeline(seed=SEED, chi_buoi=["thu_3"])
    r = api.export_ket_qua("")
    ghi(4, "chạy lại riêng một buổi (thu_3)", kiem_so(api, r["data"]["path"]) if r["ok"] else ["hong"])

    api = dung(30, 4)
    with api._ket_noi_doc() as cur:
        cid = cur.execute("SELECT club_id FROM match_results WHERE club_id IS NOT NULL").fetchone()[0]
    xoa = api.delete_club(cid)
    r = api.export_ket_qua("")
    ghi(4, "xoá một CLB đã có người sau khi chạy (xoá %s)" % ("được" if xoa["ok"] else "bị chặn"),
        kiem_so(api, r["data"]["path"]) if r["ok"] else ["hong: %r" % r.get("errors")])

    api = dung(30, 4)
    xoa = api.delete_student("S000003")
    r = api.export_ket_qua("")
    ghi(4, "xoá một học sinh sau khi chạy (xoá %s)" % ("được" if xoa["ok"] else "bị chặn"),
        kiem_so(api, r["data"]["path"]) if r["ok"] else ["hong: %r" % r.get("errors")])


# --------------------------------------------------------------------- #
# 5. HỆ THỐNG TỆP
# --------------------------------------------------------------------- #

def nhom_5():
    print("\n5. HỆ THỐNG TỆP")
    api = dung(40, 5)
    goc_tv = api.thu_muc_xuat

    api.thu_muc_xuat = "/khong/ton/tai"
    r = api.export_ket_qua("")
    loi = [] if r["ok"] else ["hong"]
    if r["ok"] and os.path.dirname(r["data"]["path"]) != os.path.dirname(os.path.abspath(api.db_path)):
        loi.append("khong lui ve canh app.db: %s" % r["data"]["path"])
    ghi(5, "thư mục Tải xuống không tồn tại → lùi về cạnh app.db", loi)
    api.thu_muc_xuat = goc_tv

    r = api.export_ket_qua("/khong/ton/tai/kq.xlsx")
    ma = r["errors"][0]["code"] if not r["ok"] else None
    ghi(5, "đường dẫn tuyệt đối vào thư mục không có",
        [] if ma in ("error_exporting_excel", "error_exporting_csv") else ["ma loi %r" % ma])

    thu_muc_gia = os.path.join(goc_tv, "la_thu_muc.xlsx")
    os.makedirs(thu_muc_gia)
    r = api.export_ket_qua(thu_muc_gia)
    ghi(5, "đường dẫn trỏ vào một THƯ MỤC",
        [] if (not r["ok"] and os.path.isdir(thu_muc_gia)) else ["khong bao loi / xoa mat thu muc"])

    for f in os.listdir(goc_tv):
        p = os.path.join(goc_tv, f)
        shutil.rmtree(p) if os.path.isdir(p) else os.remove(p)
    duong = [api.export_ket_qua("")["data"]["path"] for _ in range(50)]
    loi = []
    if len(set(duong)) != 50 or not all(os.path.isfile(p) for p in duong):
        loi.append("trung ten hoac mat tep")
    if os.path.basename(duong[-1]) != "ket_qua_phan_bo (50).xlsx":
        loi.append("ten lan 50: %s" % os.path.basename(duong[-1]))
    ghi(5, "xuất 50 lần liên tiếp", loi)

    for f in os.listdir(goc_tv):
        os.remove(os.path.join(goc_tv, f))
    with open(os.path.join(goc_tv, "ket_qua_phan_bo.csv"), "w") as f:
        f.write("TEP CU CUA NGUOI DUNG")
    with open(os.path.join(goc_tv, "ket_qua_phan_bo (2).csv"), "w") as f:
        f.write("TEP CU CUA NGUOI DUNG")
    r = api.export_ket_qua("", True)
    loi = [] if r["ok"] else ["hong"]
    for t in ("ket_qua_phan_bo.csv", "ket_qua_phan_bo (2).csv"):
        if open(os.path.join(goc_tv, t)).read() != "TEP CU CUA NGUOI DUNG":
            loi.append("ghi de %s" % t)
    if r["ok"] and os.path.basename(r["data"]["path"]) != "ket_qua_phan_bo (3).xlsx":
        loi.append("ten moi: %s" % os.path.basename(r["data"]["path"]))
    ghi(5, "kèm CSV khi đã có .csv cũ cùng tên → không ghi đè", loi)

    api.thu_muc_xuat = os.path.join(goc_tv, "Tải xuống của Trường 😀")
    os.makedirs(api.thu_muc_xuat)
    r = api.export_ket_qua("", True)
    ghi(5, "thư mục tên tiếng Việt + emoji",
        kiem_so(api, r["data"]["path"]) if r["ok"] else ["hong: %r" % r.get("errors")])

    p = os.path.join(goc_tv, "co_san.xlsx")
    with open(p, "w") as f:
        f.write("cu")
    r = api.export_ket_qua(p)
    ghi(5, "đường dẫn tuyệt đối tới tệp có sẵn → ghi đè (như tài liệu nói)",
        kiem_so(api, p) if r["ok"] and r["data"]["path"] == p else ["hong"])

    api.thu_muc_xuat = goc_tv
    ra = []

    def xuat():
        ra.append(api.export_ket_qua("", True))
    lu = [threading.Thread(target=xuat) for _ in range(4)]
    [t.start() for t in lu]
    [t.join() for t in lu]
    loi = []
    if not all(r["ok"] for r in ra):
        loi.append("co luong hong")
    else:
        ds = [r["data"]["path"] for r in ra]
        if len(set(ds)) != 4:
            loi.append("hai luong ghi cung mot tep")
        for p in ds:
            loi += kiem_so(api, p)
    ghi(5, "4 luồng xuất cùng lúc", loi)

    api2 = dung(200, 20, n_buoi=3)
    kq = {}

    def chay():
        for s in range(5):
            api2.run_pipeline(seed=s)

    def xuat2():
        for _ in range(5):
            r = api2.export_ket_qua("")
            kq.setdefault("loi", []).extend(
                [] if not r["ok"] else kiem_so(api2, r["data"]["path"]))
            kq.setdefault("hong", 0)
            kq["hong"] += 0 if r["ok"] else 1
    a, b = threading.Thread(target=chay), threading.Thread(target=xuat2)
    a.start()
    b.start()
    a.join()
    b.join()
    # kiem_so doc CSDL SAU khi xuat; neu lan chay chen vao giua, so lech la
    # do bo kiem chu khong do tep. Chi bao loi neu tep tu mau thuan noi bo.
    ghi(5, "vừa chạy sắp xếp vừa xuất (2 luồng)",
        ["%d lan xuat hong" % kq["hong"]] if kq["hong"] else [])


# --------------------------------------------------------------------- #
# 6. NHẤT QUÁN NGẪU NHIÊN
# --------------------------------------------------------------------- #

def doc_csv(p):
    with open(p, encoding="utf-8-sig", newline="") as f:
        return list(csv.reader(f))


def so_csv_voi_excel(d):
    """Bộ .csv rời nói cùng điều với sổ Excel."""
    loi = []
    wb = load_workbook(d["path"])
    td, dong = _bang_duoi_khoa(wb["Danh sách học sinh"])
    nguon = d["thoi_khoa_bieu_path"] or d["csv_path"]
    tu_csv = doc_csv(nguon)
    chuan = [["" if o is None else str(o) for o in r] for r in [td] + dong]
    # CSV chen dau nhay truoc o bat dau bang = + - @ \t \r; bo di de so.
    tu_csv = [[o[1:] if o[:2] in ("'=", "'+", "'-", "'@", "'\t", "'\r") else o for o in r]
              for r in tu_csv]
    tu_csv = [[lam_sach(o) or "" for o in r] for r in tu_csv]
    if sorted(map(tuple, chuan)) != sorted(map(tuple, tu_csv)):
        loi.append("Danh sach (Excel) khac %s" % os.path.basename(nguon))
    so_tep = len(glob.glob(os.path.join(d["per_club_dir"], "*.csv")))
    doc, clb = doc_theo_clb(wb["Theo CLB"])
    co_nguoi = {c for c, _ in doc}
    # moi CLB co nguoi = 1 tep, cong tep _chua_duoc_xep neu co em chua xep
    if so_tep < len(co_nguoi):
        loi.append("thu muc theo CLB %d tep < %d CLB co nguoi" % (so_tep, len(co_nguoi)))
    return loi


def nhom_6():
    print("\n6. NHẤT QUÁN NGẪU NHIÊN")
    rng = random.Random(SEED)
    n_bo = 6 if NHANH else 30
    for b in range(n_bo):
        n_hs = rng.randint(20, 400)
        n_buoi = rng.choice([1, 1, 2, 3, 5])
        thap = max(2, n_buoi)
        n_clb = rng.randint(thap, max(thap, min(30, n_hs // 8)))
        ty_le = rng.uniform(0.5, 1.4)
        for seed in (1, 7, 42):
            ten = "bộ %02d: %d em, %d CLB, %d buổi, chỗ×%.2f, seed %d" % (
                b, n_hs, n_clb, n_buoi, ty_le, seed)
            try:
                api = dung(n_hs, n_clb, n_buoi, seed=seed, ty_le=ty_le)
                truoc = van_tay_csdl(api)
                r = api.export_ket_qua("", True)
                if not r["ok"]:
                    ghi(6, ten, ["xuat hong: %r" % r["errors"][:1]])
                    continue
                loi = kiem_so(api, r["data"]["path"]) + so_csv_voi_excel(r["data"])
                if van_tay_csdl(api) != truoc:
                    loi.append("xuat lam thay doi CSDL")
                ghi(6, ten, loi)
                if b < 3 and seed == 42:
                    TEP_SOAT_LO.append(r["data"]["path"])
            except Exception as e:
                ghi(6, ten, ["NGOAI LE %s: %s" % (type(e).__name__, e)])


# --------------------------------------------------------------------- #
# MỞ LẠI BẰNG LIBREOFFICE
# --------------------------------------------------------------------- #

def soat_libreoffice():
    print("\nMỞ LẠI BẰNG LIBREOFFICE (%d tệp)" % len(TEP_SOAT_LO))
    lo = shutil.which("soffice")
    if not lo:
        ghi("LO", "không có LibreOffice trên máy — bỏ qua", [])
        return
    tam = tempfile.mkdtemp(prefix="lo_")
    vao = []
    for i, p in enumerate(TEP_SOAT_LO):
        q = os.path.join(tam, "t%03d.xlsx" % i)
        shutil.copy(p, q)
        vao.append(q)
    moi_truong = dict(os.environ, HOME=os.path.join(tam, "h"))
    subprocess.run([lo, "--headless", "--norestore", "--convert-to", "csv",
                    "--outdir", os.path.join(tam, "ra")] + vao,
                   env=moi_truong, capture_output=True, timeout=1800)
    for p, q in zip(TEP_SOAT_LO, vao):
        ra = os.path.join(tam, "ra", os.path.basename(q)[:-5] + ".csv")
        ok = os.path.isfile(ra) and os.path.getsize(ra) > 0
        ghi("LO", "LibreOffice mở được %s" % os.path.relpath(p, tempfile.gettempdir())[:40],
            [] if ok else ["LibreOffice khong doc duoc"])


def main():
    t = time.perf_counter()
    for ham in (nhom_3, nhom_4, nhom_5, nhom_2, nhom_6, nhom_1):
        ham()
    soat_libreoffice()

    cot = ["nhom", "kich_ban", "ket_qua", "chi_tiet", "giay_excel", "mb_dinh", "kb_tep",
           "giay_excel_kem_csv", "giay_export_csv_cu"]
    if not NHANH:
        with open(os.path.join(THU_MUC, "ket_qua_stress_xuat.csv"), "w",
                  newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, cot, extrasaction="ignore")
            w.writeheader()
            w.writerows(KQ)

    hong = [k for k in KQ if k["ket_qua"] != "ĐẠT"]
    print("\nTỔNG: %d kịch bản, %d đạt, %d hỏng — %.0f giây"
          % (len(KQ), len(KQ) - len(hong), len(hong), time.perf_counter() - t))
    for k in hong:
        print("  HỎNG [%s] %s: %s" % (k["nhom"], k["kich_ban"], k["chi_tiet"][:200]))
    return 1 if hong else 0


if __name__ == "__main__":
    sys.exit(main())
