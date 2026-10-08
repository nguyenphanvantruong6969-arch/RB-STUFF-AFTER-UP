# -*- coding: utf-8 -*-
"""Tạo tệp Google Apps Script dựng biểu mẫu "Đăng ký CLB và thi tuyển" từ một tệp Excel.

    python tao_form.py mau --ra MAU_DE_THI.xlsx       # tệp mẫu để nhà trường điền
    python tao_form.py tao --vao MAU_DE_THI.xlsx --ra ./form_moi

`tao` ghi ra thư mục --ra:

    TAO_GOOGLE_FORM.gs        dán vào script.google.com rồi chạy taoBieuMau
    HUONG_DAN_TUNG_BUOC.md    các bước cụ thể cho đúng bộ đề này

Soát trước, ghi sau: sai một điều là không ghi tệp nào, và in đủ mọi chỗ sai.
Không tự sửa dữ liệu, không bịa đề thi.

Tệp .gs dùng đúng khuôn mẫu của mau_forms_thi_clb/tao_google_forms.py, nên
biểu mẫu tạo ra xuất dữ liệu raw theo mau_forms_thi_clb/DINH_DANG_RAW.md, và
skill xu-li-raw-forms xử lí được.
"""

import argparse
import io
import os
import re
import sys

DAY = os.path.dirname(os.path.abspath(__file__))
GOC = os.path.normpath(os.path.join(DAY, "..", "..", "..", ".."))
sys.path.insert(0, os.path.join(GOC, "mau_forms_thi_clb"))
sys.path.insert(0, os.path.join(GOC, ".claude", "skills", "sinh-du-lieu-clb", "scripts"))
# Bản đóng gói (.zip) để các tệp phụ thuộc ngay cạnh script này: ưu tiên chúng.
sys.path.insert(0, DAY)

import sinh_du_lieu as sdl  # noqa: E402  (BUOI_CHUAN, trần của phần mềm)
import tao_google_forms as tg  # noqa: E402  (khuôn mẫu .gs)

CHU = "abcd"
COT_CLB = ["buoi", "ten_buoi", "club_id", "ten_clb"]
COT_DE = ["club_id", "cau_hoi", "lua_chon_a", "lua_chon_b", "lua_chon_c", "lua_chon_d", "dap_an"]
TEN_BUOI_CHUAN = {"thu_2": "Thứ Hai", "thu_3": "Thứ Ba", "thu_4": "Thứ Tư", "thu_5": "Thứ Năm",
                  "thu_6": "Thứ Sáu", "thu_7": "Thứ Bảy", "chu_nhat": "Chủ nhật"}
# Quá số câu này thì Apps Script có thể chạy quá 6 phút và bị ngắt giữa chừng.
NGUONG_SO_CAU = 400


class LoiDuLieu(ValueError):
    pass


def _o(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


# ------------------------------------------------------------------ tệp mẫu

HUONG_DAN_DIEN = [
    "CÁCH ĐIỀN TỆP NÀY",
    "",
    "Trang Cai_dat: tiêu đề biểu mẫu, điểm mỗi câu, số CLB thi tối đa mỗi buổi (1 tới 5).",
    "",
    "Trang CLB: mỗi CLB một dòng, theo đúng thứ tự muốn hiện trên biểu mẫu.",
    "  buoi      mã buổi: thu_2, thu_3, thu_4, thu_5, thu_6, thu_7, chu_nhat",
    "  ten_buoi  tên hiện cho học sinh, ví dụ Thứ Hai",
    "  club_id   mã CLB: chữ không dấu, số, gạch dưới. Không trùng. Ví dụ clb_covua",
    "  ten_clb   tên hiện cho học sinh, ví dụ CLB Cờ vua",
    "  Mỗi buổi tối đa 10 CLB.",
    "",
    "Trang De_thi: mỗi câu hỏi một dòng.",
    "  club_id     mã CLB của câu này (phải có ở trang CLB)",
    "  cau_hoi     nội dung câu hỏi",
    "  lua_chon_a tới lua_chon_d   các lựa chọn, ít nhất 2. Để trống lựa chọn không dùng.",
    "  dap_an      chữ a, b, c hoặc d",
    "  CLB không có câu nào ở trang này là CLB không tổ chức thi: vẫn có trong bảng xếp",
    "  hạng, nhưng không có trong câu tick và không có đề.",
    "",
    "Điền xong, đưa tệp cho Claude và nói: tạo Google Form từ tệp này.",
]


def ghi_mau(duong):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Cai_dat"
    for r in (["khoa", "gia_tri"], ["tieu_de", tg.TIEU_DE_MAC_DINH],
              ["diem_moi_cau", tg.DIEM_CAU], ["toi_da_clb_thi_moi_buoi", tg.TRAN_THI]):
        ws.append(r)

    clb = wb.create_sheet("CLB")
    de = wb.create_sheet("De_thi")
    clb.append(COT_CLB)
    de.append(COT_DE)
    for d in tg.du_lieu_gs():
        for c in d["clubs"]:
            clb.append([d["ma"], d["ten"], c["id"], c["name"]])
            for q in c["de"]:
                ops = q["opts"] + [""] * (4 - len(q["opts"]))
                de.append([c["id"], q["q"]] + ops + [CHU[q["opts"].index(q["dung"])]])

    hd = wb.create_sheet("Huong_dan")
    for dong in HUONG_DAN_DIEN:
        hd.append([dong])
    for sh in (ws, clb, de):
        for o in sh[1]:
            o.font = Font(bold=True)
        sh.freeze_panes = "A2"
    for sh, rong in ((ws, [26, 40]), (clb, [10, 12, 18, 26]), (de, [18, 50, 24, 24, 24, 24, 8]),
                     (hd, [100])):
        for i, w in enumerate(rong):
            sh.column_dimensions[chr(65 + i)].width = w
    wb.save(duong)


# ------------------------------------------------------------------ đọc và soát

def doc(duong):
    import openpyxl
    wb = openpyxl.load_workbook(duong, read_only=True, data_only=True)
    loi = []

    def trang(ten, cot):
        if ten not in wb.sheetnames:
            loi.append("thiếu trang %s" % ten)
            return []
        dong = [[_o(x) for x in r] for r in wb[ten].iter_rows(values_only=True)]
        dong = [r for r in dong if any(r)]
        if not dong:
            loi.append("trang %s trống" % ten)
            return []
        td = dong[0]
        thieu = [c for c in cot if c not in td]
        if thieu:
            loi.append("trang %s thiếu cột: %s" % (ten, ", ".join(thieu)))
            return []
        return [(i + 2, dict(zip(td, r + [""] * (len(td) - len(r))))) for i, r in enumerate(dong[1:])]

    cai_dat = {r["khoa"]: r["gia_tri"] for _, r in trang("Cai_dat", ["khoa", "gia_tri"])}
    clb = trang("CLB", COT_CLB)
    de = trang("De_thi", COT_DE)
    wb.close()
    if loi:
        raise LoiDuLieu("\n  - ".join([""] + loi).lstrip("\n"))
    return cai_dat, clb, de


def soat(cai_dat, clb, de):
    """Trả (du_lieu theo dạng BUOI của .gs, tiêu đề, điểm, trần, cảnh báo mềm)."""
    loi, canh_bao = [], []

    tieu_de = cai_dat.get("tieu_de", "").strip()
    if not tieu_de:
        loi.append("Cai_dat: tieu_de đang trống")
    try:
        diem = int(float(cai_dat.get("diem_moi_cau", "")))
        if diem <= 0:
            raise ValueError
    except ValueError:
        loi.append("Cai_dat: diem_moi_cau phải là số nguyên lớn hơn 0")
        diem = 1
    try:
        tran = int(float(cai_dat.get("toi_da_clb_thi_moi_buoi", "")))
        if not 1 <= tran <= sdl.TRAN_CLB_THI_MOI_BUOI:
            raise ValueError
    except ValueError:
        loi.append("Cai_dat: toi_da_clb_thi_moi_buoi phải từ 1 tới %d (trần của phần mềm xếp CLB)"
                   % sdl.TRAN_CLB_THI_MOI_BUOI)
        tran = 1

    buoi, ten_buoi, theo_ma = [], {}, {}
    for so, r in clb:
        b, cid, ten = r["buoi"], r["club_id"], r["ten_clb"]
        noi = "CLB dòng %d" % so
        if b not in sdl.BUOI_CHUAN:
            loi.append("%s: buổi '%s' không hợp lệ, dùng một trong %s" % (noi, b, ", ".join(sdl.BUOI_CHUAN)))
            continue
        tb = r["ten_buoi"] or TEN_BUOI_CHUAN[b]
        if b in ten_buoi and ten_buoi[b] != tb:
            loi.append("%s: buổi %s đã có tên '%s', dòng này ghi '%s'" % (noi, b, ten_buoi[b], tb))
        if b not in ten_buoi:
            buoi.append(b)
            ten_buoi[b] = tb
        if not re.match(r"^[A-Za-z0-9_]+$", cid):
            loi.append("%s: club_id '%s' chỉ được có chữ không dấu, số, gạch dưới" % (noi, cid))
            continue
        if cid in theo_ma:
            loi.append("%s: club_id '%s' bị trùng" % (noi, cid))
            continue
        if not ten:
            loi.append("%s: ten_clb đang trống" % noi)
        theo_ma[cid] = {"id": cid, "name": ten or cid, "buoi": b, "de": []}

    for so, r in de:
        noi = "De_thi dòng %d" % so
        c = theo_ma.get(r["club_id"])
        if c is None:
            loi.append("%s: club_id '%s' không có ở trang CLB" % (noi, r["club_id"]))
            continue
        if not r["cau_hoi"]:
            loi.append("%s: cau_hoi đang trống" % noi)
        ops = [r["lua_chon_" + x] for x in CHU]
        while ops and not ops[-1]:
            ops.pop()
        if len(ops) < 2 or "" in ops:
            loi.append("%s: cần ít nhất 2 lựa chọn, điền liền từ lua_chon_a, không để trống ở giữa" % noi)
            continue
        if len({" ".join(o.split()).casefold() for o in ops}) != len(ops):
            loi.append("%s: có hai lựa chọn giống nhau (chấm theo nội dung nên không phân biệt được)" % noi)
            continue
        da = r["dap_an"].strip().lower()
        if da not in CHU[:len(ops)]:
            loi.append("%s: dap_an '%s' phải là một chữ trong %s" % (noi, r["dap_an"], ", ".join(CHU[:len(ops)])))
            continue
        c["de"].append({"q": r["cau_hoi"], "opts": ops, "dung": ops[CHU.index(da)]})

    du_lieu = []
    for b in buoi:
        cs = [c for c in theo_ma.values() if c["buoi"] == b]
        if len(cs) > sdl.TRAN_NGUYEN_VONG_MOI_BUOI:
            loi.append("buổi %s có %d CLB, tối đa %d (trần nguyện vọng của phần mềm)"
                       % (b, len(cs), sdl.TRAN_NGUYEN_VONG_MOI_BUOI))
        if not any(c["de"] for c in cs):
            canh_bao.append("buổi %s không có CLB nào tổ chức thi" % ten_buoi[b])
        du_lieu.append({"ma": b, "ten": ten_buoi[b], "clubs": [
            {"id": c["id"], "name": c["name"], "de": c["de"]} for c in cs]})
    if not du_lieu:
        loi.append("trang CLB chưa có CLB nào hợp lệ")

    khong_thi = [c["name"] for c in theo_ma.values() if not c["de"]]
    if khong_thi:
        canh_bao.append("%d CLB không có đề, coi là không tổ chức thi: %s"
                        % (len(khong_thi), ", ".join(khong_thi)))
    so_cau = dem_cau(du_lieu)
    if so_cau > NGUONG_SO_CAU:
        canh_bao.append("biểu mẫu có %d câu hỏi; trên %d câu thì Apps Script có thể chạy quá 6 phút. "
                        "Nếu bị ngắt, chia thành hai biểu mẫu" % (so_cau, NGUONG_SO_CAU))
    if loi:
        raise LoiDuLieu("\n  - ".join([""] + loi).lstrip("\n"))
    return du_lieu, tieu_de, diem, tran, canh_bao


def dem_cau(du_lieu):
    """Số câu hỏi đúng như taoBieuMau tạo."""
    n = 2
    for d in du_lieu:
        thi = [c for c in d["clubs"] if c["de"]]
        n += 1 + len(d["clubs"])  # câu đăng ký buổi, mỗi Top một câu thả xuống
        if thi:
            n += 2 + len(thi) - 1  # câu tick, "thi CLB nào trước", các trang "tiếp theo"
            n += sum(len(c["de"]) for c in thi)
    return n


# ------------------------------------------------------------------ hướng dẫn

def huong_dan(du_lieu, tieu_de, diem, tran, canh_bao):
    n_clb = sum(len(d["clubs"]) for d in du_lieu)
    n_thi = sum(1 for d in du_lieu for c in d["clubs"] if c["de"])
    n_de = sum(len(c["de"]) for d in du_lieu for c in d["clubs"])
    vd = next((d for d in du_lieu if len([c for c in d["clubs"] if c["de"]]) >= 2), None)
    L = []
    w = L.append
    w("# Hướng dẫn từng bước: biểu mẫu \"%s\"" % tieu_de)
    w("")
    w("Biểu mẫu này có **%d buổi, %d CLB, %d CLB tổ chức thi, %d câu đề thi** (mỗi câu %d điểm), "
      "tổng **%d câu hỏi**. Mỗi buổi học sinh thi tối đa **%d CLB**."
      % (len(du_lieu), n_clb, n_thi, n_de, diem, dem_cau(du_lieu), tran))
    w("")
    w("Các buổi: " + "; ".join("%s (%d CLB)" % (d["ten"], len(d["clubs"])) for d in du_lieu) + ".")
    w("")
    if canh_bao:
        w("**Lưu ý trước khi làm:**")
        w("")
        for c in canh_bao:
            w("- " + c)
        w("")
    w("## Bước 1. Tạo biểu mẫu (làm một lần, khoảng 5 phút)")
    w("")
    w("1. Mở **script.google.com** bằng tài khoản Google của trường, bấm **Dự án mới**.")
    w("   Cách khác: mở một biểu mẫu Google Forms trống, bấm dấu ba chấm góc trên bên phải, chọn **Apps Script**.")
    w("2. Xoá hết chữ có sẵn trong khung soạn thảo.")
    w("3. Mở tệp `TAO_GOOGLE_FORM.gs`, chọn tất cả (Ctrl + A), sao chép (Ctrl + C), dán vào khung soạn thảo (Ctrl + V).")
    w("4. Bấm biểu tượng đĩa mềm để lưu. Đổi tên dự án (bấm vào tên ở góc trên bên trái), ví dụ \"Đăng ký CLB\".")
    w("5. Ở thanh trên, chọn hàm **taoBieuMau**, bấm **Chạy**.")
    w("6. Google hỏi quyền: bấm **Xem lại quyền**, chọn tài khoản.")
    w("7. Nếu hiện **\"Google chưa xác minh ứng dụng này\"**: bấm **Nâng cao**, rồi **Đi tới ... (không an toàn)**. "
      "Cảnh báo này hiện vì đoạn mã do chính em tạo, chưa qua Google duyệt.")
    w("8. Trang chọn quyền: tick **Chọn tất cả** (Drive, Trang tính, Biểu mẫu, Kết nối dịch vụ bên ngoài), bấm **Tiếp tục**.")
    w("9. Đợi 1 tới 3 phút. Mở **Nhật ký thực thi** ở dưới, chép hai đường dẫn: đường dẫn **SỬA** biểu mẫu "
      "và đường dẫn gửi **HỌC SINH**.")
    w("")
    w("Chạy taoBieuMau lần nữa là tạo thêm một biểu mẫu mới. Chỉ chạy một lần.")
    w("")
    w("## Bước 2. Cài đặt và làm thử (làm một lần)")
    w("")
    w("Mở đường dẫn SỬA biểu mẫu:")
    w("")
    w("1. Nếu nút **Xuất bản** ở góc trên còn màu tím, bấm để xuất bản.")
    w("2. Tab **Cài đặt** → **Bài kiểm tra**: mục **Công bố điểm** chọn **Sau khi xem xét thủ công**. "
      "Bỏ tick \"Câu hỏi bị trả lời sai\" và \"Câu trả lời đúng\".")
    w("3. Tab **Cài đặt** → **Câu trả lời**: bật **Giới hạn 1 câu trả lời**.")
    if vd:
        thi = [c for c in vd["clubs"] if c["de"]]
        w("4. Bấm biểu tượng con mắt (Xem trước) và làm thử: ở **%s** chọn Có, chọn vài CLB ở Top 1, Top 2 (các Top còn lại để (Bỏ trống)), "
          "tick **%s** và **%s**, chọn làm %s trước, rồi %s, rồi \"Em không thi thêm CLB nào buổi này\". "
          "Kiểm tra chỉ hiện đề của đúng hai CLB đó. Các buổi khác chọn \"Không, em bận buổi này\"."
          % (vd["ten"], thi[0]["name"], thi[1]["name"], thi[0]["name"], thi[1]["name"]))
    else:
        w("4. Bấm biểu tượng con mắt (Xem trước) và làm thử một lượt.")
    w("")
    w("## Bước 3. Gửi cho học sinh")
    w("")
    w("1. Gửi đường dẫn HỌC SINH qua nhóm lớp, email hoặc mã QR.")
    w("2. Nhắc học sinh: buổi bận thì chọn \"Không\"; bấm vào Top 1 và chọn CLB thích nhất trong danh sách, rồi Top 2, ...; "
      "mỗi CLB chỉ chọn một lần, Top không dùng thì để (Bỏ trống); "
      "tick tối đa %d CLB muốn thi; chọn làm bài lần lượt theo thứ tự trong danh sách. "
      "Chỉ bài của CLB đã tick và đã chọn ở một Top mới được tính điểm." % tran)
    w("3. Hết hạn: tab **Câu trả lời**, tắt **Chấp nhận câu trả lời**.")
    w("")
    w("## Bước 4. Chấm điểm và xuất Excel")
    w("")
    w("1. Mở lại dự án Apps Script ở Bước 1. Chọn hàm **chamTheoCLB**, bấm **Chạy**.")
    w("2. Nhật ký thực thi hiện đường dẫn **Thư mục kết quả** trên Google Drive, gồm:")
    w("   - `BANG_DIEM.xlsx`: điểm từng CLB và trang cảnh báo.")
    w("   - `SO_NHAP_CLB.xlsx`: Sổ nhập CLB, tệp DUY NHẤT nạp vào phần mềm.")
    w("   - `RAW_PHIEU.xlsx`: mọi câu trả lời, đã chấm từng câu. `DAP_AN.xlsx`: đáp án.")
    w("3. Bấm chuột phải vào thư mục, chọn **Tải xuống**, giải nén.")
    w("")
    w("## Bước 5. Nạp vào phần mềm xếp CLB")
    w("")
    w("1. Mở `SO_NHAP_CLB.xlsx`, trang **1. CLB**: điền cột **Chỉ tiêu** (lớn hơn 0) cho mọi CLB. Lưu lại.")
    w("2. Cách nhanh: kéo sổ vào phần mềm, xem tóm tắt, bấm **Nhập sổ**. Đọc hết cảnh báo sau khi nạp.")
    w("3. Cách kỹ hơn: đưa thư mục kết quả và sổ đã điền chỉ tiêu cho Claude, nói **\"xử lí raw\"**. "
      "Skill `xu-li-raw-forms` chấm lại, soát bất thường, rồi ra một sổ nhập mới và một báo cáo.")
    w("")
    w("## Lỗi thường gặp")
    w("")
    w("| Hiện tượng | Cách xử lí |")
    w("|---|---|")
    w("| \"Chưa biết chấm biểu mẫu nào\" | Dán đường dẫn SỬA biểu mẫu vào dòng `const LINK_BIEU_MAU = '...'` ở đầu mã, lưu, chạy lại chamTheoCLB |")
    w("| \"Không xuất được ... (mã lỗi ...)\" | Mạng hoặc Google bận. Chạy lại chamTheoCLB |")
    w("| Chạy quá 6 phút, bị ngắt | Biểu mẫu quá dài. Xoá biểu mẫu dở dang, chia đề thành hai tệp Excel, tạo hai biểu mẫu |")
    w("| Có hai biểu mẫu giống nhau | Đã chạy taoBieuMau hai lần. Xoá một cái, chỉ gửi học sinh một đường dẫn |")
    w("| Học sinh không mở được | Biểu mẫu chưa xuất bản, hoặc đã tắt nhận câu trả lời, hoặc học sinh chưa đăng nhập Google |")
    w("")
    return "\n".join(L) + "\n"


# ------------------------------------------------------------------ chạy

def tao(vao, ra):
    cai_dat, clb, de = doc(vao)
    du_lieu, tieu_de, diem, tran, canh_bao = soat(cai_dat, clb, de)
    os.makedirs(ra, exist_ok=True)
    with io.open(os.path.join(ra, "TAO_GOOGLE_FORM.gs"), "w", encoding="utf-8") as f:
        f.write(tg.noi_dung(du_lieu, tieu_de=tieu_de, diem=diem, tran=tran))
    with io.open(os.path.join(ra, "HUONG_DAN_TUNG_BUOC.md"), "w", encoding="utf-8") as f:
        f.write(huong_dan(du_lieu, tieu_de, diem, tran, canh_bao))
    return du_lieu, canh_bao


def main(argv=None):
    # Cua so lenh / ong dan tren Windows mac dinh la cp1252: in tieng Viet
    # se nem UnicodeEncodeError va script thoat voi ma 1.
    for _luong in (sys.stdout, sys.stderr):
        if hasattr(_luong, "reconfigure"):
            _luong.reconfigure(encoding="utf-8")
    p = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    sub = p.add_subparsers(dest="lenh", required=True)
    m = sub.add_parser("mau", help="ghi tệp Excel mẫu (điền sẵn bộ đề ví dụ)")
    m.add_argument("--ra", required=True)
    t = sub.add_parser("tao", help="đọc tệp Excel, ghi TAO_GOOGLE_FORM.gs và hướng dẫn")
    t.add_argument("--vao", required=True)
    t.add_argument("--ra", required=True)
    a = p.parse_args(argv)
    if a.lenh == "mau":
        ghi_mau(a.ra)
        print("Đã ghi tệp mẫu %s. Điền trang CLB và De_thi rồi chạy lệnh tao." % a.ra)
        return 0
    try:
        du_lieu, canh_bao = tao(a.vao, a.ra)
    except LoiDuLieu as e:
        print("KHÔNG GHI TỆP NÀO. Tệp Excel có chỗ sai:\n%s" % e, file=sys.stderr)
        return 1
    print("Đã ghi TAO_GOOGLE_FORM.gs và HUONG_DAN_TUNG_BUOC.md vào %s" % a.ra)
    print("%d buổi, %d CLB, %d câu hỏi." % (len(du_lieu), sum(len(d["clubs"]) for d in du_lieu),
                                          dem_cau(du_lieu)))
    for c in canh_bao:
        print("Lưu ý: " + c)
    return 0


if __name__ == "__main__":
    sys.exit(main())
