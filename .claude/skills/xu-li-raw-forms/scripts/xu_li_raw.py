# -*- coding: utf-8 -*-
"""Xử lí dữ liệu raw của biểu mẫu Google Forms thành Sổ nhập CLB của phần mềm xếp CLB.

    python xu_li_raw.py --raw RAW_PHIEU.xlsx --dap-an DAP_AN.xlsx \\
                        --clb SO_NHAP_CLB.xlsx --ra ./nap

Đầu vào theo định dạng mau_forms_thi_clb/DINH_DANG_RAW.md (do hàm chamTheoCLB
của TAO_GOOGLE_FORM.gs xuất ra). Đầu ra:

    SO_NHAP_CLB.xlsx         MỘT tệp kéo thẳng vào phần mềm: sheet "1. CLB"
                             (chỉ tiêu lấy từ --clb; không có --clb thì cột Chỉ
                             tiêu để trống cho nhà trường điền), sheet
                             "2. Học sinh" (nguyện vọng, điểm chấm lại theo đáp án)
    BAO_CAO_BAT_THUONG.md    mọi điều bất thường, kèm mã học sinh

NGUYÊN TẮC (giống skill sinh-du-lieu-clb):
- Không tin mù quáng các cột đã tính sẵn trong raw. Tự chấm lại, tự quyết lại
  phiếu nào giữ, bài nào tính. Chỗ nào khác với raw thì báo.
- Ô sai lẻ (một bài không tick, một phiếu trùng): bỏ riêng ô đó và báo.
- Đầu vào hỏng (thiếu cột, mã CLB không có trong danh sách...): không ghi tệp nào.
"""

import argparse
import csv
import io
import os
import re
import sys
from collections import Counter, OrderedDict

DAY = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(DAY, "..", "..", "sinh-du-lieu-clb", "scripts"))
# Bản đóng gói (.zip) để sinh_du_lieu.py ngay cạnh script này: ưu tiên nó.
sys.path.insert(0, DAY)

import sinh_du_lieu as sdl  # noqa: E402  (soat, ghi_bo, TEN_SO, trần)
import so_nhap  # noqa: E402  (sinh_du_lieu đã đặt sẵn đường tìm)

COT_RAW = ["phieu_so", "thoi_gian_nop", "email", "student_id", "trang_thai_phieu",
           "ma_cau", "loai_cau", "buoi", "club_id", "cau_so", "cau_hoi", "tra_loi",
           "dap_an_dung", "dung_sai", "diem", "diem_toi_da", "diem_forms", "duoc_tinh", "ghi_chu"]
COT_DAP_AN = ["buoi", "ten_buoi", "club_id", "ten_clb", "ma_cau", "cau_so", "cau_hoi",
              "dap_an_dung", "diem"]
CO_BUOI = "Có"
CO_THI = "Có, em thi CLB này"
TRAN_THI = sdl.TRAN_CLB_THI_MOI_BUOI


class LoiDauVao(ValueError):
    """Đầu vào không dùng được. Không ghi tệp nào."""


# ------------------------------------------------------------------- đọc tệp

def _o(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def doc_bang(duong):
    """Đọc .xlsx (trang đầu) hoặc .csv thành [dict]."""
    if duong.lower().endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(duong, read_only=True, data_only=True)
        dong = [[_o(x) for x in r] for r in wb[wb.sheetnames[0]].iter_rows(values_only=True)]
        wb.close()
    else:
        with io.open(duong, encoding="utf-8-sig", newline="") as f:
            dong = [[_o(x) for x in r] for r in csv.reader(f)]
    dong = [r for r in dong if any(r)]
    if not dong:
        raise LoiDauVao("%s: tệp trống" % os.path.basename(duong))
    tieu_de = dong[0]
    return tieu_de, [dict(zip(tieu_de, r + [""] * (len(tieu_de) - len(r)))) for r in dong[1:]]


def _can_cot(ten, tieu_de, can):
    thieu = [c for c in can if c not in tieu_de]
    if thieu:
        raise LoiDauVao("%s thiếu cột: %s" % (ten, ", ".join(thieu)))


def chuan(s):
    return " ".join(str(s or "").split()).casefold()


def so(s, mac_dinh=0):
    try:
        return int(float(s))
    except (TypeError, ValueError):
        return mac_dinh


# ------------------------------------------------------------------ xử lí

def xu_li(raw, dap_an=None, clb_truong=None):
    """Trả (clb, hs, bao_cao). Không ghi gì.

    raw:        [dict] theo COT_RAW
    dap_an:     [dict] theo COT_DAP_AN, hoặc None
    clb_truong: [dict] danh sách CLB có chỉ tiêu, hoặc None
    """
    bc = OrderedDict((k, []) for k in (
        "phieu_trung", "phieu_khong_ma", "ma_sai_dang", "email_hai_ma",
        "dap_an_forms_lech", "diem_raw_lech", "lam_khong_tick", "tick_khong_lam",
        "qua_tran", "thieu_xep_hang", "trung_hang", "bo_trong_top", "clb_nhieu_top", "bai_khong_hop_le",
        "cau_la", "khong_dang_ky_buoi", "ghi_chu"))

    # ---- đáp án và danh sách CLB theo đáp án
    khoa, diem_cau, thu_tu_clb, ten_clb, buoi_clb = {}, {}, [], {}, {}
    if dap_an:
        for d in dap_an:
            khoa[d["ma_cau"]] = d["dap_an_dung"]
            diem_cau[d["ma_cau"]] = so(d.get("diem"), 0)
            if d["club_id"] not in ten_clb:
                thu_tu_clb.append(d["club_id"])
                ten_clb[d["club_id"]] = d.get("ten_clb") or d["club_id"]
                buoi_clb[d["club_id"]] = d["buoi"]
    else:
        bc["ghi_chu"].append("Không có DAP_AN.xlsx: dùng cột diem có sẵn trong raw, không chấm lại.")

    # Bổ sung CLB chỉ thấy trong raw (CLB không có đề thi, hoặc thiếu đáp án).
    for r in raw:
        cid, b = r["club_id"], r["buoi"]
        if cid and r["loai_cau"] in ("xep_hang", "chon_thi", "cong_thi", "de_thi"):
            if cid not in buoi_clb:
                thu_tu_clb.append(cid)
                buoi_clb[cid] = b
            if cid not in ten_clb or ten_clb[cid] == cid:
                m = re.match(r"^(.*)\(([A-Za-z0-9_]+)\)\s*$", r["tra_loi"])
                ten_clb[cid] = m.group(1).strip() if (m and r["loai_cau"] == "chon_thi") else ten_clb.get(cid, cid)
    thu_tu = {c: i for i, c in enumerate(thu_tu_clb)}

    # ---- gom theo phiếu
    phieu = OrderedDict()
    for r in raw:
        phieu.setdefault(so(r["phieu_so"]), []).append(r)
    phieu = OrderedDict(sorted(phieu.items()))

    # ---- chọn phiếu giữ: phiếu đầu của mỗi mã
    giu, dau_tien, email_ma = OrderedDict(), {}, {}
    for ps, rows in phieu.items():
        sid = next((r["student_id"] for r in rows if r["student_id"]), "")
        email = next((r["email"] for r in rows if r["email"]), "")
        if email and sid:
            email_ma.setdefault(email.lower(), set()).add(sid)
        if not sid:
            bc["phieu_khong_ma"].append("Phiếu %d" % ps)
            continue
        if sid in dau_tien:
            bc["phieu_trung"].append("%s: phiếu %d bị bỏ, giữ phiếu %d" % (sid, ps, dau_tien[sid]))
            continue
        dau_tien[sid] = ps
        giu[ps] = (sid, rows)
    for e, ds in email_ma.items():
        if len(ds) > 1:
            bc["email_hai_ma"].append("%s dùng %d mã: %s" % (e, len(ds), ", ".join(sorted(ds))))

    # ---- mã học sinh sai dạng: so với dạng chiếm đa số
    dang = lambda s: re.sub(r"[0-9]", "9", re.sub(r"[A-Za-zÀ-ỹ]", "A", s))  # noqa: E731
    cac_ma = [sid for sid, _ in giu.values()]
    if cac_ma:
        pho_bien = Counter(dang(s) for s in cac_ma).most_common(1)[0][0]
        for s in cac_ma:
            if dang(s) != pho_bien or s != s.strip() or " " in s:
                bc["ma_sai_dang"].append("%s (đa số có dạng %s)" % (s, pho_bien))

    # ---- chấm và dựng từng học sinh
    lech_forms, lech_raw = Counter(), Counter()
    hs = []
    for ps, (sid, rows) in giu.items():
        ten = next((r["tra_loi"] for r in rows if r["ma_cau"] == "name"), "")
        dang_ky = {r["buoi"] for r in rows if r["loai_cau"] == "dang_ky_buoi" and r["tra_loi"] == CO_BUOI}
        if not dang_ky:
            bc["khong_dang_ky_buoi"].append(sid)

        # nguyện vọng
        nv = {}
        for b in sorted(dang_ky):
            hang = [(so(re.sub(r"\D", "", r["tra_loi"]), 0), r["club_id"]) for r in rows
                    if r["loai_cau"] == "xep_hang" and r["buoi"] == b and r["tra_loi"]]
            if not hang:
                bc["thieu_xep_hang"].append("%s · %s" % (sid, b))
                continue
            dem = Counter(h for h, _ in hang)
            if any(v > 1 for v in dem.values()):
                bc["trung_hang"].append("%s · %s" % (sid, b))
            if sorted(dem) != list(range(1, len(dem) + 1)):
                bc["bo_trong_top"].append("%s · %s: đã chọn %s" % (
                    sid, b, ", ".join("Top %d" % h for h in sorted(dem))))
            # Một CLB chọn ở nhiều Top (danh sách thả xuống không chặn được): giữ Top cao nhất.
            ds, top_cua = [], {}
            for h, c in sorted(hang, key=lambda x: (x[0], thu_tu.get(x[1], 99))):
                if c in top_cua:
                    bc["clb_nhieu_top"].append("%s · %s: %s ở Top %d và Top %d, giữ Top %d"
                                               % (sid, b, c, top_cua[c], h, top_cua[c]))
                    continue
                top_cua[c] = h
                ds.append(c)
            nv[b] = ds

        tick = {r["club_id"] for r in rows if r["loai_cau"] == "chon_thi" and r["club_id"]}
        lam = OrderedDict()
        for r in rows:
            if r["loai_cau"] != "de_thi":
                if r["loai_cau"] == "khac":
                    bc["cau_la"].append("%s: [%s]" % (sid, r["ma_cau"]))
                continue
            ma = r["ma_cau"]
            if khoa:
                if ma not in khoa:
                    bc["bai_khong_hop_le"].append("%s: câu %s không có trong đáp án" % (sid, ma))
                    continue
                diem = diem_cau[ma] if chuan(r["tra_loi"]) == chuan(khoa[ma]) else 0
                if r["diem"] != "" and so(r["diem"], -1) != diem:
                    lech_raw[ma] += 1
            else:
                diem = so(r["diem"], 0)
            if r["diem_forms"] != "" and so(r["diem_forms"], -1) != diem:
                lech_forms[ma] += 1
            lam.setdefault(r["club_id"], 0)
            lam[r["club_id"]] += diem

        diem_hs, dem_buoi = {}, Counter()
        for cid in sorted(lam, key=lambda c: thu_tu.get(c, 99)):
            b = buoi_clb.get(cid, "")
            if cid not in tick:
                bc["lam_khong_tick"].append("%s: %s" % (sid, cid))
                continue
            if b not in nv or cid not in nv[b]:
                bc["bai_khong_hop_le"].append("%s: %s không nằm trong nguyện vọng %s" % (sid, cid, b or "?"))
                continue
            if dem_buoi[b] >= TRAN_THI:
                bc["qua_tran"].append("%s · %s: bỏ bài %s" % (sid, b, cid))
                continue
            dem_buoi[b] += 1
            diem_hs[cid] = lam[cid]
        for cid in sorted(tick - set(lam), key=lambda c: thu_tu.get(c, 99)):
            bc["tick_khong_lam"].append("%s: %s" % (sid, cid))

        hs.append({"student_id": sid, "name": ten, "reserve_group": "",
                   "nguyen_vong": nv, "diem": diem_hs})

    for ma, n in sorted(lech_forms.items()):
        bc["dap_an_forms_lech"].append("%s: %d bài Forms chấm khác đáp án gốc" % (ma, n))
    for ma, n in sorted(lech_raw.items()):
        bc["diem_raw_lech"].append("%s: %d bài điểm trong raw khác khi chấm lại" % (ma, n))

    # ---- danh sách CLB
    if clb_truong is not None:
        clb = clb_truong
        co = {c["club_id"] for c in clb}
        thieu = [c for c in thu_tu_clb if c not in co]
        if thieu:
            raise LoiDauVao("Tệp danh sách CLB thiếu %d CLB có trong raw: %s"
                            % (len(thieu), ", ".join(thieu)))
        for c in clb:
            if c["club_id"] in buoi_clb and (c.get("buoi") or "") != buoi_clb[c["club_id"]]:
                raise LoiDauVao("%s: tệp danh sách CLB ghi buổi %s, biểu mẫu ghi buổi %s"
                                % (c["club_id"], c.get("buoi") or "(trống)", buoi_clb[c["club_id"]]))
    else:
        clb = [{"club_id": c, "name": ten_clb.get(c, c), "capacity": "", "reserve_capacity": 0,
                "reserve_group": "", "buoi": buoi_clb.get(c, "")} for c in thu_tu_clb]
    return clb, hs, bc


# ------------------------------------------------------------------ báo cáo

TIEU_DE_BC = OrderedDict([
    ("phieu_trung", "Phiếu nộp trùng (giữ phiếu đầu)"),
    ("phieu_khong_ma", "Phiếu không có mã học sinh (bỏ cả phiếu)"),
    ("ma_sai_dang", "Mã học sinh khác dạng với số đông (kiểm tra có gõ sai không)"),
    ("email_hai_ma", "Một email nộp bằng nhiều mã học sinh"),
    ("dap_an_forms_lech", "Đáp án trên Google Forms lệch đáp án gốc (sửa đáp án trên Forms)"),
    ("diem_raw_lech", "Điểm trong raw khác khi chấm lại theo DAP_AN.xlsx"),
    ("lam_khong_tick", "Làm bài mà không tick CLB đó (bài không được tính)"),
    ("tick_khong_lam", "Tick CLB mà không làm bài"),
    ("qua_tran", "Quá %d CLB trong một buổi (bài thừa bị bỏ)" % TRAN_THI),
    ("thieu_xep_hang", "Đăng ký buổi mà không có xếp hạng (bỏ nguyện vọng buổi đó)"),
    ("trung_hang", "Hai CLB cùng một hạng trong một buổi"),
    ("bo_trong_top", "Bỏ trống một Top ở giữa (đã dồn nguyện vọng lên)"),
    ("clb_nhieu_top", "Chọn một CLB ở nhiều Top (giữ Top cao nhất)"),
    ("bai_khong_hop_le", "Bài thi không hợp lệ (bị bỏ)"),
    ("cau_la", "Câu hỏi lạ, không thuộc mẫu (biểu mẫu có thể bị sửa tay)"),
    ("khong_dang_ky_buoi", "Học sinh không đăng ký buổi nào"),
    ("ghi_chu", "Ghi chú"),
])


def viet_bao_cao(clb, hs, bc, co_clb):
    L = ["# Báo cáo xử lí dữ liệu raw", ""]
    n_thi = sum(len(h["diem"]) for h in hs)
    L += ["- Học sinh được xuất: **%d**" % len(hs),
          "- Bài thi được tính điểm: **%d**" % n_thi,
          "- Số loại bất thường có mặt: **%d**" % sum(1 for v in bc.values() if v), ""]
    if not co_clb:
        L += ["> **Chưa có chỉ tiêu CLB.** Trong SO_NHAP_CLB.xlsx, cột Chỉ tiêu của sheet "
              "\"1. CLB\" đang trống. Điền chỉ tiêu (số lớn hơn 0) cho mọi CLB trước khi nạp.", ""]
    for k, ten in TIEU_DE_BC.items():
        ds = bc[k]
        if not ds:
            continue
        L += ["## %s: %d" % (ten, len(ds)), ""]
        L += ["- %s" % x for x in ds[:200]]
        if len(ds) > 200:
            L.append("- ... và %d mục nữa" % (len(ds) - 200))
        L.append("")
    if not any(bc.values()):
        L += ["Không có điều gì bất thường.", ""]
    L += ["## Nạp vào phần mềm", "",
          "Kéo SO_NHAP_CLB.xlsx vào ô nạp ở màn hình Vận hành, xem phần tóm tắt "
          "rồi bấm Nhập sổ.", "",
          "Sau khi nạp, đọc hết cảnh báo của phần mềm."]
    return "\n".join(L) + "\n"


# ------------------------------------------------------------------ chạy

def doc_clb_tu_so(duong):
    """Đọc sheet "1. CLB" của một Sổ nhập CLB. None nếu tệp không phải sổ."""
    if not str(duong).lower().endswith((".xlsx", ".xlsm")):
        return None
    import openpyxl

    wb = openpyxl.load_workbook(duong, read_only=True, data_only=True)
    try:
        if not so_nhap.la_so_nhap(wb.sheetnames):
            return None
        du = so_nhap.doc_so_nhap(wb)
    finally:
        wb.close()
    loi = [e for e in du["loi"] if e["params"].get("sheet") == so_nhap.SHEET_CLB]
    if loi:
        import i18n_errors
        raise LoiDauVao("Sheet \"1. CLB\" của --clb có lỗi:\n  - "
                        + "\n  - ".join(i18n_errors.format_all(loi)))
    return [{"club_id": c["club_id"], "name": c["name"], "capacity": str(c["capacity"]),
             "reserve_capacity": str(c["reserve_capacity"]),
             "reserve_group": c["reserve_group"], "buoi": c["buoi"]} for c in du["clubs"]]


def chay(duong_raw, thu_muc_ra, duong_dap_an=None, duong_clb=None):
    td, raw = doc_bang(duong_raw)
    _can_cot("RAW_PHIEU", td, COT_RAW)
    dap_an = None
    if duong_dap_an:
        td, dap_an = doc_bang(duong_dap_an)
        _can_cot("DAP_AN", td, COT_DAP_AN)
    clb_truong = None
    if duong_clb:
        # Nhận sổ nhập (sheet "1. CLB" đã điền chỉ tiêu), hoặc tệp danh sách
        # CLB kiểu cũ (club_id, name, capacity...).
        clb_truong = doc_clb_tu_so(duong_clb)
    if duong_clb and clb_truong is None:
        td, clb_truong = doc_bang(duong_clb)
        _can_cot("Danh sách CLB", td, ["club_id", "name", "capacity"])
        for c in clb_truong:
            c.setdefault("reserve_capacity", "0")
            c.setdefault("reserve_group", "")
            c.setdefault("buoi", "")
            c["capacity"] = str(so(c["capacity"], c["capacity"]))
            c["reserve_capacity"] = str(so(c["reserve_capacity"], 0))

    clb, hs, bc = xu_li(raw, dap_an, clb_truong)

    # Soát bằng chính bộ soát của skill sinh-du-lieu-clb. Chưa có chỉ tiêu thì
    # soát với chỉ tiêu tạm 1, rồi ghi sổ với cột Chỉ tiêu trống.
    clb_soat = clb if clb_truong is not None else [dict(c, capacity=1) for c in clb]
    loi = sdl.soat(clb_soat, hs)
    if loi:
        raise LoiDauVao("Dữ liệu sau xử lí chưa hợp lệ:\n  - " + "\n  - ".join(loi))

    os.makedirs(thu_muc_ra, exist_ok=True)
    clubs, students = sdl.thanh_so(clb_soat, hs)
    if clb_truong is None:
        for c in clubs:
            c["capacity"] = ""
    # co_nhom=False: biểu mẫu không hỏi nhóm ưu tiên. Không có cột thì phần
    # mềm GIỮ nhóm đã gán; có cột mà trống thì phần mềm sẽ BỎ nhóm của em.
    so_nhap.ghi_so_nhap(os.path.join(thu_muc_ra, sdl.TEN_SO), clubs, students,
                        co_nhom=False)
    with io.open(os.path.join(thu_muc_ra, "BAO_CAO_BAT_THUONG.md"), "w", encoding="utf-8") as f:
        f.write(viet_bao_cao(clb, hs, bc, clb_truong is not None))
    return clb, hs, bc


def main(argv=None):
    # Cua so lenh / ong dan tren Windows mac dinh la cp1252: in tieng Viet
    # se nem UnicodeEncodeError va script thoat voi ma 1.
    for _luong in (sys.stdout, sys.stderr):
        if hasattr(_luong, "reconfigure"):
            _luong.reconfigure(encoding="utf-8")
    p = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    p.add_argument("--raw", required=True, help="RAW_PHIEU.xlsx (hoặc .csv)")
    p.add_argument("--dap-an", help="DAP_AN.xlsx; có thì chấm lại theo đáp án gốc")
    p.add_argument("--clb", help="SO_NHAP_CLB.xlsx (sheet 1. CLB) hoặc danh sách CLB "
                                 "(.xlsx/.csv) đã điền chỉ tiêu")
    p.add_argument("--ra", required=True, help="thư mục ghi kết quả")
    a = p.parse_args(argv)
    try:
        clb, hs, bc = chay(a.raw, a.ra, a.dap_an, a.clb)
    except LoiDauVao as e:
        print("KHÔNG GHI TỆP NÀO. %s" % e, file=sys.stderr)
        return 1
    print("Đã xuất %d học sinh, %d bài thi -> %s"
          % (len(hs), sum(len(h["diem"]) for h in hs), a.ra))
    for k, ten in TIEU_DE_BC.items():
        if bc[k]:
            print("  %s: %d" % (ten, len(bc[k])))
    if not a.clb:
        print("CHƯA CÓ CHỈ TIÊU: điền cột Chỉ tiêu ở sheet \"1. CLB\" của SO_NHAP_CLB.xlsx trước khi nạp.")
    print("Chi tiết: BAO_CAO_BAT_THUONG.md")
    return 0


if __name__ == "__main__":
    sys.exit(main())
