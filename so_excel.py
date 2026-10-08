"""Viết một trang tính Excel cho NGƯỜI ĐỌC — tiêu đề, bảng, in ấn.

Dùng cho sổ kết quả (`_ghi_so_excel` trong api_xuat.py). Tách khỏi phần lấy
dữ liệu để mỗi trang chỉ còn là "trang này có những bảng nào", không lẫn
chuyện tô màu, giãn cột hay khổ giấy.

Mọi ô đều đi qua `Trang.o`, nơi DUY NHẤT chặn công thức: openpyxl coi mọi
chuỗi bắt đầu bằng "=" là công thức, nên một học sinh tên "=HYPERLINK(...)"
sẽ thành liên kết sống trong tệp giáo viên mở. Ép về kiểu chuỗi thì Excel
hiện NGUYÊN VĂN, không tính, không thêm dấu nháy như bản CSV phải làm.
Cũng ở đó, `lam_sach` bỏ ký tự Excel cấm và cắt ô quá dài — xem hàm ấy.
"""

from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.properties import PageSetupProperties

NEN_TIEU_DE_COT = PatternFill("solid", fgColor="DDE5F0")
NEN_KHOI = PatternFill("solid", fgColor="C9D7EA")
NEN_SOC = PatternFill("solid", fgColor="F5F7FA")
NEN_DO = PatternFill("solid", fgColor="FDE2E1")
CHU_TIEU_DE = Font(bold=True, size=14)
CHU_PHU = Font(italic=True, color="666666")
CHU_MUC = Font(bold=True, size=12)
CHU_DAM = Font(bold=True)
CHU_CANH_BAO = Font(bold=True, color="C00000")
CHU_LIEN_KET = Font(color="0563C1", underline="single")
CANH_GIUA = Alignment(horizontal="center")

RONG_TOI_DA = 60
# Gioi han cung cua Excel. Vuot qua la tep phai "sua chua" khi mo.
O_TOI_DA = 32767           # ky tu trong mot o
NGAT_TRANG_TOI_DA = 1026   # ngat trang thu cong trong mot trang tinh


def lam_sach(gia_tri):
    """Chuỗi mà Excel nhận được: bỏ ký tự điều khiển, cắt ở 32.767 ký tự.

    openpyxl từ chối \\x00-\\x08, \\x0b, \\x0c, \\x0e-\\x1f bằng
    IllegalCharacterError, mà nạp CSV vẫn để lọt chúng (dán từ Word hay
    Forms). Không lọc thì MỘT ký tự vô hình trong tên MỘT em làm hỏng cả
    lần xuất. Tab và xuống dòng vẫn giữ — Excel nhận được.
    """
    if not isinstance(gia_tri, str):
        return gia_tri
    return ILLEGAL_CHARACTERS_RE.sub("", gia_tri)[:O_TOI_DA]


class Trang:
    """Một trang tính, ghi từ trên xuống: `hang` là dòng trống kế tiếp."""

    def __init__(self, ws):
        self.ws = ws
        self.hang = 1
        self.rong: dict[int, int] = {}
        self.so_cot = 1

    # ---- ô ----
    def o(self, hang, cot, gia_tri, do_rong=True):
        gia_tri = lam_sach(gia_tri)
        if gia_tri == "":
            gia_tri = None
        c = self.ws.cell(row=hang, column=cot, value=gia_tri)
        if c.data_type == "f":
            c.data_type = "s"
            c.quotePrefix = True
        if do_rong and gia_tri is not None:
            self.rong[cot] = max(self.rong.get(cot, 0), len(str(gia_tri)))
        self.so_cot = max(self.so_cot, cot)
        return c

    # ---- khối chữ ----
    def tieu_de(self, chinh: str, phu: str = ""):
        """Dòng 1 tên trang, dòng 2 một câu nói trang dùng để làm gì."""
        self.o(1, 1, chinh, do_rong=False).font = CHU_TIEU_DE
        if phu:
            self.o(2, 1, phu, do_rong=False).font = CHU_PHU
        self.hang = 4

    def muc(self, chu: str):
        """Tên một phần trong trang, cách phần trước một dòng trống."""
        if self.hang > 4:
            self.hang += 1
        self.o(self.hang, 1, chu, do_rong=False).font = CHU_MUC
        self.hang += 1

    def ghi_chu(self, chu: str, font=CHU_PHU):
        self.o(self.hang, 1, chu, do_rong=False).font = font
        self.hang += 1

    def dong(self, *cac_o, dam=False):
        """Một dòng nhãn — giá trị. Căn trái để số đứng sát nhãn của nó."""
        for i, v in enumerate(cac_o, 1):
            c = self.o(self.hang, i, v)
            c.alignment = Alignment(horizontal="left")
            if dam:
                c.font = CHU_DAM
        self.hang += 1

    def lien_ket(self, ten_trang: str, mo_ta: str):
        """Một dòng mục lục: bấm tên trang là nhảy tới trang đó."""
        c = self.o(self.hang, 1, ten_trang)
        c.hyperlink = Hyperlink(ref=c.coordinate, location="'%s'!A1" % ten_trang)
        c.font = CHU_LIEN_KET
        self.o(self.hang, 2, mo_ta)
        self.hang += 1

    # ---- bảng ----
    def bang(self, cot, dong, to_do=None, khoa=False, loc=False, rong="", bo_soc=False):
        """Tiêu đề cột + các dòng. Trả về (dòng tiêu đề, dòng cuối).

        `to_do(dòng)` True -> tô đỏ nhạt cả dòng (em chưa có chỗ...).
        `khoa`: khoá cuộn ngay dưới tiêu đề — cuộn xuống dòng 300 vẫn biết
        cột nào là cột nào. `loc`: bật bộ lọc ở dòng tiêu đề. `rong`: chữ
        hiện ở dòng đầu khi bảng không có dòng nào.
        """
        dau = self.hang
        for i, ten in enumerate(cot, 1):
            c = self.o(dau, i, ten)
            c.font = CHU_DAM
            c.fill = NEN_TIEU_DE_COT
            c.alignment = Alignment(vertical="center", wrap_text=True)
        self.hang += 1
        if not dong and rong:
            self.o(self.hang, 1, rong, do_rong=False).font = CHU_PHU
            self.hang += 1
        for n, d in enumerate(dong):
            nen = NEN_DO if (to_do and to_do(d)) else (
                NEN_SOC if (n % 2 and not bo_soc) else None)
            for i, v in enumerate(d, 1):
                c = self.o(self.hang, i, v)
                if nen:
                    c.fill = nen
                # So canh phai dinh lien chu cot ben ("1HS006") — canh giua.
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    c.alignment = CANH_GIUA
            self.hang += 1
        cuoi = self.hang - 1
        if khoa:
            self.ws.freeze_panes = "A%d" % (dau + 1)
        if loc and cot:
            self.ws.auto_filter.ref = "A%d:%s%d" % (
                dau, get_column_letter(len(cot)), max(cuoi, dau))
        return dau, cuoi

    def khoi(self, chu: str, so_cot: int):
        """Dòng đầu một khối (một CLB trên trang Theo CLB)."""
        for i in range(1, so_cot + 1):
            self.ws.cell(row=self.hang, column=i).fill = NEN_KHOI
        self.o(self.hang, 1, chu, do_rong=False).font = CHU_MUC
        self.hang += 1

    def ngat_trang(self):
        """Ngắt trang ngay sau dòng vừa ghi — khối sau in sang trang mới.

        Quá 1.026 ngắt (giới hạn của Excel) thì các khối sau in nối tiếp.
        """
        if len(self.ws.row_breaks.brk) < NGAT_TRANG_TOI_DA:
            self.ws.row_breaks.append(Break(id=self.hang - 1))

    # ---- hoàn tất ----
    def xong(self, ngang=None, lap_hang=None, rong_co_dinh=None):
        """Giãn cột theo nội dung bảng và đặt khổ in A4 vừa một chiều ngang.

        Chỉ đo ô trong BẢNG (tiêu đề, ghi chú dài không tính) — không thì một
        câu ghi chú kéo cột A rộng ra cả màn hình.
        """
        for i in range(1, self.so_cot + 1):
            r = (rong_co_dinh or {}).get(i) or min(self.rong.get(i, 8) + 2, RONG_TOI_DA)
            self.ws.column_dimensions[get_column_letter(i)].width = r

        ps = self.ws.page_setup
        ps.paperSize = self.ws.PAPERSIZE_A4
        if ngang is None:
            ngang = self.so_cot > 7
        ps.orientation = "landscape" if ngang else "portrait"
        self.ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ps.fitToWidth = 1
        ps.fitToHeight = 0       # 0 = bao nhieu trang doc cung duoc
        if lap_hang:
            self.ws.print_title_rows = "%d:%d" % (lap_hang, lap_hang)
        self.ws.oddFooter.left.text = "&A"
        self.ws.oddFooter.center.text = "Trang &P / &N"
