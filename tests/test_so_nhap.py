# -*- coding: utf-8 -*-
"""Canh Sổ nhập CLB — một tệp Excel thay cho ba tệp nạp (so_nhap.py).

Điều quan trọng nhất: nạp SỔ và nạp BỘ BA CSV cũ phải ra CÙNG MỘT CSDL và
CÙNG MỘT KẾT QUẢ phân bổ. Sổ chỉ là lớp vỏ dễ dùng hơn — nó không được đổi
dù một em học sinh vào CLB nào.

Còn lại: sổ lỗi thì KHÔNG ghi gì, lỗi chỉ đúng sheet và dòng; tên CLB và mã
CLB đều dùng được trong ô NV; chữ `thi` nghĩa là đã thi, chấm sau.
"""

import base64
import csv
import io
import os
import sqlite3

import openpyxl
import pytest

import so_nhap

GOC = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MAU = os.path.join(GOC, "mau_csv")
TEP_CSV = ["01_danh_sach_CLB.csv", "02_chon_CLB_muon_thi.csv",
           "03_xep_hang_nguyen_vong.csv"]


def b64_tep(duong_dan) -> str:
    with open(duong_dan, "rb") as f:
        return base64.b64encode(f.read()).decode()


def b64_sach(clubs, students, **kw) -> str:
    buf = io.BytesIO()
    so_nhap.ghi_so_nhap(buf, clubs, students, **kw)
    return base64.b64encode(buf.getvalue()).decode()


def b64_wb(wb) -> str:
    buf = io.BytesIO()
    wb.save(buf)
    return base64.b64encode(buf.getvalue()).decode()


def chup_csdl(api) -> dict:
    """Mọi bảng ĐẦU VÀO, sắp cố định — để so hai CSDL."""
    con = sqlite3.connect(api.db_path)
    try:
        return {
            "clubs": con.execute(
                "SELECT club_id, name, capacity, reserve_capacity, "
                "COALESCE(reserve_group,''), COALESCE(buoi,'') FROM clubs ORDER BY 1").fetchall(),
            "students": con.execute(
                "SELECT student_id, name, COALESCE(reserve_group,'') FROM students "
                "ORDER BY 1").fetchall(),
            "preferences": con.execute(
                "SELECT student_id, club_id, rank FROM preferences ORDER BY 1, 3").fetchall(),
            "tick": con.execute(
                "SELECT student_id, club_id FROM club_test_selection ORDER BY 1, 2").fetchall(),
            "scores": con.execute(
                "SELECT student_id, club_id, score FROM club_scores ORDER BY 1, 2").fetchall(),
        }
    finally:
        con.close()


def chup_ket_qua(api) -> list:
    con = sqlite3.connect(api.db_path)
    try:
        return con.execute(
            "SELECT * FROM match_results ORDER BY 1, 2").fetchall()
    finally:
        con.close()


CLB = [
    {"club_id": "", "name": "CLB Tin học", "capacity": 3, "reserve_capacity": 0,
     "reserve_group": "", "buoi": "thu_2"},
    {"club_id": "", "name": "CLB Bóng rổ", "capacity": 3, "reserve_capacity": 1,
     "reserve_group": "chinh_sach", "buoi": "thu_2"},
    {"club_id": "clb_robot", "name": "CLB Robotics", "capacity": 2, "reserve_capacity": 0,
     "reserve_group": "", "buoi": "thu_4"},
]


def hs(sid, *nv, nhom=""):
    """hs('HS1', ('CLB Tin học', 9), 'CLB Bóng rổ') — tuple là (CLB, điểm)."""
    ds = []
    for x in nv:
        ten, diem = (x if isinstance(x, tuple) else (x, None))
        ds.append({"club_id": ten, "thi": diem is not None,
                   "diem": "" if diem in (None, "thi") else diem})
    return {"student_id": sid, "name": "Em " + sid, "reserve_group": nhom, "nv": ds}


# ------------------------------------------------------------------ #
# 1. SỔ VÀ BỘ BA CSV RA CÙNG MỘT CSDL, CÙNG MỘT KẾT QUẢ
# ------------------------------------------------------------------ #

@pytest.mark.parametrize("bo", ["vi_du_day_du", "vi_du_ca_tuan"])
def test_so_vi_du_ra_dung_csdl_va_ket_qua_nhu_bo_ba_csv(api_factory, bo):
    api_csv = api_factory()
    for ten in TEP_CSV:
        with io.open(os.path.join(MAU, bo, ten), encoding="utf-8-sig") as f:
            assert api_csv.import_csv_auto(f.read())["ok"]

    api_so = api_factory()
    kq = api_so.import_so_nhap(b64_tep(os.path.join(MAU, bo, "SO_NHAP_CLB_vi_du.xlsx")))
    assert kq["ok"], kq
    assert kq["data"]["warnings"] == [], kq["data"]["warnings"]

    assert chup_csdl(api_so) == chup_csdl(api_csv)

    assert api_csv.run_pipeline(seed=42)["ok"]
    assert api_so.run_pipeline(seed=42)["ok"]
    assert chup_ket_qua(api_so) == chup_ket_qua(api_csv)


DLT = os.path.join(GOC, "du_lieu_test")


@pytest.mark.parametrize("thu_muc,tien_to,duoi,ten_so", [
    (os.path.join(DLT, "vi_du_huong_dan"), "VIDU_", ".csv", "SO_NHAP_VIDU.xlsx"),
    (os.path.join(DLT, "bo_sach"), "SACH_", ".csv", "SO_NHAP_SACH.xlsx"),
    (DLT, "TEST_", ".xlsx", "SO_NHAP_TEST.xlsx"),
])
def test_so_du_lieu_test_ra_dung_csdl_nhu_bo_cu(api_factory, thu_muc, tien_to, duoi, ten_so):
    """Các bộ chạy thử trong du_lieu_test/ có bản sổ nhập: nạp sổ phải ra
    đúng CSDL như nạp bộ ba tệp cũ (mau_csv/tao_so_nhap.py sinh sổ từ đó)."""
    api_cu = api_factory()
    for t in TEP_CSV:
        p = os.path.join(thu_muc, tien_to + t.replace(".csv", duoi))
        if duoi == ".csv":
            with io.open(p, encoding="utf-8-sig") as f:
                text = f.read()
        else:
            text = api_cu.xlsx_to_csv_text(b64_tep(p))["data"]["csv_text"]
        assert api_cu.import_csv_auto(text)["ok"]
    api_so = api_factory()
    kq = api_so.import_so_nhap(b64_tep(os.path.join(thu_muc, ten_so)))
    assert kq["ok"], kq
    assert chup_csdl(api_so) == chup_csdl(api_cu)


def test_so_co_loi_co_y_bao_dung_nam_loi_va_khong_ghi_gi(api):
    import importlib.util
    spec = importlib.util.spec_from_file_location(
        "tao_so_nhap", os.path.join(MAU, "tao_so_nhap.py"))
    tsn = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(tsn)

    kq = api.import_so_nhap(b64_tep(os.path.join(DLT, "SO_NHAP_CO_LOI_CO_Y.xlsx")))
    assert not kq["ok"]
    assert sorted(e["code"] for e in kq["errors"]) == sorted(x[3] for x in tsn.LOI_CO_Y)
    assert chup_csdl(api)["clubs"] == []


def test_so_vi_du_khop_bo_csv_cung_thu_muc():
    """Sổ ví dụ sinh từ CSV (mau_csv/tao_so_nhap.py). Ai sửa CSV mà quên
    sinh lại sổ thì test này đỏ."""
    for bo in ("vi_du_day_du", "vi_du_ca_tuan"):
        def doc(ten):
            with io.open(os.path.join(MAU, bo, ten), encoding="utf-8-sig", newline="") as f:
                return list(csv.DictReader(f))
        clubs, students, _ = so_nhap.tu_csv(*(doc(t) for t in TEP_CSV))
        wb = openpyxl.load_workbook(os.path.join(MAU, bo, "SO_NHAP_CLB_vi_du.xlsx"),
                                    read_only=True)
        du = so_nhap.doc_so_nhap(wb)
        assert du["loi"] == []
        assert [c["club_id"] for c in du["clubs"]] == [c["club_id"] for c in clubs]
        assert [(s["student_id"], [x["club_id"] for x in s["nv"]]) for s in du["students"]] \
            == [(s["student_id"], [x["club_id"] for x in s["nv"]]) for s in students]


def test_mau_trong_doc_duoc_va_dung_bo_cuc():
    wb = openpyxl.load_workbook(os.path.join(MAU, "SO_NHAP_CLB.xlsx"))
    assert wb.sheetnames[:2] == [so_nhap.SHEET_CLB, so_nhap.SHEET_HS]
    assert so_nhap.SHEET_HD in wb.sheetnames
    clb, hs_ = wb[so_nhap.SHEET_CLB], wb[so_nhap.SHEET_HS]
    assert [c.value for c in clb[1]] == so_nhap.COT_CLB
    tieu_de = [c.value for c in hs_[1]]
    assert tieu_de[:5] == ["Mã HS", "Họ tên", "Nhóm ưu tiên", "NV1", "Điểm 1"]
    assert tieu_de.count("NV%d" % so_nhap.SO_NV_MAC_DINH) == 1
    # Ô NV là danh sách thả xuống lấy tên CLB từ sheet 1.
    nguon = [dv.formula1 for dv in hs_.data_validations.dataValidation if dv.type == "list"]
    assert any(so_nhap.SHEET_CLB in f for f in nguon), nguon
    assert any(dv.type == "list" for dv in clb.data_validations.dataValidation)


# ------------------------------------------------------------------ #
# 2. ĐỌC Ô: tên hay mã, buổi, điểm, chữ "thi"
# ------------------------------------------------------------------ #

def test_nv_nhan_ca_ten_lan_ma_clb_va_khong_phan_biet_dau(api):
    students = [hs("HS1", ("CLB Tin học", 9.5), "clb_robot"),
                hs("HS2", "clb tin hoc", ("CLB BÓNG RỔ", "8,5"))]
    kq = api.import_so_nhap(b64_sach(CLB, students))
    assert kq["ok"], kq
    db = chup_csdl(api)
    assert [c[0] for c in db["clubs"]] == ["clb_bong_ro", "clb_robot", "clb_tin_hoc"]
    assert db["preferences"] == [("HS1", "clb_tin_hoc", 1), ("HS1", "clb_robot", 2),
                                 ("HS2", "clb_tin_hoc", 1), ("HS2", "clb_bong_ro", 2)]
    assert db["tick"] == [("HS1", "clb_tin_hoc"), ("HS2", "clb_bong_ro")]
    assert db["scores"] == [("HS1", "clb_tin_hoc", 9.5), ("HS2", "clb_bong_ro", 8.5)]


def test_chu_thi_la_da_thi_chua_co_diem(api):
    """`thi` = em đã thi, chưa có điểm. Điểm chấm sau ở thẻ 05; nạp lại một
    sổ còn ghi `thi` thì sổ thắng (xem test_thi_de_len_diem_cu_...)."""
    kq = api.import_so_nhap(b64_sach(CLB, [hs("HS1", ("CLB Tin học", "thi"))]))
    assert kq["ok"], kq
    db = chup_csdl(api)
    assert db["tick"] == [("HS1", "clb_tin_hoc")]
    assert db["scores"] == []


def test_nap_lai_so_da_bo_thi_thi_xoa_luon_tick_cu(api):
    assert api.import_so_nhap(b64_sach(CLB, [hs("HS1", ("CLB Tin học", 9))]))["ok"]
    assert api.import_so_nhap(b64_sach(CLB, [hs("HS1", "CLB Tin học")]))["ok"]
    db = chup_csdl(api)
    assert db["tick"] == [] and db["scores"] == []
    assert db["preferences"] == [("HS1", "clb_tin_hoc", 1)]


@pytest.mark.parametrize("nhan,ma", [
    ("Thứ 2", "thu_2"), ("thu 3", "thu_3"), ("T4", "thu_4"), ("Chủ nhật", "chu_nhat"),
    ("CN", "chu_nhat"), ("", ""), ("thu_3_tiet_9", "thu_3_tiet_9"),
    ("Sáng thứ 7", "sáng_thứ_7"),
])
def test_nhan_buoi(nhan, ma):
    assert so_nhap.ma_buoi(nhan) == ma


def test_truong_mot_buoi_de_trong_cot_buoi(api):
    clubs = [dict(c, buoi="") for c in CLB]
    kq = api.import_so_nhap(b64_sach(clubs, [hs("HS1", "CLB Tin học")]))
    assert kq["ok"], kq
    assert kq["data"]["n_buoi"] == 1
    assert {c[5] for c in chup_csdl(api)["clubs"]} == {""}


def test_ma_tu_sinh_trung_ma_nguoi_dung_dat_thi_bao_loi_khong_tu_danh_so(api):
    """Tự thêm _2 thì số phụ thuộc thứ tự dòng: sắp lại dòng là hai CLB đổi
    mã cho nhau. Đụng thì bắt người dùng điền Mã CLB, ở cả hai thứ tự."""
    a = {"club_id": "clb_tin_hoc", "name": "Tin học nâng cao", "capacity": 2}
    b = {"club_id": "", "name": "CLB Tin học", "capacity": 2}
    for clubs in ([a, b], [b, a]):
        kq = api.import_so_nhap(b64_sach(clubs, []))
        assert not kq["ok"]
        assert [e["code"] for e in kq["errors"]] == ["so_nhap_ma_tu_sinh_trung"]
    assert chup_csdl(api)["clubs"] == []


def test_chua_co_hoc_sinh_thi_nap_rieng_clb(api):
    kq = api.import_so_nhap(b64_sach(CLB, []))
    assert kq["ok"], kq
    assert kq["data"]["n_clb"] == 3 and kq["data"]["n_hoc_sinh"] == 0
    assert len(chup_csdl(api)["clubs"]) == 3


# ------------------------------------------------------------------ #
# 3. SỔ LỖI: không ghi gì, chỉ đúng sheet và dòng
# ------------------------------------------------------------------ #

def _so_co(sua):
    """Sổ sạch, rồi `sua(wb)` làm hỏng một chỗ."""
    buf = io.BytesIO()
    so_nhap.ghi_so_nhap(buf, CLB, [hs("HS1", ("CLB Tin học", 9)), hs("HS2", "CLB Bóng rổ")])
    wb = openpyxl.load_workbook(io.BytesIO(buf.getvalue()))
    sua(wb)
    return b64_wb(wb)


@pytest.mark.parametrize("sua,ma_loi,tham_so", [
    (lambda wb: wb[so_nhap.SHEET_HS].__setitem__("D3", "CLB Không Có"),
     "so_nhap_clb_khong_co", {"dong": 3, "k": 1, "gia_tri": "CLB Không Có"}),
    (lambda wb: wb[so_nhap.SHEET_HS].__setitem__("F2", "CLB Tin học"),
     "so_nhap_nv_trung", {"dong": 2, "k": 2, "k_truoc": 1}),
    (lambda wb: wb[so_nhap.SHEET_HS].__setitem__("E2", "tám"),
     "so_nhap_diem_sai", {"dong": 2, "k": 1}),
    (lambda wb: wb[so_nhap.SHEET_HS].__setitem__("G3", 8),
     "so_nhap_diem_khong_nv", {"dong": 3, "k": 2}),
    (lambda wb: wb[so_nhap.SHEET_HS].__setitem__("A3", "HS1"),
     "so_nhap_hs_trung", {"dong": 3, "dong_truoc": 2}),
    (lambda wb: wb[so_nhap.SHEET_HS].__setitem__("A3", None),
     "so_nhap_thieu_ma_hs", {"dong": 3}),
    # Dòng 4 (Robotics): không em nào chọn, nên lỗi trùng tên đứng một mình.
    (lambda wb: wb[so_nhap.SHEET_CLB].__setitem__("A4", "CLB Tin học"),
     "so_nhap_clb_trung_ten", {"dong": 4, "dong_truoc": 2}),
    (lambda wb: wb[so_nhap.SHEET_CLB].__setitem__("C2", 0),
     "so_nhap_chi_tieu_sai", {"dong": 2}),
    (lambda wb: wb[so_nhap.SHEET_CLB].__setitem__("D3", 9),
     "so_nhap_chi_tieu_sai", {"dong": 3}),
    (lambda wb: wb[so_nhap.SHEET_HS].__setitem__("A1", "Số báo danh"),
     "so_nhap_thieu_cot", {"cot": "Mã HS"}),
])
def test_so_loi_khong_ghi_gi_va_chi_dung_cho(api, sua, ma_loi, tham_so):
    file_b64 = _so_co(sua)
    xem = api.xem_truoc_so_nhap(file_b64)
    assert xem["ok"], xem
    loi = xem["data"]["loi"]
    assert [e["code"] for e in loi] == [ma_loi], loi
    for k, v in tham_so.items():
        assert loi[0]["params"][k] == v, loi

    kq = api.import_so_nhap(file_b64)
    assert not kq["ok"]
    assert [e["code"] for e in kq["errors"]] == [ma_loi]
    db = chup_csdl(api)
    assert db["clubs"] == [] and db["students"] == []


def test_tep_excel_khong_phai_so_nhap(api):
    wb = openpyxl.Workbook()
    wb.active.append(["club_id", "name", "capacity"])
    wb.active.append(["clb_a", "A", 3])
    kq = api.import_so_nhap(b64_wb(wb))
    assert not kq["ok"]
    assert kq["errors"][0]["code"] == "so_nhap_khong_phai_so"
    assert not api.xem_truoc_so_nhap(b64_wb(wb))["ok"]


def test_tep_khong_phai_excel(api):
    kq = api.import_so_nhap(base64.b64encode(b"student_id,name\nHS1,A\n").decode())
    assert not kq["ok"]
    assert kq["errors"][0]["code"] == "xlsx_read_failed"


def test_tieu_de_go_lai_khong_dau_van_nhan(api):
    def sua(wb):
        ws = wb[so_nhap.SHEET_HS]
        ws["A1"], ws["B1"], ws["D1"], ws["E1"] = "MA HOC SINH", "ho va ten", "nv 1", "DIEM 1"
        wb[so_nhap.SHEET_CLB].title = "CLB"
    kq = api.import_so_nhap(_so_co(sua))
    assert kq["ok"], kq
    assert chup_csdl(api)["scores"] == [("HS1", "clb_tin_hoc", 9.0)]


# ------------------------------------------------------------------ #
# 4. LUẬT CŨ VẪN ÁP DỤNG (trần nguyện vọng mỗi buổi…)
# ------------------------------------------------------------------ #

def test_luat_cua_ham_nap_cu_van_bao_qua_so(api):
    """Luật soát của hàm nạp cũ vẫn chạy với sổ: nhãn nhóm ưu tiên gõ sai
    (không CLB nào nhận) vẫn được cảnh báo kèm gợi ý."""
    kq = api.import_so_nhap(b64_sach(CLB, [hs("HS1", "CLB Bóng rổ", nhom="chinh_sac")]))
    assert kq["ok"], kq
    assert any(w["code"] == "csv_reserve_group_unknown" for w in kq["data"]["warnings"]), \
        kq["data"]["warnings"]


def test_tran_trong_so_nhap_bang_tran_cua_phan_mem():
    import rbda_priority_pipeline as rp
    assert so_nhap.TRAN_NV_MOI_BUOI == rp.TRAN_NGUYEN_VONG_MOI_BUOI
    assert so_nhap.TRAN_THI_MOI_BUOI == rp.TRAN_CLB_THI_MOI_BUOI


# ------------------------------------------------------------------ #
# 5. CÁC LỖI ĐÃ TÌM QUA CODE REVIEW
# ------------------------------------------------------------------ #

def test_ten_chi_khac_ngoac_hay_ky_hieu_la_hai_clb_khac_nhau(api):
    clubs = [{"club_id": "", "name": n, "capacity": 3} for n in
             ("Bóng đá (nam)", "Bóng đá (nữ)", "C++", "C#")]
    kq = api.import_so_nhap(b64_sach(clubs, [hs("HS1", "Bóng đá (nữ)", "C#")]))
    assert kq["ok"], kq
    ma = {c[1]: c[0] for c in chup_csdl(api)["clubs"]}
    assert len(set(ma.values())) == 4
    assert ma["Bóng đá (nữ)"] == "clb_bong_da_nu" and ma["C#"] == "clb_c_sharp"
    assert chup_csdl(api)["preferences"] == [
        ("HS1", ma["Bóng đá (nữ)"], 1), ("HS1", ma["C#"], 2)]


def test_go_thieu_dau_chi_nhan_khi_ra_dung_mot_clb(api):
    clubs = [{"club_id": "", "name": n, "capacity": 3} for n in
             ("Bóng đá (nam)", "Bóng đá (nữ)", "CLB Tin học")]
    xem = api.xem_truoc_so_nhap(b64_sach(clubs, [hs("HS1", "clb tin hoc", "bong da")]))
    assert [e["code"] for e in xem["data"]["loi"]] == ["so_nhap_clb_khong_co"]
    assert xem["data"]["loi"][0]["params"]["gia_tri"] == "bong da"


@pytest.mark.parametrize("so_nv,so_thi,ma_loi", [
    (so_nhap.TRAN_NV_MOI_BUOI + 1, 0, "so_nhap_qua_tran_nv"),
    (so_nhap.TRAN_THI_MOI_BUOI + 1, so_nhap.TRAN_THI_MOI_BUOI + 1, "so_nhap_qua_tran_thi"),
])
def test_qua_tran_moi_buoi_bi_chan_tu_luc_doc_thu(api, so_nv, so_thi, ma_loi):
    clubs = [{"club_id": "", "name": "CLB %02d" % i, "capacity": 5, "buoi": "thu_2"}
             for i in range(so_nv)]
    em = hs("HS1", *[("CLB %02d" % i, 5) if i < so_thi else "CLB %02d" % i
                     for i in range(so_nv)])
    xem = api.xem_truoc_so_nhap(b64_sach(clubs, [em]))
    assert [e["code"] for e in xem["data"]["loi"]] == [ma_loi]
    assert xem["data"]["loi"][0]["params"]["buoi"] == "Thứ 2"
    assert not api.import_so_nhap(b64_sach(clubs, [em]))["ok"]
    assert chup_csdl(api)["clubs"] == []


def test_qua_tran_tinh_rieng_tung_buoi(api):
    """11 nguyện vọng chia hai buổi (6 + 5) là hợp lệ."""
    clubs = [{"club_id": "", "name": "CLB %02d" % i, "capacity": 5,
              "buoi": "thu_2" if i < 6 else "thu_4"} for i in range(11)]
    kq = api.import_so_nhap(b64_sach(clubs, [hs("HS1", *["CLB %02d" % i for i in range(11)])]))
    assert kq["ok"], kq


def test_buoi_khai_bang_form_va_bang_so_la_mot(api):
    assert api.create_or_update_club("clb_x", "CLB X", 5, 0, "", buoi="Thứ 3")["ok"]
    kq = api.import_so_nhap(b64_sach(
        [{"club_id": "clb_y", "name": "CLB Y", "capacity": 5, "buoi": "thu_3"}], []))
    assert kq["ok"], kq
    assert {c[5] for c in chup_csdl(api)["clubs"]} == {"thu_3"}


def test_csdl_cu_con_nhan_thu_co_dau_thi_so_dung_lai_nhan_do(api):
    con = sqlite3.connect(api.db_path)
    con.execute("INSERT INTO clubs (club_id, name, capacity, reserve_capacity, buoi) "
                "VALUES ('clb_x', 'CLB X', 5, 0, 'thứ_3')")
    con.commit()
    con.close()
    kq = api.import_so_nhap(b64_sach(
        [{"club_id": "clb_y", "name": "CLB Y", "capacity": 5, "buoi": "thu_3"}], []))
    assert kq["ok"], kq
    assert {c[5] for c in chup_csdl(api)["clubs"]} == {"thứ_3"}
    assert api.create_or_update_club("clb_z", "CLB Z", 5, 0, "", buoi="Thứ 3")["ok"]
    assert {c[5] for c in chup_csdl(api)["clubs"]} == {"thứ_3"}


def test_thi_de_len_diem_cu_thi_xoa_diem_va_bao_truoc(api):
    assert api.import_so_nhap(b64_sach(CLB, [hs("HS1", ("CLB Tin học", 9))]))["ok"]
    so = b64_sach(CLB, [hs("HS1", ("CLB Tin học", "thi"))])
    xem = api.xem_truoc_so_nhap(so)["data"]
    assert [(w["code"], w["params"]["n"]) for w in xem["canh_bao"]] == [
        ("so_nhap_canh_bao_xoa_diem", 1)]
    kq = api.import_so_nhap(so)
    assert kq["ok"], kq
    assert chup_csdl(api)["scores"] == []
    assert chup_csdl(api)["tick"] == [("HS1", "clb_tin_hoc")]


def test_bo_trong_nhom_uu_tien_thi_bo_nhom_va_bao_truoc(api):
    assert api.import_so_nhap(b64_sach(CLB, [hs("HS1", "CLB Bóng rổ", nhom="chinh_sach")]))["ok"]
    assert chup_csdl(api)["students"] == [("HS1", "Em HS1", "chinh_sach")]
    so = b64_sach(CLB, [hs("HS1", "CLB Bóng rổ")])
    xem = api.xem_truoc_so_nhap(so)["data"]
    assert [w["code"] for w in xem["canh_bao"]] == ["so_nhap_canh_bao_bo_nhom"]
    assert api.import_so_nhap(so)["ok"]
    assert chup_csdl(api)["students"] == [("HS1", "Em HS1", "")]


def test_ma_trong_dung_lai_clb_cung_ten_va_bao_clb_ngoai_so(api):
    """Nạp lại sổ bỏ trống Mã CLB không đẻ CLB thứ hai; CLB có trong phần
    mềm mà sổ không nhắc tới (vd vừa đổi tên) thì được báo, không bị xoá."""
    assert api.create_or_update_club("tin", "CLB Tin học", 5, 0, "")["ok"]
    assert api.create_or_update_club("cu", "CLB Cũ", 5, 0, "")["ok"]
    so = b64_sach([{"club_id": "", "name": "clb tin học", "capacity": 7}], [])
    xem = api.xem_truoc_so_nhap(so)["data"]
    assert [w["code"] for w in xem["canh_bao"]] == ["so_nhap_canh_bao_clb_ngoai_so"]
    assert "CLB Cũ" in xem["canh_bao"][0]["params"]["ten"]
    assert api.import_so_nhap(so)["ok"]
    db = {c[0]: c for c in chup_csdl(api)["clubs"]}
    assert set(db) == {"tin", "cu"} and db["tin"][2] == 7


def test_doi_ten_giu_ma_clb_thi_khong_sinh_clb_ma(api):
    assert api.import_so_nhap(b64_sach([{"club_id": "", "name": "CLB Cờ", "capacity": 4}], []))["ok"]
    mau = api.tao_so_nhap_mau()["data"]["path"]      # sổ mẫu ghi sẵn Mã CLB
    wb = openpyxl.load_workbook(mau)
    wb[so_nhap.SHEET_CLB]["A2"] = "CLB Cờ vua"
    assert api.import_so_nhap(b64_wb(wb))["ok"]
    assert chup_csdl(api)["clubs"] == [("clb_co", "CLB Cờ vua", 4, 0, "", "")]


# ------------------------------------------------------------------ #
# 6. CODE REVIEW LẦN CUỐI
# ------------------------------------------------------------------ #

@pytest.mark.parametrize("cach_hong", ["tra_loi", "nem_loi"])
def test_buoc_sau_hong_thi_khong_giu_lai_gi_ca(api, monkeypatch, cach_hong):
    """Nạp nửa sổ còn tệ hơn không nạp: CLB đã ghi xong mà bước nguyện vọng
    hỏng (CSDL bị khoá, lỗi bất ngờ) thì CLB cũng phải được huỷ."""
    import i18n_errors

    def hong(*a, **k):
        if cach_hong == "nem_loi":
            raise RuntimeError("database is locked")
        return i18n_errors.phan_hoi_loi(i18n_errors.err("error_importing_preferences_csv"))
    monkeypatch.setattr(api, "import_preferences_csv", hong)

    kq = api.import_so_nhap(b64_sach(CLB, [hs("HS1", ("CLB Tin học", 9))]))
    assert not kq["ok"]
    db = chup_csdl(api)
    assert db == {"clubs": [], "students": [], "preferences": [], "tick": [], "scores": []}


def test_nap_lai_hong_giu_nguyen_du_lieu_cu(api, monkeypatch):
    assert api.import_so_nhap(b64_sach(CLB, [hs("HS1", ("CLB Tin học", 9))]))["ok"]
    truoc = chup_csdl(api)
    monkeypatch.setattr(api, "import_preferences_csv",
                        lambda *a, **k: (_ for _ in ()).throw(RuntimeError("hong")))
    assert not api.import_so_nhap(b64_sach(CLB, [hs("HS1", "CLB Bóng rổ")]))["ok"]
    assert chup_csdl(api) == truoc


@pytest.mark.parametrize("chi_tieu", ["inf", "1e999", "nan", "-inf"])
def test_chi_tieu_vo_cuc_bao_dung_dong(api, chi_tieu):
    clubs = [dict(c) for c in CLB]
    clubs[0]["capacity"] = chi_tieu
    xem = api.xem_truoc_so_nhap(b64_sach(clubs, []))
    assert xem["ok"], xem
    assert [(e["code"], e["params"]["dong"]) for e in xem["data"]["loi"]] == [
        ("so_nhap_chi_tieu_sai", 2)]


def test_ghi_so_tu_doi_ten_clb_trung_o_hai_buoi(api, tmp_path):
    """Trường mở "CLB Cờ vua" cả Thứ 2 lẫn Thứ 5: mọi bộ sinh sổ đều đi qua
    ghi_so_nhap, nên sổ nào ra cũng nạp được."""
    clubs = [{"club_id": "covua_t2", "name": "CLB Cờ vua", "capacity": 4, "buoi": "thu_2"},
             {"club_id": "covua_t5", "name": "CLB Cờ vua", "capacity": 4, "buoi": "thu_5"}]
    students = [hs("HS1", "covua_t2", "covua_t5")]
    kq = api.import_so_nhap(b64_sach(clubs, students))
    assert kq["ok"], kq
    db = chup_csdl(api)
    assert sorted(c[1] for c in db["clubs"]) == ["CLB Cờ vua", "CLB Cờ vua (Thứ 5)"]
    assert db["preferences"] == [("HS1", "covua_t2", 1), ("HS1", "covua_t5", 2)]
    assert clubs[1]["name"] == "CLB Cờ vua", "khong duoc sua du lieu cua ben goi"


def test_so_khong_co_cot_nhom_thi_giu_nhom_da_gan(api):
    """Sổ sinh từ biểu mẫu không có cột Nhóm ưu tiên: phần mềm phải giữ nhóm
    đã gán, không báo bỏ nhóm. Có cột mà ô trống thì vẫn bỏ nhóm."""
    assert api.import_so_nhap(b64_sach(CLB, [hs("HS1", "CLB Bóng rổ", nhom="chinh_sach")]))["ok"]
    so = b64_sach(CLB, [hs("HS1", "CLB Tin học")], co_nhom=False)
    wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(so)))
    assert "Nhóm ưu tiên" not in [c.value for c in wb[so_nhap.SHEET_HS][1]]
    assert api.xem_truoc_so_nhap(so)["data"]["canh_bao"] == []
    assert api.import_so_nhap(so)["ok"]
    db = chup_csdl(api)
    assert db["students"] == [("HS1", "Em HS1", "chinh_sach")]
    assert db["preferences"] == [("HS1", "clb_tin_hoc", 1)]


# ------------------------------------------------------------------ #
# 7. CODE REVIEW THEO MỤC TIÊU: không lặng lẽ đặt dữ liệu sai chỗ
# ------------------------------------------------------------------ #

def test_buoi_sang_chieu_cung_thu_la_hai_buoi(api):
    clubs = [{"club_id": "a", "name": "CLB A", "capacity": 3, "buoi": "Thứ 3 (sáng)"},
             {"club_id": "b", "name": "CLB B", "capacity": 3, "buoi": "Thứ 3 (chiều)"}]
    kq = api.import_so_nhap(b64_sach(clubs, []))
    assert kq["ok"], kq
    assert len({c[5] for c in chup_csdl(api)["clubs"]}) == 2
    assert api.create_or_update_club("c", "CLB C", 3, 0, "", buoi="Thứ 3 (sáng)")["ok"]
    assert len({c[5] for c in chup_csdl(api)["clubs"]}) == 2


def test_nhieu_nhan_cu_cung_ngay_thi_dung_nhan_nhieu_clb_nhat(api):
    """CSDL cũ có hai nhãn mà chuẩn hoá ra cùng một buổi: không chọn bừa."""
    con = sqlite3.connect(api.db_path)
    con.executemany("INSERT INTO clubs (club_id, name, capacity, reserve_capacity, buoi) "
                    "VALUES (?, ?, 5, 0, ?)",
                    [("x1", "X1", "thứ_3"), ("x2", "X2", "thứ_3"), ("y1", "Y1", "thứ 3")])
    con.commit()
    con.close()
    assert api.create_or_update_club("z", "Z", 5, 0, "", buoi="Thứ 3")["ok"]
    buoi = {c[0]: c[5] for c in chup_csdl(api)["clubs"]}
    assert buoi["x1"] == "thứ_3" and buoi["y1"] == "thứ 3"
    # Cùng là thứ Ba: dùng nhãn nhiều CLB nhất, không đẻ nhãn thứ ba.
    assert buoi["z"] == "thứ_3"


def test_nv_go_ten_khong_co_khong_roi_vao_clb_gan_giong(api):
    clubs = [{"club_id": "", "name": n, "capacity": 3} for n in ("Bóng đá (nam)", "C++")]
    xem = api.xem_truoc_so_nhap(b64_sach(clubs, [hs("HS1", "Bóng đá (nữ)", "C#")]))
    loi = xem["data"]["loi"]
    assert [e["code"] for e in loi] == ["so_nhap_clb_khong_co"] * 2, loi
    assert [e["params"]["gia_tri"] for e in loi] == ["Bóng đá (nữ)", "C#"]


@pytest.mark.parametrize("nhan", ["Thứ-2", "T.2", "Thứ 2.", "thu 2", "THỨ 2", "T2"])
def test_dau_cau_trong_nhan_thu_van_la_mot_buoi(nhan):
    assert so_nhap.ma_buoi(nhan) == "thu_2"


def test_nv_go_ma_clb_kieu_khac_van_khop(api):
    clubs = [{"club_id": "clb_tin_hoc", "name": "Tin học nâng cao", "capacity": 3}]
    kq = api.import_so_nhap(b64_sach(clubs, [hs("HS1", "clb tin hoc"), hs("HS2", "CLB-Tin-Hoc")]))
    assert kq["ok"], kq
    assert chup_csdl(api)["preferences"] == [("HS1", "clb_tin_hoc", 1), ("HS2", "clb_tin_hoc", 1)]


# ------------------------------------------------------------------ #
# 8. VÒNG REVIEW 1
# ------------------------------------------------------------------ #

def test_ma_tu_sinh_trung_ma_clb_khac_da_co_thi_bao_loi_khong_de(api):
    """CSDL có clb_bong_da tên "Bóng đá nam"; sổ thêm "CLB Bóng đá" bỏ trống
    Mã CLB -> mã tự tạo cũng là clb_bong_da. Không được lặng lẽ đè CLB cũ."""
    assert api.create_or_update_club("clb_bong_da", "Bóng đá nam", 20, 0, "", buoi="Thứ 2")["ok"]
    so = b64_sach([{"club_id": "", "name": "CLB Bóng đá", "capacity": 5, "buoi": "thu_3"}], [])
    xem = api.xem_truoc_so_nhap(so)["data"]
    assert [e["code"] for e in xem["loi"]] == ["so_nhap_ma_tu_sinh_trung_clb_da_co"]
    assert xem["loi"][0]["params"]["ten_cu"] == "Bóng đá nam"
    assert not api.import_so_nhap(so)["ok"]
    assert chup_csdl(api)["clubs"] == [("clb_bong_da", "Bóng đá nam", 20, 0, "", "thu_2")]


def test_csv_nguyen_vong_theo_buoi_dung_nhan_ngan_van_khop(api):
    """CLB khai buổi "t3" được lưu thành thu_3; tệp nguyện vọng cột t3_pref_1
    vẫn là cùng buổi, không bị coi là lệch buổi."""
    assert api.import_clubs_csv(
        "club_id,name,capacity,reserve_capacity,reserve_group,buoi\nclb_a,A,5,0,,t3\n")["ok"]
    kq = api.import_preferences_csv("student_id,name,t3_pref_1\nHS1,An,clb_a\n")
    assert kq["ok"], kq
    assert kq["data"]["n_nguyen_vong_lech_buoi"] == 0
    assert chup_csdl(api)["preferences"] == [("HS1", "clb_a", 1)]


# ------------------------------------------------------------------ #
# 9. VÒNG REVIEW 3
# ------------------------------------------------------------------ #

def test_sheet_clb_thieu_cot_tuy_chon_thi_giu_gia_tri_cu(api):
    """Sheet CLB chỉ có Tên CLB + Chỉ tiêu: CLB đã có giữ buổi, suất ưu tiên,
    nhóm ưu tiên — không lặng lẽ biến trường nhiều buổi thành một buổi."""
    assert api.create_or_update_club("clb_a", "CLB A", 5, 2, "chinh_sach", buoi="Thứ 2")["ok"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = so_nhap.SHEET_CLB
    ws.append(["Tên CLB", "Chỉ tiêu"])
    ws.append(["CLB A", 9])
    ws.append(["CLB Mới", 4])
    wb.create_sheet(so_nhap.SHEET_HS).append(["Mã HS", "Họ tên", "NV1", "Điểm 1"])
    kq = api.import_so_nhap(b64_wb(wb))
    assert kq["ok"], kq
    db = {c[0]: c for c in chup_csdl(api)["clubs"]}
    assert db["clb_a"] == ("clb_a", "CLB A", 9, 2, "chinh_sach", "thu_2")
    assert db["clb_moi"][2:] == (4, 0, "", "")


def test_cot_co_ma_o_trong_van_la_xoa(api):
    assert api.create_or_update_club("clb_a", "CLB A", 5, 2, "chinh_sach", buoi="Thứ 2")["ok"]
    kq = api.import_so_nhap(b64_sach([{"club_id": "clb_a", "name": "CLB A", "capacity": 5}], []))
    assert kq["ok"], kq
    assert chup_csdl(api)["clubs"] == [("clb_a", "CLB A", 5, 0, "", "")]


@pytest.mark.parametrize("nhan,ma", [("TBA", "tba"), ("thai", "thai"), ("Thứ Ba", "thu_3"),
                                     ("thu hai", "thu_2"), ("t 7", "thu_7")])
def test_chi_nhan_dung_thu_trong_tuan(nhan, ma):
    assert so_nhap.ma_buoi(nhan) == ma


# ------------------------------------------------------------------ #
# 10. VÒNG REVIEW 5
# ------------------------------------------------------------------ #

def test_nv_ghi_ma_dung_tung_ky_tu_khong_lan_sang_ma_khac_hoa_thuong(api):
    clubs = [{"club_id": "clb_a", "name": "CLB A thường", "capacity": 3},
             {"club_id": "CLB_A", "name": "CLB A hoa", "capacity": 3}]
    kq = api.import_so_nhap(b64_sach(clubs, [hs("HS1", "CLB_A"), hs("HS2", "clb_a")]))
    assert kq["ok"], kq
    assert chup_csdl(api)["preferences"] == [("HS1", "CLB_A", 1), ("HS2", "clb_a", 1)]


def test_nv_ghi_ten_clb_thang_ma_trung_chu_cua_clb_khac(api, tmp_path):
    """Tên một CLB trùng chữ với mã của CLB khác: ô NV (ghi tên) phải ra CLB
    mang tên đó, kể cả khi xuất sổ rồi nạp lại."""
    clubs = [{"club_id": "robot", "name": "Lego", "capacity": 3},
             {"club_id": "r2", "name": "robot", "capacity": 3}]
    buf = io.BytesIO()
    so_nhap.ghi_so_nhap(buf, clubs, [])
    wb = openpyxl.load_workbook(io.BytesIO(buf.getvalue()))
    ws = wb[so_nhap.SHEET_HS]
    cot_nv = [c.column for c in ws[1] if c.value == "NV1"][0]
    for i, (sid, o) in enumerate([("HS1", "robot"), ("HS2", "Lego")], start=2):
        ws.cell(i, 1, sid)
        ws.cell(i, cot_nv, o)
    kq = api.import_so_nhap(b64_wb(wb))
    assert kq["ok"], kq
    assert chup_csdl(api)["preferences"] == [("HS1", "r2", 1), ("HS2", "robot", 1)]
    # Xuất sổ rồi nạp lại: không đổi CLB nào.
    xuat = api.export_du_lieu_dau_vao(str(tmp_path / "so.xlsx"))
    assert xuat["ok"], xuat
    with open(xuat["data"]["path"], "rb") as f:
        assert api.import_so_nhap(base64.b64encode(f.read()).decode())["ok"]
    assert chup_csdl(api)["preferences"] == [("HS1", "r2", 1), ("HS2", "robot", 1)]


def test_khong_co_cot_buoi_thi_nap_giu_dung_nhan_doc_thu_da_soat(api):
    """CSDL cũ có hai nhãn cùng thứ Ba (thứ_3, thu_3). Sổ không có cột Buổi:
    đọc thử soát trần theo nhãn cũ, nên nạp cũng phải giữ đúng nhãn cũ."""
    con = sqlite3.connect(api.db_path)
    con.executemany("INSERT INTO clubs (club_id, name, capacity, reserve_capacity, buoi) "
                    "VALUES (?, ?, 5, 0, ?)",
                    [("a%d" % i, "A%d" % i, "thứ_3") for i in range(3)]
                    + [("b%d" % i, "B%d" % i, "thu_3") for i in range(2)])
    con.commit()
    con.close()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = so_nhap.SHEET_CLB
    ws.append(["Tên CLB", "Chỉ tiêu", "Mã CLB"])
    for i in range(3):
        ws.append(["A%d" % i, 6, "a%d" % i])
    wb.create_sheet(so_nhap.SHEET_HS).append(["Mã HS", "Họ tên", "NV1", "Điểm 1"])
    assert api.import_so_nhap(b64_wb(wb))["ok"]
    buoi = {c[0]: c[5] for c in chup_csdl(api)["clubs"]}
    assert [buoi["a%d" % i] for i in range(3)] == ["thứ_3"] * 3
    assert buoi["b0"] == "thu_3"


def test_nap_lai_so_sua_ho_ten_thi_cap_nhat_ten(api):
    assert api.import_so_nhap(b64_sach(CLB, [hs("HS1", "CLB Tin học")]))["ok"]
    em = hs("HS1", "CLB Tin học")
    em["name"] = "Nguyễn Văn An"
    assert api.import_so_nhap(b64_sach(CLB, [em]))["ok"]
    assert chup_csdl(api)["students"] == [("HS1", "Nguyễn Văn An", "")]


def test_ma_clb_trong_lay_lai_ma_cu_chi_khac_hoa_thuong_van_nap_duoc(api):
    """CSDL có CLB_A "Alpha" và clb_a "Beta". Sổ để trống Mã CLB (sổ từ biểu
    mẫu): mỗi dòng lấy lại đúng mã cũ theo tên, không báo trùng mã."""
    con = sqlite3.connect(api.db_path)
    con.executemany("INSERT INTO clubs (club_id, name, capacity, reserve_capacity) "
                    "VALUES (?, ?, 5, 0)", [("CLB_A", "Alpha"), ("clb_a", "Beta")])
    con.commit()
    con.close()
    clubs = [{"club_id": "", "name": "Alpha", "capacity": 5},
             {"club_id": "", "name": "Beta", "capacity": 5}]
    kq = api.import_so_nhap(b64_sach(clubs, [hs("HS1", "Alpha", "Beta")]))
    assert kq["ok"], kq
    assert chup_csdl(api)["preferences"] == [("HS1", "CLB_A", 1), ("HS1", "clb_a", 2)]


def test_o_nhom_chi_co_dau_gach_la_bo_nhom_va_duoc_bao_truoc(api):
    assert api.import_so_nhap(b64_sach(CLB, [hs("HS1", "CLB Tin học", nhom="chinh_sach")]))["ok"]
    assert chup_csdl(api)["students"][0][2] == "chinh_sach"
    so = b64_sach(CLB, [hs("HS1", "CLB Tin học", nhom="-")])
    xem = api.xem_truoc_so_nhap(so)
    assert xem["ok"], xem
    assert any("nhóm" in str(c).lower() or "bo_nhom" in str(c)
               for c in xem["data"]["canh_bao"]), xem
    assert api.import_so_nhap(so)["ok"]
    assert chup_csdl(api)["students"][0][2] == ""


def test_ma_tu_sinh_trung_ma_cu_khac_hoa_thuong_goi_y_dung_ma_that(api):
    """CSDL có CLB_TIN "Tin cu". Sổ đổi tên thành "Tin", bỏ trống Mã CLB: lỗi
    phải gợi ý điền CLB_TIN (mã thật), không phải clb_tin (sẽ tạo CLB mới)."""
    assert api.create_or_update_club("CLB_TIN", "Tin cu", 5, 0, "")["ok"]
    so = b64_sach([{"club_id": "", "name": "CLB Tin", "capacity": 5}], [])
    loi = api.xem_truoc_so_nhap(so)["data"]["loi"]
    assert [e["code"] for e in loi] == ["so_nhap_ma_tu_sinh_trung_clb_da_co"]
    assert loi[0]["params"]["ma_cu"] == "CLB_TIN"
    assert loi[0]["params"]["ten_cu"] == "Tin cu"
    # Làm theo gợi ý: điền CLB_TIN thì đổi tên đúng CLB cũ, không đẻ CLB mới.
    so = b64_sach([{"club_id": "CLB_TIN", "name": "CLB Tin", "capacity": 5}], [])
    assert api.import_so_nhap(so)["ok"]
    assert [c[:2] for c in chup_csdl(api)["clubs"]] == [("CLB_TIN", "CLB Tin")]
