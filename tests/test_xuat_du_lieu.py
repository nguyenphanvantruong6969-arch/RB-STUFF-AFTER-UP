# -*- coding: utf-8 -*-
"""Xuất "Dữ liệu đầu vào hiện tại" — một Sổ nhập CLB phải NẠP LẠI ĐƯỢC y nguyên.

Lời hứa của tính năng: sửa dữ liệu trong phần mềm (tick, điểm, nguyện
vọng, CLB) rồi xuất ra thì sổ ấy dựng lại ĐÚNG CSDL đó. Test mạnh nhất là
vòng tròn: xuất -> xoá sạch -> nạp lại qua đúng cửa giao diện dùng
(`import_so_nhap`) -> so từng bảng -> chạy sắp xếp cùng hạt giống -> so
từng dòng kết quả.
"""

import base64
import os
import sqlite3

import openpyxl

import so_nhap


def dung_truong(api):
    """Trường hai buổi, có nhãn dự trữ, có tick chưa chấm, có tên khó."""
    clb = [
        ("clb_a", "CLB A", 2, 1, "chinh_sach", "thu_2"),
        ("clb_b", 'CLB "Vì cộng đồng"', 3, 0, "", "thu_2"),
        ("clb_c", "CLB C, nấu ăn", 2, 0, "", "thu_4"),
    ]
    for c in clb:
        assert api.create_or_update_club(*c)["ok"]
    hs = [
        ("HS01", "Ngô Văn An", "", ["clb_a", "clb_c"], ["clb_a", "clb_b", "clb_c"]),
        ("HS02", "Lê Thị Bình", "chinh_sach", ["clb_a"], ["clb_a", "clb_c"]),
        ("HS03", "=1+1", "", ["clb_b", "clb_c"], ["clb_b", "clb_a"]),
        ("HS04", "Đỗ Minh Khoa", "", [], ["clb_c"]),
        ("HS05", "Chưa nhập gì", "", [], []),
    ]
    for ma, ten, nhom, thi, nv in hs:
        assert api.create_student_if_missing(ma, ten)["ok"]
        if nhom:
            assert api.set_student_reserve_group(ma, nhom)["ok"]
        if thi:
            assert api.submit_test_selection(ma, thi)["ok"]
        if nv:
            assert api.submit_preferences(ma, nv)["ok"]
    # clb_c: HS03 được chấm, HS01 tick nhưng CHƯA chấm.
    assert api.submit_club_scores("clb_a", [
        {"student_id": "HS01", "score": 8.5}, {"student_id": "HS02", "score": 7}])["ok"]
    assert api.submit_club_scores("clb_b", [{"student_id": "HS03", "score": 9}])["ok"]
    assert api.submit_club_scores("clb_c", [{"student_id": "HS03", "score": 6.25}])["ok"]


def chup(api):
    conn = sqlite3.connect(api.db_path)
    try:
        q = lambda sql: sorted(conn.execute(sql).fetchall())
        return {
            "students": q("SELECT student_id, name, COALESCE(reserve_group,'') FROM students "
                          "WHERE student_id IN (SELECT student_id FROM club_test_selection "
                          "UNION SELECT student_id FROM preferences)"),
            "clubs": q("SELECT club_id, name, capacity, reserve_capacity, "
                       "COALESCE(reserve_group,''), COALESCE(buoi,'') FROM clubs"),
            "tick": q("SELECT * FROM club_test_selection"),
            "diem": q("SELECT * FROM club_scores"),
            "nv": q("SELECT * FROM preferences"),
        }
    finally:
        conn.close()


def ket_qua(api):
    return sorted(
        (r["student_id"], r["buoi"], r["club_id"], r["matched_tier"])
        for r in api.get_match_results("")["data"])


def nap(api, path):
    with open(path, "rb") as f:
        return api.import_so_nhap(base64.b64encode(f.read()).decode())


def test_vong_tron_xuat_xoa_nap_lai_ra_dung_csdl_va_dung_ket_qua(api):
    dung_truong(api)
    assert api.run_pipeline(42)["ok"]
    truoc, kq_truoc = chup(api), ket_qua(api)

    res = api.export_du_lieu_dau_vao()
    assert res["ok"], res["errors"]
    d = res["data"]
    assert d["path"].endswith(".xlsx") and os.path.isfile(d["path"])
    assert (d["n_clubs"], d["n_tick"], d["n_scores"], d["n_preferences"]) == (3, 4, 3, 8)
    # HS01 tick clb_a, clb_c và xếp nguyện vọng cả hai: vào sổ. HS02 tick
    # clb_a, xếp clb_a: vào sổ. HS03 tick clb_c nhưng KHÔNG xếp clb_c.
    assert d["n_thi_ngoai_nv"] == 1

    assert api.reset_data("tat_ca", "XOA")["ok"]
    r = nap(api, d["path"])
    assert r["ok"], r["errors"]
    assert r["data"]["warnings"] == [], r["data"]["warnings"]

    sau = chup(api)
    # Khác DUY NHẤT ở lượt thi HS03–clb_c (không có trong nguyện vọng), đúng
    # con số n_thi_ngoai_nv đã báo trước.
    bo = lambda rows: [x for x in rows if x[:2] != ("HS03", "clb_c")]  # noqa: E731
    assert sau["clubs"] == truoc["clubs"] and sau["nv"] == truoc["nv"]
    assert sau["students"] == truoc["students"]
    assert sau["tick"] == bo(truoc["tick"]) and sau["diem"] == bo(truoc["diem"])

    assert api.run_pipeline(42)["ok"]
    # HS03 không xếp clb_c nên lượt thi đó không đổi được kết quả. HS05 (chưa
    # nhập gì) nay CÓ trong sổ, nên kết quả trùng khít.
    assert ket_qua(api) == kq_truoc


def test_so_xuat_ra_dung_bo_cuc_so_nhap(api):
    dung_truong(api)
    p = api.export_du_lieu_dau_vao()["data"]["path"]
    wb = openpyxl.load_workbook(p, read_only=True)
    assert wb.sheetnames[:2] == [so_nhap.SHEET_CLB, so_nhap.SHEET_HS]
    du = so_nhap.doc_so_nhap(wb)
    assert du["loi"] == []
    # Mã CLB được ghi ra, nên nạp lại giữ đúng mã cũ.
    assert [c["club_id"] for c in du["clubs"]] == ["clb_a", "clb_b", "clb_c"]
    assert {c["buoi"] for c in du["clubs"]} == {"thu_2", "thu_4"}


def test_tick_chua_cham_ghi_chu_thi_va_ten_bi_chan_cong_thuc(api):
    dung_truong(api)
    p = api.export_du_lieu_dau_vao()["data"]["path"]
    ws = openpyxl.load_workbook(p)[so_nhap.SHEET_HS]
    dong = {r[0].value: r for r in ws.iter_rows(min_row=2) if r[0].value}
    hs01 = [o.value for o in dong["HS01"]]
    # HS01 xếp clb_a, clb_b, clb_c; tick clb_a (8.5) và clb_c (chưa chấm).
    assert hs01[3:9] == ["CLB A", 8.5, 'CLB "Vì cộng đồng"', None, "CLB C, nấu ăn", "thi"]
    o_ten = dong["HS03"][1]
    assert o_ten.value == "=1+1" and o_ten.data_type == "s"   # chữ, không phải công thức


def test_em_chua_nhap_gi_van_co_dong_trong_so(api):
    dung_truong(api)
    d = api.export_du_lieu_dau_vao()["data"]
    assert d["n_hoc_sinh"] == 5            # kể cả HS05


def test_diem_khong_con_tick_khong_duoc_xuat(api):
    dung_truong(api)
    conn = sqlite3.connect(api.db_path)
    conn.execute("DELETE FROM club_test_selection WHERE student_id='HS02'")
    conn.commit()
    conn.close()
    p = api.export_du_lieu_dau_vao()["data"]["path"]
    du = so_nhap.doc_so_nhap(openpyxl.load_workbook(p, read_only=True))
    hs02 = next(s for s in du["students"] if s["student_id"] == "HS02")
    assert not any(x["thi"] for x in hs02["nv"])


def test_xuat_khong_ghi_de(api):
    dung_truong(api)
    a = api.export_du_lieu_dau_vao()["data"]["path"]
    b = api.export_du_lieu_dau_vao()["data"]["path"]
    assert a != b and os.path.isfile(a) and os.path.isfile(b)


def test_tai_so_nhap_mau_dien_san_clb_dang_co(api):
    trong = api.tao_so_nhap_mau()
    assert trong["ok"] and trong["data"]["n_clubs"] == 0
    dung_truong(api)
    r = api.tao_so_nhap_mau()
    assert r["ok"], r
    du = so_nhap.doc_so_nhap(openpyxl.load_workbook(r["data"]["path"], read_only=True))
    assert du["loi"] == [] and du["students"] == []
    assert [c["name"] for c in du["clubs"]] == ["CLB A", 'CLB "Vì cộng đồng"', "CLB C, nấu ăn"]


def test_ten_clb_dung_nhau_van_xuat_ra_so_nap_lai_duoc(api):
    """Form CLB cho phép "Toán" và "toán"; sổ thì không (ô NV chọn theo tên).
    Sổ xuất ra phải đổi tên CLB sau để vẫn nạp lại được, và nói ra điều đó."""
    for ma, ten in (("toan_a", "Toán"), ("toan_b", "toán"), ("rong", "")):
        assert api.create_or_update_club(ma, ten, 3, 0, "")["ok"]
    assert api.create_student_if_missing("HS1", "An")["ok"]
    assert api.submit_preferences("HS1", ["toan_b", "toan_a"])["ok"]
    d = api.export_du_lieu_dau_vao()["data"]
    assert d["n_ten_doi"] == 1

    assert api.reset_data("tat_ca", "XOA")["ok"]
    r = nap(api, d["path"])
    assert r["ok"], r["errors"]
    conn = sqlite3.connect(api.db_path)
    try:
        assert sorted(conn.execute("SELECT club_id FROM clubs")) == [
            ("rong",), ("toan_a",), ("toan_b",)]
        assert conn.execute("SELECT club_id FROM preferences ORDER BY rank").fetchall() == [
            ("toan_b",), ("toan_a",)]
    finally:
        conn.close()


def test_ma_chi_khac_hoa_thuong_van_xuat_va_nap_lai_duoc(api):
    """Phần mềm phân biệt hoa/thường ở mã: hs01/HS01, clb_a/CLB_A là bốn thứ
    khác nhau. Sổ xuất từ một CSDL như thế phải dựng lại đúng CSDL đó."""
    for ma, ten in (("clb_a", "CLB A thường"), ("CLB_A", "CLB A hoa")):
        assert api.create_or_update_club(ma, ten, 3, 0, "")["ok"]
    for sid, nv in (("hs01", ["clb_a"]), ("HS01", ["CLB_A"])):
        assert api.create_student_if_missing(sid, sid)["ok"]
        assert api.submit_preferences(sid, nv)["ok"]
    truoc = chup(api)
    d = api.export_du_lieu_dau_vao()["data"]
    assert api.reset_data("tat_ca", "XOA")["ok"]
    r = nap(api, d["path"])
    assert r["ok"], r["errors"]
    sau = chup(api)
    assert (sau["clubs"], sau["nv"], sau["students"]) == (truoc["clubs"], truoc["nv"], truoc["students"])
