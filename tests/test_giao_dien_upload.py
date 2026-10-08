"""Test GIAO DIỆN nạp Sổ nhập CLB bằng trình duyệt thật (Playwright + Chromium).

Vùng nạp tệp là chỗ người dùng thao tác nhiều nhất. Giờ nó chỉ nhận MỘT
tệp: Sổ nhập CLB (so_nhap.py). Luồng phải là: thả sổ -> phần mềm đọc thử,
hiện tóm tắt (hoặc danh sách lỗi kèm sheet, dòng) -> bấm Nhập sổ -> xong.

File này chạy giao diện thật để bảo đảm luồng đó không vỡ trong im
lặng. Máy nào không có `playwright` hoặc không có Chromium thì tự bỏ
qua — nó là lớp kiểm tra thêm, không phải điều kiện bắt buộc để chạy
được bộ test.
"""

import io
import os

import openpyxl
import pytest

import so_nhap

pytest.importorskip("playwright", reason="chưa cài playwright")
from playwright.sync_api import sync_playwright  # noqa: E402

import browser_host  # noqa: E402

GOC = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MAU = os.path.join(GOC, "mau_csv")

# Chromium do Playwright cài sẵn; đường dẫn này là của môi trường phát
# triển. Không có thì bỏ qua cả file.
from _trinh_duyet import CHROMIUM  # noqa: E402


@pytest.fixture
def trang(api):
    """Mở index.html thật, backend là PipelineAPI thật qua browser_host."""
    if CHROMIUM is None:
        pytest.skip("khong tim thay Chromium")
    url = browser_host.serve(api, GOC, "index.html", open_browser=False)
    with sync_playwright() as pw:
        br = pw.chromium.launch(executable_path=CHROMIUM, args=["--no-sandbox"])
        page = br.new_page(bypass_csp=True)  # CSP chan eval cua wait_for_function; test CSP rieng o test_giao_dien_sua_loi_p1
        loi = []
        page.on("pageerror", lambda e: loi.append(str(e)))
        page.goto(url)
        page.wait_for_selector("#dropZone")
        yield page, loi
        br.close()


def mau(ten):
    return os.path.join(MAU, ten)


SO_VI_DU = os.path.join(MAU, "vi_du_day_du", "SO_NHAP_CLB_vi_du.xlsx")
CHO = ("!document.getElementById('importActions').hidden"
       " && !document.getElementById('btnImportAll').hidden")


def tha(page, duong_dan):
    page.locator("#fileAny").set_input_files([str(duong_dan)])


def test_tha_so_hien_tom_tat_roi_nhap_vao_dung_bang(trang, api):
    page, loi = trang
    tha(page, SO_VI_DU)
    page.wait_for_function(
        "document.querySelector('.queue-row .queue-detail')"
        " && document.querySelector('.queue-row .queue-detail').innerText.indexOf('Sẵn sàng') === 0")
    chi_tiet = page.locator(".queue-row .queue-detail").inner_text()
    assert "9 CLB" in chi_tiet and "3 buổi" in chi_tiet and "24 học sinh" in chi_tiet
    # Đọc thử KHÔNG ghi gì.
    assert api.list_clubs_admin()["data"] == []

    page.locator("#btnImportAll").click()
    page.wait_for_selector(".queue-row.is-done", timeout=20000)
    assert "Xong: 9 CLB, 24 học sinh" in page.locator(".queue-row .queue-detail").inner_text()
    assert page.locator("#btnImportAll").is_hidden()   # không nhập lại lần hai

    st = api.get_student_entry_state("HS001")["data"]
    assert st["ranked_clubs"][:2] == ["clb_tinhoc", "clb_bongro"]
    assert sorted(st["tested_clubs"]) == ["clb_robotics", "clb_tienganh", "clb_tinhoc"]
    assert len(api.list_clubs_admin()["data"]) == 9
    assert not loi, loi


def test_so_co_loi_liet_ke_tung_loi_va_khoa_nut_nhap(trang, api, tmp_path):
    wb = openpyxl.load_workbook(SO_VI_DU)
    wb[so_nhap.SHEET_HS]["D2"] = "CLB Không Có"
    wb[so_nhap.SHEET_HS]["E3"] = "chín"
    hong = tmp_path / "so_loi.xlsx"
    wb.save(hong)

    page, loi = trang
    tha(page, hong)
    page.wait_for_selector(".queue-row.is-unknown")
    assert "2 lỗi" in page.locator(".queue-row .queue-detail").inner_text()
    ds = page.locator(".so-nhap-loi li").all_inner_texts()
    assert len(ds) == 2
    assert "dòng 2" in ds[0] and "CLB Không Có" in ds[0]
    assert "dòng 3" in ds[1] and "chín" in ds[1]
    assert page.locator("#btnImportAll").is_hidden()
    assert api.list_clubs_admin()["data"] == []
    assert not loi, loi


def test_so_sach_nhung_doi_du_lieu_cu_thi_bao_truoc_ma_van_cho_nhap(trang, api, tmp_path):
    """Nhóm ưu tiên để trống sẽ bỏ nhóm của em: báo vàng, không chặn."""
    api.create_or_update_club("clb_a", "CLB A", 5, 1, "chinh_sach")
    api.create_student_if_missing("HS1", "An")
    api.set_student_reserve_group("HS1", "chinh_sach")
    p = tmp_path / "so.xlsx"
    so_nhap.ghi_so_nhap(str(p), [{"club_id": "clb_a", "name": "CLB A", "capacity": 5,
                                  "reserve_capacity": 1, "reserve_group": "chinh_sach"}],
                        [{"student_id": "HS1", "name": "An", "reserve_group": "",
                          "nv": [{"club_id": "clb_a", "thi": False, "diem": ""}]}])
    page, loi = trang
    tha(page, p)
    page.wait_for_selector(".so-nhap-canh-bao li")
    assert "bỏ nhóm" in page.locator(".so-nhap-canh-bao").inner_text()
    assert page.locator("#btnImportAll").is_visible()
    assert not loi, loi


def test_loi_phia_phan_mem_thi_giu_tep_de_nhap_lai(trang, api, monkeypatch):
    """CSDL bận không phải lỗi của sổ: không bảo người dùng sửa Excel, giữ
    tệp và nút Nhập để thử lại."""
    import i18n_errors
    lan = {"n": 0}
    goc = api.import_so_nhap

    def lan_dau_hong(b64):
        lan["n"] += 1
        if lan["n"] == 1:
            return i18n_errors.phan_hoi_loi(i18n_errors.err("loi_nap_so_nhap", detail="locked"))
        return goc(b64)
    monkeypatch.setattr(api, "import_so_nhap", lan_dau_hong)
    page, loi = trang
    tha(page, SO_VI_DU)
    page.wait_for_function(CHO)
    page.locator("#btnImportAll").click()
    page.wait_for_selector(".toast.is-error")
    assert page.locator(".so-nhap-loi").count() == 0
    assert page.locator("#btnImportAll").is_visible()
    page.locator("#btnImportAll").click()
    page.wait_for_selector(".queue-row.is-done", timeout=20000)
    assert not loi, loi


def test_tep_csv_bi_tu_choi_kem_loi_chi_duong(trang, api):
    page, loi = trang
    tha(page, mau("05_danh_sach_club.csv"))
    page.wait_for_selector(".queue-row.is-unknown")
    assert "Tải sổ nhập mẫu" in page.locator(".queue-row .queue-detail").inner_text()
    assert page.locator("#btnImportAll").is_hidden()
    assert api.list_clubs_admin()["data"] == []
    assert not loi, loi


def test_excel_khong_phai_so_nhap_bi_tu_choi(trang, tmp_path):
    wb = openpyxl.Workbook()
    wb.active.append(["club_id", "name", "capacity"])
    wb.active.append(["clb_a", "A", 3])
    p = tmp_path / "ba_tep_cu.xlsx"
    wb.save(p)
    page, loi = trang
    tha(page, p)
    page.wait_for_selector(".queue-row.is-unknown")
    assert "không phải Sổ nhập CLB" in page.locator(".queue-row .queue-detail").inner_text()
    assert not loi, loi


def test_bo_tep_xoa_luon_phan_hoi_cua_lan_nhap_truoc(trang):
    """Để lại kết quả cũ thì người dùng tưởng đó là của tệp đang có."""
    page, loi = trang
    tha(page, SO_VI_DU)
    page.wait_for_function("!document.getElementById('btnImportAll').hidden"
                           " && !document.getElementById('importActions').hidden")
    page.locator("#btnImportAll").click()
    page.wait_for_selector(".queue-row.is-done", timeout=20000)
    assert page.locator("#feedbackImportAll").inner_text().strip()

    page.locator("#btnClearQueue").click()
    page.wait_for_function(
        "document.getElementById('feedbackImportAll').textContent.trim() === ''"
    )
    assert page.locator("#importQueue").is_hidden()
    assert not loi, loi


def test_tai_so_nhap_mau_tao_tep_trong_thu_muc_tai_ve(trang, api, monkeypatch):
    mo = []
    monkeypatch.setattr(api, "mo_thu_muc", lambda p: mo.append(p) or {"ok": True})
    api.create_or_update_club("clb_a", "CLB A", 5, 0, "")
    page, loi = trang
    page.locator("#btnTaiSoMau").click()
    page.wait_for_function(
        "document.getElementById('feedbackImportAll').textContent.indexOf('SO_NHAP_CLB') !== -1")
    cau = page.locator("#feedbackImportAll").inner_text()
    assert "1 CLB" in cau
    duong = cau.split("→")[-1].strip()
    wb = openpyxl.load_workbook(duong, read_only=True)
    assert so_nhap.doc_so_nhap(wb)["clubs"][0]["name"] == "CLB A"
    assert not loi, loi


def test_file_excel_hong_bao_loi_chu_khong_sap_giao_dien(trang, tmp_path):
    hong = tmp_path / "khong_phai_excel.xlsx"
    hong.write_bytes(b"day khong phai file excel")
    page, loi = trang
    tha(page, hong)
    page.wait_for_selector(".queue-row.is-unknown")
    assert page.locator(".queue-row .queue-detail").inner_text().strip()
    assert not loi, loi


def test_so_trong_bo_nho_doc_duoc(tmp_path):
    """Phép thử tầng Python cho tiện ích tạo sổ dùng trong các test giao diện."""
    buf = io.BytesIO()
    so_nhap.ghi_so_nhap(buf, [{"name": "CLB A", "capacity": 2}], [])
    assert so_nhap.doc_so_nhap(openpyxl.load_workbook(buf))["loi"] == []
