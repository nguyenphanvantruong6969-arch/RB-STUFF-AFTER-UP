# -*- coding: utf-8 -*-
"""Skill `tao-google-form-clb`: tệp Excel -> TAO_GOOGLE_FORM.gs + hướng dẫn.

Ba phép kiểm:
  1. vòng tròn: tệp mẫu (bộ đề ví dụ) -> .gs giống hệt tệp .gs trong kho;
  2. soát: điền sai thì không ghi tệp nào, và báo đúng dòng;
  3. chạy thật đoạn mã sinh ra trên bản giả lập FormApp (cần Node): CLB không
     thi không có trong câu tick, rẽ nhánh chỉ hiện đề CLB được chọn.
"""

import importlib.util
import json
import os
import shutil
import subprocess

import pytest

openpyxl = pytest.importorskip("openpyxl")

GOC = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCRIPT = os.path.join(GOC, ".claude", "skills", "tao-google-form-clb", "scripts", "tao_form.py")


@pytest.fixture(scope="module")
def tf():
    spec = importlib.util.spec_from_file_location("tao_form", SCRIPT)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


def _sua(duong, trang, sua):
    """sua(ws): chỉnh một trang của tệp Excel rồi lưu."""
    wb = openpyxl.load_workbook(duong)
    sua(wb[trang])
    wb.save(duong)


def test_vong_tron_ra_dung_tep_gs_trong_kho(tf, tmp_path):
    mau = str(tmp_path / "MAU.xlsx")
    assert tf.main(["mau", "--ra", mau]) == 0
    assert tf.main(["tao", "--vao", mau, "--ra", str(tmp_path / "ra")]) == 0
    with open(os.path.join(GOC, "mau_forms_thi_clb", "TAO_GOOGLE_FORM.gs"), encoding="utf-8") as f:
        goc = f.read()
    with open(str(tmp_path / "ra" / "TAO_GOOGLE_FORM.gs"), encoding="utf-8") as f:
        assert f.read() == goc
    with open(str(tmp_path / "ra" / "HUONG_DAN_TUNG_BUOC.md"), encoding="utf-8") as f:
        hd = f.read()
    assert "—" not in hd and "–" not in hd
    assert "212 câu hỏi" in hd and "taoBieuMau" in hd and "chamTheoCLB" in hd


def test_dien_sai_thi_khong_ghi_gi(tf, tmp_path):
    mau = str(tmp_path / "MAU.xlsx")
    tf.main(["mau", "--ra", mau])

    def hong(ws):
        ws["C3"] = "clb_covua"          # CLB dòng 3 trùng mã với dòng 2
        ws["A4"] = "thu_9"              # buổi không hợp lệ

    def hong_de(ws):
        ws["G2"] = "e"                  # đáp án ngoài a..d
        ws["D3"] = ws["C3"].value       # hai lựa chọn giống nhau
    _sua(mau, "CLB", hong)
    _sua(mau, "De_thi", hong_de)
    ra = tmp_path / "ra"
    with pytest.raises(tf.LoiDuLieu) as e:
        tf.tao(mau, str(ra))
    s = str(e.value)
    assert "CLB dòng 3" in s and "bị trùng" in s
    assert "CLB dòng 4" in s and "thu_9" in s
    assert "De_thi dòng 2" in s and "dap_an" in s
    assert "De_thi dòng 3" in s and "giống nhau" in s
    assert not ra.exists()
    assert tf.main(["tao", "--vao", mau, "--ra", str(ra)]) == 1


def test_qua_10_clb_mot_buoi_va_tran_qua_5(tf, tmp_path):
    mau = str(tmp_path / "MAU.xlsx")
    tf.main(["mau", "--ra", mau])

    def them(ws):
        ws.append(["thu_2", "Thứ Hai", "clb_them", "CLB Thêm"])

    def cai(ws):
        ws["B4"] = 6
    _sua(mau, "CLB", them)
    _sua(mau, "Cai_dat", cai)
    with pytest.raises(tf.LoiDuLieu) as e:
        tf.tao(mau, str(tmp_path / "ra"))
    assert "thu_2 có 11 CLB" in str(e.value) and "toi_da_clb_thi_moi_buoi" in str(e.value)


def _bo_de_nho(tf, tmp_path, tran=2):
    """Hai buổi: Thứ Hai có 3 CLB (1 CLB không thi), Thứ Tư chỉ có CLB không thi."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Cai_dat"
    for r in (["khoa", "gia_tri"], ["tieu_de", "Form thử"], ["diem_moi_cau", 4],
              ["toi_da_clb_thi_moi_buoi", tran]):
        ws.append(r)
    c = wb.create_sheet("CLB")
    c.append(tf.COT_CLB)
    for r in (["thu_2", "", "clb_a", "CLB A"], ["thu_2", "", "clb_b", "CLB B (mới)"],
              ["thu_2", "", "clb_c", "CLB C"], ["thu_4", "", "clb_d", "CLB D"]):
        c.append(r)
    d = wb.create_sheet("De_thi")
    d.append(tf.COT_DE)
    d.append(["clb_a", "1 + 1 = ?", "2", "3", "", "", "a"])
    d.append(["clb_c", "Thủ đô Việt Nam?", "Huế", "Hà Nội", "Đà Nẵng", "", "B"])
    d.append(["clb_c", "5 x 5 = ?", "10", "25", "", "", "b"])
    p = str(tmp_path / "nho.xlsx")
    wb.save(p)
    return p


def test_clb_khong_thi_va_buoi_khong_co_ai_thi(tf, tmp_path):
    du_lieu, canh_bao = tf.tao(_bo_de_nho(tf, tmp_path), str(tmp_path / "ra"))
    assert [d["ten"] for d in du_lieu] == ["Thứ Hai", "Thứ Tư"]      # tên buổi tự điền
    assert [len(c["de"]) for c in du_lieu[0]["clubs"]] == [1, 0, 2]
    assert du_lieu[0]["clubs"][2]["de"][0]["dung"] == "Hà Nội"         # "B" viết hoa vẫn nhận
    noi = " | ".join(canh_bao)
    assert "CLB B (mới)" in noi and "Thứ Tư không có CLB nào tổ chức thi" in noi
    # 2 + Thứ Hai (đăng ký + 3 Top + tick + trước + 1 tiếp theo + 3 câu đề) + Thứ Tư (đăng ký + 1 Top)
    assert tf.dem_cau(du_lieu) == 2 + (1 + 3 + 2 + 1 + 3) + (1 + 1)


MOCK = r"""
const fs = require('fs');
const items = []; const SUBMIT = 'SUBMIT';
function base(type) { const o = {type, title: '', choices: [], values: [],
  setTitle(t) { this.title = t; return this; }, setRequired() { return this; }, setHelpText() { return this; },
  setPoints(p) { this.points = p; return this; }, setRows(r) { this.rows = r; return this; },
  setColumns(c) { this.cols = c; return this; }, setChoiceValues(v) { this.values = v; return this; },
  setValidation(v) { this.val = v; return this; },
  createChoice(v, x) { return {v, nav: (x === true || x === false) ? undefined : x, correct: x === true}; },
  setChoices(c) { this.choices = c; return this; }}; items.push(o); return o; }
const form = {setIsQuiz() { return this; }, setDescription() { return this; }, setShuffleQuestions() { return this; },
  setProgressBar() { return this; }, setShowLinkToRespondAgain() { return this; }, setPublishingSummary() { return this; },
  setConfirmationMessage() { return this; }, setPublished() { return this; }, getId: () => 'F',
  addTextItem: () => base('text'), addPageBreakItem: () => base('page'), addMultipleChoiceItem: () => base('mc'),
  addGridItem: () => base('grid'), addCheckboxItem: () => base('cb'), addListItem: () => base('list'), getEditUrl: () => 'e', getPublishedUrl: () => 'p'};
function vb() { return {setHelpText() { return this; }, requireLimitOneResponsePerColumn() { return this; },
  requireSelectAtMost(n) { this.max = n; return this; }, build() { return {max: this.max}; }}; }
let title = null;
global.FormApp = {create: (t) => { title = t; return form; }, PageNavigationType: {SUBMIT},
  createGridValidation: vb, createCheckboxValidation: vb};
global.Logger = {log() {}};
global.PropertiesService = {getScriptProperties: () => ({setProperty() {}})};
eval(fs.readFileSync(process.argv[2], 'utf8').replace(/^const /gm, 'var ')); taoBieuMau();
const pages = [[]], pageOf = new Map();
items.forEach(it => { if (it.type === 'page') { pages.push([]); pageOf.set(it, pages.length - 1); } else pages[pages.length - 1].push(it); });
function walk(pick) { let p = 0, seen = []; while (p < pages.length) { let nav = null;
  for (const q of pages[p]) { seen.push(q.title.match(/^\[([^\]]+)\]/)[1]);
    if (q.type === 'mc') { const v = pick(q); const c = q.choices.find(x => x.v === v); if (!c) throw new Error(v); if (c.nav) nav = c.nav; } }
  if (nav === SUBMIT) break; p = nav ? pageOf.get(nav) : p + 1; } return seen; }
const cb = items.filter(i => i.type === 'cb');
const firstSel = items.find(i => i.title.startsWith('[thi-thu_2]'));
const s = walk(q => { const t = q.title;
  if (t.startsWith('[di-')) return CO_BUOI;
  if (t.startsWith('[thi-thu_2]')) return 'CLB C (clb_c)';
  return q.choices[0].v; });
const lists = items.filter(i => i.type === 'list');
console.log(JSON.stringify({title, n: items.filter(i => i.type !== 'page').length,
  lists: lists.map(l => [l.title, l.values]),
  cb: cb.map(c => [c.values, c.val.max]), first: firstSel.choices.map(c => c.v),
  points: items.filter(i => i.points).map(i => i.points), walk: s}));
"""


@pytest.mark.skipif(shutil.which("node") is None, reason="cần Node.js để chạy bản giả lập FormApp")
def test_ma_sinh_ra_chay_dung_tren_ban_gia_lap(tf, tmp_path):
    du_lieu, _ = tf.tao(_bo_de_nho(tf, tmp_path), str(tmp_path / "ra"))
    js = tmp_path / "mock.js"
    js.write_text(MOCK, encoding="utf-8")
    out = subprocess.run(["node", str(js), str(tmp_path / "ra" / "TAO_GOOGLE_FORM.gs")],
                         capture_output=True, text=True, encoding="utf-8", check=True).stdout
    kq = json.loads(out)
    assert kq["title"] == "Form thử"
    assert kq["n"] == tf.dem_cau(du_lieu)
    # CLB B không thi: không có trong câu tick và trang chọn; Thứ Tư không có câu tick.
    assert kq["cb"] == [[["CLB A (clb_a)", "CLB C (clb_c)"], 2]]
    assert kq["first"] == ["CLB A (clb_a)", "CLB C (clb_c)", "Em không thi CLB nào buổi này"]
    assert set(kq["points"]) == {4}
    # Chọn CLB C trước: không thấy đề CLB A, làm xong CLB C (cuối) là sang Thứ Tư.
    assert kq["walk"] == ["student_id", "name", "di-thu_2", "thu_2-top-1", "thu_2-top-2", "thu_2-top-3",
                          "chon-thu_2", "thi-thu_2", "clb_c-1", "clb_c-2", "di-thu_4", "thu_4-top-1"]
    # Mỗi Top một danh sách thả xuống, "(Bỏ trống)" đứng đầu.
    assert kq["lists"][0] == ["[thu_2-top-1] Top 1",
                              ["(Bỏ trống)", "CLB A (clb_a)", "CLB B (mới) (clb_b)", "CLB C (clb_c)"]]
    assert len(kq["lists"]) == 3 + 1


CHAM = r"""// Bản giả lập tối thiểu để chạy chamTheoCLB trên các phiếu cho trước.
const fs=require('fs'); const [,, gsPath, outPath, kichBan]=process.argv;
function base(){ const o={setTitle(){return this},setRequired(){return this},setHelpText(){return this},setPoints(){return this},setRows(){return this},setColumns(){return this},setChoiceValues(){return this},setValidation(){return this},createChoice(){return {}},setChoices(){return this}}; return o; }
const form={setIsQuiz(){return this},setDescription(){return this},setShuffleQuestions(){return this},setProgressBar(){return this},setShowLinkToRespondAgain(){return this},setPublishingSummary(){return this},setConfirmationMessage(){return this},setPublished(){return this},getId:()=>'F',addTextItem:base,addPageBreakItem:base,addMultipleChoiceItem:base,addGridItem:base,addCheckboxItem:base,addListItem:base,getEditUrl:()=>'e',getPublishedUrl:()=>'p'};
let responses=[];
function vb(){ return {setHelpText(){return this},requireLimitOneResponsePerColumn(){return this},requireSelectAtMost(){return this},build(){return {}}}; }
global.FormApp={create:()=>form,openById:()=>({getResponses:()=>responses}),PageNavigationType:{SUBMIT:'S'},createGridValidation:vb,createCheckboxValidation:vb};
global.Logger={log:m=>{}}; global.PropertiesService={getScriptProperties:()=>({getProperty:()=>'F',setProperty(){}})};
global.Utilities={formatDate:()=>'T'};
const sheets=[]; function sheet(n){const s={n,setName(x){this.n=x;return this},getRange:()=>({setNumberFormat(){return this},setValues(v){s.data=v;return this},setFontWeight(){return this}}),setFrozenRows(){}};sheets.push(s);return s;}
global.SpreadsheetApp={flush(){},create:()=>{const f=sheet('S');return {getId:()=>'x',getSheets:()=>[f],insertSheet:n=>sheet(n),getUrl:()=>'u'};}};
global.DriveApp={createFolder:()=>({getUrl:()=>'f',createFile(){}}),getFileById:()=>({moveTo(){},setTrashed(){}})};
global.UrlFetchApp={fetch:()=>({getResponseCode:()=>200,getBlob:()=>({setName(){return this}})})}; global.ScriptApp={getOAuthToken:()=>'t'};
eval(fs.readFileSync(gsPath,'utf8').replace(/^const /gm,'var '));
function resp(a){ return {getTimestamp:()=>new Date(0),getRespondentEmail:()=>'',getItemResponses:()=>Object.entries(a).map(([k,v])=>({getItem:()=>({getTitle:()=>'['+k+'] x'}),getResponse:()=>v,getScore:()=>null}))}; }
const d=BUOI[0], C=d.clubs, lab=c=>c.name+' ('+c.id+')';
const off={'di-thu_3':KHONG_BUOI,'di-thu_4':KHONG_BUOI,'di-thu_5':KHONG_BUOI,'di-thu_6':KHONG_BUOI};
function lam(a,c,n){ c.de.forEach((q,j)=>a[c.id+'-'+(j+1)]= j<n?q.dung:'sai'); }
// Phiếu kiểu thả xuống: Top1 Nhiếp ảnh, Top2 (Bỏ trống), Top3 Cờ vua, Top4 Nhiếp ảnh (trùng), Top5 Văn học
const a=Object.assign({},off,{student_id:'HS1',name:'An','di-thu_2':CO_BUOI,
  'thu_2-top-1':lab(C[2]),'thu_2-top-2':BO_TRONG,'thu_2-top-3':lab(C[0]),'thu_2-top-4':lab(C[2]),'thu_2-top-5':lab(C[1]),
  'chon-thu_2':[lab(C[0]),lab(C[2])],'thi-thu_2':lab(C[0])});
a['thi-sau-'+C[0].id]=lab(C[2]); a['thi-sau-'+C[2].id]=DUNG; lam(a,C[0],1); lam(a,C[2],2);
// Phiếu kiểu lưới cũ (dòng = CLB, "Hạng n") và câu Có/Không cũ
const b=Object.assign({},off,{student_id:'HS2',name:'Bình','di-thu_2':CO_BUOI,thu_2:C.map((c,i)=>'Hạng '+(i+1)),'chon-thu_2':[lab(C[0])]});
b['thi-'+C[0].id]=CO; lam(b,C[0],2);
responses=[resp(a),resp(b)];
chamTheoCLB();
const by=n=>sheets.filter(s=>s.n===n).pop().data;
console.log('HS', JSON.stringify(by('2. Học sinh').map(r=>r.filter(x=>x!==''))));
console.log('CLB', JSON.stringify(by('1. CLB').slice(0,2)));
console.log('Cảnh báo', JSON.stringify(by('Cảnh báo').slice(1).map(r=>r[2])));
console.log('raw xep_hang HS1', JSON.stringify(by('RAW_PHIEU').filter(r=>r[6]==='xep_hang'&&r[3]==='HS1').map(r=>[r[5],r[8],r[11]])));

"""


@pytest.mark.skipif(shutil.which("node") is None, reason="cần Node.js để chạy bản giả lập FormApp")
def test_cham_top_tha_xuong_va_bieu_mau_cu(tmp_path):
    """Chấm phiếu kiểu thả xuống (bỏ trống Top 2, chọn trùng Nhiếp ảnh ở Top 1 và Top 4)
    và phiếu kiểu lưới cũ bằng chamTheoCLB trong tệp .gs của kho."""
    js = tmp_path / "cham.js"
    js.write_text(CHAM, encoding="utf-8")
    gs = os.path.join(GOC, "mau_forms_thi_clb", "TAO_GOOGLE_FORM.gs")
    out = subprocess.run(["node", str(js), gs, str(tmp_path / "x.json")],
                         capture_output=True, text=True, encoding="utf-8", check=True).stdout
    dong = {l.split(" ", 1)[0]: l.split(" ", 1)[1] for l in out.strip().splitlines()}
    # Sổ nhập CLB: NV theo hạng em chọn, điểm ngay cạnh CLB em đã thi.
    hs = json.loads(dong["HS"])
    # Không có cột Nhóm ưu tiên: biểu mẫu không hỏi nhóm, và cột vắng mặt
    # nghĩa là phần mềm giữ nhóm đã gán (cột có mà ô trống thì bỏ nhóm).
    assert hs[0][:4] == ["Mã HS", "Họ tên", "NV1", "Điểm 1"]
    assert hs[1] == ["HS1", "An", "CLB Nhiếp ảnh", 10, "CLB Cờ vua", 5, "CLB Văn học"]
    assert hs[2][:6] == ["HS2", "Bình", "CLB Cờ vua", 10, "CLB Văn học", "CLB Nhiếp ảnh"]
    assert json.loads(dong["CLB"]) == [
        ["Tên CLB", "Buổi", "Chỉ tiêu", "Suất ưu tiên", "Nhóm ưu tiên", "Mã CLB"],
        ["CLB Cờ vua", "thu_2", "", "", "", "clb_covua"]]
    canh_bao = " | ".join(json.loads(dong["Cảnh"].split(" ", 1)[1]))
    assert "bỏ trống Top 2" in canh_bao and "clb_nhiepanh ở cả Top 1 và Top 4, giữ Top 1" in canh_bao


@pytest.mark.skipif(shutil.which("node") is None, reason="cần Node.js để chạy bản giả lập FormApp")
def test_clb_cung_ten_o_hai_buoi_ra_hai_ten_rieng_trong_so(tmp_path):
    """Trường mở "CLB Cờ vua" cả hai buổi: sổ chọn CLB theo tên, nên CLB sau
    phải mang thêm tên buổi, không thì phần mềm từ chối cả sổ."""
    js = tmp_path / "cham.js"
    kich_ban = CHAM.replace(
        "responses=[resp(a),resp(b)];",
        "BUOI[1].clubs[0].name=BUOI[0].clubs[0].name; responses=[];")
    kich_ban += "console.log('TEN', JSON.stringify(by('1. CLB').slice(1).map(r=>r[0])));\n"
    js.write_text(kich_ban, encoding="utf-8")
    gs = os.path.join(GOC, "mau_forms_thi_clb", "TAO_GOOGLE_FORM.gs")
    out = subprocess.run(["node", str(js), gs, str(tmp_path / "x.json")],
                         capture_output=True, text=True, encoding="utf-8", check=True).stdout
    ten = json.loads(next(l for l in out.splitlines() if l.startswith("TEN ")).split(" ", 1)[1])
    assert len(ten) == len({t.lower() for t in ten}), ten
    assert any(t.endswith("(Thứ Ba)") for t in ten), ten
